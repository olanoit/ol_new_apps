# -*- coding: utf-8 -*-
"""Resumen mensual del CONAFOVICER (D.L. 21067).

La retención se descuenta en cada planilla semanal, pero el pago al
Banco de la Nación es **mensual y hasta el día 15 del mes siguiente**.
Este modelo consolida las boletas del mes —incluidas las de todas sus
semanas— y deja el detalle por trabajador que exige la entidad.
"""
import base64
import io

from dateutil.relativedelta import relativedelta

from odoo import _, api, fields, models
from odoo.exceptions import UserError

from odoo.addons.al_hr_pe.tools import custom_round

#: Columnas del detalle que acompaña al pago.
COLUMNS = [
    'Tipo doc.', 'N° documento', 'Apellidos y nombres', 'Categoría',
    'Días', 'Jornal básico', 'D.S.O.', 'Base', 'Retenido',
]


class L10nPeHrConafovicer(models.Model):
    _name = 'l10n_pe.hr.conafovicer'
    _description = 'Resumen mensual de CONAFOVICER'
    _order = 'date_end desc, id desc'
    _inherit = ['mail.thread']
    _check_company_auto = True

    name = fields.Char(string='Referencia', compute='_compute_name',
                       store=True)
    company_id = fields.Many2one(
        'res.company', string='Compañía', required=True, readonly=True,
        default=lambda self: self.env.company)
    period_id = fields.Many2one(
        'hr.period', string='Periodo', required=True, check_company=True,
        domain="[('period_type', '=', 'monthly')]",
        help='Mes que se liquida. Se toman las boletas del mes y las de '
             'todas sus semanas.')
    date_start = fields.Date(related='period_id.date_start', store=True)
    date_end = fields.Date(related='period_id.date_end', store=True)
    date_due = fields.Date(
        string='Vence', compute='_compute_date_due', store=True,
        help='Hasta el día 15 del mes siguiente.')
    state = fields.Selection(
        selection=[('draft', 'Borrador'), ('computed', 'Calculado'),
                   ('paid', 'Pagado')],
        string='Estado', default='draft', required=True, tracking=True,
        copy=False)
    line_ids = fields.One2many(
        'l10n_pe.hr.conafovicer.line', 'summary_id', string='Detalle',
        copy=False, readonly=True)
    currency_id = fields.Many2one(
        related='company_id.currency_id', string='Moneda')
    amount_base = fields.Monetary(
        string='Base', compute='_compute_amounts', store=True)
    amount_total = fields.Monetary(
        string='Retenido', compute='_compute_amounts', store=True)
    employee_count = fields.Integer(
        string='Trabajadores', compute='_compute_amounts', store=True)
    payment_reference = fields.Char(string='Constancia de pago')

    _period_company_uniq = models.Constraint(
        'UNIQUE(period_id, company_id)',
        'Ya existe el resumen de CONAFOVICER de ese periodo.')

    @api.depends('period_id', 'company_id')
    def _compute_name(self):
        for summary in self:
            summary.name = 'CONAFOVICER %s' % (
                summary.period_id.name or summary.period_id.code or '/')

    @api.depends('date_end')
    def _compute_date_due(self):
        for summary in self:
            summary.date_due = (
                (summary.date_end + relativedelta(months=1)).replace(day=15)
                if summary.date_end else False)

    @api.depends('line_ids.amount', 'line_ids.base')
    def _compute_amounts(self):
        for summary in self:
            summary.amount_base = sum(summary.line_ids.mapped('base'))
            summary.amount_total = sum(summary.line_ids.mapped('amount'))
            summary.employee_count = len(
                summary.line_ids.mapped('employee_id'))

    # ------------------------------------------------------------------
    # Cálculo
    # ------------------------------------------------------------------
    def _get_payslips(self):
        """Boletas confirmadas del mes y de todas sus semanas."""
        self.ensure_one()
        periods = self.period_id | self.period_id.child_ids
        return self.env['hr.payslip'].search([
            ('periodo_id', 'in', periods.ids),
            ('company_id', '=', self.company_id.id),
            ('state', 'in', ('validated', 'paid')),
            ('l10n_pe_is_construction', '=', True),
        ])

    def action_compute(self):
        for summary in self:
            if summary.state == 'paid':
                raise UserError(_(
                    'El resumen %s ya está pagado.', summary.name))
            summary.line_ids.unlink()
            vals = []
            for payslip in summary._get_payslips():
                amount = payslip._l10n_pe_construction_conafovicer()
                if not amount:
                    continue
                vals.append(fields.Command.create({
                    'employee_id': payslip.employee_id.id,
                    'payslip_id': payslip.id,
                    'days': payslip._l10n_pe_construction_days(),
                    'wage_amount': payslip._l10n_pe_construction_amount(
                        'jornal'),
                    'dso_amount': payslip._l10n_pe_construction_amount('dso'),
                    'base': payslip._l10n_pe_construction_conafovicer_base(),
                    'amount': amount,
                }))
            summary.line_ids = vals
            summary.state = 'computed'
        return True

    def action_mark_paid(self):
        for summary in self:
            if summary.state != 'computed':
                raise UserError(_(
                    'Calcule el resumen %s antes de marcarlo pagado.',
                    summary.name))
            summary.state = 'paid'
        return True

    def action_draft(self):
        self.write({'state': 'draft'})
        return True

    # ------------------------------------------------------------------
    # Detalle para el pago
    # ------------------------------------------------------------------
    def action_export_xlsx(self):
        """Detalle por trabajador que acompaña el depósito."""
        self.ensure_one()
        if not self.line_ids:
            raise UserError(_('Calcule el resumen antes de exportarlo.'))
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:  # pragma: no cover
            raise UserError(_(
                'Falta la librería openpyxl para generar el Excel.')) from exc

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'CONAFOVICER'
        sheet.append(['Empleador', self.company_id.name])
        sheet.append(['RUC', self.company_id.vat or ''])
        sheet.append(['Periodo', self.period_id.name or ''])
        sheet.append(['Vence', self.date_due and
                      self.date_due.strftime('%d/%m/%Y') or ''])
        sheet.append([])
        sheet.append(COLUMNS)
        header_row = sheet.max_row
        for cell in sheet[header_row]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='4F6228')
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
        for line in self.line_ids:
            employee = line.employee_id
            sheet.append([
                (employee.l10n_latam_identification_type_id
                 .l10n_pe_hr_sunat_code or ''),
                employee.identification_id or '',
                employee.name or '',
                line.category_id.name or '',
                line.days,
                line.wage_amount, line.dso_amount, line.base, line.amount,
            ])
        sheet.append([])
        sheet.append(['', '', '', '', '', '', 'TOTAL',
                      self.amount_base, self.amount_total])
        for index, column in enumerate(COLUMNS, start=1):
            sheet.column_dimensions[
                sheet.cell(row=header_row, column=index).column_letter
            ].width = max(12, min(32, len(column) + 6))

        stream = io.BytesIO()
        workbook.save(stream)
        attachment = self.env['ir.attachment'].create({
            'name': 'conafovicer_%s_%s.xlsx' % (
                self.company_id.vat or self.company_id.id,
                self.period_id.code or ''),
            'type': 'binary',
            'datas': base64.b64encode(stream.getvalue()),
            'res_model': self._name,
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'self',
        }


class L10nPeHrConafovicerLine(models.Model):
    _name = 'l10n_pe.hr.conafovicer.line'
    _description = 'Detalle de CONAFOVICER'
    _order = 'employee_id, payslip_id'
    _check_company_auto = True

    summary_id = fields.Many2one(
        'l10n_pe.hr.conafovicer', string='Resumen', required=True,
        ondelete='cascade', index=True)
    company_id = fields.Many2one(
        related='summary_id.company_id', store=True, index=True)
    currency_id = fields.Many2one(related='summary_id.currency_id')
    employee_id = fields.Many2one(
        'hr.employee', string='Trabajador', required=True,
        check_company=True)
    payslip_id = fields.Many2one(
        'hr.payslip', string='Boleta', check_company=True)
    category_id = fields.Many2one(
        'l10n_pe.hr.construction.category', string='Categoría',
        related='payslip_id.l10n_pe_construction_category_id', store=True)
    days = fields.Float(string='Días')
    wage_amount = fields.Monetary(string='Jornal básico')
    dso_amount = fields.Monetary(string='D.S.O.')
    base = fields.Monetary(
        string='Base',
        help='Jornal básico más el descanso semanal obligatorio.')
    amount = fields.Monetary(string='Retenido')

    @api.constrains('base', 'wage_amount', 'dso_amount')
    def _check_base(self):
        for line in self:
            expected = custom_round(line.wage_amount + line.dso_amount)
            if custom_round(line.base) != expected:
                raise UserError(_(
                    'La base de %(name)s no es el jornal más el D.S.O.',
                    name=line.employee_id.display_name))
