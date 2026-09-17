# -*- coding: utf-8 -*-
"""Excel de revisión de los libros PLE de la localización oficial.

El Excel se arma con el TXT oficial: los tests comprueban que cada línea del
TXT quede en su fila, con los encabezados del Anexo 2 de SUNAT.
"""
import base64
import io
import subprocess
import sys
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from odoo.tests import tagged
from odoo.addons.account.tests.common import AccountTestInvoicingCommon

from odoo.addons.al_l10n_pe_ple.models.ple_mixin import PLE_EXPECTED_FIELDS
from odoo.addons.al_l10n_pe_ple.models.ple_official_headers import PLE_OFFICIAL_HEADERS
from odoo.addons.al_l10n_pe_ple.models.ple_xlsx import PLE_XLSX_HEADERS, PLE_XLSX_TITLES

# Formatos que solo genera la localización oficial.
NATIVE_BOOKS = ('010100', '010200', '030100', '030200', '030300', '030400',
                '030500', '030600', '030700', '031100', '031200', '031300',
                '031400', '031500', '031601', '031602', '031700', '031800',
                '032000', '032400', '032500', '050100', '050300', '060100',
                '120100', '130100')
GENERATOR = Path(__file__).resolve().parents[2] / 'docs' / 'ple' / 'oficial' / 'generar_encabezados.py'


def txt(rows):
    """TXT como lo escribe la localización oficial: «|» al final de línea."""
    return ''.join('|'.join(row) + '|\n' for row in rows).encode()


@tagged('post_install', '-at_install')
class TestNativeXlsx(AccountTestInvoicingCommon):

    @classmethod
    @AccountTestInvoicingCommon.setup_country('pe')
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.company_data['company']
        cls.company.vat = '20512528458'
        cls.options = {'date': {'date_from': '2026-03-01', 'date_to': '2026-03-31'}}

    @staticmethod
    def _sheet_rows(content, sheet=None):
        workbook = load_workbook(io.BytesIO(content))
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        return [list(row) for row in worksheet.iter_rows(values_only=True)]

    # ------------------------------------------------------------------
    # Encabezados oficiales
    # ------------------------------------------------------------------
    def test_every_ple_format_has_headers(self):
        """Todo formato que valida el módulo tiene encabezados, con una
        columna por campo del TXT."""
        for code, count in PLE_EXPECTED_FIELDS.items():
            if code in ('080100', '080200', '140100'):
                continue  # reemplazados por el SIRE (8.4, 8.5 y 14.4)
            with self.subTest(book=code):
                headers = self.env['l10n_pe.ple.mixin']._ple_xlsx_headers(code)
                self.assertEqual(len(headers), count)
                self.assertTrue(all(headers))

    def test_native_books_use_official_headers(self):
        for code in NATIVE_BOOKS:
            with self.subTest(book=code):
                self.assertNotIn(code, PLE_XLSX_HEADERS)
                self.assertIn(code, PLE_OFFICIAL_HEADERS)
                self.assertIn(code, PLE_XLSX_TITLES)
                self.assertLessEqual(len(PLE_XLSX_TITLES[code]), 31, 'límite de Excel')
        self.assertEqual(PLE_OFFICIAL_HEADERS['050100'][1], 'Código Único de la Operación (CUO)')
        self.assertEqual(PLE_OFFICIAL_HEADERS['010200'][-1], 'Indica el estado de la operación')

    def test_generated_headers_are_up_to_date(self):
        """El módulo lleva los encabezados del Excel oficial que hay en docs."""
        if not GENERATOR.exists():
            self.skipTest('sin docs/ple/oficial (módulo instalado fuera del repositorio)')
        result = subprocess.run([sys.executable, str(GENERATOR), '--check'],
                                capture_output=True, text=True, check=False)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_txt_rows_drop_the_closing_pipe(self):
        rows = self.env['l10n_pe.ple.mixin']._ple_txt_rows(b'a|b||\r\nc|d|e\n\n')
        self.assertEqual(rows, [['a', 'b', ''], ['c', 'd', 'e']])

    # ------------------------------------------------------------------
    # Libro Mayor: 5.1, 5.3, 6.1 y Libro 3
    # ------------------------------------------------------------------
    def _gl(self):
        return self.env['account.general.ledger.report.handler'].with_company(self.company)

    def test_general_ledger_formats(self):
        line = ['20260300', 'CUO1', 'M1', '1011000', '', '', 'PEN', '', '', '00', '', '0', '',
                '01/03/2026', '', '01/03/2026', 'Venta', '', '100.00', '0.00', '', '1']
        cases = (
            ('l10n_pe_export_ple_51_to_txt', 'l10n_pe_export_ple_51_to_xlsx', 'PLE 5.1 Libro Diario'),
            ('l10n_pe_export_ple_53_to_txt', 'l10n_pe_export_ple_53_to_xlsx', 'PLE 5.3 Plan Contable'),
            ('l10n_pe_export_ple_61_to_txt', 'l10n_pe_export_ple_61_to_xlsx', 'PLE 6.1 Libro Mayor'),
        )
        handler_class = type(self.env['account.general.ledger.report.handler'])
        for txt_method, xlsx_method, sheet in cases:
            with self.subTest(method=xlsx_method), patch.object(
                    handler_class, txt_method, autospec=True, return_value={
                        'file_name': 'LE2051252845820260300050100001111',
                        'file_content': txt([line[:21]]), 'file_type': 'txt'}):
                result = getattr(self._gl(), xlsx_method)(self.options)
                self.assertEqual(result['file_type'], 'xlsx')
                self.assertEqual(result['file_name'], 'LE2051252845820260300050100001111')
                rows = self._sheet_rows(result['file_content'], sheet)
                self.assertIn(self.company.name, rows[0][0])
                self.assertEqual(rows[1][:7], ['RUC', '20512528458', None, 'Período', '2026', 'Mes', '03'])
                self.assertEqual(rows[5][0], 1)
                self.assertEqual(rows[5][1], '20260300')
                self.assertEqual(rows[5][2], 'CUO1')

    def test_general_ledger_buttons(self):
        report = self.env.ref('account_reports.general_ledger_report')
        options = report.with_company(self.company).get_options({})
        params = {button.get('action_param') for button in options['buttons']}
        for method in ('l10n_pe_export_ple_51_to_xlsx', 'l10n_pe_export_ple_53_to_xlsx',
                       'l10n_pe_export_ple_61_to_xlsx', 'l10n_pe_export_lib_to_xlsx'):
            self.assertIn(method, params)

    def test_inventory_and_balance_one_sheet_per_format(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            archive.writestr('LE2051252845820260331030100011111.txt',
                             txt([['20260331', '01', '1D01ST', '1500.00', '1']]))
            archive.writestr('LE2051252845820260331031700011111.txt', b'')
        handler_class = type(self.env['account.general.ledger.report.handler'])
        with patch.object(handler_class, 'l10n_pe_export_lib_to_txt', autospec=True,
                          return_value={'file_name': 'Inventory_and_balance_reports_20260331',
                                        'file_content': buffer.getvalue(), 'file_type': 'zip'}):
            result = self._gl().l10n_pe_export_lib_to_xlsx(self.options)
        workbook = load_workbook(io.BytesIO(result['file_content']))
        self.assertEqual(workbook.sheetnames,
                         ['PLE 3.1 Situación financiera', 'PLE 3.17 Balance comprobación'])
        rows = self._sheet_rows(result['file_content'], 'PLE 3.1 Situación financiera')
        self.assertEqual(rows[3][1:6], PLE_OFFICIAL_HEADERS['030100'])
        self.assertEqual(rows[5][1:6], ['20260331', '01', '1D01ST', '1500.00', '1'])
        empty = self._sheet_rows(result['file_content'], 'PLE 3.17 Balance comprobación')
        self.assertEqual(len(empty), 5, 'sin datos: solo las filas de encabezado')

    # ------------------------------------------------------------------
    # Flujo de caja: 1.1 y 1.2
    # ------------------------------------------------------------------
    def test_cash_and_bank(self):
        handler = self.env['account.cash.flow.report.handler'].with_company(self.company)
        handler_class = type(handler)
        row = ['20260300', 'CUO9', 'M9', '1041000', '01/03/2026', '001', 'Pago',
               '6', '20131312955', 'PROVEEDOR SAC', '00012', '0.00', '250.00', '1']
        with patch.object(handler_class, 'l10n_pe_export_ple_12_to_txt', autospec=True,
                          return_value={'file_name': 'LE2051252845820260300010200001111',
                                        'file_content': txt([['20260300', '02'] + row[:13]]),
                                        'file_type': 'txt'}):
            result = handler.l10n_pe_export_ple_12_to_xlsx(self.options)
        rows = self._sheet_rows(result['file_content'], 'PLE 1.2 Bancos')
        self.assertEqual(rows[3][1:16], PLE_OFFICIAL_HEADERS['010200'])
        self.assertEqual(len([value for value in rows[5][1:] if value is not None]), 15)

        report = self.env.ref('account_reports.cash_flow_report')
        options = report.with_company(self.company).get_options({})
        params = {button.get('action_param') for button in options['buttons']}
        self.assertIn('l10n_pe_export_ple_11_to_xlsx', params)
        self.assertIn('l10n_pe_export_ple_12_to_xlsx', params)

    # ------------------------------------------------------------------
    # Inventario permanente: 12.1 y 13.1
    # ------------------------------------------------------------------
    def test_stock_wizard(self):
        wizard = self.env['l10n_pe.stock.ple.wizard'].with_company(self.company).create({
            'date_from': '2026-03-01', 'date_to': '2026-03-31'})
        line = ['20260300', 'CUO1', 'M1', '0000', '9', '01', 'P001', '', '',
                '05/03/2026', '00', '0', '0', '01', 'Producto', 'NIU', '1',
                '10.00', '100.00', '1000.00', '0.00', '0.00', '0.00', '10.00',
                '100.00', '1000.00', '1']
        wizard_class = type(wizard)
        with patch.object(wizard_class, '_get_ple_report_content', autospec=True,
                          return_value=txt([line]).decode()):
            action = wizard.get_ple_xlsx_13_1()
        self.assertEqual(action['type'], 'ir.actions.act_url')
        self.assertEqual(wizard.report_filename, 'LE2051252845820260300130100001111.xlsx')
        content = base64.b64decode(wizard.report_data)
        rows = self._sheet_rows(content, 'PLE 13.1 Inventario valorizado')
        # Los encabezados largos del Anexo 2 caben: la fila crece sobre los 40
        # puntos del formato v18.
        worksheet = load_workbook(io.BytesIO(content)).active
        self.assertGreater(worksheet.row_dimensions[5].height, 40)
        self.assertEqual(rows[3][1:28], PLE_OFFICIAL_HEADERS['130100'])
        # Las celdas vacías del TXT quedan vacías en la hoja.
        self.assertEqual(rows[5][1:28], [value or None for value in line])

        with patch.object(wizard_class, '_get_ple_report_content', autospec=True,
                          return_value=''):
            wizard.get_ple_xlsx_12_1()
        self.assertEqual(wizard.report_filename, 'LE2051252845820260300120100001011.xlsx')
