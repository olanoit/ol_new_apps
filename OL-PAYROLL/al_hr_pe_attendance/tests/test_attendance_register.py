# -*- coding: utf-8 -*-
"""Registro de control de asistencia (D.S. 004-2006-TR)."""
import base64
import io
from datetime import date, datetime

from odoo.tests.common import TransactionCase, tagged


@tagged('post_install', '-at_install', 'al_hr_pe_attendance')
class TestAttendanceRegister(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        Param = cls.env['hr.main.parameter']
        Param.search([('company_id', '=', cls.env.company.id)], limit=1) \
            or Param.create({'company_id': cls.env.company.id})
        calendar = cls.env['resource.calendar'].create({
            'name': 'Diurno registro',
            'tz': 'America/Lima',
            'attendance_ids': [(5, 0, 0)] + [
                (0, 0, {'name': 'Día %s' % dia, 'dayofweek': str(dia),
                        'hour_from': 8.0, 'hour_to': 16.0,
                        'day_period': 'morning'})
                for dia in range(0, 5)],
        })
        cls.employee = cls.env['hr.employee'].create({
            'name': 'Trabajador registro',
            'identification_id': '45678912',
            'tz': 'America/Lima',
            'resource_calendar_id': calendar.id,
        })
        cls.employee.version_id.l10n_pe_is_overtime = True
        # Lunes 02/11/2026: 08:00 → 19:00 hora de Lima (UTC-5).
        cls.env['hr.attendance'].create({
            'employee_id': cls.employee.id,
            'check_in': datetime(2026, 11, 2, 13, 0),
            'check_out': datetime(2026, 11, 3, 0, 0),
        })

    def _wizard(self):
        return self.env['l10n_pe.hr.attendance.register.wizard'].create({
            'date_from': date(2026, 11, 1), 'date_to': date(2026, 11, 30),
            'employee_ids': [(6, 0, self.employee.ids)],
        })

    def test_row_has_the_legal_fields(self):
        """Ingreso, salida y sobretiempo (inicio y fin) del día."""
        rows = self._wizard()._register_rows()
        self.assertEqual(len(rows), 1)
        doc, name, day, check_in, check_out, ot_from, ot_to, worked, note = \
            rows[0]
        self.assertEqual((doc, day), ('45678912', date(2026, 11, 2)))
        self.assertEqual((check_in, check_out), ('08:00', '19:00'))
        self.assertEqual((ot_from, ot_to), ('16:00', '19:00'))
        self.assertEqual(worked, '11:00')

    def test_excel_has_employer_header(self):
        from openpyxl import load_workbook
        action = self._wizard().action_export()
        attachment = self.env['ir.attachment'].browse(
            int(action['url'].split('/')[3].split('?')[0]))
        sheet = load_workbook(io.BytesIO(base64.b64decode(attachment.datas))).active
        self.assertIn('004-2006-TR', sheet['A1'].value)
        self.assertEqual(sheet['A2'].value, 'Empleador')
        self.assertEqual(sheet['A3'].value, 'RUC')
        self.assertEqual(sheet.cell(row=6, column=4).value, 'Ingreso')
