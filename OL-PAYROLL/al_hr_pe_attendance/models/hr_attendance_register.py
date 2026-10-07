# -*- coding: utf-8 -*-
"""Registro permanente de control de asistencia (D.S. 004-2006-TR).

El empleador debe llevar un registro con, por trabajador y día (art. 3):

a) nombre o razón social del empleador y b) su RUC;
c) nombre y documento de identidad del trabajador;
d) fecha, hora y minutos de ingreso y salida de la jornada;
e) hora y minutos de inicio y fin del sobretiempo.

Debe estar a disposición de la inspección (SUNAFIL) y conservarse cinco
años. Este asistente lo genera en Excel a partir de las marcaciones
(``hr.attendance``), con la misma consolidación y el mismo cálculo de
sobretiempo que el tareaje, para que el registro cuadre con la boleta.
"""
import base64
import io

from odoo import _, api, fields, models
from odoo.exceptions import UserError

COLUMNS = [
    'Documento', 'Trabajador', 'Fecha', 'Ingreso', 'Salida',
    'Inicio sobretiempo', 'Fin sobretiempo', 'Horas en la jornada',
    'Observación',
]


def _hhmm(hours):
    """Horas flotantes → «HH:MM» (más de 24 = día siguiente)."""
    if hours is None:
        return ''
    minutes = int(round(hours * 60))
    days, minutes = divmod(minutes, 24 * 60)
    text = '%02d:%02d' % divmod(minutes, 60)
    return text + (' (+%d)' % days if days else '')


class HrAttendanceRegisterWizard(models.TransientModel):
    _name = 'l10n_pe.hr.attendance.register.wizard'
    _description = 'Registro de control de asistencia (D.S. 004-2006-TR)'

    company_id = fields.Many2one(
        'res.company', string='Compañía', required=True,
        default=lambda self: self.env.company)
    date_from = fields.Date(
        string='Desde', required=True,
        default=lambda self: fields.Date.context_today(self).replace(day=1))
    date_to = fields.Date(
        string='Hasta', required=True, default=fields.Date.context_today)
    employee_ids = fields.Many2many(
        'hr.employee', string='Trabajadores',
        domain="[('company_id', '=', company_id)]",
        help='Vacío: todos los trabajadores con marcaciones en el periodo.')

    @api.constrains('date_from', 'date_to')
    def _check_dates(self):
        for wizard in self:
            if wizard.date_from > wizard.date_to:
                raise UserError(_('La fecha inicial es posterior a la final.'))

    # ------------------------------------------------------------------
    # Filas
    # ------------------------------------------------------------------
    def _register_rows(self):
        """Una fila por trabajador y día con marcación."""
        self.ensure_one()
        Tareaje = self.env['hr.tareaje.manager']
        # Tareaje en memoria: solo para reutilizar la consolidación de
        # marcaciones, los feriados y el cálculo de sobretiempo.
        tareaje = Tareaje.new({
            'name': 'Registro de asistencia',
            'company_id': self.company_id.id,
            'date_start': self.date_from,
            'date_end': self.date_to,
        })
        param = self.env['hr.main.parameter'].get_main_parameter(
            self.company_id)
        marks_by_employee = tareaje._get_marks_by_employee_day()
        holidays = tareaje._get_public_holidays()
        rows = []
        employees = sorted(
            (emp for emp in marks_by_employee
             if not self.employee_ids or emp in self.employee_ids),
            key=lambda emp: emp.name or '')
        for employee in employees:
            calendar = employee.resource_calendar_id \
                or self.company_id.resource_calendar_id
            version = employee._get_version(self.date_to)
            for day in sorted(marks_by_employee[employee]):
                marks_in, marks_out = marks_by_employee[employee][day]
                schedule = tareaje._get_day_schedule(calendar, day)
                if day in holidays.get(calendar.id, set()) \
                        or day in holidays.get(False, set()):
                    day_kind = 'feriado'
                elif schedule is None:
                    day_kind = 'descanso'
                else:
                    day_kind = 'workday'
                sched_in, sched_out, break_hours = \
                    schedule or (None, None, 0.0)
                values = tareaje._classify_day(
                    marks_in, marks_out,
                    sched_in=sched_in, sched_out=sched_out,
                    break_hours=break_hours, day_kind=day_kind,
                    night_from=param.tareaje_night_from,
                    night_to=param.tareaje_night_to,
                    he25_limit=param.tareaje_he25_hours,
                    round_minutes=param.tareaje_round_minutes,
                    compute_overtime=bool(version.l10n_pe_is_overtime))
                overtime = values['he25'] + values['he35'] + values['he100']
                note = ''
                if marks_out is None:
                    note = _('Sin marcación de salida')
                elif day_kind == 'feriado':
                    note = _('Feriado laborado')
                elif day_kind == 'descanso':
                    note = _('Descanso laborado')
                out = marks_out
                if out is not None and out <= marks_in:
                    out += 24.0
                worked = max(0.0, out - marks_in - break_hours) \
                    if out is not None else None
                rows.append([
                    employee.identification_id or '',
                    employee.name or '',
                    day,
                    _hhmm(marks_in),
                    _hhmm(out),
                    _hhmm(out - overtime) if overtime and out else '',
                    _hhmm(out) if overtime and out else '',
                    _hhmm(worked),
                    note,
                ])
        return rows

    # ------------------------------------------------------------------
    # Excel
    # ------------------------------------------------------------------
    def action_export(self):
        self.ensure_one()
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:  # pragma: no cover
            raise UserError(_(
                'Falta la librería openpyxl para generar el Excel.')) from exc
        rows = self._register_rows()
        if not rows:
            raise UserError(_('No hay marcaciones en el periodo.'))
        company = self.company_id
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = _('Registro')
        sheet.append([_('REGISTRO DE CONTROL DE ASISTENCIA (D.S. 004-2006-TR)')])
        sheet.append([_('Empleador'), company.name or ''])
        sheet.append([_('RUC'), company.vat or ''])
        sheet.append([_('Periodo'), '%s al %s' % (
            self.date_from.strftime('%d/%m/%Y'),
            self.date_to.strftime('%d/%m/%Y'))])
        sheet.append([])
        sheet['A1'].font = Font(bold=True, size=12)
        sheet.append(COLUMNS)
        header_row = sheet.max_row
        for cell in sheet[header_row]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='4F6228')
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
        for row in rows:
            sheet.append(row)
            sheet.cell(row=sheet.max_row, column=3).number_format = 'DD/MM/YYYY'
        widths = (14, 34, 12, 10, 12, 12, 12, 12, 24)
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[
                sheet.cell(row=header_row, column=index).column_letter
            ].width = width
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

        stream = io.BytesIO()
        workbook.save(stream)
        filename = 'registro_asistencia_%s_%s_%s.xlsx' % (
            company.vat or company.id,
            self.date_from.strftime('%Y%m%d'), self.date_to.strftime('%Y%m%d'))
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(stream.getvalue()),
            # Sin res_model/res_id: lleva documentos de identidad y horarios;
            # solo lo descarga quien lo generó.
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'self',
        }
