# -*- coding: utf-8 -*-
"""Fase 8: tests del framework de importación Excel (al_hr_pe_import).

Cubren el mixin ``al.import.payroll.mixin`` (parseo openpyxl, detección
de hojas, lotes, contadores, reporte y plantilla descargable) y los
wizards concretos de asistencias, reglas salariales y récord vacacional,
más la regla multicompañía del historial de progreso.

Los tests llaman a ``_process_all_rows`` de forma síncrona (camino
documentado para tests): ``action_run_import`` lanza un hilo con cursor
real, que no es apto para la transacción de prueba.
"""
import base64
import io
from datetime import date, datetime

from odoo.exceptions import UserError, ValidationError
from odoo.tests import TransactionCase, tagged


def _xlsx(sheets):
    """Construye un xlsx en memoria: ``{nombre_hoja: [filas]}`` → b64."""
    import openpyxl

    wb = openpyxl.Workbook()
    default = wb.active
    for idx, (name, rows) in enumerate(sheets.items()):
        ws = default if idx == 0 else wb.create_sheet()
        ws.title = name
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue())


@tagged('post_install', '-at_install')
class TestFase8Import(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.pen = cls.env.ref('base.PEN')
        cls.pen.active = True
        Company = cls.env['res.company']
        cls.company = Company.create({
            'name': 'PE Import SAC',
            'country_id': cls.env.ref('base.pe').id,
            'currency_id': cls.pen.id,
        })
        cls.company_b = Company.create({
            'name': 'PE Import Beta SAC',
            'country_id': cls.env.ref('base.pe').id,
            'currency_id': cls.pen.id,
        })
        cls.env = cls.env(context=dict(
            cls.env.context,
            allowed_company_ids=[cls.company.id, cls.company_b.id]))
        cls.env.user.company_ids |= cls.company | cls.company_b
        cls.env.user.group_ids |= cls.env.ref(
            'hr_payroll.group_hr_payroll_manager')

        cls.structure = cls.env.ref('al_hr_pe.base_structure')
        cls.employee = cls.env['hr.employee'].with_company(
            cls.company).create({
                'names': 'Rosa',
                'last_name': 'Huamán',
                'm_last_name': 'Cruz',
                'company_id': cls.company.id,
                'date_version': date(2025, 1, 1),
                'contract_date_start': date(2025, 1, 1),
                'wage': 3000.0,
                'structure_type_id': cls.structure.type_id.id,
            })
        cls.employee.version_id.identification_id = '46271883'

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------
    def _wizard_asistencias(self, sheets, **vals):
        wizard = self.env['al.import.hr.attendance.wizard'].create(dict({
            'file_data': _xlsx(sheets),
            'file_name': 'asistencias.xlsx',
            'company_id': self.company.id,
            'tz': 'America/Lima',
        }, **vals))
        wizard.action_load_file()
        return wizard

    @staticmethod
    def _correr(wizard):
        rows = wizard._preprocess_rows(wizard._iter_data_rows())
        return wizard._process_all_rows(rows)

    # ------------------------------------------------------------------
    # Mixin: carga de archivo y detección de hojas
    # ------------------------------------------------------------------
    def test_01_carga_y_deteccion_de_hojas(self):
        wizard = self._wizard_asistencias({
            'RESUMEN': [['solo texto']],
            'DETALLE': [
                ['EMPLEADO', 'ENTRADA', 'SALIDA'],
                [self.employee.name, '2026-07-01 08:00:00', ''],
            ],
        })
        self.assertEqual(wizard.state, 'configure')
        self.assertEqual(wizard.sheet_names_csv, 'RESUMEN||DETALLE')
        # Sin keyword el mixin autoselecciona la primera hoja.
        self.assertEqual(wizard.sheet_id.name, 'RESUMEN')
        # El selector de hoja es reasignable a cualquier hoja detectada.
        detalle = self.env['al.import.payroll.sheet'].search([
            ('res_model', '=', wizard._name),
            ('res_id', '=', wizard.id),
            ('name', '=', 'DETALLE'),
        ])
        self.assertEqual(len(detalle), 1)
        wizard.sheet_id = detalle
        rows = wizard._iter_data_rows()
        self.assertEqual(len(rows), 1)

        # Extensión inválida: la constraint rechaza el archivo.
        with self.assertRaises(ValidationError):
            self.env['al.import.hr.attendance.wizard'].create({
                'file_data': base64.b64encode(b'no es excel'),
                'file_name': 'archivo.txt',
                'company_id': self.company.id,
            })
        # Sin archivo: no se puede lanzar la carga.
        vacio = self.env['al.import.hr.attendance.wizard'].create({
            'company_id': self.company.id})
        with self.assertRaises(UserError):
            vacio.action_load_file()

    # ------------------------------------------------------------------
    # Asistencias: creación con zona horaria, actualización y errores
    # ------------------------------------------------------------------
    def test_02_importar_asistencias(self):
        wizard = self._wizard_asistencias({
            'ASISTENCIAS': [
                ['EMPLEADO', 'ENTRADA', 'SALIDA'],
                [self.employee.name, '2026-07-01 08:00:00',
                 '2026-07-01 17:00:00'],
                ['', '2026-07-02 08:00:00', ''],          # sin empleado
                ['NO EXISTE SAC', '2026-07-02 08:00:00', ''],
                [self.employee.name, 'no-es-fecha', ''],
            ],
        })
        results, counts, log_lines, created_ids = self._correr(wizard)
        self.assertEqual(counts, {
            'created': 1, 'updated': 0, 'skipped': 1, 'error': 2})
        self.assertEqual(len(log_lines), 4)
        self.assertEqual(len(created_ids), 1)

        att = self.env['hr.attendance'].browse(created_ids[0])
        # 08:00 America/Lima (UTC-5) → 13:00 UTC almacenado.
        self.assertEqual(att.employee_id, self.employee)
        self.assertEqual(att.check_in, datetime(2026, 7, 1, 13, 0, 0))
        self.assertEqual(att.check_out, datetime(2026, 7, 1, 22, 0, 0))

        # Los errores llevan sugerencia de corrección específica.
        errores = [r for r in results if r['status'] == 'error']
        self.assertTrue(all(r['fix'] for r in errores))

        # Segunda pasada con la misma entrada: actualiza el check_out.
        wizard2 = self._wizard_asistencias({
            'ASISTENCIAS': [
                ['EMPLEADO', 'ENTRADA', 'SALIDA'],
                [self.employee.name, '2026-07-01 08:00:00',
                 '2026-07-01 18:00:00'],
            ],
        })
        _res, counts2, _log, _ids = self._correr(wizard2)
        self.assertEqual(counts2['updated'], 1)
        self.assertEqual(counts2['created'], 0)
        self.assertEqual(att.check_out, datetime(2026, 7, 1, 23, 0, 0))

        # Con "Actualizar existentes" desmarcado la fila se omite.
        wizard3 = self._wizard_asistencias({
            'ASISTENCIAS': [
                ['EMPLEADO', 'ENTRADA', 'SALIDA'],
                [self.employee.name, '2026-07-01 08:00:00',
                 '2026-07-01 19:00:00'],
            ],
        }, update_existing=False)
        _res, counts3, _log, _ids = self._correr(wizard3)
        self.assertEqual(counts3['skipped'], 1)
        self.assertEqual(att.check_out, datetime(2026, 7, 1, 23, 0, 0))

    def test_03_asistencias_multicompania(self):
        # El empleado es de la compañía A: importarlo desde B falla por
        # fila (no encontrado), sin tocar datos de A.
        wizard = self._wizard_asistencias({
            'ASISTENCIAS': [
                ['EMPLEADO', 'ENTRADA', 'SALIDA'],
                [self.employee.name, '2026-07-01 08:00:00', ''],
            ],
        }, company_id=self.company_b.id)
        _res, counts, log_lines, _ids = self._correr(wizard)
        self.assertEqual(counts['error'], 1)
        self.assertIn('no encontrado', log_lines[0])
        self.assertFalse(self.env['hr.attendance'].search([
            ('employee_id', '=', self.employee.id)]))

    # ------------------------------------------------------------------
    # Lotes: los contadores cierran con cualquier tamaño de lote
    # ------------------------------------------------------------------
    def test_04_procesamiento_por_lotes(self):
        filas = [['EMPLEADO', 'ENTRADA', 'SALIDA']]
        for day in range(1, 6):
            filas.append([self.employee.name,
                          '2026-06-%02d 08:00:00' % day,
                          '2026-06-%02d 17:00:00' % day])
        wizard = self._wizard_asistencias(
            {'ASISTENCIAS': filas}, use_batching=True, batch_size=2)
        _res, counts, log_lines, created_ids = self._correr(wizard)
        self.assertEqual(counts['created'], 5)
        self.assertEqual(len(created_ids), 5)
        self.assertEqual(len(log_lines), 5)

    # ------------------------------------------------------------------
    # Reglas salariales
    # ------------------------------------------------------------------
    def test_05_importar_reglas_salariales(self):
        struct = self.env['hr.payroll.structure'].create({
            'name': 'Estructura import F8',
            'type_id': self.structure.type_id.id,
        })
        contenido = {
            'REGLAS': [
                ['CATEGORÍA', 'COMPAÑÍA', 'CÓDIGO', 'NOMBRE', 'SECUENCIA',
                 'CÓDIGO PYTHON', 'CONDICIÓN PYTHON'],
                ['ING', '', 'F8BONO', 'Bono de prueba F8', 45,
                 'result = 100.0', 'result = True'],
                ['NO-EXISTE', '', 'F8MAL', 'Regla sin categoría', 50,
                 'result = 0.0', ''],
            ],
        }
        wizard = self.env['al.import.hr.salary.rule.wizard'].create({
            'file_data': _xlsx(contenido),
            'file_name': 'reglas.xlsx',
            'company_id': self.company.id,
            'struct_id': struct.id,
        })
        wizard.action_load_file()
        _res, counts, _log, created_ids = self._correr(wizard)
        self.assertEqual(counts['created'], 1)
        self.assertEqual(counts['error'], 1)

        regla = self.env['hr.salary.rule'].browse(created_ids[0])
        self.assertEqual(regla.code, 'F8BONO')
        self.assertEqual(regla.struct_id, struct)
        self.assertEqual(regla.category_id,
                         self.env.ref('al_hr_pe.ING'))
        self.assertEqual(regla.amount_select, 'code')
        self.assertEqual(regla.amount_python_compute, 'result = 100.0')
        self.assertEqual(regla.condition_select, 'python')

        # Reimportar el mismo código sobre la misma estructura actualiza.
        contenido['REGLAS'][1][3] = 'Bono renombrado F8'
        wizard2 = self.env['al.import.hr.salary.rule.wizard'].create({
            'file_data': _xlsx(contenido),
            'file_name': 'reglas.xlsx',
            'company_id': self.company.id,
            'struct_id': struct.id,
        })
        wizard2.action_load_file()
        _res, counts2, _log, _ids = self._correr(wizard2)
        self.assertEqual(counts2['updated'], 1)
        self.assertEqual(regla.name, 'Bono renombrado F8')

    def test_05b_saneo_codigo_v18(self):
        """El export de v18 se adapta al localdict de v19.

        Los exports reales traen (a) la condición por defecto de Odoo en
        TODAS las filas aunque la regla sea "Siempre verdadero" y (b)
        código que usa ``contract``/``payslip.wage``, inexistentes en v19.
        """
        struct = self.env['hr.payroll.structure'].create({
            'name': 'Estructura saneo F8',
            'type_id': self.structure.type_id.id,
        })
        boilerplate = (
            "\n# Available variables:\n#--------------------\n"
            "# payslip: hr.payslip object\n\n"
            "result = rules['NET']['total'] > categories['NET'] * 0.10")
        codigo_v18 = (
            "if contract.wage_type == 'hourly':\n"
            "    result = payslip.wage / contract.resource_calendar_id"
            ".hours_per_day\n"
            "else:\n"
            "    result = payslip.wage / 30")
        contenido = {
            'REGLAS': [
                ['CATEGORÍA', 'COMPAÑÍA', 'CÓDIGO', 'NOMBRE', 'SECUENCIA',
                 'CÓDIGO PYTHON', 'CONDICIÓN PYTHON'],
                ['ING', '', 'F8V18', 'Regla exportada de v18', 10,
                 codigo_v18, boilerplate],
            ],
        }
        wizard = self.env['al.import.hr.salary.rule.wizard'].create({
            'file_data': _xlsx(contenido),
            'file_name': 'reglas_v18.xlsx',
            'company_id': self.company.id,
            'struct_id': struct.id,
        })
        wizard.action_load_file()
        _res, counts, log, created_ids = self._correr(wizard)
        self.assertEqual(counts['created'], 1)

        regla = self.env['hr.salary.rule'].browse(created_ids[0])
        self.assertEqual(regla.condition_select, 'none',
                         'la condición por defecto debe descartarse')
        self.assertNotIn('contract', regla.amount_python_compute)
        self.assertNotIn('payslip.wage', regla.amount_python_compute)
        self.assertIn('version.wage', regla.amount_python_compute)
        self.assertIn('version.resource_calendar_id',
                      regla.amount_python_compute)
        self.assertIn('adaptado v19', '\n'.join(log))

        # El código adaptado es Python válido y evaluable.
        compile(regla.amount_python_compute, '<regla F8V18>', 'exec')

        # Sin saneo, el código entra literal y la condición se respeta.
        wizard2 = self.env['al.import.hr.salary.rule.wizard'].create({
            'file_data': _xlsx(contenido),
            'file_name': 'reglas_v18.xlsx',
            'company_id': self.company.id,
            'struct_id': struct.id,
            'sanitize_v18': False,
        })
        wizard2.action_load_file()
        self._correr(wizard2)
        self.assertEqual(regla.condition_select, 'python')
        self.assertIn('contract', regla.amount_python_compute)

    def test_02b_asistencias_por_documento(self):
        """Los relojes biométricos exportan el DNI, no el nombre."""
        contenido = {
            'ASISTENCIAS': [
                ['EMPLEADO', 'ENTRADA', 'SALIDA'],
                # DNI numérico tal y como lo entrega openpyxl.
                [46271883, '2026-04-02 08:00:00', '2026-04-02 17:00:00'],
                ['99999999', '2026-04-03 08:00:00', '2026-04-03 17:00:00'],
            ],
        }
        wizard = self.env['al.import.hr.attendance.wizard'].create({
            'file_data': _xlsx(contenido),
            'file_name': 'asistencias.xlsx',
            'company_id': self.company.id,
            'tz': 'America/Lima',
        })
        wizard.action_load_file()
        _res, counts, _log, created_ids = self._correr(wizard)
        self.assertEqual(counts['created'], 1,
                         'el DNI debe identificar al empleado')
        self.assertEqual(counts['error'], 1,
                         'un documento inexistente sigue siendo error')
        marca = self.env['hr.attendance'].browse(created_ids[0])
        self.assertEqual(marca.employee_id, self.employee)

    # ------------------------------------------------------------------
    # Récord vacacional: saldo inicial que se reemplaza, no se acumula
    # ------------------------------------------------------------------
    def test_06_importar_record_vacacional(self):
        contenido = {
            'VACACIONES': [
                ['FECHA', 'DOCUMENTO', 'DÍAS', 'IMPORTE'],
                # DNI numérico a propósito: openpyxl lo entrega como
                # número y _clean_code debe volverlo texto sin decimales.
                ['2026-01-01', 46271883, 12.5, 1250.00],
                ['2026-01-01', 99999999, 5, 0],   # documento inexistente
            ],
        }
        wizard = self.env['al.import.vacation.rest.wizard'].create({
            'file_data': _xlsx(contenido),
            'file_name': 'vacaciones.xlsx',
            'company_id': self.company.id,
        })
        wizard.action_load_file()
        _res, counts, _log, created_ids = self._correr(wizard)
        self.assertEqual(counts['created'], 1)
        self.assertEqual(counts['error'], 1)

        rest = self.env['hr.vacation.rest'].browse(created_ids[0])
        self.assertEqual(rest.employee_id, self.employee)
        self.assertEqual(rest.internal_motive, 'rest')
        self.assertEqual(rest.days_rest, 12.5)
        self.assertEqual(rest.amount_rest, 1250.00)
        self.assertEqual(rest.motive, 'Saldo acumulado anterior')

        # Reimportar reemplaza el saldo (paridad v18: unlink + create) y
        # el signo negativo cambia el motivo.
        contenido['VACACIONES'] = [
            ['FECHA', 'DOCUMENTO', 'DÍAS', 'IMPORTE'],
            ['2026-02-01', '46271883', -3, 0],
        ]
        wizard2 = self.env['al.import.vacation.rest.wizard'].create({
            'file_data': _xlsx(contenido),
            'file_name': 'vacaciones.xlsx',
            'company_id': self.company.id,
        })
        wizard2.action_load_file()
        _res, counts2, _log, ids2 = self._correr(wizard2)
        self.assertEqual(counts2['updated'], 1)
        saldos = self.env['hr.vacation.rest'].search([
            ('employee_id', '=', self.employee.id),
            ('internal_motive', '=', 'rest'),
        ])
        self.assertEqual(len(saldos), 1)
        self.assertEqual(saldos.id, ids2[0])
        self.assertEqual(saldos.days_rest, -3)
        self.assertEqual(saldos.motive, 'Saldo Ajuste Adelantos')

    # ------------------------------------------------------------------
    # Plantilla y reporte Excel
    # ------------------------------------------------------------------
    def test_07_plantilla_descargable(self):
        import openpyxl

        wizard = self.env['al.import.hr.attendance.wizard'].create({
            'company_id': self.company.id})
        data, filename = wizard._build_xlsx_template()
        self.assertEqual(filename, 'plantilla_hr_attendance.xlsx')
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb.active
        self.assertEqual(ws.title, 'ASISTENCIAS')
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers, wizard._template_headers())
        # La fila de ejemplo acompaña a la cabecera.
        self.assertEqual(ws.cell(row=2, column=1).value,
                         wizard._template_example_rows()[0][0])

        # action_download_template persiste el binario y devuelve la URL.
        action = wizard.action_download_template()
        self.assertEqual(action['type'], 'ir.actions.act_url')
        self.assertTrue(wizard.template_file)

    def test_08_reporte_de_resultados(self):
        import openpyxl

        wizard = self._wizard_asistencias({
            'ASISTENCIAS': [
                ['EMPLEADO', 'ENTRADA', 'SALIDA'],
                [self.employee.name, '2026-05-04 08:00:00', ''],
                ['NO EXISTE SAC', '2026-05-04 08:00:00', ''],
                ['', '', ''],
                ['', '2026-05-05 08:00:00', ''],  # skipped
            ],
        })
        results, counts, _log, _ids = self._correr(wizard)
        data, filename = wizard._build_xlsx_report(results, counts)
        self.assertTrue(filename.startswith('import_hr_attendance_'))
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames,
                         ['Resultado', 'Errores', 'Omitidos'])
        ws = wb['Resultado']
        estados = [ws.cell(row=r, column=2).value
                   for r in range(2, 2 + len(results))]
        self.assertCountEqual(estados, ['CREADO', 'ERROR', 'OMITIDO'])
        # La hoja de errores solo contiene los errores.
        ws_err = wb['Errores']
        self.assertEqual(ws_err.cell(row=2, column=2).value, 'ERROR')
        self.assertIsNone(ws_err.cell(row=3, column=2).value)

    # ------------------------------------------------------------------
    # Progreso: datos para el widget OWL y regla multicompañía
    # ------------------------------------------------------------------
    def test_09_progreso_y_multicompania(self):
        Progress = self.env['al.import.payroll.progress']
        prog_a = Progress.create({
            'name': 'Import A',
            'wizard_model': 'al.import.hr.attendance.wizard',
            'wizard_id': 1,
            'target_model': 'hr.attendance',
            'company_id': self.company.id,
            'status': 'running',
            'total': 40,
            'current': 10,
            'created': 8,
            'errors': 2,
        })
        prog_b = Progress.create({
            'name': 'Import B',
            'wizard_model': 'al.import.hr.attendance.wizard',
            'wizard_id': 2,
            'target_model': 'hr.attendance',
            'company_id': self.company_b.id,
            'status': 'done',
            'total': 10,
            'current': 10,
        })
        data = prog_a.get_progress_data()
        self.assertEqual(data['percent'], 25)
        self.assertEqual(data['created'], 8)
        self.assertEqual(data['errors'], 2)
        self.assertFalse(data['has_report'])

        prog_a.created_res_ids_csv = '5,7, 9'
        self.assertEqual(prog_a._created_ids_list(), [5, 7, 9])

        # Usuario de nómina restringido a la compañía B: la ir.rule del
        # historial solo le muestra sus importaciones.
        user_b = self.env['res.users'].create({
            'name': 'Nómina Beta F8',
            'login': 'fase8_import_beta',
            'company_id': self.company_b.id,
            'company_ids': [(6, 0, [self.company_b.id])],
            'group_ids': [(6, 0, [
                self.env.ref('base.group_user').id,
                self.env.ref('hr_payroll.group_hr_payroll_user').id,
            ])],
        })
        visibles = Progress.with_user(user_b).with_context(
            allowed_company_ids=self.company_b.ids).search([
                ('id', 'in', (prog_a | prog_b).ids)])
        self.assertEqual(visibles, prog_b)
