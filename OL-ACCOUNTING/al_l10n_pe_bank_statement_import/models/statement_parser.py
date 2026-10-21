"""Lectura de los archivos de movimientos de la banca por internet peruana.

Cada banco exporta su propio formato y lo cambia sin aviso, por eso no se leen
posiciones fijas: se busca la fila de encabezados (puede haber títulos y datos
de la cuenta arriba) y cada columna se reconoce por su nombre.

Funciones puras (sin ORM) para poder probarlas con archivos de ejemplo.
"""
import csv
import datetime
import hashlib
import io
import re
import unicodedata

#: Nombres de columna reconocidos, normalizados (minúsculas, sin tildes ni signos).
#: La primera coincidencia de cada lista, en orden, es la que se usa.
COLUMNS = {
    'date': ['fecha operacion', 'fecha de operacion', 'f operacion', 'fec operacion', 'fecha', 'fecha movimiento',
             'fecha de movimiento', 'fecha transaccion', 'fecha de transaccion', 'fecha proceso',
             'fecha de proceso', 'fecha contable'],
    'description': ['descripcion operacion', 'descripcion de la operacion', 'descripcion', 'concepto', 'detalle',
                    'glosa', 'descripcion del movimiento', 'movimiento', 'transaccion', 'operacion'],
    'ref': ['operacion numero', 'n operacion', 'no operacion', 'nro operacion', 'nro de operacion',
            'numero de operacion', 'num operacion', 'n de operacion', 'no de operacion', 'n doc', 'no doc',
            'nro doc', 'nro documento', 'no documento',
            'numero de documento', 'documento', 'referencia', 'referencia2', 'nro de referencia'],
    'amount': ['monto', 'importe', 'monto s', 'importe s', 'monto us', 'importe us', 'valor'],
    'debit': ['cargo', 'cargos', 'debe', 'retiro', 'retiros', 'debito', 'debitos', 'cargo s', 'cargos s'],
    'credit': ['abono', 'abonos', 'haber', 'deposito', 'depositos', 'credito', 'creditos', 'abono s', 'abonos s'],
    'balance': ['saldo', 'saldo contable', 'saldo disponible', 'saldo s', 'saldo final'],
}
DATE_FORMATS = ['%d/%m/%Y', '%d/%m/%y', '%d-%m-%Y', '%d-%m-%y', '%Y-%m-%d', '%d.%m.%Y', '%Y/%m/%d']
HEADER_SCAN_ROWS = 40


class StatementParseError(ValueError):
    """El archivo no es un extracto reconocible."""


def normalize(text):
    text = unicodedata.normalize('NFKD', str(text or '')).encode('ascii', 'ignore').decode().lower()
    return ' '.join(re.sub(r'[^a-z0-9]+', ' ', text).split())


# ----------------------------------------------------------------------
# Lectura del archivo
# ----------------------------------------------------------------------

def read_rows(raw):
    """Filas del archivo como listas de celdas (texto, número o fecha)."""
    if raw[:4] == b'PK\x03\x04':
        return _read_xlsx(raw)
    if raw[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
        return _read_xls(raw)
    head = raw[:4096].lower()
    if b'<table' in raw[:500000].lower() or head.lstrip().startswith((b'<html', b'<!doctype')):
        return _read_html(raw)
    return _read_text(raw)


def _read_xlsx(raw):
    import openpyxl  # noqa: PLC0415 - dependencia externa declarada en el manifiesto
    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _read_xls(raw):
    import xlrd  # noqa: PLC0415 - dependencia externa declarada en el manifiesto
    book = xlrd.open_workbook(file_contents=raw)
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
            else:
                row.append(cell.value)
        rows.append(row)
    return rows


def _read_html(raw):
    from lxml import html  # noqa: PLC0415
    document = html.fromstring(_decode(raw))
    rows = []
    for tr in document.iter('tr'):
        rows.append([' '.join(cell.text_content().split()) for cell in tr if cell.tag in ('td', 'th')])
    return rows


def _decode(raw):
    for encoding in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode('latin-1', errors='replace')


def _read_text(raw):
    text = _decode(raw)
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        raise StatementParseError('Archivo vacío.')
    sample = '\n'.join(lines[:50])
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=';,\t|').delimiter
    except csv.Error:
        delimiter = max(';,\t|', key=sample.count)
    return list(csv.reader(lines, delimiter=delimiter))


# ----------------------------------------------------------------------
# Valores
# ----------------------------------------------------------------------

def parse_amount(value):
    """Importe de una celda, o ``None`` si no es un número."""
    if value is None or value == '':
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r'(?i)s/\.?|us\$|\$|pen|usd|\s', '', str(value))
    negative = (text.startswith('(') and text.endswith(')')) or text.startswith('-') or text.endswith('-')
    text = text.strip('()-+')
    if not text or not re.fullmatch(r'[\d.,]+', text):
        return None
    if ',' in text and '.' in text:
        decimal = ',' if text.rfind(',') > text.rfind('.') else '.'
    elif ',' in text:
        decimal = ',' if text.count(',') == 1 and re.search(r',\d{1,2}$', text) else None
    elif '.' in text:
        decimal = '.' if text.count('.') == 1 else None
    else:
        decimal = None
    if decimal == ',':
        text = text.replace('.', '').replace(',', '.')
    else:
        text = text.replace(',', '') if decimal == '.' else text.replace(',', '').replace('.', '')
    try:
        amount = float(text)
    except ValueError:
        return None
    return -amount if negative else amount


def parse_date(value):
    """Fecha de una celda, o ``None``."""
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value or '').strip()
    if not text:
        return None
    token = text.split()[0]
    for fmt in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(token, fmt).date()
        except ValueError:
            continue
    return None


# ----------------------------------------------------------------------
# Encabezados y movimientos
# ----------------------------------------------------------------------

def find_header(rows):
    """(índice de la fila de encabezados, {campo: columna}) o error."""
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        names = [normalize(cell) for cell in row]
        columns = {}
        for field, synonyms in COLUMNS.items():
            for synonym in synonyms:
                if synonym in names:
                    column = names.index(synonym)
                    if column not in columns.values():
                        columns[field] = column
                        break
        if 'date' in columns and ('amount' in columns or ('debit' in columns and 'credit' in columns)):
            return index, columns
    raise StatementParseError('No se encontró la fila de encabezados (fecha y monto, o cargo y abono).')


def _cell(row, column):
    return row[column] if column is not None and column < len(row) else None


def parse_statement(raw):
    """Lee el archivo y devuelve ``{'transactions': [...], 'balance_start', 'balance_end_real',
    'date_from', 'date_to'}`` con los movimientos en orden cronológico."""
    rows = read_rows(raw)
    header, columns = find_header(rows)
    transactions = []
    for row in rows[header + 1:]:
        day = parse_date(_cell(row, columns['date']))
        if not day:
            continue  # filas en blanco, subtotales o pie
        if 'amount' in columns:
            amount = parse_amount(_cell(row, columns['amount']))
        else:
            debit = parse_amount(_cell(row, columns['debit'])) or 0.0
            credit = parse_amount(_cell(row, columns['credit'])) or 0.0
            amount = abs(credit) - abs(debit)
        if not amount:
            continue
        description = ' '.join(str(_cell(row, columns.get('description')) or '').split())
        ref = ' '.join(str(_cell(row, columns.get('ref')) or '').split())
        if ref.endswith('.0') and ref[:-2].isdigit():
            ref = ref[:-2]  # número de operación leído como número en Excel
        transactions.append({
            'date': day,
            'amount': round(amount, 2),
            'payment_ref': description or ref or '/',
            'ref': ref,
            'balance': parse_amount(_cell(row, columns.get('balance'))),
        })
    if not transactions:
        raise StatementParseError('El archivo no tiene movimientos.')
    # Los bancos suelen listar del más reciente al más antiguo.
    if transactions[0]['date'] > transactions[-1]['date']:
        transactions.reverse()
    result = {
        'transactions': transactions,
        'date_from': transactions[0]['date'],
        'date_to': transactions[-1]['date'],
        'columns': columns,
    }
    first, last = transactions[0], transactions[-1]
    if first['balance'] is not None and last['balance'] is not None:
        result['balance_start'] = round(first['balance'] - first['amount'], 2)
        result['balance_end_real'] = last['balance']
    _set_unique_ids(transactions)
    return result


def _set_unique_ids(transactions):
    """Identificador estable de cada movimiento: volver a importar el mismo
    archivo (o uno que se superpone) no duplica movimientos. Dos movimientos
    idénticos el mismo día se distinguen por su orden."""
    seen = {}
    for transaction in transactions:
        key = '%s|%.2f|%s|%s' % (transaction['date'], transaction['amount'], transaction['ref'],
                                 transaction['payment_ref'])
        seen[key] = seen.get(key, 0) + 1
        digest = hashlib.sha1(('%s|%d' % (key, seen[key])).encode()).hexdigest()[:16]
        transaction['unique_import_id'] = 'pe-%s' % digest
