# -*- coding: utf-8 -*-
"""Importación masiva de constancias de depósito de detracciones.

Lee el archivo de la consulta de constancias de depósito de SUNAT Operaciones
en Línea (Excel o CSV; columnas por su encabezado, en cualquier orden) y, por
cada constancia, busca el comprobante:

- compras: el proveedor de la fila es el tercero de la factura de compra;
- ventas: el proveedor de la fila es la compañía y el adquiriente, el cliente.

Llena el número y la fecha de la constancia (PLE 8.1, campos 32-33) y, si se
pide, registra el depósito con el mismo asistente del botón «Registrar
depósito». Primero muestra una vista previa sin escribir nada.
"""
import base64
import csv
import io
import re
import unicodedata
from datetime import date, datetime

from markupsafe import Markup, escape

from odoo import api, fields, models
from odoo.exceptions import UserError

# Encabezado normalizado → clave interna. Varias redacciones por columna.
HEADERS = {
    'numero constancia': 'constancy', 'numero de constancia': 'constancy',
    'nro constancia': 'constancy', 'constancia': 'constancy',
    'ruc proveedor': 'supplier_vat', 'numero de ruc del proveedor': 'supplier_vat',
    'numero de documento adquiriente': 'buyer_vat', 'numero documento adquiriente': 'buyer_vat',
    'numero de documento del adquiriente': 'buyer_vat',
    'fecha pago': 'date', 'fecha de pago': 'date', 'fecha deposito': 'date',
    'monto deposito': 'amount', 'monto del deposito': 'amount', 'monto': 'amount',
    'tipo de comprobante': 'doc_type', 'tipo comprobante': 'doc_type',
    'serie de comprobante': 'serie', 'serie comprobante': 'serie', 'serie': 'serie',
    'numero de comprobante': 'number', 'numero comprobante': 'number',
}
REQUIRED = ('constancy', 'supplier_vat', 'date', 'serie', 'number')


def _norm(text):
    text = unicodedata.normalize('NFKD', str(text or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', ' ', text.lower()).strip()


def _digits(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r'\D', '', str(value or ''))


def _doc_key(serie, number):
    """«F001-00000123» y «F001-123» son el mismo comprobante."""
    serie = str(serie or '').strip().upper()
    number = _digits(number).lstrip('0') or '0'
    return '%s-%s' % (serie, number)


def _move_keys(move):
    """Serie y número del comprobante, con el mismo formato que la fila, desde
    el número del documento, el nombre o la referencia del proveedor."""
    keys = set()
    for number in (move.l10n_latam_document_number, move.name, move.ref):
        number = (number or '').strip()
        if '-' not in number:
            continue
        serie, correlative = number.rsplit('-', 1)
        keys.add(_doc_key(serie.split()[-1] if serie.split() else serie, correlative))
    return keys


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or '').strip()
    for fmt in ('%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d', '%d-%m-%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return False


def _parse_amount(value):
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value or '').strip().replace(',', '')
    try:
        return float(text)
    except ValueError:
        return 0.0


class L10nPeDetractionImportWizard(models.TransientModel):
    _name = 'l10n_pe.detraction.import.wizard'
    _description = 'Importar constancias de detracción'
    _check_company_auto = True

    company_id = fields.Many2one('res.company', string='Compañía', required=True,
                                 default=lambda self: self.env.company)
    file = fields.Binary(string='Archivo', required=True, attachment=False)
    filename = fields.Char(string='Nombre del archivo')
    register_deposit = fields.Boolean(
        string='Registrar el depósito',
        help='Además de la constancia, registra el pago de la detracción (compras) o el cobro '
             '(ventas) desde el diario indicado, como el botón «Registrar depósito».')
    journal_id = fields.Many2one(
        'account.journal', string='Diario del depósito', check_company=True,
        domain="[('type', 'in', ('bank', 'cash'))]")
    preview_html = fields.Html(string='Vista previa', compute='_compute_preview', sanitize=False)

    # ------------------------------------------------------------------
    # Lectura
    # ------------------------------------------------------------------
    def _read_rows(self):
        """Filas del archivo como dicts por clave interna."""
        self.ensure_one()
        content = base64.b64decode(self.with_context(bin_size=False).file or b'')
        rows = self._read_excel(content)
        if rows is None:
            rows = self._read_csv(content)
        if not rows:
            raise UserError(self.env._('El archivo está vacío o no se pudo leer.'))
        header_index = next((i for i, row in enumerate(rows[:10])
                             if sum(1 for cell in row if _norm(cell) in HEADERS) >= 3), None)
        if header_index is None:
            raise UserError(self.env._(
                'No se reconocen las columnas. El archivo debe traer al menos: número de '
                'constancia, RUC del proveedor, fecha de pago, serie y número de comprobante.'))
        columns = {i: HEADERS[_norm(cell)] for i, cell in enumerate(rows[header_index])
                   if _norm(cell) in HEADERS}
        missing = [key for key in REQUIRED if key not in columns.values()]
        if missing:
            raise UserError(self.env._('Faltan columnas en el archivo: %s.', ', '.join(missing)))
        result = []
        for number, row in enumerate(rows[header_index + 1:], start=header_index + 2):
            values = {key: row[i] if i < len(row) else '' for i, key in columns.items()}
            if not any(str(v).strip() for v in values.values()):
                continue
            values['_row'] = number
            result.append(values)
        return result

    def _read_excel(self, content):
        try:
            from openpyxl import load_workbook
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        except Exception:  # noqa: BLE001 - no es un .xlsx: se prueba .xls y CSV
            pass
        try:
            import xlrd
            sheet = xlrd.open_workbook(file_contents=content).sheet_by_index(0)
            return [sheet.row_values(i) for i in range(sheet.nrows)]
        except Exception:  # noqa: BLE001
            return None

    def _read_csv(self, content):
        text = content.decode('utf-8-sig', errors='replace')
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=',;|\t') if text.strip() else csv.excel
        return [row for row in csv.reader(io.StringIO(text), dialect)]

    # ------------------------------------------------------------------
    # Emparejamiento
    # ------------------------------------------------------------------
    def _match(self):
        """``[(fila, comprobante o vacío, estado, detalle)]``."""
        self.ensure_one()
        company_vat = _digits(self.company_id.vat)
        Move = self.env['account.move']
        result = []
        for row in self._read_rows():
            constancy = _digits(row.get('constancy'))
            day = _parse_date(row.get('date'))
            supplier = _digits(row.get('supplier_vat'))
            buyer = _digits(row.get('buyer_vat'))
            key = _doc_key(row.get('serie'), row.get('number'))
            if not constancy or not day:
                result.append((row, Move, 'error', self.env._('Constancia o fecha no válida')))
                continue
            if supplier == company_vat:
                move_type, partner_vat = 'out_invoice', buyer
            else:
                move_type, partner_vat = 'in_invoice', supplier
            candidates = Move.search([
                ('company_id', '=', self.company_id.id), ('move_type', '=', move_type),
                ('state', '=', 'posted'), ('commercial_partner_id.vat', '=', partner_vat)])
            move = candidates.filtered(lambda m: key in _move_keys(m))[:1]
            if not move:
                result.append((row, Move, 'missing', self.env._('Comprobante no encontrado')))
            elif move.l10n_pe_detraction_number and move.l10n_pe_detraction_number != constancy:
                result.append((row, move, 'conflict', self.env._(
                    'Ya tiene la constancia %s', move.l10n_pe_detraction_number)))
            elif move.l10n_pe_detraction_number:
                result.append((row, move, 'done', self.env._('Ya registrada')))
            else:
                amount = _parse_amount(row.get('amount'))
                expected = move.l10n_pe_detraction_amount
                note = ''
                if amount and expected and move.company_currency_id.compare_amounts(amount, expected):
                    note = self.env._('Monto %(file)s distinto del calculado %(calc)s',
                                      file=amount, calc=expected)
                result.append((row, move, 'ok', note))
        return result

    @api.depends('file', 'company_id')
    def _compute_preview(self):
        labels = {
            'ok': self.env._('Por registrar'), 'done': self.env._('Ya registrada'),
            'conflict': self.env._('Otra constancia'), 'missing': self.env._('No encontrado'),
            'error': self.env._('Fila no válida'),
        }
        for wizard in self:
            if not wizard.file:
                wizard.preview_html = False
                continue
            try:
                matches = wizard._match()
            except UserError as exc:
                wizard.preview_html = Markup('<div class="alert alert-danger">%s</div>') % exc.args[0]
                continue
            rows = Markup('').join(
                Markup('<tr><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>') % (
                    row['_row'], _digits(row.get('constancy')),
                    move.display_name if move else _doc_key(row.get('serie'), row.get('number')),
                    labels[state], note)
                for row, move, state, note in matches)
            counts = {state: sum(1 for m in matches if m[2] == state) for state in labels}
            summary = ' · '.join('%s: %s' % (labels[s], c) for s, c in counts.items() if c)
            wizard.preview_html = Markup(
                '<p>%s</p><table class="table table-sm o_main_table"><thead><tr><th>%s</th>'
                '<th>%s</th><th>%s</th><th>%s</th><th>%s</th></tr></thead><tbody>%s</tbody>'
                '</table>') % (
                summary, self.env._('Fila'), self.env._('Constancia'), self.env._('Comprobante'),
                self.env._('Estado'), self.env._('Detalle'), rows)

    # ------------------------------------------------------------------
    # Importación
    # ------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        if self.register_deposit and not self.journal_id:
            raise UserError(self.env._('Indique el diario del depósito.'))
        done = self.env['account.move']
        for row, move, state, _note in self._match():
            if state != 'ok':
                continue
            constancy = _digits(row.get('constancy'))
            day = _parse_date(row.get('date'))
            if self.register_deposit:
                amount = _parse_amount(row.get('amount')) or move.l10n_pe_detraction_amount
                self.env['l10n_pe.detraction.deposit.wizard'].create({
                    'move_id': move.id, 'journal_id': self.journal_id.id, 'amount': amount,
                    'payment_date': day, 'constancy_number': constancy,
                }).action_confirm()
            else:
                move.write({'l10n_pe_detraction_number': constancy,
                            'l10n_pe_detraction_date': day})
            done |= move
        if not done:
            raise UserError(self.env._('No hay constancias por registrar en el archivo.'))
        return {
            'type': 'ir.actions.act_window',
            'name': self.env._('Constancias registradas'),
            'res_model': 'account.move',
            'view_mode': 'list,form',
            'domain': [('id', 'in', done.ids)],
        }
