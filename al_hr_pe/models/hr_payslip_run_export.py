# -*- coding: utf-8 -*-
"""Exportadores PLAME / AFPNet del lote de nóminas (``hr.payslip.run``).

Port v18 → v19 de los 5 exportadores de
``al_hr_payroll/hr_fields/models/hr_payslip_run.py``:

=========================  =========  =============================================
Método                     Archivo    Contenido
=========================  =========  =============================================
``export_plame``           ``.rem``   Remuneraciones por concepto (Tabla 22 SUNAT)
``export_plame_hours``     ``.jor``   Jornada laboral (horas ordinarias y extras)
``export_plame_suspencion``  ``.snl``   Suspensiones de labores (Tabla 21 SUNAT)
``export_plame_other_conditions``  ``.toc``  Otras condiciones (SCTR / Vida Ley)
``afp_net``                ``.xlsx``  Plantilla de importación AFPNet
=========================  =========  =============================================

Cambios estructurales respecto a v18 (plan §6.5 «attachments, no filesystem»):

* Se elimina por completo ``hr.main.parameter.dir_create_file`` y los
  ``open()`` a disco: cada método construye el contenido **en memoria**
  y lo publica como ``ir.attachment`` vinculado al lote, devolviendo una
  acción de descarga ``/web/content/<id>?download=true``.
* Las consultas SQL crudas con ``.format()`` (inyección + acoplamiento al
  esquema) se reescriben con el ORM sobre ``hr.payslip`` /
  ``hr.payslip.line`` / ``hr.payslip.worked_days``.
* ``hr.contract`` → ``hr.version`` (``payslip.version_id``); los campos PE
  viven en la versión (``membership_id``, ``situation_id``,
  ``l10n_pe_cuspp``, ``l10n_pe_exception``, ``l10n_pe_work_type``, …).
* ``hr.type.document`` → ``l10n_latam.identification.type`` con
  ``l10n_pe_hr_sunat_code`` / ``l10n_pe_hr_afp_code`` (plan §6.4).
* AFPNet se genera con ``openpyxl`` (antes ``xlsxwriter`` + archivo en
  disco).

El FORMATO DE SALIDA (separador ``|``, orden de columnas, fin de línea
``\\r\\n``, extensiones y nombres de archivo ``0601AAAAMM<RUC>.<ext>``)
se conserva idéntico al v18: es el formato oficial de importación de
PLAME/AFPNet.
"""
import base64
import io
from math import modf

from odoo import models
from odoo.exceptions import UserError


class HrPayslipRun(models.Model):
    _inherit = 'hr.payslip.run'

    # ------------------------------------------------------------------
    # Helpers comunes
    # ------------------------------------------------------------------
    def _l10n_pe_check_single(self, message):
        """Valida que solo haya un lote seleccionado (mismo control v18)."""
        if len(self.ids) > 1:
            raise UserError(message)
        self.ensure_one()

    def _l10n_pe_plame_filename(self, extension):
        """Nombre oficial del archivo PLAME: ``0601AAAAMM<RUC>.<ext>``.

        v18 componía ``'0601' + año + mes`` a partir de ``date_end`` del
        lote y el RUC de la compañía (``company_id.vat``). Se conserva.
        """
        self.ensure_one()
        if not self.date_end:
            raise UserError(self.env._(
                'El lote no tiene Fecha de fin: no se puede componer el '
                'nombre del archivo PLAME (0601AAAAMM<RUC>).'))
        return '0601%s%s%s.%s' % (
            self.date_end.strftime('%Y'),          # AAAA del periodo
            self.date_end.strftime('%m'),          # MM del periodo
            self.company_id.vat or '',              # RUC del empleador
            extension,
        )

    def _l10n_pe_download_attachment(self, filename, content):
        """Publica ``content`` (bytes) como adjunto del lote y lo descarga.

        Sustituye al par v18 «escribir en ``dir_create_file`` + popup.it»:
        el archivo queda trazable en el chatter del lote y la descarga
        funciona en despliegues multi-worker/cloud.
        """
        self.ensure_one()
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(content),
            'res_model': 'hr.payslip.run',
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'self',
        }

    @staticmethod
    def _l10n_pe_doc_type(employee):
        """Código SUNAT del tipo de documento del empleado (Tabla 3).

        v18: ``he.type_document_id.sunat_code`` (LEFT JOIN, podía imprimir
        ``None``). v19: extensión latam nativa. Se devuelve ``''`` cuando
        falta el dato en lugar del ``None`` literal que colaba el SQL.
        """
        return (employee.l10n_latam_identification_type_id
                .l10n_pe_hr_sunat_code or '')

    def _l10n_pe_get_dlabs(self, slip, param):
        """Días efectivamente laborados = DLAB − subsidios − vacaciones.

        Port literal de ``hr.payslip.get_dlabs()`` v18 (aún no existe en
        el ``hr.payslip`` v19; se define aquí como helper privado).
        # TODO(fase2-revisar): cuando el motor de nómina (Fase 2) porte
        # get_dlabs() a hr.payslip, borrar este helper y delegar en él.
        """
        dlab_codes = param.wd_dlab.mapped('code')
        dsub_codes = param.wd_dsub.mapped('code')
        dvac_codes = param.wd_dvac.mapped('code')
        dlab = slip.worked_days_line_ids.filtered(
            lambda wd: wd.code in dlab_codes)
        dsub = slip.worked_days_line_ids.filtered(
            lambda wd: wd.code in dsub_codes)
        dvac = slip.worked_days_line_ids.filtered(
            lambda wd: wd.code in dvac_codes)
        return (sum(dlab.mapped('number_of_days'))
                - sum(dsub.mapped('number_of_days'))
                - sum(dvac.mapped('number_of_days')))

    # ------------------------------------------------------------------
    # PLAME .rem — Remuneraciones por concepto
    # ------------------------------------------------------------------
    def export_plame(self):
        """Genera el ``.rem`` de PLAME (ingresos/descuentos por concepto).

        Estructura de cada línea (idéntica a v18)::

            tipo_doc|nro_doc|codigo_sunat|monto_devengado|monto_pagado|\\r\\n

        1. ``tipo_doc``: código SUNAT del tipo de documento (Tabla 3).
        2. ``nro_doc``: número de documento del trabajador.
        3. ``codigo_sunat``: concepto remunerativo (Tabla 22, campo
           ``sunat_code`` de la regla salarial).
        4. ``monto_devengado``: suma de los totales de las líneas.
        5. ``monto_pagado``: en v18 era el mismo total (devengado =
           pagado); se conserva.

        Reglas de inclusión (traducción de la UNION ALL v18):

        * Rama 1 — cualquier concepto con ``sunat_code`` distinto de
          ``0804/0607/0605/0601`` y total ≠ 0.
        * Rama 2 — conceptos previsionales ``0605``/``0601``: si la
          afiliación es ONP o «SIN REGIMEN» solo entra ``0605``; para AFP
          entran ``0605`` y ``0601``. (Sin filtro de total ≠ 0, como en
          v18.) Requiere afiliación en la versión (era INNER JOIN).

        Ambas ramas son disjuntas por código, así que no hay doble conteo.
        Se agrupa por (nro_doc, codigo_sunat) y se ordena igual que el
        ``ORDER BY a1.dni, a1.sunat`` original.
        """
        self._l10n_pe_check_single(self.env._(
            'Solo se puede procesar una planilla a la vez, '
            'seleccione una sola nómina'))

        # Códigos previsionales tratados aparte en la rama 2 del v18.
        pension_codes = ('0605', '0601')
        excluded_branch1 = ('0804', '0607', '0605', '0601')
        # Afiliaciones que solo reportan 0605 (nombres literales v18).
        # TODO(fase2-revisar): v18 comparaba hm.name IN ('ONP','SIN REGIMEN')
        # — depende del nombre exacto del registro hr.membership; valorar
        # sustituirlo por un flag is_afp (ONP/sin régimen = not is_afp).
        onp_names = ('ONP', 'SIN REGIMEN')

        grouped = {}  # (nro_doc, codigo_sunat) -> dict acumulador
        for slip in self.slip_ids:
            employee = slip.employee_id
            dni = employee.identification_id or ''
            doc_type = self._l10n_pe_doc_type(employee)
            membership = slip.version_id.membership_id
            for line in slip.line_ids:
                code = line.salary_rule_id.sunat_code
                if not code:
                    continue
                if code not in excluded_branch1:
                    # Rama 1: conceptos generales con total distinto de 0.
                    if not line.total:
                        continue
                elif code in pension_codes and membership:
                    # Rama 2: previsionales según afiliación.
                    if membership.name in onp_names and code != '0605':
                        continue
                else:
                    # 0804/0607 fuera de rama 1 y sin rama 2 aplicable,
                    # o previsional sin afiliación (INNER JOIN v18).
                    continue
                key = (dni, code)
                vals = grouped.setdefault(key, {
                    'doc_type': doc_type,
                    'amount_earn': 0.0,
                    'amount_paid': 0.0,
                })
                # v18: amount_earn = amount_paid = hpl.total.
                vals['amount_earn'] += line.total
                vals['amount_paid'] += line.total

        # ORDER BY dni, sunat del SQL original.
        output = io.StringIO()
        for (dni, code) in sorted(grouped):
            vals = grouped[(dni, code)]
            # Los montos se emiten con 2 decimales: el SQL v18 devolvía
            # NUMERIC(x,2) que psycopg2 renderizaba como '1500.00'.
            output.write('%s|%s|%s|%.2f|%.2f|\r\n' % (
                vals['doc_type'],
                dni,
                code,
                vals['amount_earn'],
                vals['amount_paid'],
            ))

        filename = self._l10n_pe_plame_filename('rem')
        return self._l10n_pe_download_attachment(
            filename, output.getvalue().encode('utf-8'))

    # ------------------------------------------------------------------
    # PLAME .jor — Jornada laboral
    # ------------------------------------------------------------------
    def export_plame_hours(self):
        """Genera el ``.jor`` de PLAME (jornada por trabajador).

        Estructura de cada línea (idéntica a v18)::

            tipo_doc|nro_doc|horas_ordinarias|0|horas_extras|0|\\r\\n

        1. ``tipo_doc``: código SUNAT del tipo de documento.
        2. ``nro_doc``: número de documento del trabajador.
        3. ``horas_ordinarias``: parte entera de
           «días laborados efectivos × horas por día del calendario»
           (v18: ``modf(get_dlabs() * hours_per_day)`` truncado con %d).
        4. Minutos ordinarios: siempre ``0`` (v18 hardcodeado).
        5. ``horas_extras``: suma (truncada a entero) de las horas de las
           work entries clasificadas como sobretiempo (``wd_ext``).
        6. Minutos extras: siempre ``0`` (v18 hardcodeado).

        Solo se emiten los payslips con al menos una línea de días
        trabajados cuyos tipos estén clasificados en Parámetros
        Principales como días no laborados (``wd_dnlab``), sobretiempo
        (``wd_ext``) o ausencias/vacaciones (``wd_dvac``): es la
        traducción del INNER JOIN + WHERE del SQL v18. (El SQL también
        sumaba faltas y vacaciones, pero nunca los volcaba al archivo;
        aquí no se calculan.) El orden de salida es por número de
        documento, como el ``ORDER BY he.identification_id``.
        """
        self._l10n_pe_check_single(self.env._(
            'Solo se puede procesar una planilla a la vez, '
            'seleccione una sola nómina'))
        param = self.env['hr.main.parameter'].get_main_parameter()
        # v18 validaba la configuración de work entries dentro de
        # get_dlabs() vía check_voucher_values(); se conserva el control.
        param.check_voucher_values()

        gating_types = param.wd_dnlab | param.wd_ext | param.wd_dvac
        ext_types = param.wd_ext

        output = io.StringIO()
        slips = self.slip_ids.sorted(
            key=lambda s: s.employee_id.identification_id or '')
        for slip in slips:
            relevant = slip.worked_days_line_ids.filtered(
                lambda wd: wd.work_entry_type_id in gating_types)
            if not relevant:
                # Sin líneas en fal/ext/vac: el INNER JOIN v18 lo excluía.
                continue
            hext = sum(relevant.filtered(
                lambda wd: wd.work_entry_type_id in ext_types
            ).mapped('number_of_hours'))
            # v18: horas por día del calendario del contrato.
            hours_per_day = slip.version_id.resource_calendar_id.hours_per_day
            dlab = self._l10n_pe_get_dlabs(slip, param)
            # v18: modf() separa la parte entera y %d la trunca.
            hlab = modf(dlab * hours_per_day)
            output.write('%s|%s|%d|0|%d|0|\r\n' % (
                self._l10n_pe_doc_type(slip.employee_id),
                slip.employee_id.identification_id or '',
                hlab[1],
                hext,
            ))

        filename = self._l10n_pe_plame_filename('jor')
        return self._l10n_pe_download_attachment(
            filename, output.getvalue().encode('utf-8'))

    # ------------------------------------------------------------------
    # PLAME .snl — Suspensiones de labores
    # ------------------------------------------------------------------
    def export_plame_suspencion(self):
        """Genera el ``.snl`` de PLAME (suspensiones — Tabla 21 SUNAT).

        Estructura de cada línea (idéntica a v18)::

            tipo_doc|nro_doc|tipo_suspension|nro_dias|\\r\\n

        1. ``tipo_doc``: código SUNAT del tipo de documento, rellenado a
           2 posiciones con ceros a la izquierda (``rjust(2, '0')``).
        2. ``nro_doc``: número de documento del trabajador.
        3. ``tipo_suspension``: código Tabla 21 (``hr.suspension.type``).
        4. ``nro_dias``: suma de días de todas las suspensiones de ese
           tipo del trabajador en el periodo del lote.

        v18 recorría ``contract.work_suspension_ids`` filtrado por
        ``periodo_id`` y emitía una línea por tipo (deduplicando con la
        lista «memoria»), sumando los días de todos los registros del
        mismo tipo en el periodo.
        """
        self._l10n_pe_check_single(self.env._(
            'Solo se puede procesar una planilla a la vez, '
            'seleccione una sola nómina'))

        # TODO(fase2-revisar): el modelo hr.work.suspension (v18 colgaba de
        # hr.contract) aún no está migrado a v19; el plan lo ubica en la
        # Fase 2 asociado a hr.version. Este export asume que existirá con
        # los campos employee_id, periodo_id, suspension_type_id y days.
        # Revisar también si el M2O final es version_id en vez de
        # employee_id y ajustar el domain.
        if 'hr.work.suspension' not in self.env:
            raise UserError(self.env._(
                'El registro de suspensiones de labores '
                '(hr.work.suspension) aún no está disponible en esta '
                'versión: se migra con el motor de nómina (Fase 2).'))
        Suspension = self.env['hr.work.suspension']

        output = io.StringIO()
        for slip in self.slip_ids:
            employee = slip.employee_id
            code = self._l10n_pe_doc_type(employee)
            # v18: sunat_code.rjust(2, '0') o '' si no hay tipo de doc.
            tdoc = code.rjust(2, '0') if code else ''
            ndoc = employee.identification_id or ''
            # Suspensiones del trabajador en el periodo del lote.
            lineas = Suspension.search([
                ('periodo_id', '=', self.periodo_id.id),
                ('employee_id', '=', employee.id),
            ])
            memoria = []  # tipos ya emitidos (dedupe, como v18)
            for line in lineas:
                tipo = line.suspension_type_id
                if tipo.code in memoria:
                    continue
                # v18 relanzaba un search por (periodo, contrato, tipo)
                # para sumar los días; aquí basta filtrar el recordset.
                total_dias = sum(lineas.filtered(
                    lambda l, t=tipo: l.suspension_type_id == t
                ).mapped('days'))
                output.write('%s|%s|%s|%s|\r\n' % (
                    tdoc,
                    ndoc,
                    tipo.code,
                    total_dias,
                ))
                memoria.append(tipo.code)

        filename = self._l10n_pe_plame_filename('snl')
        return self._l10n_pe_download_attachment(
            filename, output.getvalue().encode('utf-8'))

    # ------------------------------------------------------------------
    # PLAME .toc — Otras condiciones (Vida Ley)
    # ------------------------------------------------------------------
    def export_plame_other_conditions(self):
        """Genera el ``.toc`` de PLAME (otras condiciones del trabajador).

        Estructura de cada línea (idéntica a v18, incluido el campo
        vacío entre el 4º y el 6º)::

            tipo_doc|nro_doc|0|1||condicion|\\r\\n

        1. ``tipo_doc``: código SUNAT del tipo de documento.
        2. ``nro_doc``: número de documento del trabajador.
        3. Constante ``0`` (v18 hardcodeado).
        4. Constante ``1`` (v18 hardcodeado — indicador Vida Ley).
        5. Campo vacío (v18 emitía ``||``).
        6. ``condicion``: ``2`` si el empleado es «no domiciliado»,
           ``1`` en caso contrario.

        Se emite una línea por cada línea de boleta cuya regla salarial
        tenga código ``SVLEY`` (Seguro Vida Ley) con total ≠ 0, como el
        SQL v18 (que seleccionaba el monto pero no lo volcaba).
        """
        self._l10n_pe_check_single(self.env._(
            'Solo se puede procesar una planilla a la vez, '
            'seleccione una sola nómina'))

        output = io.StringIO()
        for slip in self.slip_ids:
            employee = slip.employee_id
            for line in slip.line_ids:
                if line.salary_rule_id.code != 'SVLEY' or not line.total:
                    continue
                # v18: CASE WHEN he.condition = 'not_domiciled'
                #      THEN '2' ELSE '1' END
                condition = '2' if employee.condition == 'not_domiciled' \
                    else '1'
                output.write('%s|%s|%s|%s||%s|\r\n' % (
                    self._l10n_pe_doc_type(employee),
                    employee.identification_id or '',
                    '0',
                    '1',
                    condition,
                ))

        filename = self._l10n_pe_plame_filename('toc')
        return self._l10n_pe_download_attachment(
            filename, output.getvalue().encode('utf-8'))

    # ------------------------------------------------------------------
    # AFPNet — plantilla XLSX de declaración de aportes AFP
    # ------------------------------------------------------------------
    def _l10n_pe_get_first_version(self, employee, last_version):
        """Primera «versión-contrato» del vínculo laboral vigente.

        Port de ``hr.contract.get_first_contract()`` v18, que recorría
        los contratos del empleado (de más reciente a más antiguo) usando
        la situación (Tabla 15; código ``'0'`` = baja/inactivo) para
        detectar dónde empieza el vínculo laboral actual.

        En v19 los contratos son versiones: se toma como «contrato» cada
        ``contract_date_start`` distinto, representado por su versión más
        reciente (que refleja la situación final de ese contrato).

        # TODO(fase2-revisar): validar la equivalencia contrato ≈ grupo de
        # versiones por contract_date_start cuando la Fase 2 defina cómo
        # se registran los ceses/realtas en hr.version, y confirmar que el
        # código de situación '0' sigue marcando la baja en la Tabla 15
        # migrada.
        """
        Version = self.env['hr.version']
        domain = [('employee_id', '=', employee.id)]
        if last_version and last_version.contract_date_start:
            domain.append(('contract_date_start', '<=',
                           last_version.contract_date_start))
        versions = Version.search(
            domain, order='contract_date_start desc, date_version desc')
        # Una versión representativa por fecha de inicio de contrato.
        contracts = []
        seen = set()
        for version in versions:
            key = version.contract_date_start
            if key in seen:
                continue
            seen.add(key)
            contracts.append(version)

        # Algoritmo v18 tal cual (aux / roll_back / delimiter).
        aux, roll_back = None, None
        delimiter = len(contracts)
        if delimiter > 1:
            for c, contract in enumerate(contracts):
                code = contract.situation_id.code
                if code == '0' and c == 0:
                    aux = (contract, c)
                    continue
                if code == '0' and aux and c - aux[1] == 1:
                    return aux[0]
                if code == '0' and aux and not c - aux[1] == 1:
                    return roll_back
                if code == '0' and not aux:
                    return roll_back
                if code != '0' and delimiter - 1 == c:
                    return contract
                roll_back = contract
            # v18 caía a None implícito; devolvemos recordset vacío para
            # que el llamador pueda operar sin AttributeError.
            return Version
        return contracts[0] if contracts else Version

    def afp_net(self):
        """Genera la plantilla XLSX de importación de AFPNet.

        Una fila por boleta cuyo trabajador esté afiliado a una AFP
        (``membership_id.is_afp``), sin cabecera, hoja «AFP NET» con
        pestaña azul. Columnas (mismo orden y anchos que v18):

        =====  =========================================================
        Col    Contenido
        =====  =========================================================
        A      Correlativo: índice de la boleta dentro del lote (v18
               usaba el índice del ``enumerate`` sobre TODOS los slips,
               con lo que los no-AFP dejan huecos en la numeración; se
               conserva esa peculiaridad por fidelidad de formato).
        B      CUSPP del afiliado (``l10n_pe_cuspp``).
        C      Código AFP del tipo de documento
               (``l10n_pe_hr_afp_code``).
        D      Número de documento.
        E      Apellido paterno.
        F      Apellido materno.
        G      Nombres.
        H      ¿Se devenga remuneración? S/N: «S» si el contrato sigue
               vigente al cierre (fin de contrato ≥ inicio del lote) o,
               sin fecha fin, si la situación no es baja (código '0').
        I      ¿Inicio de relación laboral en el periodo? S/N: «S» si el
               primer contrato del vínculo empieza dentro del lote.
        J      ¿Cese en el periodo? S/N: «S» si la situación es baja
               ('0') y la fecha fin de contrato cae dentro del lote.
        K      Excepción de jornada (L/U/J/I/P/O) o vacío.
        L      Remuneración asegurable: total de la línea de la regla
               «R.S. Ingresos afectos AFP» de Parámetros Principales.
        M-O    Ceros (aporte voluntario del afiliado / del empleador /
               comisión — v18 emitía 0.00 fijo).
        P      Tipo de labor (N/C/M/P; «N» por defecto).
        =====  =========================================================
        """
        # openpyxl en lugar de xlsxwriter (regla de la migración); import
        # local como en v18 para no cargar la librería en cada request.
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter

        self._l10n_pe_check_single(self.env._(
            'No se puede seleccionar más de un registro para este proceso'))
        param = self.env['hr.main.parameter'].get_main_parameter()
        insurable_rule = param.insurable_remuneration

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = 'AFP NET'
        worksheet.sheet_properties.tabColor = '0000FF'  # v18: 'blue'

        # Equivalente openpyxl del formato v18 «numberdosespecial»:
        # num_format '0.00', alineado derecha/centro, Times New Roman 11.
        number_font = Font(name='Times New Roman', size=11)
        number_align = Alignment(horizontal='right', vertical='center')

        def write_number(row, col, value):
            """Escribe un monto con el formato numérico de 2 decimales."""
            cell = worksheet.cell(row=row, column=col, value=value)
            cell.number_format = '0.00'
            cell.font = number_font
            cell.alignment = number_align

        x = 0  # fila de salida (solo avanza con afiliados AFP, como v18)
        for c, slip in enumerate(self.slip_ids):
            version = slip.version_id
            if not version.membership_id.is_afp:
                continue
            employee = slip.employee_id
            first_version = self._l10n_pe_get_first_version(
                employee, version)
            # v18: hr.payslip.line de la regla de remuneración asegurable.
            # (v18 hacía search y accedía a .total: con >1 línea habría
            # reventado; aquí se toma la primera por seguridad.)
            ir_line = slip.line_ids.filtered(
                lambda l: l.salary_rule_id == insurable_rule)[:1] \
                if insurable_rule else slip.line_ids.browse()

            # Columna H — «¿devenga remuneración?» (lógica v18 literal):
            # con fecha fin de contrato: S si fin ≥ inicio del lote;
            # sin fecha fin: N solo si la situación es baja ('0').
            date_end = version.contract_date_end
            if date_end:
                if date_end >= self.date_end:
                    resul = 'S'
                elif self.date_start <= date_end <= self.date_end:
                    resul = 'S'
                else:
                    resul = 'N'
            else:
                resul = 'N' if version.situation_id.code == '0' else 'S'

            # Columna I — inicio del vínculo dentro del periodo del lote.
            # (v18 crasheaba si get_first_contract devolvía None; aquí un
            # vacío/sin fecha produce 'N'.)
            first_start = first_version.contract_date_start \
                if first_version else False
            alta = 'S' if (first_start
                           and self.date_start <= first_start
                           <= self.date_end) else 'N'

            # Columna J — cese dentro del periodo del lote.
            cese = 'S' if (version.situation_id.code == '0'
                           and date_end
                           and self.date_start <= date_end
                           <= self.date_end) else 'N'

            row = x + 1  # openpyxl es 1-indexado; v18 escribía desde 0
            worksheet.cell(row=row, column=1, value=c)
            worksheet.cell(row=row, column=2,
                           value=version.l10n_pe_cuspp or '')
            worksheet.cell(
                row=row, column=3,
                value=(employee.l10n_latam_identification_type_id
                       .l10n_pe_hr_afp_code or ''))
            worksheet.cell(row=row, column=4,
                           value=employee.identification_id or '')
            worksheet.cell(row=row, column=5,
                           value=employee.last_name or '')
            worksheet.cell(row=row, column=6,
                           value=employee.m_last_name or '')
            worksheet.cell(row=row, column=7, value=employee.names or '')
            worksheet.cell(row=row, column=8, value=resul)
            worksheet.cell(row=row, column=9, value=alta)
            worksheet.cell(row=row, column=10, value=cese)
            worksheet.cell(row=row, column=11,
                           value=version.l10n_pe_exception or '')
            write_number(row, 12, ir_line.total if ir_line.total else 0.00)
            write_number(row, 13, 0.00)
            write_number(row, 14, 0.00)
            write_number(row, 15, 0.00)
            worksheet.cell(row=row, column=16,
                           value=version.l10n_pe_work_type or 'N')
            x += 1

        # Anchos de columna v18: [2,15,2,12,20,20,20,2,2,2,2,8,8,8,8,2].
        widths = [2, 15, 2, 12, 20, 20, 20, 2, 2, 2, 2, 8, 8, 8, 8, 2]
        for i, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(i)].width = width

        buffer = io.BytesIO()
        workbook.save(buffer)
        return self._l10n_pe_download_attachment(
            'AFP_NET.xlsx', buffer.getvalue())
