# -*- coding: utf-8 -*-
"""Derechohabientes: vigencia, asignación familiar y exportación."""
import io
from datetime import date

from dateutil.relativedelta import relativedelta

from odoo import fields
from odoo.exceptions import UserError, ValidationError
from odoo.tests import tagged
from odoo.tests.common import TransactionCase


@tagged('post_install', '-at_install')
class TestDependents(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env['res.company'].create({
            'name': 'Derechohabientes Test S.A.C.',
            'vat': '20512528458',
        })
        cls.env.user.company_ids = [(4, cls.company.id)]
        cls.env = cls.env(context=dict(
            cls.env.context, allowed_company_ids=cls.company.ids))
        cls.dni = cls.env.ref('l10n_pe.it_DNI', raise_if_not_found=False) \
            or cls.env['l10n_latam.identification.type'].search(
                [('country_id.code', '=', 'PE')], limit=1)
        cls.employee = cls.env['hr.employee'].create({
            'name': 'Quispe Mamani Juan',
            'last_name': 'Quispe', 'm_last_name': 'Mamani', 'names': 'Juan',
            'company_id': cls.company.id,
            'identification_id': '45678901',
            'l10n_latam_identification_type_id': cls.dni.id,
        })
        cls.type_spouse = cls.env.ref('al_hr_pe.dependent_type_02')
        cls.type_partner = cls.env.ref('al_hr_pe.dependent_type_03')
        cls.type_child = cls.env.ref('al_hr_pe.dependent_type_05')
        cls.type_unborn = cls.env.ref('al_hr_pe.dependent_type_04')
        cls.reason_age = cls.env.ref('al_hr_pe.dependent_end_07')
        cls.today = fields.Date.context_today(cls.env['hr.employee'])

    def _dependent(self, type_id=None, years=None, **kwargs):
        vals = {
            'employee_id': self.employee.id,
            'type_id': (type_id or self.type_child).id,
            'last_name': 'Quispe', 'm_last_name': 'Flores',
            'names': 'Ana', 'gender': 'female',
            'l10n_latam_identification_type_id': self.dni.id,
            'identification_id': kwargs.pop('identification_id', '71234567'),
            'date_start': self.today - relativedelta(years=1),
        }
        if years is not None:
            vals['birthday'] = self.today - relativedelta(years=years)
        else:
            vals['birthday'] = self.today - relativedelta(years=5)
        vals.update(kwargs)
        return self.env['l10n_pe.hr.dependent'].create(vals)

    # ------------------------------------------------------------------
    # Datos básicos
    # ------------------------------------------------------------------
    def test_name_and_age(self):
        dependent = self._dependent(years=7)
        self.assertEqual(dependent.name, 'Quispe Flores Ana',
                         'apellidos y nombres en el orden de SUNAT')
        self.assertEqual(dependent.age, 7)

    def test_employee_counts_only_current(self):
        self._dependent(years=5)
        ended = self._dependent(years=9, identification_id='71234568')
        self.assertEqual(self.employee.l10n_pe_dependent_count, 2)
        ended.date_end = self.today - relativedelta(days=1)
        self.employee.invalidate_recordset()
        self.assertEqual(self.employee.l10n_pe_dependent_count, 1)

    def test_state(self):
        dependent = self._dependent()
        self.assertEqual(dependent.state, 'draft', 'aún no declarado')
        dependent.is_declared = True
        self.assertEqual(dependent.state, 'current')
        dependent.date_end = self.today - relativedelta(days=1)
        self.assertEqual(dependent.state, 'ended')

    # ------------------------------------------------------------------
    # Asignación familiar (Ley 25129)
    # ------------------------------------------------------------------
    def test_child_under_18_gives_allowance(self):
        dependent = self._dependent(years=10)
        self.assertTrue(dependent.gives_family_allowance)
        self.assertTrue(self.employee._l10n_pe_has_family_allowance())

    def test_child_over_18_does_not(self):
        dependent = self._dependent(years=19)
        self.assertFalse(dependent.gives_family_allowance)
        self.assertFalse(self.employee._l10n_pe_has_family_allowance())

    def test_student_up_to_24_does(self):
        dependent = self._dependent(years=22, is_studying=True)
        self.assertTrue(dependent.gives_family_allowance,
                        'hasta los 24 si cursa estudios superiores')
        dependent.is_studying = False
        self.assertFalse(dependent.gives_family_allowance)

    def test_student_over_24_does_not(self):
        dependent = self._dependent(years=25, is_studying=True)
        self.assertFalse(dependent.gives_family_allowance)

    def test_spouse_does_not_give_allowance(self):
        dependent = self._dependent(type_id=self.type_spouse, years=35)
        self.assertFalse(dependent.gives_family_allowance,
                         'la asignación nace de los hijos, no del cónyuge')

    def test_allowance_expires_by_itself(self):
        """El derecho caduca el día que el hijo cumple 18, sin tocar nada."""
        dependent = self._dependent(years=17)
        self.assertTrue(dependent._is_family_allowance_source(self.today))
        birthday_18 = dependent.birthday + relativedelta(years=18)
        self.assertFalse(
            dependent._is_family_allowance_source(birthday_18),
            'a los 18 cumplidos ya no corresponde')
        self.assertTrue(
            dependent._is_family_allowance_source(
                birthday_18 - relativedelta(days=1)))

    def test_ended_link_does_not_give_allowance(self):
        dependent = self._dependent(years=10)
        dependent.date_end = self.today - relativedelta(days=1)
        self.assertFalse(dependent._is_family_allowance_source(self.today))

    def test_falls_back_to_children_when_no_dependents(self):
        """Sin derechohabientes cargados se conserva el criterio anterior."""
        self.assertFalse(self.employee.l10n_pe_dependent_ids)
        self.assertFalse(self.employee._l10n_pe_has_family_allowance())
        self.employee.children = 2
        self.assertTrue(self.employee._l10n_pe_has_family_allowance())
        # En cuanto hay derechohabientes, mandan ellos.
        self._dependent(years=20)
        self.employee.invalidate_recordset()
        self.assertFalse(self.employee._l10n_pe_has_family_allowance(),
                         'el hijo mayor de edad no da derecho aunque '
                         'children siga en 2')

    # ------------------------------------------------------------------
    # Validaciones
    # ------------------------------------------------------------------
    def test_only_one_partner(self):
        self._dependent(type_id=self.type_spouse, years=35)
        with self.assertRaises(ValidationError):
            self._dependent(type_id=self.type_partner, years=33,
                            identification_id='71234569')

    def test_partner_allowed_after_ending_the_previous(self):
        first = self._dependent(type_id=self.type_spouse, years=35)
        first.date_end = self.today - relativedelta(days=1)
        second = self._dependent(type_id=self.type_partner, years=33,
                                 identification_id='71234569')
        self.assertTrue(second.id)

    def test_identity_required(self):
        with self.assertRaises(ValidationError):
            self._dependent(years=5, identification_id=False)

    def test_unborn_needs_no_identity_but_needs_due_date(self):
        with self.assertRaises(ValidationError):
            self.env['l10n_pe.hr.dependent'].create({
                'employee_id': self.employee.id,
                'type_id': self.type_unborn.id,
                'names': 'Madre gestante',
                'date_start': self.today,
            })
        dependent = self.env['l10n_pe.hr.dependent'].create({
            'employee_id': self.employee.id,
            'type_id': self.type_unborn.id,
            'names': 'Madre gestante',
            'date_start': self.today,
            'gestation_due_date': self.today + relativedelta(months=3),
        })
        self.assertFalse(dependent.birthday)
        self.assertFalse(dependent.gives_family_allowance)

    def test_end_before_start_rejected(self):
        dependent = self._dependent()
        with self.assertRaises(ValidationError):
            dependent.date_end = dependent.date_start - relativedelta(days=1)

    def test_duplicate_document_rejected(self):
        self._dependent(years=5)
        with self.assertRaises(Exception):
            self._dependent(years=8)
            self.env.flush_all()

    def test_action_set_end(self):
        dependent = self._dependent()
        dependent.action_set_end()
        self.assertEqual(dependent.date_end, self.today)

    # ------------------------------------------------------------------
    # Exportación para el T-Registro
    # ------------------------------------------------------------------
    def test_export_xlsx(self):
        from openpyxl import load_workbook

        alta = self._dependent(years=6)
        baja = self._dependent(years=19, identification_id='71234570',
                               date_end=self.today,
                               end_reason_id=self.reason_age.id)
        action = (alta | baja).action_export_tregistro_xlsx()
        self.assertEqual(action['type'], 'ir.actions.act_url')

        attachment = self.env['ir.attachment'].search(
            [('res_model', '=', 'res.company'),
             ('res_id', '=', self.company.id)], order='id desc', limit=1)
        self.assertIn('derechohabientes_20512528458', attachment.name)

        import base64
        workbook = load_workbook(io.BytesIO(base64.b64decode(attachment.datas)))
        self.assertEqual(workbook.sheetnames, ['Altas', 'Bajas'])

        altas = list(workbook['Altas'].values)
        self.assertEqual(len(altas), 2, 'cabecera + un alta')
        fila = altas[1]
        self.assertEqual(fila[1], '45678901', 'documento del trabajador')
        self.assertEqual(fila[4], '05', 'T19: hijo menor de edad')
        self.assertEqual(fila[6], '71234567', 'documento del derechohabiente')
        self.assertEqual(fila[12], '2', 'sexo femenino = 2')
        # El T-Registro exige el tipo de documento con 2 dígitos, aunque el
        # catálogo lo guarde en la forma corta que usa el PLAME.
        expected_code = (self.dni.l10n_pe_hr_sunat_code or '').zfill(2)
        self.assertEqual(fila[0], expected_code)
        self.assertEqual(len(fila[0]), 2)
        self.assertEqual(fila[5], expected_code)

        bajas = list(workbook['Bajas'].values)
        self.assertEqual(len(bajas), 2)
        self.assertEqual(bajas[1][9], '07', 'T20: hijo adquiere mayoría de edad')

        self.assertTrue(alta.is_declared, 'el alta queda marcada')
        self.assertFalse(baja.is_declared)

    def test_export_blocks_incomplete_data(self):
        """SUNAT rechaza el alta si el trabajador no tiene documento."""
        undocumented = self.env['hr.employee'].create({
            'name': 'Sin Documento Aún',
            'company_id': self.company.id,
        })
        dependent = self.env['l10n_pe.hr.dependent'].create({
            'employee_id': undocumented.id,
            'type_id': self.type_child.id,
            'last_name': 'Sin', 'names': 'Hija',
            'l10n_latam_identification_type_id': self.dni.id,
            'identification_id': '73333333',
            'birthday': self.today - relativedelta(years=3),
            'date_start': self.today,
        })
        self.assertFalse(undocumented.identification_id)
        with self.assertRaises(UserError):
            dependent.action_export_tregistro_xlsx()

    def test_export_rejects_two_companies(self):
        other = self.env['res.company'].create({'name': 'Otra S.A.C.'})
        other_employee = self.env['hr.employee'].create({
            'name': 'Otro Trabajador', 'company_id': other.id,
            'identification_id': '11111111',
            'l10n_latam_identification_type_id': self.dni.id,
        })
        first = self._dependent(years=5)
        second = self.env['l10n_pe.hr.dependent'].create({
            'employee_id': other_employee.id,
            'type_id': self.type_child.id,
            'names': 'Hijo', 'last_name': 'Otro',
            'l10n_latam_identification_type_id': self.dni.id,
            'identification_id': '72222222',
            'birthday': self.today - relativedelta(years=4),
            'date_start': self.today,
        })
        with self.assertRaises(UserError):
            (first | second).action_export_tregistro_xlsx()
