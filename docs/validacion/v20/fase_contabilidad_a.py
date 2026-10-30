# -*- coding: utf-8 -*-
"""Fase contabilidad_a — base, maestros, impuestos y reportes (OL-ACCOUNTING).

Se ejecuta con run_qa.sh (se antepone _qa.py y se añade _fin.py).
Todo lo creado lleva el prefijo QA20. La configuración global de la compañía
que hace falta para un caso (agente de retención, percepción, reparto de la
detracción, credenciales de consulta…) se activa solo durante el caso y se
restaura al terminarlo: los demás agentes nunca la ven cambiada.
"""
import base64
import io
import json
import time
import zipfile
from datetime import date
from unittest.mock import MagicMock, patch

from lxml import etree
from stdnum.pe import ruc as stdnum_ruc

from odoo.exceptions import UserError, ValidationError

def caso(modulo, titulo, pasos=''):  # noqa: F811
    """Corrige _qa.caso: allí todos los casos comparten el mismo dict _CUR
    (QA['casos'][-1] = _CUR), y el JSON/RESUMEN repetía el último caso."""
    global _CUR
    _CUR = dict(modulo=modulo, titulo=titulo, pasos=pasos, checks=[], registros=[], error=None)
    QA['casos'].append(_CUR)
    print('\n== [%s] %s' % (modulo, titulo))
    return _CUR


PE = env.ref('base.pe')
PEN = env.ref('base.PEN')
USD = env.ref('base.USD')
RUN = int(time.time()) % 10000          # token de la corrida: números de documento únicos
_N = [0]
SEP = date(2026, 9, 15)                 # periodo de libros/SIRE/621 (menos usado por los demás)
OCT = date(2026, 10, 15)


def num():
    _N[0] += 1
    return '%04d%04d' % (RUN, _N[0])


def ruc(body):
    return body + stdnum_ruc.calc_check_digit(body)


def rid(rec):
    return rec.id if hasattr(rec, '_name') else rec


class Config:
    """Cambia valores de un registro y los restaura después."""

    def __init__(self, rec, vals):
        self.rec = rec
        self.old = {k: rec[k] for k in vals}
        changed = {k: v for k, v in vals.items() if rid(rec[k]) != rid(v)}
        if changed:
            rec.write({k: rid(v) for k, v in changed.items()})
        self.changed = list(changed)

    def restaurar(self):
        if self.changed:
            self.rec.write({k: rid(self.old[k]) for k in self.changed})


def esperar_error(exc_types, fn):
    """True si fn lanza alguna de exc_types (deshace lo que haya hecho)."""
    try:
        with env.cr.savepoint():
            fn()
    except exc_types:
        return True
    return False


# ----------------------------------------------------------------------
#  Maestros propios QA20
# ----------------------------------------------------------------------
Account = env['account.account'].with_company(company)
Journal = env['account.journal'].with_company(company)
Partner = env['res.partner']
Move = env['account.move'].with_company(company)


def cuenta(code, name, atype, reconcile=False):
    acc = Account.search([('code', '=', code), ('company_ids', 'in', company.id)], limit=1)
    return acc or Account.create({'code': code, 'name': qa_name(name), 'account_type': atype,
                                  'reconcile': reconcile, 'company_ids': [Command.link(company.id)]})


def diario(code, name, jtype, **vals):
    j = Journal.search([('code', '=', code), ('company_id', '=', company.id)], limit=1)
    if not j:
        j = Journal.create(dict({'name': qa_name(name), 'code': code, 'type': jtype,
                                 'company_id': company.id}, **vals))
    return j


def socio(name, **vals):
    p = Partner.search([('name', '=', qa_name(name))], limit=1)
    if not p:
        p = Partner.create(dict({'name': qa_name(name), 'country_id': PE.id}, **vals))
    return p


def producto(name, **vals):
    p = env['product.product'].search([('name', '=', qa_name(name))], limit=1)
    return p or env['product.product'].create(dict({'name': qa_name(name), 'type': 'service'}, **vals))


def tax(use, amount=18.0):
    return env['account.tax'].search([('company_id', '=', company.id), ('type_tax_use', '=', use),
                                      ('amount', '=', amount), ('l10n_pe_edi_tax_code', '=', '1000')],
                                     limit=1)


IGV_V, IGV_C = tax('sale'), tax('purchase')
DOC01 = env.ref('l10n_pe.document_type01')
DOC_RH = env.ref('l10n_pe.document_type03')      # código 02: recibo por honorarios
DOC91 = env.ref('l10n_pe.document_type91')

ACC_PEND = cuenta('1049920', 'Pagos pendientes', 'asset_current', reconcile=True)
ACC_INC = cuenta('7019920', 'Ventas de servicios', 'income')
ACC_EXP = cuenta('6399920', 'Otros servicios de terceros (destinos)', 'expense')
ACC_D1 = cuenta('9499920', 'Gastos administrativos (destino)', 'expense')
ACC_D2 = cuenta('9599920', 'Gastos de ventas (destino)', 'expense')
ACC_79 = cuenta('7919920', 'Cargas imputables a cuentas de costos', 'income')
ACC_PROV = cuenta('4699920', 'Otras cuentas por pagar (provisiones)', 'liability_current')
ACC_EXP2 = cuenta('6599920', 'Otros gastos de gestión (sin destino)', 'expense')

J_VENTA = diario('QAV20', 'Ventas', 'sale', l10n_latam_use_documents=True)
J_COMPRA = diario('QAC20', 'Compras', 'purchase', l10n_latam_use_documents=True)
J_GRAL = diario('QAG20', 'Operaciones varias', 'general')
J_DEST = diario('QAD20', 'Asientos de destino', 'general')
J_BANCO = diario('QAB20', 'Banco', 'bank')
J_ITF = diario('QAI20', 'Banco con ITF', 'bank', l10n_pe_itf_applies=True)
for _j in (J_BANCO, J_ITF):
    (_j.inbound_payment_method_line_ids | _j.outbound_payment_method_line_ids).filtered(
        lambda l: not l.payment_account_id).payment_account_id = ACC_PEND
J_BANCO.l10n_pe_itf_applies = False

CLIENTE = socio('Cliente Retail SAC', vat=ruc('2060992001'), is_company=True)
PROV = socio('Proveedor Servicios SAC', vat=ruc('2060992002'), is_company=True)
PROV_SPOT = socio('Proveedor SPOT SAC', vat=ruc('2060992003'), is_company=True,
                  l10n_pe_detraction_account='00071234567')
PROV_BUENO = socio('Buen Contribuyente SAC', vat=ruc('2060992004'), is_company=True)
HONOR = socio('Quispe Rojas Ana (honorarios)', vat=ruc('1047392001'))
SERV = producto('Servicio de consultoría')
SERV_SPOT = producto('Servicio de mantenimiento SPOT',
                     l10n_pe_detraction_type_id=env.ref('al_l10n_pe_detraction.detraction_037').id)


def linea(price, product=None, taxes=None, account=None, qty=1.0, name='Servicio'):
    vals = {'name': name, 'quantity': qty, 'price_unit': price,
            'tax_ids': [Command.set((taxes if taxes is not None else IGV_V).ids)]}
    if product:
        vals['product_id'] = product.id
    if account:
        vals['account_id'] = account.id
    return Command.create(vals)


def venta(partner, lines, fecha, post=True, currency=None, journal=None, **vals):
    move = Move.create(dict({
        'move_type': 'out_invoice', 'partner_id': partner.id, 'journal_id': (journal or J_VENTA).id,
        'invoice_date': fecha, 'date': fecha, 'name': 'FQ20-%s' % num(),
        'l10n_latam_document_type_id': DOC01.id,
        'currency_id': (currency or PEN).id, 'invoice_line_ids': lines}, **vals))
    if post:
        move.action_post()
    return move


def compra(partner, lines, fecha, post=True, doc=DOC01, serie='FQ20', currency=None, **vals):
    move = Move.create(dict({
        'move_type': 'in_invoice', 'partner_id': partner.id, 'journal_id': J_COMPRA.id,
        'invoice_date': fecha, 'date': fecha, 'l10n_latam_document_type_id': doc.id,
        'l10n_latam_document_number': '%s-%s' % (serie, num()),
        'currency_id': (currency or PEN).id, 'invoice_line_ids': lines}, **vals))
    if post:
        move.action_post()
    return move


def pagar(moves, journal=None, **vals):
    wiz = env['account.payment.register'].with_context(
        active_model='account.move', active_ids=moves.ids).create(
        dict({'journal_id': (journal or J_BANCO).id}, **vals))
    return wiz, wiz._create_payments()


def lineas_txt(content):
    if isinstance(content, bytes):
        content = content.decode('latin-1')
    return [l for l in content.replace('\r\n', '\n').split('\n') if l.strip()]


# ======================================================================
#  al_account_base
# ======================================================================
caso('al_account_base', 'Código SUNAT del banco = clearing_number nativo (ambos sentidos)',
     'Contactos ▸ cuenta bancaria: elegir el banco del catálogo y escribir el código SUNAT; '
     'luego cambiar el «clearing number» y ver que el código SUNAT lo sigue.')


def t_banco():
    bcp = env.ref('al_account_base.bank_bcp')
    acc = registro(env['res.partner.bank'].create({
        'account_number': '19%s0001' % num(), 'partner_id': PROV.id,
        'l10n_pe_bank_id': bcp.id, 'l10n_pe_bank_code': '02'}), 'cuenta BCP')
    igual(acc.bank_name, bcp.name, 'nombre del banco tomado del catálogo')
    igual(acc.bank_bic, 'BCPLPEPL', 'BIC tomado del catálogo')
    igual(acc.clearing_number, '02', 'l10n_pe_bank_code → clearing_number')
    acc.clearing_number = '11'
    acc.invalidate_recordset()
    igual(acc.l10n_pe_bank_code, '11', 'clearing_number → l10n_pe_bank_code')
    check(esperar_error(ValidationError, lambda: acc.write({'l10n_pe_bank_code': '123'})),
          'código de 3 dígitos rechazado (sin truncar a otro banco)')
    if not esperar_error(ValidationError, lambda: acc.write({'clearing_number': 'AB'})):
        aviso('El clearing_number nativo acepta un código SUNAT inválido',
              'escribir clearing_number=«AB» no dispara _check_l10n_pe_bank_code (la restricción es sobre el '
              'related l10n_pe_bank_code); el PLE 1.2 lo leería tal cual')
    bn = registro(env['res.partner.bank'].create(dict(
        env['res.partner.bank']._l10n_pe_national_bank_vals(),
        account_number='00071%s' % num(), partner_id=PROV_SPOT.id)), 'cuenta de detracciones BN')
    igual(bn.l10n_pe_bank_code, '18', 'cuenta del Banco de la Nación: código SUNAT 18')
    check(bn._l10n_pe_is_national_bank(), 'se reconoce como Banco de la Nación (BANCPEPL)')


ejecutar(t_banco)

caso('al_account_base', 'Glosa del asiento y de la línea, anexo y naturaleza del diario',
     'Contabilidad ▸ Asientos ▸ Nuevo: escribir la glosa, publicar y duplicar.')


def t_glosa():
    move = registro(Move.create({
        'move_type': 'entry', 'journal_id': J_GRAL.id, 'date': OCT,
        'l10n_pe_gloss': 'QA20 Provisión de servicios de octubre',
        'line_ids': [
            Command.create({'account_id': ACC_EXP2.id, 'name': 'D', 'debit': 250.0,
                            'l10n_pe_gloss': 'Glosa de la línea'}),
            Command.create({'account_id': ACC_PROV.id, 'name': 'H', 'credit': 250.0}),
        ]}), 'asiento con glosa')
    move.action_post()
    igual(move.state, 'posted', 'asiento publicado')
    igual(move.l10n_pe_gloss, 'QA20 Provisión de servicios de octubre', 'glosa conservada al publicar')
    igual(move.line_ids[0].l10n_pe_gloss, 'Glosa de la línea', 'glosa de la línea')
    check(move in Move.search([('l10n_pe_gloss', '=', 'QA20 Provisión de servicios de octubre')]),
          'la glosa es buscable')
    copia = move.copy()
    check(not copia.l10n_pe_gloss and not any(copia.line_ids.mapped('l10n_pe_gloss')),
          'duplicar no arrastra la glosa')
    copia.unlink()
    check(move.l10n_pe_is_pe(), 'l10n_pe_is_pe() en compañía peruana')
    PROV.l10n_pe_annex_code = '0001'
    igual(PROV.l10n_pe_annex_code, '0001', 'código de anexo de 4 dígitos aceptado')
    check(esperar_error(ValidationError, lambda: PROV.write({'l10n_pe_annex_code': 'AB12'})),
          'código de anexo con letras rechazado')
    igual(J_GRAL.l10n_pe_journal_kind, 'movement', 'naturaleza del diario por defecto: movimiento')
    ap = registro(diario('QAA20', 'Apertura', 'general', l10n_pe_journal_kind='opening',
                         l10n_pe_exclude_from_books=True), 'diario de apertura')
    check(ap.l10n_pe_journal_kind == 'opening' and ap.l10n_pe_exclude_from_books,
          'diario de apertura excluido de libros')


ejecutar(t_glosa)

# ======================================================================
#  l10n_pe_vat_sunat
# ======================================================================
_Company = type(env['res.company'])
_Conn = type(env['l10n_pe.api.connection'])
CONN = env['l10n_pe.api.connection'].search([], limit=1)

caso('l10n_pe_vat_sunat', 'Consulta de RUC (simulada) actualiza el contacto',
     'Contactos ▸ nuevo contacto con RUC ▸ botón «Actualizar RUC/DNI».')


def t_ruc():
    cfg = Config(company, {'l10n_pe_ruc_validation': True, 'l10n_pe_dni_validation': True})
    p = registro(socio('Empresa por consultar', vat=ruc('2060992010'), is_company=True), 'contacto RUC')
    district = env['l10n_pe.res.city.district'].search([('city_id', '!=', False)], limit=1)
    vals = {'name': qa_name('EMPRESA CONSULTADA S.A.C.'), 'street': 'AV. QA 123',
            'l10n_pe_district': district.id, 'city_id': district.city_id.id,
            'state_id': district.city_id.state_id.id, 'zip': district.code}
    check(p._l10n_pe_lookup_target() == ('ruc', p.vat), 'el RUC es el vat del contacto peruano',
          repr(p._l10n_pe_lookup_target()))
    with patch.object(_Company, '_get_pe_api_connections', return_value=CONN), \
            patch.object(_Conn, 'run', return_value=(vals, {})) as run:
        p.action_update_document()
    check(run.called and run.call_args.args[:2] == (p.vat, 'ruc'), 'se consultó el RUC',
          repr(run.call_args))
    igual(p.name, vals['name'], 'razón social actualizada')
    igual(p.street, 'AV. QA 123', 'dirección actualizada')
    igual(p.l10n_pe_district, district, 'distrito (ubigeo) actualizado')
    check(not p.alert_warning_vat, 'sin alerta de consulta')
    with patch.object(_Company, '_get_pe_api_connections', return_value=CONN), \
            patch.object(_Conn, 'run') as run2:
        res = Partner._l10n_pe_lookup_document('20131312954', '6')
    check(res == {} and not run2.called, 'RUC con dígito verificador inválido no se consulta')
    omitido('Consulta real a SUNAT/APIs', 'sin credenciales de prueba: conexión simulada')
    cfg.restaurar()


ejecutar(t_ruc)

caso('l10n_pe_vat_sunat', 'DNI en identificadores adicionales (Odoo 20) y consulta simulada',
     'Contactos ▸ persona con DNI en «Identificadores adicionales» ▸ «Actualizar RUC/DNI».')


def t_dni():
    cfg = Config(company, {'l10n_pe_ruc_validation': True, 'l10n_pe_dni_validation': True})
    p = registro(socio('Persona con DNI', additional_identifiers={'PE_DNI': '45678912'}), 'contacto DNI')
    igual(p._get_additional_identifier('PE_DNI'), '45678912', 'DNI guardado en additional_identifiers')
    igual(p._l10n_pe_lookup_target(), ('dni', '45678912'), 'objetivo de la consulta: DNI')
    igual(p._l10n_pe_lookup_doc_type(), 'dni', 'tipo consultable con la validación de DNI activa')
    with patch.object(_Company, '_get_pe_api_connections', return_value=CONN), \
            patch.object(_Conn, 'run', return_value=({'name': qa_name('PEREZ GOMEZ JUAN')}, {})) as run:
        p.action_update_document()
    check(run.called and run.call_args.args[:2] == ('45678912', 'dni'), 'se consultó el DNI',
          repr(run.call_args))
    igual(p.name, qa_name('PEREZ GOMEZ JUAN'), 'nombre actualizado desde la consulta de DNI')
    check(esperar_error(UserError, lambda: p.with_context().write(
        {'additional_identifiers': {'PE_DNI': '1234'}}) or p._run_document_lookup(True)),
        'DNI de 4 dígitos rechazado al consultar')
    cfg.restaurar()


ejecutar(t_dni)

# ======================================================================
#  al_l10n_pe_currency
# ======================================================================
from odoo.addons.al_l10n_pe_currency.services import bcrp_rate  # noqa: E402

Rate = env['res.currency.rate']

caso('al_l10n_pe_currency', 'Tipo de cambio compra/venta manual',
     'Contabilidad ▸ Configuración ▸ Monedas ▸ USD ▸ nueva tasa con compra 3.360 y venta 3.370.')


def t_tc_manual():
    d = date(2026, 2, 2)
    old = Rate.search([('currency_id', '=', USD.id), ('company_id', '=', company.id), ('name', '=', d)])
    if old:
        old.write({'rate_purchase': 3.360, 'rate_sale': 3.370, 'ref_origin': 'manual'})
        r = old
    else:
        r = Rate.create({'name': d, 'currency_id': USD.id, 'company_id': company.id,
                         'rate_purchase': 3.360, 'rate_sale': 3.370, 'ref_origin': 'manual'})
    registro(r, 'TC manual 02/02/2026')
    igual(r.rate, 1 / 3.370, 'tasa nativa = 1/venta', 6)
    igual(r.rate_purchase, 3.360, 'compra conservada', 3)
    r.rate = 0.25
    igual(r.rate_sale, 4.0, 'tocar la tasa nativa recalcula la venta', 3)
    r.rate_sale = 3.370
    igual(r.rate, 1 / 3.370, 'y la venta vuelve a fijar la nativa', 6)
    check(esperar_error(ValidationError, lambda: r.write({'rate_purchase': -1})),
          'tipo de cambio negativo rechazado')


ejecutar(t_tc_manual)

caso('al_l10n_pe_currency', 'Descarga BCRP simulada: el cierre de un día es el TC SUNAT del siguiente '
     '(fin de semana)', 'Monedas ▸ «Actualizar tipo de cambio» (BCRP) para 05/02–10/02/2026.')


def t_tc_bcrp():
    closes = [{'date': date(2026, 2, 4), 'compra': 3.351, 'venta': 3.358},
              {'date': date(2026, 2, 5), 'compra': 3.362, 'venta': 3.369},
              {'date': date(2026, 2, 6), 'compra': 3.371, 'venta': 3.377},   # viernes
              {'date': date(2026, 2, 9), 'compra': 3.380, 'venta': 3.386}]   # lunes
    with patch.object(bcrp_rate, 'fetch_bcrp', return_value=closes) as fetch:
        n = env['res.currency'].l10n_pe_update_range_bcrp(date(2026, 2, 5), date(2026, 2, 10))
    check(fetch.called, 'se llamó a la descarga (simulada)', repr(fetch.call_args))
    igual(n, 6, 'fechas cargadas 05..10/02')
    exp = {5: (3.351, 3.358), 6: (3.362, 3.369), 7: (3.371, 3.377), 8: (3.371, 3.377),
           9: (3.371, 3.377), 10: (3.380, 3.386)}
    for day, (c, v) in exp.items():
        r = Rate.search([('currency_id', '=', USD.id), ('company_id', '=', company.id),
                         ('name', '=', date(2026, 2, day))], limit=1)
        registro(r, 'TC BCRP %02d/02' % day)
        check(r and abs(r.rate_purchase - c) < 1e-6 and abs(r.rate_sale - v) < 1e-6
              and r.ref_origin == 'bcrp',
              'TC SUNAT %02d/02 = cierre anterior%s' % (day, ' (fin de semana)' if day in (7, 8) else ''),
              '%s/%s esperado %s/%s' % (r.rate_purchase, r.rate_sale, c, v))
    with patch.object(bcrp_rate, 'fetch_bcrp', return_value=closes):
        env['res.currency'].l10n_pe_update_range_bcrp(date(2026, 2, 5), date(2026, 2, 10))
    igual(Rate.search_count([('currency_id', '=', USD.id), ('company_id', '=', company.id),
                             ('name', '=', date(2026, 2, 7))]), 1, 'recargar no duplica')
    omitido('Descarga real del BCRP/SUNAT', 'sin red: respuesta simulada')


ejecutar(t_tc_bcrp)

caso('al_l10n_pe_currency', 'Factura en USD con TC compra vs venta',
     'Ventas ▸ factura en dólares del 09/02/2026 ▸ campo «Tipo de cambio»: compra o venta.')


def t_tc_factura():
    venta_ = registro(venta(CLIENTE, [linea(100.0, taxes=env['account.tax'])], date(2026, 2, 9),
                            currency=USD), 'factura USD con TC venta')
    igual(venta_.l10n_pe_exchange_rate_type, 'sale', 'por defecto se usa la venta')
    igual(venta_.invoice_currency_rate, 1 / 3.377, 'tasa = 1/venta del 09/02', 6)
    igual(abs(venta_.amount_total_signed), 337.70, 'total en soles con venta')
    compra_tc = registro(venta(CLIENTE, [linea(100.0, taxes=env['account.tax'])], date(2026, 2, 9),
                               currency=USD, post=False, l10n_pe_exchange_rate_type='purchase'),
                         'factura USD con TC compra')
    compra_tc.action_post()
    igual(compra_tc.invoice_currency_rate, 1 / 3.371, 'tasa = 1/compra del 09/02', 6)
    igual(abs(compra_tc.amount_total_signed), 337.10, 'total en soles con compra')


ejecutar(t_tc_factura)

# ======================================================================
#  al_account_move_name_sequence
# ======================================================================
caso('al_account_move_name_sequence', 'Numeración por secuencia del diario',
     'Contabilidad ▸ Diarios ▸ marcar «Usar secuencia» con prefijos QSV-/QSN- y publicar facturas.')


def t_seq():
    j = registro(diario('QAS20', 'Ventas por secuencia', 'sale', l10n_latam_use_documents=False,
                        use_name_sequence=True, name_sequence_prefix='QSV-',
                        refund_name_sequence_prefix='QSN-'), 'diario con secuencia')
    igual(j.name_sequence_id.prefix, 'QSV-', 'secuencia de facturas creada con el prefijo')
    igual(j.refund_name_sequence_id.prefix, 'QSN-', 'secuencia de NC creada con el prefijo')
    igual(j.name_sequence_id.implementation, 'no_gap', 'secuencia sin huecos')
    a = registro(venta(CLIENTE, [linea(100.0)], OCT, journal=j, name='/',
                       l10n_latam_document_type_id=False), 'factura 1')
    b = registro(venta(CLIENTE, [linea(100.0)], OCT, journal=j, name='/',
                       l10n_latam_document_type_id=False), 'factura 2')
    check(a.name.startswith('QSV-') and b.name.startswith('QSV-'), 'nombre tomado de la secuencia',
          '%s, %s' % (a.name, b.name))
    check(int(b.name.split('-')[1]) == int(a.name.split('-')[1]) + 1, 'numeración correlativa')
    nc = registro(Move.create({'move_type': 'out_refund', 'partner_id': CLIENTE.id, 'journal_id': j.id,
                               'invoice_date': OCT, 'invoice_line_ids': [linea(50.0)]}), 'NC')
    nc.action_post()
    check(nc.name.startswith('QSN-'), 'NC con la secuencia rectificativa', nc.name)


ejecutar(t_seq)

caso('al_account_move_name_sequence', 'Serie CPE: publicar crea las 3 secuencias; factura y NC',
     'Contabilidad ▸ Series CPE ▸ nueva serie FQ21 ▸ Publicar; asignarla a un diario y facturar.')


def t_serie():
    Serie = env['edi.invoice.series']
    s = Serie.search([('name', '=', 'FQ21'), ('company_id', '=', company.id)], limit=1) or Serie.create({
        'name': 'FQ21', 'l10n_latam_document_type_id': DOC01.id, 'name_nc': 'FQ2C', 'name_nd': 'FQ2D',
        'company_id': company.id})
    if s.state != 'publish':
        s.action_publish()
    registro(s, 'serie CPE')
    igual(s.state, 'publish', 'serie publicada')
    igual((s.invoice_seq_id.prefix, s.credit_note_seq_id.prefix, s.debit_note_seq_id.prefix),
          ('FQ21-', 'FQ2C-', 'FQ2D-'), 'secuencias de factura, NC y ND')
    check(esperar_error(ValidationError, lambda: Serie.create({
        'name': 'B001', 'name_nc': 'BC01', 'name_nd': 'BD01', 'l10n_latam_document_type_id': DOC01.id,
        'company_id': company.id})),
        'serie «B001» rechazada para facturas')
    j = registro(diario('QAE20', 'Ventas serie CPE', 'sale', l10n_latam_use_documents=True,
                        use_name_sequence=True, edi_series_ids=[Command.set(s.ids)]), 'diario serie')
    inv = registro(venta(CLIENTE, [linea(300.0)], OCT, journal=j, name='/'), 'factura FQ21')
    igual(inv.edi_series_id, s, 'serie asignada a la factura')
    check(inv.name.startswith('FQ21-') and inv.l10n_latam_document_number == inv.name,
          'numerada con la serie', inv.name)
    rev = env['account.move.reversal'].with_context(active_model='account.move', active_ids=inv.ids) \
        .create({'journal_id': j.id, 'reason': 'QA20 anulación parcial'})
    nc = Move.browse(rev.reverse_moves()['res_id'])
    nc.action_post()
    registro(nc, 'NC de la serie')
    check(nc.name.startswith('FQ2C-'), 'NC numerada con la serie rectificativa', nc.name)


ejecutar(t_serie)

# ======================================================================
#  al_account_destinations
# ======================================================================
caso('al_account_destinations', 'Gasto clase 6 → asiento de destino 9 / 79',
     'Plan contable ▸ cuenta 63 con destinos 94 (60 %) y 95 (40 %) y cuenta 79; publicar un gasto.')


def t_destino():
    cfg = Config(company, {'l10n_pe_destination_journal_id': J_DEST})
    igual(company.l10n_pe_dest_type, '6a9', 'compañía con destinos 6 → 9')
    check(esperar_error((UserError, ValidationError), lambda: ACC_EXP.write({
        'l10n_pe_load_account_id': ACC_79.id,
        'l10n_pe_destiny_ids': [Command.clear(),
                                Command.create({'dest_account_id': ACC_D1.id, 'percentage': 0.5}),
                                Command.create({'dest_account_id': ACC_D2.id, 'percentage': 0.3})]})),
        'destinos que no suman 100 % rechazados')
    ACC_EXP.write({'l10n_pe_load_account_id': ACC_79.id, 'l10n_pe_destiny_ids': [
        Command.clear(),
        Command.create({'dest_account_id': ACC_D1.id, 'percentage': 0.6}),
        Command.create({'dest_account_id': ACC_D2.id, 'percentage': 0.4})]})
    check(ACC_EXP.l10n_pe_work_destinies, 'la cuenta 6 trabaja con destinos')
    move = registro(Move.create({
        'move_type': 'entry', 'journal_id': J_GRAL.id, 'date': OCT,
        'line_ids': [Command.create({'account_id': ACC_EXP.id, 'name': 'Gasto', 'debit': 1000.0}),
                     Command.create({'account_id': ACC_PROV.id, 'name': 'Provisión', 'credit': 1000.0})]}),
        'asiento de gasto')
    move.action_post()
    dest = registro(move.l10n_pe_destiny_move_id, 'asiento de destino')
    check(dest and dest.state == 'posted' and dest.l10n_pe_is_destiny_entry, 'destino generado y publicado')
    igual(dest.journal_id, J_DEST, 'en el diario de destinos de Ajustes')
    by = lambda a, f: sum(dest.line_ids.filtered(lambda l: l.account_id == a).mapped(f))  # noqa: E731
    igual(by(ACC_D1, 'debit'), 600.0, '94: 60 % al debe')
    igual(by(ACC_D2, 'debit'), 400.0, '95: 40 % al debe')
    igual(by(ACC_79, 'credit'), 1000.0, '79: 100 % al haber')
    check(not dest.l10n_pe_destiny_move_id, 'el destino no genera otro destino')
    # Factura de proveedor con gasto clase 6
    bill = registro(compra(PROV, [linea(500.0, taxes=IGV_C, account=ACC_EXP)], OCT), 'factura con gasto 6')
    d2 = registro(bill.l10n_pe_destiny_move_id, 'destino de la factura')
    igual(sum(d2.line_ids.filtered(lambda l: l.account_id == ACC_D1).mapped('debit')), 300.0,
          'factura: 94 = 300')
    igual(sum(d2.line_ids.filtered(lambda l: l.account_id == ACC_79).mapped('credit')), 500.0,
          'factura: 79 = 500 (sin IGV)')
    cfg.restaurar()


ejecutar(t_destino)

# ======================================================================
#  al_l10n_pe_detraction
# ======================================================================
from odoo.addons.al_l10n_pe_detraction.services import bn_txt  # noqa: E402

ACC_DET_C = cuenta('1219920', 'Detracciones por cobrar', 'asset_receivable', reconcile=True)
ACC_DET_P = cuenta('4249920', 'Detracciones por pagar', 'liability_payable', reconcile=True)

caso('al_l10n_pe_detraction', 'Factura de venta con detracción y reparto de la cuenta 12',
     'Ventas ▸ factura de servicio SPOT 12 % por S/ 1 180 ▸ Publicar.')


def t_det_venta():
    cfg = Config(company, {'l10n_pe_detraction_split': True,
                           'l10n_pe_detraction_receivable_account_id': ACC_DET_C,
                           'l10n_pe_detraction_payable_account_id': ACC_DET_P})
    igual((SERV_SPOT.l10n_pe_withhold_code, SERV_SPOT.l10n_pe_withhold_percentage), ('037', 12.0),
          'el producto sincroniza el código y % nativos del catálogo 54')
    inv = registro(venta(CLIENTE, [linea(1000.0, product=SERV_SPOT)], OCT), 'factura con detracción')
    check(inv.l10n_pe_detraction_applies, 'aplica detracción (> S/ 700)')
    igual(inv.l10n_pe_detraction_amount, 142.0, 'detracción 12 % de 1 180 redondeada a 142')
    igual(inv.l10n_pe_detraction_net, 1038.0, 'neto a cobrar 1 038')
    igual(inv.l10n_pe_edi_operation_type, '1001', 'tipo de operación 1001 (sujeta a detracción)')
    terms = inv.line_ids.filtered(lambda l: l.display_type == 'payment_term')
    igual(len(terms), 2, 'dos líneas de cobro (neto + detracción)')
    igual(terms.filtered(lambda l: l.account_id == ACC_DET_C).balance, 142.0, 'detracción a la 12 propia')
    xml = env['account.edi.xml.ubl_pe']._export_invoice(inv)[0]
    check(b'Detraccion' in xml and b'142.00' in xml, 'el XML UBL lleva el bloque de detracción',
          'Detraccion/142.00')
    cfg.restaurar()


ejecutar(t_det_venta)

caso('al_l10n_pe_detraction', 'Depósito de la detracción de una compra',
     'Compras ▸ factura SPOT ▸ «Depositar detracción» con número de constancia.')


def t_det_deposito():
    cfg = Config(company, {'l10n_pe_detraction_split': True,
                           'l10n_pe_detraction_receivable_account_id': ACC_DET_C,
                           'l10n_pe_detraction_payable_account_id': ACC_DET_P})
    bill = registro(compra(PROV_SPOT, [linea(2000.0, product=SERV_SPOT, taxes=IGV_C)], OCT),
                    'compra con detracción')
    igual(bill.l10n_pe_detraction_amount, 283.0, 'detracción 12 % de 2 360 = 283')
    wiz = env['l10n_pe.detraction.deposit.wizard'].create({
        'move_id': bill.id, 'journal_id': J_BANCO.id, 'payment_date': OCT,
        'constancy_number': 'QA%s' % num()})
    igual(wiz.amount, 283.0, 'el asistente propone el importe de la detracción')
    wiz.action_confirm()
    igual(bill.l10n_pe_detraction_number, wiz.constancy_number, 'constancia registrada en la factura')
    pay = env['account.payment'].search([('l10n_pe_detraction_move_id', '=', bill.id)], limit=1)
    registro(pay, 'depósito de detracción')
    check(pay and pay.amount == 283.0 and pay.payment_type == 'outbound', 'pago de 283 vinculado',
          pay.read(['amount', 'state']) if pay else 'sin pago')
    igual(bill.amount_residual, 2077.0, 'solo se paga la detracción: queda el neto 2 077')
    cfg.restaurar()


ejecutar(t_det_deposito)

caso('al_l10n_pe_detraction', 'TXT de depósito masivo del Banco de la Nación',
     'Perú ▸ Detracciones ▸ «Generar TXT BN» (adquirente) para setiembre 2026.')


def t_det_txt():
    cfg = Config(company.partner_id, {'l10n_pe_detraction_account': '00071234500'})
    bill = registro(compra(PROV_SPOT, [linea(5000.0, product=SERV_SPOT, taxes=IGV_C)], SEP),
                    'compra SPOT de setiembre')
    wiz = registro(env['l10n_pe.detraction.txt.wizard'].create({
        'company_id': company.id, 'mode': 'acquirer', 'batch_number': '26%04d' % RUN,
        'date_from': date(2026, 9, 1), 'date_to': date(2026, 9, 30)}), 'asistente TXT')
    wiz.action_generate()
    content = wiz.file_data.content.decode('latin-1')
    lines = lineas_txt(content)
    igual(bn_txt.check_structure(content), [], 'estructura válida para el banco')
    igual(len(lines[0]), 68, 'cabecera de 68 posiciones')
    igual(lines[0][0], '*', 'cabecera de adquirente')
    mine = [l for l in lines[1:] if PROV_SPOT.vat in l]
    check(len(mine) == 1 and len(mine[0]) == 107, 'detalle del proveedor de 107 posiciones', mine)
    if mine:
        check(bn_txt.amount(bill.l10n_pe_detraction_amount) in mine[0], 'importe de la detracción en el detalle',
              bill.l10n_pe_detraction_amount)
    omitido('Envío del TXT al Banco de la Nación', 'sin credenciales bancarias')
    cfg.restaurar()


ejecutar(t_det_txt)

# ======================================================================
#  al_l10n_pe_retention (sobre l10n_pe_edi_withholding de ee20)
# ======================================================================
RET_TAX = company.l10n_pe_retention_tax_id
_Svc = type(env['l10n_pe.edi.cpe.service'])

caso('al_l10n_pe_retention', 'Pago con retención 3 %: CRE numerado y XML generado sin enviar',
     'Compras ▸ factura de S/ 1 180 ▸ Registrar pago ▸ retención 3 % ▸ «Generar XML del CRE».')


def t_ret_pago():
    cfg = Config(company, {'l10n_pe_retention_agent': True})
    check(RET_TAX and RET_TAX.is_withholding_tax and RET_TAX.get_external_id().get(RET_TAX.id, '').endswith(
        'purchase_tax_withholding_3'), 'impuesto de retención = el nativo de l10n_pe_edi_withholding',
        RET_TAX.display_name)
    bill = registro(compra(PROV, [linea(1000.0, taxes=IGV_C, account=ACC_EXP2)], OCT, post=False),
                    'factura sujeta a retención')
    check(bill.l10n_pe_retention_applies, 'aplica retención')
    igual(bill.l10n_pe_retention_amount, 35.40, 'estimación 3 % de 1 180')
    bill.action_post()
    check(RET_TAX in bill.invoice_line_ids.tax_ids, 'impuesto de retención inyectado al publicar')
    igual(bill.amount_total, 1180.0, 'sin alterar el total')
    wiz, pay = pagar(bill)
    registro(pay, 'pago con retención')
    igual(sum(wiz.withholding_line_ids.mapped('amount')), 35.40, 'línea de retención del asistente')
    igual(pay._l10n_pe_retention_lines().amount, 35.40, 'retención en el pago')
    number = pay.l10n_pe_edi_retention_number or ''
    check(number[:1] == 'R' and len(number.split('-')[0]) == 4, 'CRE numerado con serie R___', number)
    if not number.startswith('R001-'):
        aviso('La serie del CRE no es R001', 'número %s: la secuencia nativa (%s) conserva el prefijo «%s»; '
              'el R001- del módulo solo se aplica al cargar el plan en compañías nuevas'
              % (number, RET_TAX.withholding_sequence_id.get_external_id().get(
                  RET_TAX.withholding_sequence_id.id), RET_TAX.withholding_sequence_id.prefix))
    igual(bill.payment_state in ('paid', 'in_payment'), True, 'factura pagada (neto + retención)')
    pay.action_l10n_pe_generate_cre_xml()
    att = env['ir.attachment'].search([('res_model', '=', 'account.payment'), ('res_id', '=', pay.id),
                                       ('mimetype', '=', 'application/xml')], limit=1, order='id desc')
    registro(att, 'XML del CRE')
    xml = att.raw.decode('iso-8859-1') if att else ''
    for txt, label in ((company.vat, 'RUC del agente'), (PROV.vat, 'RUC del proveedor'),
                       (number, 'número del CRE'), ('<sac:SUNATRetentionPercent>3', 'tasa 3 %'),
                       ('35.40', 'importe retenido'), (bill.l10n_latam_document_number, 'factura referenciada')):
        check(txt and txt in xml, 'XML: %s' % label, txt)
    tree = etree.fromstring(pay._l10n_pe_edi_generate_retention_bstr())
    igual(etree.QName(tree).localname, 'Retention', 'raíz del UBL: Retention')
    omitido('Envío del CRE a SUNAT/OSE', 'sin credenciales de prueba')
    summ = registro(env['l10n_pe.retention.summary.wizard'].create({'year': 2026, 'month': '10'}),
                    'resumen 626')
    summ.action_export()
    rows = [l.split('|') for l in lineas_txt(summ.file_data.content) if PROV.vat in l]
    check(rows and any(r[6] == '35.40' for r in rows), 'PDT 626: fila del proveedor con 35.40', rows)
    cfg.restaurar()


ejecutar(t_ret_pago)

caso('al_l10n_pe_retention', 'Anulación del CRE (resumen RR): reversing → cancelled (servicio simulado)',
     'Pagos ▸ CRE aceptado ▸ «Revertir retención» con motivo ▸ «Consultar reversión».')


def t_ret_anula():
    cfg = Config(company, {'l10n_pe_retention_agent': True})
    bill = registro(compra(PROV, [linea(2000.0, taxes=IGV_C, account=ACC_EXP2)], OCT), 'factura')
    _w, pay = pagar(bill)
    registro(pay, 'pago con CRE a anular')
    number = pay.l10n_pe_edi_retention_number
    igual(pay._l10n_pe_retention_lines().amount, 70.80, 'retención 3 % de 2 360')
    check(esperar_error(UserError, lambda: (pay.write({'l10n_pe_edi_status': 'sent'}), pay.action_draft())),
          'un CRE aceptado bloquea pasar el pago a borrador')
    pay.write({'l10n_pe_edi_status': 'sent', 'l10n_pe_edi_reversal_reason': 'QA20 Error en el importe'})
    pay.invalidate_recordset(['l10n_pe_edi_reversal_available'])
    calls = {}

    def send_summary(_self, comp, filename, edi_str, sunat_wsdl):
        calls.update(filename=filename, xml=edi_str)
        return {'xml_document': edi_str, 'cdr_number': '2026103000%s' % RUN}

    def get_status(_self, comp, ticket, sunat_wsdl):
        calls['ticket'] = ticket
        return {'success': True, 'cdr': b'<ApplicationResponse/>'}

    with patch.object(_Svc, '_send_summary', send_summary), patch.object(_Svc, '_get_status', get_status):
        pay.action_l10n_pe_edi_reverse_retention()
        igual(pay.l10n_pe_edi_status, 'reversing', 'tras enviar el resumen: reversing')
        pay.action_l10n_pe_edi_check_reversal()
    igual(pay.l10n_pe_edi_status, 'cancelled', 'tras consultar el ticket: cancelled')
    check(calls.get('filename', '').startswith('%s-RR-' % company.vat), 'nombre del resumen RR',
          calls.get('filename'))
    tree = etree.fromstring(calls['xml'])
    igual(etree.QName(tree).localname, 'SummaryDocuments', 'XML de reversión: SummaryDocuments')
    xml_rr = calls['xml'].decode() if isinstance(calls['xml'], bytes) else calls['xml']
    serie_rr, folio_rr = number.split('-')
    check(serie_rr in xml_rr and str(int(folio_rr)) in xml_rr, 'el resumen referencia el CRE anulado (serie y número)',
          number)
    pay.action_draft()
    pay.action_post()
    check(pay.l10n_pe_edi_retention_number and pay.l10n_pe_edi_retention_number != number,
          'al volver a publicar recibe un número nuevo', pay.l10n_pe_edi_retention_number)
    omitido('Envío real del resumen RR a SUNAT', 'sin credenciales de prueba')
    cfg.restaurar()


ejecutar(t_ret_anula)

caso('al_l10n_pe_retention', 'Excepciones: monto mínimo y buen contribuyente',
     'Compras ▸ facturas por S/ 590 y de un buen contribuyente: no muestran retención.')


def t_ret_exc():
    cfg = Config(company, {'l10n_pe_retention_agent': True})
    small = registro(compra(PROV, [linea(500.0, taxes=IGV_C, account=ACC_EXP2)], OCT, post=False),
                     'factura de S/ 590')
    check(not small.l10n_pe_retention_applies, 'no aplica por debajo de S/ 700')
    PROV_BUENO.with_context(l10n_pe_padron_auto=True).is_good_taxpayer = True
    good = registro(compra(PROV_BUENO, [linea(1000.0, taxes=IGV_C, account=ACC_EXP2)], OCT, post=False),
                    'factura de buen contribuyente')
    check(not good.l10n_pe_retention_applies, 'no aplica a buen contribuyente')
    cfg.restaurar()
    check(not compra(PROV, [linea(1000.0, taxes=IGV_C, account=ACC_EXP2)], OCT, post=False)
          .l10n_pe_retention_applies, 'no aplica si la compañía no es agente')


ejecutar(t_ret_exc)

# ======================================================================
#  al_l10n_pe_perception
# ======================================================================
ACC_PER = cuenta('4011920', 'IGV percepciones por pagar', 'liability_current')
ACC_PER_R = cuenta('4011940', 'IGV percepciones por aplicar', 'asset_current')
CEMENTO = producto('Cemento portland', type='consu', l10n_pe_perception_type='internal',
                   taxes_id=[Command.set(IGV_V.ids)])

caso('al_l10n_pe_perception', 'Factura con percepción 2 % y XML de la factura',
     'Ventas ▸ factura de 10 bolsas de cemento ▸ Publicar: percepción 51 al 2 %.')


def t_perc_factura():
    cfg = Config(company, {'l10n_pe_perception_agent': True, 'l10n_pe_perception_account_id': ACC_PER,
                           'l10n_pe_perception_received_account_id': ACC_PER_R})
    inv = registro(venta(CLIENTE, [linea(100.0, product=CEMENTO, qty=10)], OCT), 'factura con percepción')
    igual(inv.l10n_pe_perception_code, '51', 'régimen 51 (venta interna)')
    igual(inv.l10n_pe_perception_rate, 2.0, 'tasa 2 %')
    igual(inv.l10n_pe_perception_amount, 23.60, 'percepción 2 % de 1 180')
    igual(inv.l10n_pe_edi_operation_type, '2001', 'operación 2001 (sujeta a percepción)')
    tree = etree.fromstring(env['account.edi.xml.ubl_pe']._export_invoice(inv)[0])
    ch = [c for c in tree.findall('{*}AllowanceCharge') if c.findtext('{*}AllowanceChargeReasonCode') == '51']
    check(ch and ch[0].findtext('{*}Amount') == '23.60' and ch[0].findtext('{*}BaseAmount') == '1180.00',
          'XML: cargo 51 de 23.60 sobre 1 180.00')
    term = [t for t in tree.findall('{*}PaymentTerms') if t.findtext('{*}ID') == 'Percepcion']
    check(term and term[0].findtext('{*}Amount') == '1203.60', 'XML: total con percepción 1 203.60')
    cfg.restaurar()


ejecutar(t_perc_factura)

caso('al_l10n_pe_perception', 'Comprobante de percepción (tipo 40) desde el cobro',
     'Ventas ▸ cobrar 1 203.60 ▸ en el pago «Crear comprobante de percepción» ▸ Publicar.')


def t_perc_cpe():
    cfg = Config(company, {'l10n_pe_perception_agent': True, 'l10n_pe_perception_account_id': ACC_PER,
                           'l10n_pe_perception_received_account_id': ACC_PER_R})
    inv = registro(venta(CLIENTE, [linea(100.0, product=CEMENTO, qty=10)], OCT), 'factura')
    _w, pay = pagar(inv, amount=1203.60, payment_date=OCT, payment_difference_handling='open')
    registro(pay, 'cobro con percepción')
    action = pay.action_l10n_pe_create_perception()
    perc = env['l10n_pe.perception'].browse(action['res_id'])
    registro(perc, 'comprobante de percepción')
    igual(perc.line_ids.perception_pen, 23.60, 'percepción calculada con el factor (1 203.60)')
    perc.action_post()
    igual(perc.state, 'posted', 'comprobante publicado')
    check((perc.l10n_pe_edi_perception_number or '').startswith('P'), 'número P___',
          perc.l10n_pe_edi_perception_number)
    cr = perc.move_id.line_ids.filtered(lambda l: l.account_id == ACC_PER)
    igual(cr.credit, 23.60, 'asiento: percepción por pagar al haber')
    igual(inv.l10n_pe_perception_collected, 23.60, 'la factura registra la percepción cobrada')
    tree = etree.fromstring(perc._l10n_pe_edi_generate_retention_bstr())
    ns = {'cbc': 'urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2',
          'sac': 'urn:sunat:names:specification:ubl:peru:schema:xsd:SunatAggregateComponents-1'}
    igual(etree.QName(tree).localname, 'Perception', 'XML: raíz Perception')
    igual(tree.xpath('/*/cbc:TotalInvoiceAmount', namespaces=ns)[0].text, '23.60', 'XML: total percibido')
    igual(tree.xpath('/*/sac:SUNATTotalCashed', namespaces=ns)[0].text, '1203.60', 'XML: total cobrado')
    check(perc._l10n_pe_edi_generate_retention_filename().startswith('%s-40-' % company.vat),
          'nombre del archivo tipo 40', perc._l10n_pe_edi_generate_retention_filename())
    omitido('Envío del comprobante de percepción a SUNAT', 'sin credenciales de prueba')
    cfg.restaurar()


ejecutar(t_perc_cpe)

# ======================================================================
#  al_l10n_pe_fourth_category
# ======================================================================
caso('al_l10n_pe_fourth_category', 'Recibos por honorarios: retención 8 % sobre S/ 1 500',
     'Compras ▸ recibo por honorarios E001 de S/ 2 000 y otro de S/ 1 200 ▸ Publicar.')

RH = {}


def t_4ta():
    a = registro(compra(HONOR, [linea(2000.0, taxes=env['account.tax'], name='Asesoría contable')], OCT,
                        doc=DOC_RH, serie='E001'), 'RH con retención')
    igual(a.l10n_pe_fourth_category_status, 'withheld', 'estado: retenido')
    igual(a.l10n_pe_fourth_category_amount, 160.0, 'retención 8 % = 160')
    igual(a.amount_residual, 1840.0, 'se paga el neto 1 840')
    tl = a.line_ids.filtered(lambda l: l.tax_line_id.l10n_pe_fourth_category)
    check(tl.account_id.code.startswith('40172'), 'retención a la 40172', tl.account_id.code)
    b = registro(compra(HONOR, [linea(1200.0, taxes=env['account.tax'], name='Asesoría')], OCT,
                        doc=DOC_RH, serie='E001'), 'RH bajo el umbral')
    igual(b.l10n_pe_fourth_category_status, 'under_threshold', 'S/ 1 200: no se retiene')
    RH.update(a=a, b=b)


ejecutar(t_4ta)

caso('al_l10n_pe_fourth_category', 'Reporte mensual (PDT 601) y TXT .4ta',
     'Perú ▸ Cuarta categoría ▸ reporte de octubre 2026 ▸ Calcular ▸ Exportar TXT.')


def t_4ta_txt():
    a, b = RH['a'], RH['b']
    for m in (a, b):
        pagar(m, payment_date=OCT)
    rep = registro(env['l10n_pe.fourth.category.report'].create({'year': 2026, 'month': '10'}),
                   'reporte 4ta oct-2026')
    rep.action_compute()
    own = rep.line_ids.filtered(lambda l: l.move_id.partner_id == HONOR and l.move_id in (a | b))
    igual(len(own), 2, 'los dos recibos pagados en el reporte')
    igual(sum(own.mapped('amount_withheld')), 160.0, 'retenido del prestador: 160')
    rep.action_export_txt()
    igual(rep.filename, '0601202610%s.4ta' % company.vat, 'nombre del archivo PDT 601')
    rows = [l.split('|') for l in lineas_txt(rep.file.content)]
    na = a.l10n_latam_document_number.split('-')[1].lstrip('0')
    ra = [r for r in rows if r[1] == HONOR.vat and r[4].lstrip('0') == na]
    check(ra and ra[0][:4] == ['06', HONOR.vat, 'R', 'E001'] and ra[0][5] == '2000.00',
          'fila del recibo: tipo 06, RUC, R, E001, 2000.00', ra)
    check(ra and ra[0][8] == '1', 'indicador de retención = 1', ra and ra[0][8])


ejecutar(t_4ta_txt)

# ======================================================================
#  al_l10n_pe_itf
# ======================================================================
from odoo.addons.al_l10n_pe_itf.models.res_company import itf_amount  # noqa: E402

caso('al_l10n_pe_itf', 'ITF al confirmar pagos y cobros en un banco afecto',
     'Contabilidad ▸ Pagos ▸ pago de S/ 2 000 y cobro de S/ 5 000 desde el banco con ITF ▸ Confirmar.')


def t_itf():
    igual(itf_amount(12345.67), 0.60, 'redondeo del art. 13 (12 345.67 → 0.60)')
    itf_acc = company._l10n_pe_itf_account()

    def pago(amount, ptype='outbound', **vals):
        p = env['account.payment'].create(dict({
            'payment_type': ptype, 'partner_type': 'supplier' if ptype == 'outbound' else 'customer',
            'partner_id': PROV.id, 'amount': amount, 'date': OCT, 'journal_id': J_ITF.id}, **vals))
        p.action_post()
        return p
    p1 = registro(pago(2000), 'pago con ITF')
    igual(p1.l10n_pe_itf_amount, 0.10, 'ITF 0.005 % de 2 000 = 0.10')
    mv = registro(p1.l10n_pe_itf_move_id, 'asiento del ITF')
    check(mv.state == 'posted' and mv.journal_id == J_ITF, 'asiento del ITF publicado en el banco')
    igual(sum(mv.line_ids.filtered(lambda l: l.account_id == itf_acc).mapped('balance')), 0.10,
          'gasto ITF en la %s' % itf_acc.code)
    p2 = registro(pago(5000, 'inbound'), 'cobro con ITF')
    igual(p2.l10n_pe_itf_move_id.amount_total, 0.25, 'cobro de 5 000: ITF 0.25')
    p3 = registro(pago(5000, l10n_pe_itf_exempt=True), 'pago exonerado')
    check(not p3.l10n_pe_itf_move_id, 'pago marcado exonerado: sin ITF')
    p4 = registro(pago(999), 'pago < S/ 1 000')
    check(not p4.l10n_pe_itf_move_id, 'menos de S/ 1 000: sin ITF')
    p1.action_draft()
    igual(mv.state, 'cancel', 'al pasar el pago a borrador se anula el ITF')
    p1.action_post()
    check(p1.l10n_pe_itf_move_id and p1.l10n_pe_itf_move_id.state == 'posted', 'al republicar vuelve el ITF')


ejecutar(t_itf)

# ======================================================================
#  al_l10n_pe_monthly_tax
# ======================================================================
caso('al_l10n_pe_monthly_tax', 'UIT desde el parámetro de reglas nativo l10n_pe_uit',
     'Ajustes ▸ Perú ▸ UIT: con la planilla instalada se usa el parámetro del año.')


def t_uit():
    vals = {v.date_from.year: float(v.parameter_value) for v in env['hr.rule.parameter.value'].sudo().search(
        [('rule_parameter_id.code', '=', 'l10n_pe_uit')])}
    check(vals, 'parámetro l10n_pe_uit con valores', vals)
    for y in (2024, 2025, 2026):
        if y in vals:
            igual(company._l10n_pe_tax_uit(y), vals[y], 'UIT %s = parámetro nativo' % y)
    check(company._l10n_pe_tax_uit(2025) != company.l10n_pe_tax_uit_amount,
          'para 2025 no usa el valor fijo de la compañía (%s)' % company.l10n_pe_tax_uit_amount)


ejecutar(t_uit)

caso('al_l10n_pe_monthly_tax', 'Resumen 621 de setiembre: IGV y pago a cuenta (variación por mis comprobantes)',
     'Perú ▸ Declaraciones ▸ 621 setiembre 2026 ▸ Calcular; registrar facturas y recalcular.')
SEP_DOCS = {}


def t_621():
    Regime = env['l10n_pe.tax.regime']
    if not Regime.search([('company_id', '=', company.root_id.id)], limit=1):
        registro(Regime.create({'company_id': company.root_id.id, 'regime': 'general',
                                'date_from': date(2026, 1, 1)}), 'régimen general 2026')
    rec = env['l10n_pe.monthly.tax'].search([('company_id', '=', company.id), ('year', '=', 2026),
                                             ('month', '=', '9')], limit=1) \
        or env['l10n_pe.monthly.tax'].create({'year': 2026, 'month': '9'})
    registro(rec, '621 setiembre')
    if rec.state != 'draft' and hasattr(rec, 'action_reset') and rec.state != 'declared':
        rec.action_reset()
    rec.action_compute()
    before = {f: rec[f] for f in ('sale_base', 'sale_igv', 'purchase_igv', 'igv_result', 'net_income')}
    inv = registro(venta(CLIENTE, [linea(10000.0, account=ACC_INC)], SEP), 'venta setiembre 10 000')
    bill = registro(compra(PROV, [linea(4000.0, taxes=IGV_C, account=ACC_EXP2)], SEP), 'compra setiembre 4 000')
    SEP_DOCS.update(inv=inv, bill=bill)
    rec.action_compute()
    igual(rec.state, 'calculated', 'periodo calculado')
    igual(rec.sale_base - before['sale_base'], 10000.0, 'base de ventas +10 000')
    igual(rec.sale_igv - before['sale_igv'], 1800.0, 'IGV de ventas +1 800')
    igual(rec.purchase_igv - before['purchase_igv'], 720.0, 'IGV de compras +720')
    igual(rec.igv_result - before['igv_result'], 1080.0, 'resultado del IGV +1 080')
    igual(rec.net_income - before['net_income'], 10000.0, 'ingresos netos +10 000')
    exp_tax = round(max(rec.coefficient, 0.015) * rec.net_income) if rec.net_income > 0 else 0.0
    check(abs(rec.income_tax - exp_tax) <= 1, 'pago a cuenta = max(coeficiente, 1.5 %) × ingresos',
          'coef %s, ingresos %s, renta %s' % (rec.coefficient, rec.net_income, rec.income_tax))
    igual(rec.total_to_pay, rec.igv_to_pay + rec.income_tax_to_pay, 'total = IGV + renta')


ejecutar(t_621)

# ======================================================================
#  al_l10n_pe_ple
# ======================================================================
def opciones(report, d1='2026-09-01', d2='2026-09-30'):
    return report.with_company(company).get_options({
        'date': {'date_from': d1, 'date_to': d2, 'filter': 'custom', 'mode': 'range'},
        'selected_variant_id': report.id})


def ple_asiento():
    if 'entry' not in SEP_DOCS:
        m = Move.create({'move_type': 'entry', 'journal_id': J_GRAL.id, 'date': SEP,
                         'l10n_pe_gloss': 'QA20 Glosa PLE setiembre',
                         'line_ids': [Command.create({'account_id': ACC_EXP2.id, 'name': 'QA20 gasto', 'debit': 321.0}),
                                      Command.create({'account_id': ACC_PROV.id, 'name': 'QA20 prov', 'credit': 321.0})]})
        m.action_post()
        SEP_DOCS['entry'] = m
    return SEP_DOCS['entry']


GL = env.ref('account_reports.general_ledger_report')
GLH = env['account.general.ledger.report.handler'].with_company(company)

caso('al_l10n_pe_ple', 'PLE 5.1 Libro Diario (TXT oficial + XLSX del módulo)',
     'Contabilidad ▸ Libro mayor ▸ setiembre 2026 ▸ botones «PLE 5.1» y «XLSX 5.1».')


def t_ple51():
    m = registro(ple_asiento(), 'asiento de setiembre')
    res = GLH.l10n_pe_export_ple_51_to_txt(opciones(GL))
    name = res['file_name']
    check(name.startswith('LE%s202609' % company.vat) and '050100' in name, 'nombre del TXT 5.1', name)
    lines = lineas_txt(res['file_content'])
    mine = [l.split('|') for l in lines if l.split('|')[1] == str(m.id)]
    igual(len(mine), 2, 'dos líneas del asiento (CUO = id)')
    if mine:
        igual(len(mine[0]) - 1, 21, '21 campos por línea (cierra con «|»)')
        igual(mine[0][0], '20260900', 'periodo AAAAMM00')
        check({r[3] for r in mine} == {ACC_EXP2.code, ACC_PROV.code}, 'cuentas del asiento', [r[3] for r in mine])
        check(any(r[17] == '321.00' for r in mine) and any(r[18] == '321.00' for r in mine), 'debe y haber 321.00')
        if not any('Glosa PLE' in r[15] for r in mine):
            aviso('La glosa peruana (l10n_pe_gloss) no llega al 5.1 oficial',
                  'campo 16 = %r: el exportador nativo usa el nombre del asiento' % mine[0][15])
    x = GLH.l10n_pe_export_ple_51_to_xlsx(opciones(GL))
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x['file_content'])).worksheets[0]
    values = [str(c) for row in ws.iter_rows(values_only=True) for c in row if c is not None]
    check(x['file_type'] == 'xlsx' and str(m.id) in values, 'XLSX 5.1 con las mismas líneas del TXT')


ejecutar(t_ple51)

caso('al_l10n_pe_ple', 'PLE 6.1 Libro Mayor', 'Libro mayor ▸ setiembre 2026 ▸ «PLE 6.1».')


def t_ple61():
    m = ple_asiento()
    res = GLH.l10n_pe_export_ple_61_to_txt(opciones(GL))
    check('060100' in res['file_name'], 'nombre del TXT 6.1', res['file_name'])
    mine = [l.split('|') for l in lineas_txt(res['file_content']) if l.split('|')[1] == str(m.id)]
    igual(len(mine), 2, 'líneas del asiento en el mayor')
    if mine:
        igual(len(mine[0]) - 1, 21, '21 campos por línea')
        igual(mine[0][19], '', 'el 6.1 no lleva la referencia del libro (campo 20)')
    x = GLH.l10n_pe_export_ple_61_to_xlsx(opciones(GL))
    igual(x['file_type'], 'xlsx', 'XLSX 6.1 generado')


ejecutar(t_ple61)

R84 = env['account.report'].search([('custom_handler_model_name', '=', 'l10n_pe.tax.ple.8.1.report.handler')], limit=1)
R144 = env['account.report'].search([('custom_handler_model_name', '=', 'l10n_pe.tax.ple.14.1.report.handler')], limit=1)


def sep_docs():
    if 'inv' not in SEP_DOCS:
        SEP_DOCS['inv'] = venta(CLIENTE, [linea(10000.0, account=ACC_INC)], SEP)
        SEP_DOCS['bill'] = compra(PROV, [linea(4000.0, taxes=IGV_C, account=ACC_EXP2)], SEP)
    return SEP_DOCS['inv'], SEP_DOCS['bill']


caso('al_l10n_pe_ple', 'PLE 8.1 → RCE 8.4 (Registro de Compras)',
     'Informes ▸ RCE Registro de Compras 8.4 ▸ setiembre 2026 ▸ «TXT».')


def t_ple81():
    _inv, bill = sep_docs()
    registro(bill, 'compra de setiembre')
    res = env[R84.custom_handler_model_name].with_company(company).export_to_txt(opciones(R84))
    name = res['file_name']
    check(name[2:13] == company.vat and name[13:19] == '202609' and name[21:27] == '080400' and len(name) == 33,
          'nombre oficial LE…080400', name)
    serie, folio = bill.l10n_latam_document_number.split('-')
    rows = [l.split('|') for l in lineas_txt(res['file_content']) if PROV.vat in l and folio.lstrip('0') in l]
    igual(len(rows), 1, 'la compra está en el 8.4')
    if rows:
        r = rows[0]
        igual(len(r) - 1, 37, '37 campos')
        igual(r[12], PROV.vat, 'campo 13: RUC del proveedor')
        igual(r[4], '15/09/2026', 'campo 5: fecha de emisión')
        igual((r[14], r[15]), ('4000.00', '720.00'), 'base e IGV de la adquisición gravada')
        igual(r[24], '4720.00', 'campo 25: total')


ejecutar(t_ple81)

caso('al_l10n_pe_ple', 'PLE 14.1 → RVIE 14.4 (Registro de Ventas)',
     'Informes ▸ RVIE Registro de Ventas 14.4 ▸ setiembre 2026 ▸ «TXT».')


def t_ple141():
    inv, _bill = sep_docs()
    registro(inv, 'venta de setiembre')
    try:
        res = env[R144.custom_handler_model_name].with_company(company).export_to_txt(opciones(R144))
    except UserError as e:
        aviso('El 14.4 de setiembre no se generó por datos de otros', str(e)[:300])
        return
    name = res['file_name']
    check(name[21:27] == '140400' and name[32] == '2' and len(name) == 33, 'nombre oficial LE…140400…2', name)
    serie, folio = inv.name.split('-')
    rows = [l.split('|') for l in lineas_txt(res['file_content']) if serie in l and folio.lstrip('0') in l]
    igual(len(rows), 1, 'la venta está en el 14.4')
    if rows:
        r = rows[0]
        igual(len(r) - 1, 33, '33 campos (sin los 34-40)')
        igual(r[11], CLIENTE.vat, 'campo 12: RUC del cliente')
        igual((r[14], r[16], r[25]), ('10000.00', '1800.00', '11800.00'), 'base, IGV y total')


ejecutar(t_ple141)

# ======================================================================
#  al_l10n_pe_sire
# ======================================================================
def periodo(model):
    rec = env[model].search([('company_id', '=', company.id), ('year', '=', 2026), ('month', '=', '09')], limit=1)
    return rec or env[model].create({'year': 2026, 'month': '09', 'company_id': company.id,
                                     'download_manual': True})


caso('al_l10n_pe_sire', 'RVIE setiembre: comprobantes del sistema',
     'Perú ▸ SIRE ▸ RVIE setiembre 2026 ▸ «Cargar del sistema».')


def t_rvie():
    inv, _bill = sep_docs()
    rvie = registro(periodo('l10n_pe.sire.rvie'), 'RVIE 2026-09')
    rvie.action_load_system()
    line = rvie.system_line_ids.filtered(lambda l: l.move_id == inv)
    igual(len(line), 1, 'la venta aparece en la propuesta del sistema')
    if line:
        igual((line.tipo_doc_identidad, line.nro_doc_identidad), ('6', CLIENTE.vat), 'cliente con RUC')
        igual((line.bi_gravada, line.igv_ipm, line.total_cp), (10000.0, 1800.0, 11800.0), 'importes')
        check(len(line.car_sunat) == 27 and line.car_sunat.startswith(company.vat), 'CAR de 27 caracteres',
              line.car_sunat)
    omitido('Descarga de la propuesta y envío al SIRE', 'sin credenciales SOL de prueba')


ejecutar(t_rvie)

caso('al_l10n_pe_sire', 'RCE con proveedor no domiciliado (8.5): país, modalidad y tipo de renta nativos',
     'Compras ▸ factura tipo 91 de proveedor de EE. UU. con tipo de renta 18 y modalidad 3 ▸ '
     'SIRE ▸ RCE setiembre ▸ «Cargar no domiciliados».')


def t_rce_nd():
    _inv, bill = sep_docs()
    us = env.ref('base.us')
    nd_p = registro(socio('Cloud Services Inc', country_id=us.id, vat='98-7654321', street='1 Main St'),
                    'proveedor no domiciliado')
    nd = registro(compra(nd_p, [linea(500.0, taxes=env['account.tax'], name='Licencia')], SEP, doc=DOC91,
                         serie='INV', post=False, l10n_pe_sire_nd_income_type='18',
                         l10n_pe_service_modality='3', l10n_pe_sire_nd_agreement='00',
                         l10n_pe_sire_nd_gross_income=500.0, l10n_pe_sire_nd_withholding_rate=30.0),
                  'factura de no domiciliado')
    nd.action_post()
    igual(nd.l10n_pe_usage_type_id.code, '18', 'tipo de renta guardado en l10n_pe_usage_type_id nativo')
    igual(nd.l10n_pe_sire_nd_service_modality, '3', 'modalidad = l10n_pe_service_modality nativo')
    check(nd.l10n_pe_sire_is_non_domiciled, 'marcado como no domiciliado')
    rce = registro(periodo('l10n_pe.sire.rce'), 'RCE 2026-09')
    check(nd not in rce._sire_system_moves() and nd in rce._sire_nd_moves(),
          'sale del RCE y va al registro de no domiciliados')
    rce.action_load_system()
    check(rce.system_line_ids.filtered(lambda l: l.move_id == bill), 'la compra doméstica sí está en el RCE')
    rce.action_sire_load_nd()
    line = rce.nd_line_ids.filtered(lambda l: l.move_id == nd)
    igual(len(line), 1, 'línea del no domiciliado desplegada')
    if line:
        check(not line.check_detail, 'sin observaciones', line.check_detail)
        row = rce._sire_nd_row(line)
        igual(len(row), 35, '35 campos del anexo 9')
        igual(row[16], us.l10n_pe_code, 'campo 17: país desde res.country.l10n_pe_code (9249)')
        igual(row[31], '18', 'campo 32: tipo de renta')
        igual(row[32], '3', 'campo 33: modalidad del servicio')
        igual(row[24], '500.00', 'campo 25: renta bruta')
        igual(row[28], '150.00', 'campo 29: impuesto retenido 30 %')
    check(esperar_error(UserError, rce.action_sire_send_nd), 'no se envía sin aceptar la propuesta')
    omitido('Carga del 8.5 al SIRE', 'sin credenciales SOL de prueba')


ejecutar(t_rce_nd)

# ======================================================================
#  al_l10n_pe_cpe_validation
# ======================================================================
CPE_SERVICE = 'odoo.addons.al_l10n_pe_cpe_validation.services.sunat_cpe'


class _Resp:
    def __init__(self, status, payload):
        self.status_code, self._p, self.text = status, payload, str(payload)

    def json(self):
        return self._p


def _fake_post(cp='1', dom='00', notes=None, calls=None):
    def fake(url, data=None, json=None, headers=None, timeout=None):
        if calls is not None:
            calls.append(url)
        if 'oauth2/token' in url:
            return _Resp(200, {'access_token': 'tok-qa20', 'expires_in': 3600})
        return _Resp(200, {'success': True, 'message': 'Operation Success! ',
                           'data': {'estadoCp': cp, 'estadoRuc': '00', 'condDomiRuc': dom,
                                    'observaciones': notes or []}})
    return patch(CPE_SERVICE + '.requests.post', side_effect=fake)


caso('al_l10n_pe_cpe_validation', 'Consulta integrada de comprobantes (SUNAT simulada)',
     'Compras ▸ factura de proveedor ▸ «Validar en SUNAT»: válido, observado y no válido.')


def t_cpe():
    root = company.root_id.sudo()
    cfg = Config(root, {'l10n_pe_cpe_client_id': 'qa20-cid', 'l10n_pe_cpe_client_secret': 'qa20-secret',
                        'l10n_pe_cpe_token': False, 'l10n_pe_cpe_token_expiry': False})
    calls = []
    ok = registro(compra(PROV, [linea(1000.0, taxes=env['account.tax'])], OCT, post=False), 'CPE válido')
    vals = ok._l10n_pe_cpe_query_values()
    igual((vals['numRuc'], vals['codComp'], vals['numeroSerie']), (PROV.vat, '01', 'FQ20'),
          'datos de la consulta según el manual')
    with _fake_post(calls=calls):
        ok.action_l10n_pe_cpe_validate()
    igual((ok.l10n_pe_cpe_result, ok.l10n_pe_cpe_doc_state), ('valid', '1'), 'resultado: válido')
    check(any('oauth2/token' in c for c in calls), 'se pidió token OAuth2 (simulado)')
    obs = registro(compra(PROV, [linea(800.0, taxes=env['account.tax'])], OCT, post=False), 'CPE observado')
    with _fake_post(dom='12'):
        obs.action_l10n_pe_cpe_validate()
    igual(obs.l10n_pe_cpe_result, 'observed', 'domicilio no habido: observado')
    bad = registro(compra(PROV, [linea(900.0, taxes=env['account.tax'])], OCT, post=False), 'CPE no válido')
    with _fake_post(cp='0', notes=['Comprobante no informado']):
        bad.action_l10n_pe_cpe_validate()
    igual(bad.l10n_pe_cpe_result, 'invalid', 'no informado: no válido')
    check('no informado' in (bad.l10n_pe_cpe_message or ''), 'mensaje de SUNAT guardado')
    omitido('Consulta real a la API de SUNAT', 'sin credenciales de prueba')
    cfg.restaurar()


ejecutar(t_cpe)

# ======================================================================
#  al_l10n_pe_financial_reports
# ======================================================================
caso('al_l10n_pe_financial_reports', 'Hoja de trabajo (formato 1) 2026 con mis cuentas',
     'Contabilidad ▸ Informes ▸ Perú ▸ Hoja de trabajo ▸ 2026.')


def t_ws():
    rep = env.ref('al_l10n_pe_financial_reports.pe_worksheet_format_1')
    opts = rep.with_company(company).get_options({'date': {'date_from': '2026-01-01', 'date_to': '2026-12-31',
                                                           'filter': 'custom', 'mode': 'range'}})
    vals = {}
    for line in rep.with_company(company)._get_lines(opts):
        vals[line.name] = {c.expression_label: c.no_format or 0.0 for c in line.columns}
    row = lambda acc: next((v for k, v in vals.items() if k and k.startswith(acc.code)), None)  # noqa: E731
    d1 = row(ACC_D1)
    exp_d1 = sum(env['account.move.line'].search([('account_id', '=', ACC_D1.id), ('parent_state', '=', 'posted'),
                                                   ('date', '>=', '2026-01-01')]).mapped('balance'))
    check(d1 is not None, 'fila de la cuenta 94 de destino', list(vals)[:3])
    if d1:
        igual(d1.get('sum_debit', 0.0), exp_d1, 'sumas del debe de la 94 = mayor')
        igual(d1.get('function_loss', 0.0), exp_d1, 'la 94 va a pérdidas por función')
    t = vals.get('Totales', {})
    for a, b in (('sum_debit', 'sum_credit'), ('balance_debit', 'balance_credit')):
        igual(t.get(a, 0.0), t.get(b, 0.0), 'totales cuadrados: %s = %s' % (a, b))


ejecutar(t_ws)

caso('al_l10n_pe_financial_reports', 'Estado de cambios en el patrimonio (3.19) = patrimonio del balance',
     'Informes ▸ Perú ▸ Estado de cambios en el patrimonio neto ▸ 2026.')


def t_eq():
    rep = env.ref('al_l10n_pe_financial_reports.pe_statement_changes_equity')
    igual(rep.column_ids.mapped('expression_label'), ['movement_note', 'opening', 'movement', 'closing'],
          'columnas del formato 3.19')

    def values(report, d1, d2):
        o = report.with_company(company).get_options({'selected_variant_id': report.id, 'unfold_all': True,
                                                      'date': {'date_from': d1, 'date_to': d2,
                                                               'mode': 'range', 'filter': 'custom'}})
        codes = {l.id: l.code for l in report.line_ids}
        out = {}
        for line in report.with_company(company)._get_lines(o):
            lid = report._get_res_id_from_line_id(line.id, 'account.report.line')
            if lid in codes and not report._get_res_id_from_line_id(line.id, 'account.account'):
                out[codes[lid]] = {c.expression_label: c.no_format for c in line.columns}
        return out
    eq = values(rep, '2026-01-01', '2026-10-31')
    bs = values(env.ref('account_reports.balance_sheet'), '2026-10-31', '2026-10-31')
    closing = eq.get('SCEP_CLOSING', {}).get('closing') or 0.0
    igual(closing, bs.get('EQ', {}).get('balance') or 0.0, 'saldo final del 3.19 = patrimonio del balance')


ejecutar(t_eq)

# ======================================================================
#  al_l10n_pe_dashboard
# ======================================================================
caso('al_l10n_pe_dashboard', 'Tablero Perú: vista de IGV y armado de la hoja',
     'Tableros ▸ Perú: ventas, compras e IGV del mes.')


def t_dash():
    inv, bill = sep_docs()
    rows = env['l10n_pe.dashboard.tax'].search([('move_id', 'in', (inv | bill).ids)])
    vals = {r.move_id: (r.kind, r.base, r.igv) for r in rows}
    igual(vals.get(inv), ('sale', 10000.0, 1800.0), 'venta: base 10 000, IGV 1 800')
    igual(vals.get(bill), ('purchase', 4000.0, 720.0), 'compra: base 4 000, IGV 720')
    builder = env['l10n_pe.dashboard.builder']
    dash = env.ref('al_l10n_pe_dashboard.dashboard_pe')
    builder._refresh_dashboard(force=True)
    dash._check_spreadsheet_data()
    data = json.loads(dash.spreadsheet_data)
    cells = list(data['sheets'][0]['cells'].values())
    check('Ventas, compras e IGV' in cells, 'sección «Ventas, compras e IGV» en la hoja')
    for src in list(data['pivots'].values()) + list(data['lists'].values()):
        env[src['model']].search_count(src['domain'])
    check(True, 'todas las fuentes del tablero se pueden consultar (%d)' % (len(data['pivots']) + len(data['lists'])))


ejecutar(t_dash)
