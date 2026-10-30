# -*- coding: utf-8 -*-
"""Fase planillas_tools: OL-PAYROLL y OL-TOOLS sobre Odoo 20.

Escenario: OLANO IT S.A.C. (dev) paga su planilla mensual de jul-2025 a
oct-2026 a cuatro trabajadores QA20 (AFP, ONP, sueldo < RMV y uno que cesa
el 10/07/2026), procesa gratificaciones, CTS, vacaciones, liquidación y
renta de 5.ª, contabiliza el lote y genera los archivos (PLAME, AFPnet,
TXT del banco). Un quinto trabajador marca asistencia (importada de Excel)
y se le tarea con horas extra; un sexto es obrero de construcción civil.

Se antepone _qa.py y se cierra con _fin.py (ver run_qa.sh).
"""
import io
from calendar import monthrange
from datetime import date, datetime, timedelta

from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.tools import BinaryBytes, float_round

S = {}  # estado compartido entre casos


def correr(fn):
    """ejecutar() con confirmación por caso y reintento ante errores de
    concurrencia (otros agentes escriben en la misma base). En modo rollback
    no se confirma ni se reintenta."""
    import os
    final = not os.environ.get('QA_ROLLBACK')
    for intento in range(3):
        n = len(_CUR['checks'])
        ejecutar(fn)
        err = _CUR.get('error') or ''
        if final and intento < 2 and ('SerializationFailure' in err or 'concurrent update' in err):
            env.cr.rollback()
            env.invalidate_all()
            del _CUR['checks'][n:]
            _CUR['error'] = None
            _CUR['registros'] = []
            aviso('reintento %d de %s por concurrencia con otro agente' % (intento + 1, fn.__name__), err)
            continue
        break
    if final:
        env.cr.commit()


def importar(wiz):
    """Camino síncrono documentado del importador (el botón lanza un hilo que
    hace commit por lote): se marca current_test para que no haga commit
    dentro del savepoint del caso."""
    from odoo import modules
    prev = modules.module.current_test
    modules.module.current_test = True
    try:
        return wiz._process_all_rows(wiz._preprocess_rows(wiz._iter_data_rows()))
    finally:
        modules.module.current_test = prev


class _Deshacer(Exception):
    """Para deshacer a propósito un savepoint (pruebas hipotéticas)."""


def r2(x):
    return float_round(x, 2, rounding_method='HALF-UP')


def linea(slip, code):
    return sum(slip.line_ids.filtered(lambda l: l.code == code).mapped('total'))


def adjunto(action):
    return env['ir.attachment'].browse(int(action['url'].split('/')[3].split('?')[0]))


def xlsx(sheets):
    import openpyxl
    wb = openpyxl.Workbook()
    default = wb.active
    for idx, (name, rows) in enumerate(sheets.items()):
        ws = default if idx == 0 else wb.create_sheet()
        ws.title = name
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return BinaryBytes(buf.getvalue())


def impuesto_5ta(renta_neta, uit):
    """Escala del art. 53 LIR: 8/14/17/20/30 % por tramos de 5/20/35/45 UIT."""
    tramos = [(5, 8), (20, 14), (35, 17), (45, 20), (None, 30)]
    total, piso = 0.0, 0.0
    for limite, tasa in tramos:
        techo = limite * uit if limite else float('inf')
        if renta_neta <= piso:
            break
        total += (min(renta_neta, techo) - piso) * tasa / 100.0
        piso = techo
    return total


Rule = env['hr.salary.rule'].with_company(company)
Payslip = env['hr.payslip']
Run = env['hr.payslip.run']
STRUCT = env.ref('al_hr_pe.base_structure')
PARAM = env['hr.main.parameter'].search([('company_id', '=', company.id)], limit=1)
DNI = env.ref('al_hr_pe.l10n_pe_hr_document_type_dni')


def regla(code, struct=STRUCT):
    return Rule.search([('code', '=', code), ('struct_ids', 'in', struct.ids)], limit=1)


def lote(y, m):
    return S['lotes'][(y, m)]


def boleta(emp, y, m):
    return lote(y, m).slip_ids.filtered(lambda s: s.employee_id == emp)


def set_if_empty(rec, vals):
    """Configura solo lo que esté vacío (no pisa configuración ajena)."""
    todo = {k: v for k, v in vals.items() if not rec[k]}
    if todo:
        rec.write(todo)
    return sorted(todo)


# ======================================================================
# al_hr_pe — preparación
# ======================================================================
caso('al_hr_pe', 'Preparación: parámetros, periodos y trabajadores QA20',
     'El jefe de planillas genera los periodos 2025-2026, completa la pestaña de '
     'beneficios de Parámetros Principales y registra 4 trabajadores con su contrato.')


def preparar():
    check(PARAM, 'Parámetros Principales de la compañía existen', PARAM.display_name)
    for year in (2025, 2026):
        env['hr.period.generator'].create({'year': year, 'company_id': company.id}).action_generate()
    periodos = env['hr.period'].search([('company_id', '=', company.id), ('period_type', '!=', 'weekly'),
                                        ('code', 'in', ['202507', '202610'])])
    check(len(periodos) == 2, 'periodos mensuales 2025-07 y 2026-10 generados')
    wet = env['hr.work.entry.type']
    cambios = set_if_empty(PARAM, {
        'basic_sr_id': regla('BAS').id,
        'household_allowance_sr_id': regla('AF').id,
        'extra_hours_sr_id': regla('HE25').id,
        'commission_sr_ids': [Command.set(regla('COMI').ids)],
        'bonus_sr_ids': [Command.set(regla('BONR').ids)],
        'cts_input_id': regla('CTS').id,
        'gratification_input_id': regla('GRA').id,
        'bonus_nine_input_id': regla('BON9').id,
        'truncated_gratification_input_id': regla('GRA_TRU').id,
        'truncated_bonus_nine_input_id': regla('BON9_TRU').id,
        'truncated_cts_input_id': regla('CTS_TRU').id,
        'vacation_input_id': regla('VAC').id,
        'truncated_vacation_input_id': regla('VATRU').id,
        'lack_wd_ids': [Command.set(wet.search([('code', '=', 'FAL')]).ids)],
        'working_wd_ids': [Command.set(wet.search([('code', 'in', ('DLAB', 'DOM'))]).ids)],
        'medical_rest_wd_ids': [Command.set(wet.search([('code', '=', 'DMED')]).ids)],
        'fifth_afect_sr_id': regla('BAS').id,
        'fifth_extr_sr_id': regla('REAQ').id,
        'proy_afect_sr_id': regla('BAS').id,
        'fifth_category_input_id': regla('QUINTA').id,
        'insurable_remuneration': regla('AAFP').id,
        'net_to_pay_sr_id': regla('NETO').id,
    })
    aviso('Parámetros Principales: se completaron campos vacíos', ', '.join(cambios) or 'ninguno')
    # Asiento peruano por lote: opt-in documentado (sin diario nativo en la estructura).
    if STRUCT.with_company(company).journal_id:
        aviso('Estructura BASE: se quita el diario nativo para usar el asiento peruano por lote',
              'antes: %s (hr_payroll_account generaría un asiento por boleta)'
              % STRUCT.with_company(company).journal_id.display_name)
        STRUCT.with_company(company).journal_id = False
    PARAM.action_generate_brackets(year=2026)
    afp = {m.name: m for m in env['hr.membership'].search([('company_id', '=', False)])}
    essalud = env['hr.social.insurance'].search([('name', '=', 'EsSalud')], limit=1)
    S['essalud'] = essalud
    datos = [
        # clave, ap. paterno, ap. materno, nombres, DNI, sueldo, ingreso, afiliación
        ('afp', 'Torres', 'Ruiz', 'Ana', '79000101', 6000.0, date(2025, 1, 1), 'AFP INTEGRA'),
        ('onp', 'Díaz', 'Mora', 'Bruno', '79000102', 2500.0, date(2025, 1, 1), 'ONP'),
        ('rmv', 'Ríos', 'León', 'Carla', '79000103', 1000.0, date(2025, 6, 1), 'AFP HABITAT'),
        ('cese', 'Cárdenas', 'Paz', 'Diego', '79000104', 2400.0, date(2025, 3, 1), 'AFP PRIMA'),
    ]
    S['emp'] = {}
    for key, ap, am, nom, dni, wage, start, memb in datos:
        emp = env['hr.employee'].create({
            'name': qa_name('%s %s %s' % (ap, am, nom)),
            'names': nom, 'last_name': ap, 'm_last_name': am,
            'company_id': company.id, 'tz': 'America/Lima',
            'l10n_pe_document_type_id': DNI.id, 'identification_id': dni,
            'date_version': start, 'contract_date_start': start,
            'wage': wage, 'structure_type_id': STRUCT.type_id.id,
        })
        vals = {'membership_id': afp[memb].id, 'social_insurance_id': essalud.id}
        if afp[memb].is_afp:
            vals.update({'l10n_pe_commission_type': 'flow', 'l10n_pe_cuspp': '6%sXQA0' % dni[-6:]})
        emp.version_id.write(vals)
        S['emp'][key] = registro(emp, '%s, sueldo %.2f, %s' % (memb, wage, start))
    check(S['emp']['afp'].name.startswith('QA20'), 'nombre del trabajador con prefijo QA20',
          S['emp']['afp'].name)
    S['lotes'] = {}


def crear_lote(y, m, empleados, validar=True):
    last = monthrange(y, m)[1]
    periodo = env['hr.period'].search([('code', '=', '%04d%02d' % (y, m)),
                                       ('company_id', '=', company.id)], limit=1)
    run = Run.create({'name': qa_name('Planilla %04d-%02d' % (y, m)),
                      'date_start': date(y, m, 1), 'date_end': date(y, m, last),
                      'company_id': company.id, 'periodo_id': periodo.id})
    for emp in empleados:
        Payslip.create({'name': qa_name('Boleta %04d-%02d %s' % (y, m, emp.names)),
                        'employee_id': emp.id, 'struct_id': STRUCT.id,
                        'date_from': date(y, m, 1), 'date_to': date(y, m, last),
                        'payslip_run_id': run.id})
    run.slip_ids.compute_sheet()
    if validar:
        run.slip_ids.action_payslip_done()
    S['lotes'][(y, m)] = run
    return run


correr(preparar)
LISTO = bool(S.get('emp'))

MESES_1 = [(2025, m) for m in range(7, 13)] + [(2026, m) for m in range(1, 7)]


def planillas_hasta_junio():
    e = S['emp']
    for y, m in MESES_1:
        crear_lote(y, m, [e['afp'], e['onp'], e['rmv'], e['cese']])
    check(len(S['lotes']) == 12, '12 lotes mensuales jul-2025 → jun-2026 validados',
          ', '.join(r.name for r in S['lotes'].values()))
    registro(lote(2026, 6), 'lote junio 2026')


if LISTO:
    caso('al_hr_pe', 'Planillas mensuales jul-2025 → jun-2026 (base de los beneficios)',
         'Cada mes se crea el lote del periodo, se calculan y validan las boletas de los 4 trabajadores.')
    correr(planillas_hasta_junio)
    LISTO = len(S['lotes']) == 12

# ======================================================================
# al_hr_pe_benefits — gratificación diciembre 2025, CTS mayo, gratificación julio
# ======================================================================


def grati_diciembre():
    g = env['hr.gratification'].create({
        'company_id': company.id, 'year': 2025, 'type': '12', 'with_bonus': True,
        'payslip_run_id': lote(2025, 12).id, 'deposit_date': date(2025, 12, 15)})
    g.action_process()
    registro(g, 'gratificación Navidad 2025')
    S['grati_dic'] = g
    for key, emp in S['emp'].items():
        ln = g.line_ids.filtered(lambda l: l.employee_id == emp)
        w = emp.version_id.wage
        if not check(ln, 'línea de gratificación de %s' % emp.name):
            continue
        igual(ln.months, 6, '%s: meses computables jul-dic' % key)
        igual(ln.total_grat, w, '%s: gratificación = computable/6 × 6 meses = sueldo' % key)
        igual(ln.bonus_essalud, r2(w * 0.09), '%s: bonificación extraordinaria 9 %% (EsSalud) = sueldo × 9 %%' % key)


def cts_mayo():
    cts = env['hr.cts'].create({'company_id': company.id, 'year': 2026, 'type': '05',
                                'payslip_run_id': lote(2026, 4).id, 'deposit_date': date(2026, 5, 15)})
    cts.action_process()
    registro(cts, 'CTS mayo 2026 (nov-2025 → abr-2026)')
    S['cts_may'] = cts
    for key, emp in S['emp'].items():
        ln = cts.line_ids.filtered(lambda l: l.employee_id == emp)
        w = emp.version_id.wage
        if not check(ln, 'línea de CTS de %s' % emp.name):
            continue
        sexto = r2(w / 6.0)  # 1/6 de la gratificación de diciembre (= sueldo)
        igual(ln.sixth_of_gratification, sexto, '%s: 1/6 de la gratificación de diciembre' % key)
        igual(ln.computable_remuneration, w + sexto, '%s: computable = sueldo + 1/6 gratificación' % key)
        igual(ln.months, 6, '%s: 6 meses nov-abr' % key)
        igual(ln.total_cts, r2((w + sexto) / 12.0 * 6), '%s: CTS = computable/12 × 6' % key)


def grati_julio():
    g = env['hr.gratification'].create({
        'company_id': company.id, 'year': 2026, 'type': '07', 'with_bonus': True,
        'payslip_run_id': lote(2026, 6).id, 'deposit_date': date(2026, 7, 15)})
    g.action_process()
    registro(g, 'gratificación Fiestas Patrias 2026')
    S['grati_jul'] = g
    igual(g.name, 'Gratificación Fiestas Patrias 2026', 'nombre propuesto del documento')
    for key, emp in S['emp'].items():
        ln = g.line_ids.filtered(lambda l: l.employee_id == emp)
        w = emp.version_id.wage
        if not check(ln, 'línea de gratificación de %s' % emp.name):
            continue
        igual(ln.computable_remuneration, w, '%s: computable = sueldo (sin AF ni variables)' % key)
        igual(ln.months, 6, '%s: 6 meses ene-jun' % key)
        igual(ln.total_grat, w, '%s: gratificación = computable/6 × 6' % key)
        igual(ln.total, r2(w * 1.09), '%s: total = gratificación + bono 9 %%' % key)


if LISTO:
    caso('al_hr_pe_benefits', 'Gratificación de Navidad 2025 (lote dic-2025)',
         'El usuario crea la gratificación de diciembre 2025 sobre el lote de diciembre y pulsa «Procesar».')
    correr(grati_diciembre)
    caso('al_hr_pe_benefits', 'CTS de mayo 2026 con 1/6 de la gratificación',
         'El usuario crea la CTS semestral nov-abr sobre el lote de abril y pulsa «Procesar».')
    correr(cts_mayo)
    caso('al_hr_pe_benefits', 'Gratificación de julio 2026 con bonificación 9 %',
         'El usuario crea la gratificación de Fiestas Patrias sobre el lote de junio y pulsa «Procesar».')
    correr(grati_julio)


# ======================================================================
# Cese del 10/07/2026 y planillas jul → oct
# ======================================================================
def liquidacion():
    e = S['emp']
    cese = e['cese']
    cese.version_id.write({'contract_date_end': date(2026, 7, 10),
                           'situation_id': env.ref('al_hr_pe.situation_0').id})
    run = crear_lote(2026, 7, [e['afp'], e['onp'], e['rmv'], cese], validar=False)
    registro(run, 'lote julio 2026 (con el cese)')
    liq = env['hr.liquidation'].create({'company_id': company.id, 'year': 2026,
                                        'gratification_type': '12', 'cts_type': '11',
                                        'payslip_run_id': run.id})
    liq.action_process()
    registro(liq, 'liquidación del cese 10/07/2026')
    S['liq'] = liq
    w = cese.version_id.wage
    sj = boleta(cese, 2026, 7)
    wd = {l.code: l.number_of_days for l in sj.worked_days_line_ids if l.number_of_days}
    aviso('días de la boleta jul-2026 del cesado (contrato hasta el 10/07)', wd)
    igual(wd.get('DLAB', 0) + wd.get('DOM', 0), 10.0,
          'DLAB + DOM = 10 días de contrato en julio (1-10/07)')
    igual(linea(sj, 'BAS'), r2(w / 30 * 10), 'básico de julio = sueldo/30 × 10 días')
    check(set(liq.cts_line_ids.employee_id.ids) | set(liq.vacation_line_ids.employee_id.ids) == {cese.id},
          'solo el trabajador cesado entra en la liquidación')
    # CTS trunca may-oct: 2 meses (may, jun) + 10 días de julio; computable =
    # sueldo + 1/6 de la gratificación de diciembre 2025 (la de julio se paga
    # el 15/07, después del cese).
    cts = liq.cts_line_ids.filtered(lambda l: l.employee_id == cese)
    if check(len(cts) == 1, 'una línea de CTS trunca'):
        sexto = r2(S['grati_dic'].line_ids.filtered(lambda l: l.employee_id == cese).total_grat / 6)
        comp = w + sexto
        igual(cts.months, 2, 'CTS trunca: meses completos (may, jun)')
        igual(cts.days, 10, 'CTS trunca: días de julio hasta el cese')
        igual(cts.computable_remuneration, comp, 'CTS trunca: computable = sueldo + 1/6 grati dic-2025')
        igual(cts.total_cts, r2(comp / 12 * 2) + r2(comp / 360 * 10),
              'CTS trunca = computable/12 × 2 + computable/360 × 10')
    # Gratificación: la de ene-jun ya se pagó en la regular de julio → no hay
    # trunca de ese semestre; jul-dic solo tiene días sueltos.
    check(not liq.gratification_line_ids.filtered(lambda l: l.compute_date == date(2026, 1, 1)),
          'sin gratificación trunca ene-jun (ya incluida en la regular de julio)')
    # Vacaciones truncas: desde el último aniversario (01/03/2026): mar-jun = 4
    # meses + 10 días de julio, sobre el sueldo.
    vac = liq.vacation_line_ids.filtered(lambda l: l.employee_id == cese)
    if check(len(vac) == 1, 'una línea de vacaciones truncas'):
        igual(str(vac.compute_date), '2026-03-01', 'vacaciones: cómputo desde el último aniversario')
        igual(vac.months, 4, 'vacaciones truncas: meses mar-jun')
        igual(vac.days, 10, 'vacaciones truncas: días de julio')
        igual(vac.truncated_vacation, r2(w / 12 * 4) + r2(w / 360 * 10),
              'vacaciones truncas = sueldo/12 × 4 + sueldo/360 × 10')
        memb = cese.version_id.membership_id
        aportes = r2(vac.total_vacation * memb.retirement_fund / 100) \
            + r2(vac.total_vacation * memb.prima_insurance / 100) \
            + r2(vac.total_vacation * memb.fixed_commision / 100)
        igual(vac.total, r2(vac.total_vacation - aportes),
              'vacaciones: neto = total − fondo 10 % − prima − comisión AFP')
    liq.action_export_to_payslips()
    slip = boleta(cese, 2026, 7)
    entradas = {i.salary_rule_id.code: i.amount for i in slip.input_line_ids}
    if cts:
        igual(entradas.get('CTS_TRU', 0.0), cts.total_cts, 'exportar: input CTS_TRU en la boleta de julio')
    if vac:
        igual(entradas.get('VATRU', 0.0), vac.total_vacation, 'exportar: input VATRU en la boleta de julio')
    slip.compute_sheet()
    igual(linea(slip, 'VATRU'), entradas.get('VATRU', 0.0), 'boleta de julio: línea VATRU calculada')
    run.slip_ids.action_payslip_done()
    check(slip.state in ('validated', 'paid'), 'boleta de cese validada', slip.state)


def planillas_ago_oct():
    e = S['emp']
    for m in (8, 9):
        crear_lote(2026, m, [e['afp'], e['onp'], e['rmv']])
    crear_lote(2026, 10, [e['afp'], e['onp'], e['rmv']], validar=False)
    registro(lote(2026, 9), 'lote setiembre 2026 (validado)')
    registro(lote(2026, 10), 'lote octubre 2026 (en borrador hasta la quinta)')
    check(not lote(2026, 9).slip_ids.filtered(lambda s: s.employee_id == e['cese']),
          'el cesado ya no tiene boleta en setiembre')


if LISTO:
    caso('al_hr_pe_benefits', 'Liquidación por cese (10/07/2026): CTS y vacaciones truncas',
         'Se registra el cese en el contrato, se crea la liquidación sobre el lote de julio, '
         'se pulsa «Procesar» y «Exportar a boletas».')
    correr(liquidacion)
    caso('al_hr_pe', 'Planillas ago → oct 2026',
         'Se generan los lotes de agosto y setiembre (validados) y octubre (en borrador).')
    correr(planillas_ago_oct)
    LISTO = (2026, 10) in S['lotes']


# ======================================================================
# al_hr_pe — parámetros nativos y cálculo de la boleta
# ======================================================================
def parametros_nativos():
    P = env['hr.rule.parameter']
    Val = env['hr.rule.parameter.value']
    hist = {
        'l10n_pe_uit': {'2024-01-01': 5150.0, '2025-01-01': 5350.0, '2026-01-01': 5500.0},
        'l10n_pe_rmv': {'2022-05-01': 1025.0, '2025-01-01': 1130.0, '2026-10-01': 1230.0},
        'l10n_pe_essalud_rate': {'1997-05-01': 9.0},
    }
    for code, valores in hist.items():
        param = P.search([('code', '=', code)])
        check(len(param) == 1, 'existe el parámetro de reglas %s' % code, param.name)
        reales = {str(v.date_from): float(v.parameter_value) for v in Val.search([('code', '=', code)])}
        for d, v in valores.items():
            igual(reales.get(d), v, '%s vigente desde %s' % (code, d))
    for year, uit in ((2024, 5150.0), (2025, 5350.0), (2026, 5500.0)):
        igual(P._l10n_pe_get_uit(year), uit, '_l10n_pe_get_uit(%s)' % year)
    for d, rmv in ((date(2024, 6, 30), 1025.0), (date(2026, 9, 30), 1130.0),
                   (date(2026, 10, 1), 1230.0), (date(2022, 1, 1), 0.0)):
        igual(P._l10n_pe_get_rmv(d), rmv, '_l10n_pe_get_rmv(%s)' % d)
    try:
        P._l10n_pe_get_uit(2031)
        check(False, '_l10n_pe_get_uit(2031) debe lanzar UserError')
    except UserError as exc:
        check('2031' in str(exc) and 'UIT' in str(exc), 'sin UIT 2031 → UserError claro', str(exc))


def planilla_mensual():
    e = S['emp']
    rmv_set = env['hr.rule.parameter']._l10n_pe_get_rmv(date(2026, 9, 30))
    for key in ('afp', 'onp', 'rmv'):
        emp = e[key]
        slip = boleta(emp, 2026, 9)
        w = emp.version_id.wage
        memb = emp.version_id.membership_id
        igual(slip.rmv, 1130.0, '%s set-2026: RMV de la boleta = 1 130 (parámetro al 30/09)' % key)
        igual(linea(slip, 'BAS'), w, '%s: básico mes completo = sueldo' % key)
        afecto = w
        if memb.is_afp:
            igual(linea(slip, 'AAFP'), afecto, '%s: afecto AFP = básico' % key)
            igual(abs(linea(slip, 'A_JUB')), r2(afecto * memb.retirement_fund / 100),
                  '%s: fondo AFP = afecto × %.2f %%' % (key, memb.retirement_fund))
            igual(abs(linea(slip, 'COMFI')), r2(afecto * memb.fixed_commision / 100),
                  '%s: comisión flujo = afecto × %.2f %%' % (key, memb.fixed_commision))
            igual(abs(linea(slip, 'SEGI')), r2(min(afecto, memb.insurable_remuneration) * memb.prima_insurance / 100),
                  '%s: prima de seguro = min(afecto, tope %.2f) × %.2f %%' % (key, memb.insurable_remuneration, memb.prima_insurance))
            igual(linea(slip, 'ONP'), 0.0, '%s: sin ONP' % key)
            aportes = r2(afecto * memb.retirement_fund / 100) + r2(afecto * memb.fixed_commision / 100) \
                + r2(min(afecto, memb.insurable_remuneration) * memb.prima_insurance / 100)
        else:
            igual(abs(linea(slip, 'ONP')), r2(afecto * memb.retirement_fund / 100), '%s: ONP = afecto × 13 %%' % key)
            igual(linea(slip, 'A_JUB'), 0.0, '%s: sin aporte AFP' % key)
            aportes = r2(afecto * memb.retirement_fund / 100)
        igual(linea(slip, 'ESSALUD'), r2(max(afecto, rmv_set) * 9 / 100),
              '%s: EsSalud = max(afecto %.2f, RMV %.2f) × 9 %%' % (key, afecto, rmv_set))
        igual(linea(slip, 'NETO'), r2(w - aportes), '%s: neto = ingresos − aportes del trabajador' % key)
        registro(slip, '%s set-2026' % key)
    # Octubre: la RMV sube a 1 230 (D.S. 015-2026-TR) solo desde ese periodo.
    oct_ = boleta(e['rmv'], 2026, 10)
    igual(oct_.rmv, 1230.0, 'rmv oct-2026: RMV = 1 230')
    igual(oct_.family_allowance, 123.0, 'rmv oct-2026: asignación familiar = 10 % RMV')
    igual(linea(oct_, 'ESSALUD'), r2(1230 * 0.09), 'rmv oct-2026: EsSalud sobre la RMV nueva (1 230 × 9 %)')
    igual(boleta(e['rmv'], 2026, 9).rmv, 1130.0, 'la boleta validada de setiembre conserva RMV 1 130')


def rmv_futura():
    """RMV hipotética desde 01/01/2027 (se deshace al final)."""
    e = S['emp']['rmv']
    res = {}
    try:
        with env.cr.savepoint():
            env['hr.rule.parameter.value'].create({
                'rule_parameter_id': env.ref('al_hr_pe.rule_parameter_l10n_pe_rmv').id,
                'date_from': date(2027, 1, 1), 'parameter_value': '1300'})
            for y, m in ((2026, 12), (2027, 1)):
                s = Payslip.create({'name': qa_name('RMV futura %s-%s' % (y, m)), 'employee_id': e.id,
                                    'struct_id': STRUCT.id, 'date_from': date(y, m, 1),
                                    'date_to': date(y, m, monthrange(y, m)[1])})
                s.compute_sheet()
                res[(y, m)] = (s.rmv, linea(s, 'ESSALUD'))
            boleta(e, 2026, 10).invalidate_recordset(['rmv'])
            res['oct'] = boleta(e, 2026, 10).rmv
            raise _Deshacer()
    except _Deshacer:
        pass
    igual(res[(2026, 12)][0], 1230.0, 'dic-2026 (antes de la vigencia) sigue con RMV 1 230')
    igual(res[(2027, 1)][0], 1300.0, 'ene-2027 toma la RMV futura 1 300')
    igual(res[(2027, 1)][1], r2(1300 * 0.09), 'ene-2027: EsSalud = 1 300 × 9 % (sueldo 1 000 < RMV)')
    igual(res['oct'], 1230.0, 'la boleta de oct-2026 ya calculada no cambia')
    aviso('valor RMV 2027 creado dentro de un savepoint y deshecho', 'no queda en la base')


if LISTO:
    caso('al_hr_pe', 'Parámetros nativos UIT / RMV / EsSalud con su historial',
         'El usuario abre Nómina ▸ Configuración ▸ Parámetros de reglas y revisa los valores por fecha.')
    correr(parametros_nativos)
    caso('al_hr_pe', 'Boleta mensual: AFP, ONP y sueldo menor a la RMV',
         'El usuario abre las boletas de setiembre y octubre 2026 y revisa cada concepto.')
    correr(planilla_mensual)
    caso('al_hr_pe', 'RMV con fecha futura solo afecta a las planillas posteriores',
         'Se registra una RMV vigente desde 01/01/2027 y se calculan boletas de dic-2026 y ene-2027.')
    correr(rmv_futura)


# ======================================================================
# al_hr_pe_benefits — vacaciones, record vacacional, renta de 5.ª
# ======================================================================
def vacaciones():
    emp = S['emp']['onp']
    Rest = env['hr.vacation.rest']
    Rest.get_vacation_employee(emp, False)
    rec = Rest.search([('employee_id', '=', emp.id), ('date_from', '=', date(2025, 1, 1))])
    if check(rec, 'récord vacacional del año 2025 generado'):
        igual(rec[0].days, 30.0, 'récord 2025: 2.5 días × 12 meses = 30 días')
        registro(rec[0], 'récord vacacional 2025')
    vac = env['hr.vacation'].create({'company_id': company.id, 'year': 2026,
                                     'payslip_run_id': lote(2026, 8).id})
    ln = env['hr.vacation.line'].create({'vacation_id': vac.id, 'employee_id': emp.id,
                                         'version_id': emp.version_id.id,
                                         'wage': emp.version_id.wage, 'accrued_vacation': 15})
    ln.with_context(line_form=True).action_compute()
    registro(vac, 'vacaciones agosto 2026')
    w = emp.version_id.wage
    igual(ln.total_vacation, r2(w / 30 * 15), 'vacaciones 15 días = sueldo/30 × 15')
    igual(ln.onp, r2(ln.total_vacation * 0.13), 'ONP 13 % de la remuneración vacacional')
    igual(ln.neto_total, r2(ln.total_vacation * 0.87), 'neto vacacional = total − ONP')


def quinta():
    emp = S['emp']['afp']
    run = lote(2026, 10)
    fifth = env['hr.fifth.category'].create({'payslip_run_id': run.id, 'company_id': company.id})
    fifth.action_generate()
    registro(fifth, 'renta de 5.ª octubre 2026')
    S['fifth'] = fifth
    uit = env['hr.rule.parameter']._l10n_pe_get_uit(2026)
    ln = fifth.line_ids.filtered(lambda l: l.employee_id == emp)
    if not check(ln, 'Ana (6 000) queda afecta a 5.ª'):
        return
    w = emp.version_id.wage
    grat = S['grati_jul'].line_ids.filtered(lambda l: l.employee_id == emp)
    # Proyección anual: ene-set reales (9 × w) + oct + nov-dic (2 × w)
    # + gratificación de julio real (con bono) + diciembre proyectada × 1.09.
    total = 9 * w + w + 2 * w + (grat.total_grat + grat.bonus_essalud) + w * 1.09
    igual(ln.seven_uit, 7 * uit, '7 UIT = 7 × %.0f (UIT 2026 del parámetro)' % uit)
    igual(ln.past_rem, 9 * w, 'remuneraciones ene-set (boletas validadas)')
    igual(ln.total_proy, total, 'renta bruta proyectada = 12 sueldos + 2 gratificaciones con bono')
    igual(ln.net_rent, total - 7 * uit, 'renta neta = bruta − 7 UIT')
    tax = impuesto_5ta(total - 7 * uit, uit)
    igual(ln.tax_proy, tax, 'impuesto anual por tramos 8/14/17/20/30 % (calculado a mano)')
    igual(ln.monthly_ret, r2((tax - ln.past_months_ret) / 4),
          'retención de octubre = (impuesto − retenido) / 4 (art. 40)')
    excl = fifth.line_excluidos_ids.mapped('employee_id')
    check(S['emp']['onp'] in excl and S['emp']['rmv'] in excl,
          'Bruno (2 500) y Carla (1 000) pasan a excluidos', ', '.join(excl.mapped('name')))
    fifth.action_export_to_payslips()
    slip = boleta(emp, 2026, 10)
    slip.compute_sheet()
    igual(abs(linea(slip, 'QUINTA')), ln.monthly_ret + ln.ext_ret, 'la boleta de octubre descuenta la quinta exportada')
    S['quinta_oct'] = ln.monthly_ret
    # Mensaje claro si falta la UIT del año
    try:
        PARAM.action_generate_brackets(year=2031)
        check(False, 'generar tramos 2031 sin UIT debe fallar')
    except UserError as exc:
        check('UIT' in str(exc) and '2031' in str(exc), 'tramos 2031 sin UIT → mensaje claro', str(exc))
    PARAM.action_generate_brackets(year=2026)


if LISTO:
    caso('al_hr_pe_benefits', 'Vacaciones: récord vacacional y goce de 15 días',
         'El usuario calcula el récord vacacional de Bruno y registra 15 días de vacaciones en el lote de agosto.')
    correr(vacaciones)
    caso('al_hr_pe_benefits', 'Renta de 5.ª de octubre con la UIT 2026',
         'El usuario crea la quinta del lote de octubre, pulsa «Generar» y «Exportar a boletas».')
    correr(quinta)


# ======================================================================
# al_hr_pe_import — inputs de boleta desde Excel
# ======================================================================
def importar_inputs():
    run = lote(2026, 10)
    e = S['emp']
    wiz = env['al.import.payslip.input.wizard'].create({
        'file_data': xlsx({'INPUTS': [
            ['NRO DOCUMENTO', 'CÓDIGO DE INPUT', 'MONTO'],
            [e['onp'].identification_id, 'BONR', 300.0],
            [e['rmv'].identification_id, 'BONR', '150,50'],
            ['99999999', 'BONR', 10.0],
            [e['afp'].identification_id, 'NOEXISTE', 5.0],
        ]}),
        'file_name': 'inputs_octubre.xlsx', 'company_id': company.id, 'payslip_run_id': run.id})
    wiz.action_load_file()
    results, counts, log, ids = importar(wiz)
    igual(counts.get('created'), 2, 'filas importadas (Bruno 300, Carla 150.50)')
    igual(counts.get('error'), 2, 'filas con error (DNI sin boleta, input inexistente)', 0)
    check(all(r['fix'] for r in results if r['status'] == 'error'), 'cada error trae su sugerencia',
          ' | '.join(log))
    for key, monto in (('onp', 300.0), ('rmv', 150.5)):
        slip = boleta(e[key], 2026, 10)
        slip.compute_sheet()
        igual(linea(slip, 'BONR'), monto, '%s: línea BONR = input importado' % key)
    s = boleta(e['onp'], 2026, 10)
    igual(abs(linea(s, 'ONP')), r2((2500 + 300) * 0.13), 'onp: ONP sobre sueldo + bonificación regular')
    s = boleta(e['rmv'], 2026, 10)
    igual(linea(s, 'ESSALUD'), r2(max(1150.5, 1230) * 0.09), 'rmv: EsSalud sigue con base RMV (1 150.50 < 1 230)')


def cerrar_octubre():
    run = lote(2026, 10)
    run.slip_ids.compute_sheet()
    run.slip_ids.action_payslip_done()
    check(run.state != '01_ready', 'lote de octubre validado', run.state)


if LISTO:
    caso('al_hr_pe_import', 'Importar inputs de boleta desde Excel',
         'El usuario sube un Excel con DNI, código de input y monto al lote de octubre y lo importa.')
    correr(importar_inputs)
    caso('al_hr_pe', 'Cierre del lote de octubre',
         'Se recalculan y validan las boletas de octubre (con quinta e inputs importados).')
    correr(cerrar_octubre)


# ======================================================================
# al_hr_pe_account — asiento de planilla por lote
# ======================================================================
def asiento():
    A = env['account.account'].with_company(company)

    def cta(code):
        return A.search([('code', '=', code)], limit=1)

    c = {k: cta(v) for k, v in {'gasto': '6211000', 'neto': '4111000', 'onp': '4032000',
                               'afp': '4170000', 'ess_d': '6271000', 'ess_c': '4031000'}.items()}
    check(all(c.values()), 'cuentas PCGE de planilla existen', ', '.join(a.code for a in c.values() if a))
    cambios = []
    for code, debit, credit in (('BAS', c['gasto'], None), ('NETO', None, c['neto']),
                                ('ONP', None, c['onp']), ('ESSALUD', c['ess_d'], c['ess_c'])):
        r = regla(code)
        if debit and not r.account_debit:
            r.account_debit = debit
            cambios.append('%s debe %s' % (code, debit.code))
        if credit and not r.account_credit:
            r.account_credit = credit
            cambios.append('%s haber %s' % (code, credit.code))
    for memb in S['emp']['afp'].version_id.membership_id | S['emp']['rmv'].version_id.membership_id \
            | S['emp']['cese'].version_id.membership_id:
        if not memb.with_company(company).account_id:
            memb.with_company(company).account_id = c['afp']
            cambios.append('%s → %s' % (memb.name, c['afp'].code))
    # Diario y partner propios: evita bloqueos con otros agentes que tocan
    # el diario SLR o el partner de la compañía en paralelo.
    journal = env['account.journal'].create({'name': qa_name('Planillas'), 'code': 'QPLA', 'type': 'general',
                                             'company_id': company.id})
    partner = env['res.partner'].create({'name': qa_name('Planilla OLANO IT (contrapartida)'),
                                         'company_id': company.id})
    registro(journal, 'diario del asiento de planilla')
    registro(partner, 'partner del asiento de planilla')
    cambios += set_if_empty(PARAM, {
        'move_journal_id': journal.id, 'move_partner_id': partner.id,
        'afp_rule_ids': [Command.set((regla('A_JUB') | regla('COMFI') | regla('COMMIX') | regla('SEGI')).ids)],
    })
    aviso('configuración contable completada (solo campos vacíos)', '; '.join(cambios) or 'ninguno')
    run = lote(2026, 9)
    run._pe_generate_batch_move(adjust_account=cta('6599000') or c['gasto'])
    move = run.move_id
    if not check(move, 'el lote de setiembre tiene asiento'):
        return
    registro(move, 'asiento de planilla set-2026')
    check(move.state == 'posted', 'asiento publicado', move.state)
    igual(move.ref, 'PLA092026', 'referencia PLA + MM + AAAA')
    igual(sum(move.line_ids.mapped('debit')), sum(move.line_ids.mapped('credit')), 'asiento cuadrado')
    slips = run.slip_ids

    def suma(cuenta, campo):
        return sum(move.line_ids.filtered(lambda l: l.account_id == cuenta).mapped(campo))

    def tot(code):
        return sum(abs(linea(s, code)) for s in slips)

    igual(suma(c['gasto'], 'debit'), tot('BAS'), '6211 debe = Σ básicos del lote')
    igual(suma(c['neto'], 'credit'), tot('NETO'), '4111 haber = Σ netos')
    igual(suma(c['onp'], 'credit'), tot('ONP'), '4032 haber = Σ ONP')
    igual(suma(c['afp'], 'credit'), tot('A_JUB') + tot('COMFI') + tot('SEGI'),
          '4170 haber = Σ fondo + comisión + prima AFP (cuenta de la afiliación)')
    igual(suma(c['ess_d'], 'debit'), tot('ESSALUD'), '6271 debe = Σ EsSalud')
    igual(suma(c['ess_c'], 'credit'), tot('ESSALUD'), '4031 haber = Σ EsSalud')
    ajuste = move.line_ids.filtered(lambda l: 'redondeo' in (l.name or '').lower())
    check(not ajuste, 'sin línea de ajuste por redondeo (el lote cuadra por construcción)')


if LISTO:
    caso('al_hr_pe_account', 'Asiento de planilla del lote de setiembre',
         'El contador configura las cuentas de las reglas y pulsa «Asiento de planilla por lote».')
    correr(asiento)


# ======================================================================
# al_hr_pe_reports — boleta PDF, PLAME, AFPnet, TXT bancario
# ======================================================================
def reportes():
    run = lote(2026, 10)
    e = S['emp']
    slip = boleta(e['afp'], 2026, 10)
    Report = env['ir.actions.report']
    html = str(Report._render_qweb_html('al_hr_pe_reports.action_report_boleta_pago', slip.ids)[0])
    check('pe-doc-header' in html, 'la boleta usa la cabecera común de al_report_base')
    check(e['afp'].identification_id in html and company.vat in html, 'la boleta muestra DNI y RUC')
    pdf, fmt = Report._render_qweb_pdf('al_hr_pe_reports.action_report_boleta_pago', slip.ids)
    check(fmt == 'pdf' and pdf[:4] == b'%PDF', 'boleta PDF renderizada', '%d bytes' % len(pdf))
    att = env['ir.attachment'].create({'name': qa_name('Boleta oct-2026 Ana.pdf'), 'raw': pdf,
                                       'res_model': 'hr.payslip', 'res_id': slip.id})
    registro(att, 'PDF de la boleta')
    # PLAME .rem
    rem = adjunto(run.action_export_plame_rem())
    registro(rem, 'PLAME .rem')
    txt = rem.raw.decode('utf-8')
    filas = [f.split('|') for f in txt.split('\r\n') if f]
    check(all(len(f) == 6 and f[5] == '' for f in filas), 'formato tipo|doc|concepto|devengado|pagado|', '%d filas' % len(filas))
    idx = {(f[1], f[2]): f for f in filas}
    for key in ('afp', 'onp', 'rmv'):
        dni = e[key].identification_id
        f = idx.get((dni, '0121'))
        if check(f, '%s: concepto 0121 (básico) declarado' % key):
            igual(float(f[3]), linea(boleta(e[key], 2026, 10), 'BAS'), '%s: 0121 devengado = básico' % key)
            igual(f[0], '01', '%s: tipo de documento DNI = 01' % key)
    f = idx.get((e['afp'].identification_id, '0605'))
    if check(f, 'afp: 0605 (renta 5.ª) declarada'):
        igual(float(f[3]), S.get('quinta_oct', 0.0), 'afp: 0605 = retención de octubre')
    check((e['onp'].identification_id, '0601') not in idx, 'onp: sin comisión AFP 0601')
    check(not any(k[1] == '0804' for k in idx), 'EsSalud (0804) no se importa (lo calcula PLAME)')
    # AFPnet
    import openpyxl
    afpnet = adjunto(run.action_export_afpnet())
    registro(afpnet, 'AFPnet')
    ws = openpyxl.load_workbook(io.BytesIO(afpnet.raw)).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    igual(len(rows), 2, 'AFPnet: solo los 2 afiliados AFP (Ana, Carla)')
    if rows:
        by_doc = {r[3]: r for r in rows}
        igual([r[0] for r in rows], [1, 2], 'correlativo 1, 2 sin huecos')
        a = by_doc.get(e['afp'].identification_id)
        if check(a, 'fila de Ana en AFPnet'):
            igual(a[1], e['afp'].version_id.l10n_pe_cuspp, 'CUSPP')
            igual((a[4], a[5], a[6]), ('Torres', 'Ruiz', 'Ana'), 'apellidos y nombres')
            igual((a[7], a[8], a[9]), ('S', 'N', 'N'), 'devenga S, inicio N, cese N')
            igual(float(a[11]), linea(boleta(e['afp'], 2026, 10), 'AAFP'), 'remuneración asegurable = AAFP')
            igual(a[15], 'N', 'tipo de labor N')
    # TXT bancario BCP
    bcp = env['l10n_pe.bank'].search([('name', 'ilike', 'Crédito')], limit=1)
    if not bcp.format_bank:
        bcp.format_bank = 'bcp'
        aviso('banco BCP: se configuró format_bank = bcp')
    pen = company.currency_id
    cargo = env['res.partner.bank'].create({'account_number': '1917000123456', 'partner_id': company.partner_id.id,
                                            'l10n_pe_bank_id': bcp.id, 'company_id': company.id,
                                            'currency_id': pen.id, 'type_of_account': '0'})
    journal = env['account.journal'].create({'name': qa_name('Pagos planilla BCP'), 'type': 'bank',
                                             'code': 'QPBCP', 'company_id': company.id,
                                             'bank_account_id': cargo.id})
    registro(journal, 'diario de pago masivo')
    PARAM.journals_banks = [Command.link(journal.id)]
    cuentas = {}
    for i, key in enumerate(('afp', 'onp', 'rmv'), start=1):
        emp = e[key]
        partner = emp.work_contact_id
        acc = env['res.partner.bank'].create({'account_number': '191%011d' % (5000 + i), 'partner_id': partner.id,
                                              'l10n_pe_bank_id': bcp.id, 'company_id': company.id,
                                              'currency_id': pen.id, 'type_of_account': '1'})
        emp.bank_account_ids = [Command.link(acc.id)]
        cuentas[key] = acc
    run.action_generate_multipayment()
    pago = run.multipayment_ids.filtered(lambda p: p.journal_id == journal)
    if not check(pago, 'pago masivo BCP generado desde el lote'):
        return
    registro(pago, 'pago masivo BCP')
    igual(len(pago.line_ids), 3, 'una línea por trabajador con cuenta BCP')
    for key, acc in cuentas.items():
        ln = pago.line_ids.filtered(lambda l: l.employee_id == e[key])
        igual(ln.amount, linea(boleta(e[key], 2026, 10), 'NETO'), '%s: monto = NETO peruano' % key)
    txt = adjunto(pago.action_generate_txt())
    registro(txt, 'TXT haberes BCP')
    rows = txt.raw.decode('utf-8').split('\r\n')[:-1]
    igual(len(rows), 4, 'TXT BCP: cabecera + 3 detalles')
    if rows:
        head = rows[0]
        total = sum(pago.line_ids.mapped('amount'))
        igual(head[0], '1', 'cabecera tipo 1')
        igual(head[1:7], '000003', 'cabecera: 3 abonos')
        igual(head[16], 'C', 'tipo de cargo C')
        igual(head[17:21], '0001', 'moneda soles 0001')
        igual(head[21:41].strip(), cargo.account_number, 'cuenta de cargo')
        igual(float(head[41:58]), r2(total), 'total de la planilla en la cabecera')
        chk = sum(int(a.account_number[3:]) for a in cuentas.values()) + int(cargo.account_number[3:])
        igual(head[-15:], str(chk).rjust(15, '0'), 'checksum = Σ cuentas (desde la 4.ª posición) + cargo')
        det = [r for r in rows[1:] if e['afp'].identification_id in r]
        if check(det, 'detalle de Ana'):
            d = det[0]
            igual(d[2:22].strip(), cuentas['afp'].account_number, 'detalle: cuenta destino')
            igual(float(d[-18:-1]), linea(boleta(e['afp'], 2026, 10), 'NETO'), 'detalle: importe = neto')


if LISTO:
    caso('al_hr_pe_reports', 'Boleta PDF, PLAME .rem, AFPnet y TXT de haberes BCP del lote de octubre',
         'El usuario imprime la boleta, exporta PLAME y AFPnet, genera el pago masivo y descarga el TXT del BCP.')
    correr(reportes)


# ======================================================================
# al_hr_pe_public_holidays — feriados nativos vs catálogo propio
# ======================================================================
def feriados():
    H = env['pe.public.holiday']
    Leaves = env['resource.calendar.leaves']
    nac = H.search([('year', '=', 2026), ('labor_regime', '=', False)])
    igual(len(nac), 16, 'catálogo 2026: 16 feriados nacionales')
    sin_tipo = H.search([('year', '>=', 2026), ('work_entry_type_id', '=', False)])
    igual(len(sin_tipo), 0, 'feriados del catálogo 2026-2035 sin concepto de nómina (debería ser DOM)')
    cal = env['resource.calendar'].create({'name': qa_name('Calendario feriados'), 'company_id': company.id})
    registro(cal, 'calendario de prueba')
    previos = Leaves.search([('pe_public_holiday_id', '!=', False)])
    nac.action_apply_to_calendars()
    nac.action_apply_to_calendars()  # idempotente
    leaves = Leaves.search([('calendar_id', '=', cal.id), ('pe_public_holiday_id', 'in', nac.ids)])
    igual(len(leaves), 16, 'aplicar dos veces deja 16 descansos (sin duplicar)')
    fp = leaves.filtered(lambda l: l.pe_public_holiday_id.date == date(2026, 7, 28))
    igual(str(fp.date_from), '2026-07-28 05:00:00', '28/07 00:00 Lima = 05:00 UTC')
    igual(fp.work_entry_type_id.code, 'DOM', 'el descanso usa el concepto DOM de la planilla')
    Lc = Leaves.with_company(company).with_context(allowed_company_ids=company.ids)
    prep = Lc._prepare_public_holidays_data(date(2026, 1, 1), date(2026, 12, 31))
    nativos = prep['prepared_public_holidays'].get(company.id, [])
    cubiertos = set(H.search([('year', '=', 2026)]).mapped('date'))
    duplicados = [v for v in nativos if v['date_from'].date() in cubiertos]
    check(not duplicados, 'Odoo 20 (asistente nativo) no propone días que ya cubre el catálogo 2026',
          '%d propuestos: %s' % (len(nativos), ', '.join(v['name'] for v in nativos)))
    prep = Lc._prepare_public_holidays_data(date(2036, 1, 1), date(2036, 12, 31))
    nativos = prep['prepared_public_holidays'].get(company.id, [])
    check(H.search_count([('year', '=', 2036)]) == 0 and nativos,
          '2036 (fuera del catálogo): el nativo propone los feriados del CSV', '%d días' % len(nativos))
    t = env['hr.tareaje.manager'].create({'name': qa_name('Tareaje feriados'), 'date_start': date(2026, 10, 1),
                                          'date_end': date(2026, 10, 31)})
    hol = t._get_public_holidays()
    check(date(2026, 10, 8) in hol.get(cal.id, set()), 'el tareaje reconoce el 08/10 como feriado del calendario')
    t.unlink()
    # Efecto en la boleta: Elena (calendario QA20) en octubre, con el feriado
    # de Angamos (jueves 08/10) no laborado: se paga el mes completo.
    emp = S.get('emp_att')
    if emp:
        try:
            with env.cr.savepoint():
                slip = Payslip.create({'name': qa_name('Boleta oct-2026 Elena (feriado 08/10)'), 'employee_id': emp.id,
                                       'struct_id': STRUCT.id, 'date_from': date(2026, 10, 1),
                                       'date_to': date(2026, 10, 31)})
                slip.compute_sheet()
            registro(slip, 'boleta con feriado (borrador)')
            wd = {l.code: l.number_of_days for l in slip.worked_days_line_ids if l.number_of_days}
            aviso('días de la boleta oct-2026 con feriado', wd)
            igual(wd.get('DLAB', 0) + wd.get('DOM', 0), 31.0, 'DLAB + DOM = 31 días (el feriado no laborado es descanso)')
            igual(linea(slip, 'BAS'), 3000.0, 'el feriado no laborado no descuenta el básico')
        except Exception as exc:  # noqa: BLE001
            check(False, 'calcular la boleta de octubre de un trabajador con el feriado aplicado',
                  '%s: %s' % (type(exc).__name__, str(exc).splitlines()[0]))
            env.invalidate_all()
    # No dejar descansos en calendarios ajenos (los de la compañía los aplica su cron).
    mios = env['resource.calendar'].search([('name', '=like', 'QA20%')])
    ajenos = Leaves.search([('pe_public_holiday_id', '!=', False), ('id', 'not in', previos.ids),
                            ('calendar_id', 'not in', mios.ids)])
    ajenos.unlink()
    aviso('descansos retirados de calendarios no QA20 tras la prueba', len(ajenos))


# ======================================================================
# al_hr_pe_attendance (+ import de asistencias): turno, tardanza, HE 25/35
# ======================================================================
def asistencia():
    cal = env['resource.calendar'].create({
        'name': qa_name('Turno diurno 08-17'), 'company_id': company.id,
        'attendance_ids': [Command.clear()] + [
            Command.create({'dayofweek': str(d), 'hour_from': a, 'hour_to': b, 'day_period': p})
            for d in range(5) for a, b, p in ((8.0, 13.0, 'morning'), (14.0, 17.0, 'afternoon'))]})
    emp = env['hr.employee'].create({
        'name': qa_name('Quispe Huamán Elena'), 'names': 'Elena', 'last_name': 'Quispe',
        'm_last_name': 'Huamán', 'company_id': company.id, 'tz': 'America/Lima',
        'l10n_pe_document_type_id': DNI.id, 'identification_id': '79000105',
        'resource_calendar_id': cal.id, 'date_version': date(2026, 1, 1),
        'contract_date_start': date(2026, 1, 1), 'wage': 3000.0, 'structure_type_id': STRUCT.type_id.id})
    emp.version_id.write({'membership_id': env.ref('al_hr_pe.membership_ONP').id,
                          'social_insurance_id': S['essalud'].id, 'l10n_pe_is_overtime': True,
                          'resource_calendar_id': cal.id})
    registro(emp, 'trabajadora con marcación, sujeta a HE')
    S['emp_att'] = emp
    # Turno atípico (ciclo) y su tope legal de 48 h semanales
    ciclo = env['l10n_pe.hr.shift.cycle'].create({'name': qa_name('Ciclo 4x3'), 'days_work': 4,
                                                  'days_rest': 3, 'hours_per_day': 12.0})
    registro(ciclo, 'turno atípico 4×3 de 12 h')
    igual(ciclo.avg_weekly_hours, 48.0, 'promedio semanal 12 h × 4 × 7/7 = 48 h')
    try:
        with env.cr.savepoint():
            env['l10n_pe.hr.shift.cycle'].create({'name': qa_name('Ciclo 6x1'), 'days_work': 6,
                                                  'days_rest': 1, 'hours_per_day': 10.0})
        check(False, 'ciclo 6×1 de 10 h (60 h/sem) debe rechazarse')
    except ValidationError as exc:
        check('48' in str(exc), 'ciclo 6×1 de 10 h rechazado por superar 48 h', str(exc))
    # Excel de marcaciones de setiembre: todos los días hábiles 08:00-17:00 salvo
    # el 08 (llega 08:40), el 09 (sale 20:00) y el 10 (sale 18:30).
    filas = [['EMPLEADO', 'ENTRADA', 'SALIDA']]
    d = date(2026, 9, 1)
    especiales = {date(2026, 9, 8): ('08:40', '17:00'), date(2026, 9, 9): ('08:00', '20:00'),
                  date(2026, 9, 10): ('08:00', '18:30')}
    while d.month == 9:
        if d.weekday() < 5:
            ent, sal = especiales.get(d, ('08:00', '17:00'))
            filas.append([emp.name, '%s %s:00' % (d, ent), '%s %s:00' % (d, sal)])
        d += timedelta(days=1)
    wiz = env['al.import.hr.attendance.wizard'].create({
        'file_data': xlsx({'ASISTENCIAS': filas}), 'file_name': 'marcaciones_set2026.xlsx',
        'company_id': company.id, 'tz': 'America/Lima'})
    wiz.action_load_file()
    results, counts, log, ids = importar(wiz)
    igual(counts.get('created'), len(filas) - 1, 'al_hr_pe_import: marcaciones creadas desde Excel')
    atts = env['hr.attendance'].browse(ids)
    a9 = atts.filtered(lambda a: a.check_in == datetime(2026, 9, 9, 13, 0))
    if check(a9, 'marcación del 09/09 (08:00 Lima = 13:00 UTC)'):
        dia = a9 | a9.overtime_attendance_ids
        igual(str(max(dia.mapped('check_out'))), '2026-09-10 01:00:00',
              'el día termina 20:00 Lima aunque el sobretiempo vaya en asistencia separada')
        aviso('asistencias del 09/09', '%d registro(s): %s' % (len(dia), ', '.join(
            '%s→%s' % (x.check_in, x.check_out) for x in dia.sorted('check_in'))))
    # Tareaje de setiembre
    t = env['hr.tareaje.manager'].create({'name': qa_name('Tareaje setiembre 2026'), 'date_start': date(2026, 9, 1),
                                          'date_end': date(2026, 9, 30), 'is_compute_he': True,
                                          'time_tolerancia': 0.25})
    t.action_generate()
    registro(t, 'tareaje setiembre')
    ln = t.tareaje_line_ids.filtered(lambda l: l.employee_id == emp)
    if not check(ln, 'línea de tareaje de Elena'):
        return
    dias = {x.fecha: x for x in ln.attendance_line_ids}
    igual(dias[date(2026, 9, 8)].tar, r2(40 / 60), '08/09: tardanza 40 min (> 15 min de tolerancia) = 0.67 h')
    igual(dias[date(2026, 9, 9)].he25, 2.0, '09/09: 3 h extra → 2 h al 25 %')
    igual(dias[date(2026, 9, 9)].he35, 1.0, '09/09: … y 1 h al 35 %')
    igual(dias[date(2026, 9, 10)].he25, 1.5, '10/09: 1.5 h al 25 %')
    igual(ln.dlab, 22.0, 'días laborados = 22 días hábiles de setiembre')
    igual(ln.he25, 3.5, 'total HE 25 % = 3.5 h')
    igual(ln.he35, 1.0, 'total HE 35 % = 1 h')
    t.action_close()
    slip = Payslip.create({'name': qa_name('Boleta set-2026 Elena'), 'employee_id': emp.id, 'struct_id': STRUCT.id,
                           'date_from': date(2026, 9, 1), 'date_to': date(2026, 9, 30)})
    slip.compute_sheet()
    registro(slip, 'boleta con tareaje aplicado')
    hora = 3000 / 30 / 8
    igual(linea(slip, 'HE25'), hora * 1.25 * 3.5, 'HE25 = sueldo/30/8 × 1.25 × 3.5 h')
    igual(linea(slip, 'HE35'), hora * 1.35 * 1.0, 'HE35 = sueldo/30/8 × 1.35 × 1 h')
    tar_h = slip.worked_days_line_ids.filtered(lambda w: w.code == 'TAR').number_of_hours
    igual(linea(slip, 'TAR'), tar_h * hora, 'TAR = horas de tardanza × sueldo/30/8 (%.2f h)' % tar_h)
    igual(linea(slip, 'BAS'), 3000.0, 'básico completo (22 DLAB + 8 DOM = 30 días)')
    # Registro de control de asistencia (D.S. 004-2006-TR)
    reg = env['l10n_pe.hr.attendance.register.wizard'].create({
        'date_from': date(2026, 9, 1), 'date_to': date(2026, 9, 30), 'employee_ids': [Command.set(emp.ids)]})
    filas_r = {r[2]: r for r in reg._register_rows()}
    r9 = filas_r.get(date(2026, 9, 9))
    if check(r9, 'registro de control: fila del 09/09'):
        igual((r9[3], r9[4], r9[5], r9[6]), ('08:00', '20:00', '17:00', '20:00'),
              'ingreso, salida e inicio/fin del sobretiempo')


if LISTO:
    caso('al_hr_pe_attendance', 'Turno, marcaciones importadas, tardanza y horas extra 25 %/35 %',
         'Se define el turno, se importa el Excel de marcaciones de setiembre, se genera y aplica el '
         'tareaje y se calcula la boleta.')
    correr(asistencia)


if LISTO:
    caso('al_hr_pe_public_holidays', 'Feriados 2026 en calendarios sin duplicar los nativos de Odoo 20',
         'El usuario aplica los feriados del catálogo a los calendarios y abre el asistente nativo «Cargar feriados».')
    correr(feriados)


# ======================================================================
# al_hr_pe_construction — BUC, movilidad, CONAFOVICER
# ======================================================================
def construccion():
    struct = env.ref('al_hr_pe_construction.construction_structure')
    cat = env.ref('al_hr_pe_construction.category_operario')
    env['hr.period.generator'].create({'year': 2026, 'company_id': company.id, 'generate_weekly': True}).action_generate()
    emp = env['hr.employee'].create({
        'name': qa_name('Mamani Condori Julio'), 'names': 'Julio', 'last_name': 'Mamani',
        'm_last_name': 'Condori', 'company_id': company.id, 'l10n_pe_document_type_id': DNI.id,
        'identification_id': '79000106', 'date_version': date(2026, 1, 1),
        'contract_date_start': date(2026, 1, 1), 'wage': 0.0})
    emp.version_id.write({'l10n_pe_labor_regime': 'construccion', 'l10n_pe_construction_category_id': cat.id,
                          'membership_id': env.ref('al_hr_pe.membership_ONP').id})
    registro(emp, 'operario de construcción civil')
    week = env['hr.period'].search([('company_id', '=', company.id), ('period_type', '=', 'weekly'),
                                    ('date_start', '=', date(2026, 10, 12))], limit=1)
    check(week and week.date_end == date(2026, 10, 18), 'semana 12-18/10/2026 generada', week.code)
    slip = Payslip.create({'name': qa_name('Semana 12-18/10 Julio'), 'employee_id': emp.id, 'company_id': company.id,
                           'struct_id': struct.id, 'date_from': date(2026, 10, 12), 'date_to': date(2026, 10, 18)})
    slip.worked_days_line_ids.unlink()
    slip.worked_days_line_ids = [Command.create({
        'name': 'Días laborados', 'version_id': slip.version_id.id,
        'work_entry_type_id': env.ref('al_hr_pe.wd_DLAB').id,
        'number_of_days': 6, 'number_of_hours': 48, 'amount': 0.0})]
    slip.compute_sheet()
    registro(slip, 'boleta semanal 6 días')
    tabla = env['l10n_pe.hr.construction.wage.line'].search([('category_id', '=', cat.id),
                                                                   ('table_id.date_from', '<=', date(2026, 10, 12)),
                                                                   ('table_id.date_to', '>=', date(2026, 10, 18))], limit=1)
    jornal = tabla.daily_wage
    buc = tabla.buc_percent or cat.buc_percent
    mov = tabla.mobility_amount
    igual(slip.l10n_pe_daily_wage, jornal, 'jornal diario de la tabla del convenio (%.2f)' % jornal)
    igual(linea(slip, 'JOR'), r2(jornal * 6), 'JOR = jornal × 6 días')
    igual(linea(slip, 'DSO'), r2(jornal), 'DSO = 1 jornal por 6 días trabajados')
    igual(linea(slip, 'BUC'), r2(jornal * 6 * buc / 100), 'BUC = jornal × 6 × %.0f %%' % buc)
    igual(linea(slip, 'MOV'), r2(mov * 6), 'movilidad = %.2f × 6 días' % mov)
    rate = company.l10n_pe_conafovicer_rate
    igual(linea(slip, 'CONAF'), -r2((r2(jornal * 6) + r2(jornal)) * rate / 100),
          'CONAFOVICER = −(JOR + DSO) × %.1f %%' % rate)
    igual(linea(slip, 'TREM'), r2(jornal * 6) + r2(jornal) + r2(jornal * 6 * buc / 100),
          'afecto (TREM) = JOR + DSO + BUC; la movilidad no está afecta')
    igual(abs(linea(slip, 'ONP')), r2(linea(slip, 'TREM') * 0.13), 'ONP 13 % del afecto')


if LISTO:
    caso('al_hr_pe_construction', 'Boleta semanal de operario: BUC, movilidad y CONAFOVICER',
         'Se registra al operario, se genera la semana del 12 al 18/10 con 6 días laborados y se calcula.')
    correr(construccion)


# ======================================================================
# OL-TOOLS
# ======================================================================
def usuario_qa(login, nombre, grupos):
    U = env['res.users'].with_context(no_reset_password=True)
    user = U.search([('login', '=', login)], limit=1)
    if not user:
        user = U.create({'name': qa_name(nombre), 'login': login, 'company_id': company.id,
                         'company_ids': [Command.set(company.ids)],
                         'group_ids': [Command.set([env.ref(g).id for g in grupos])]})
    return user


def docs_help():
    Help = env['al.docs.help']
    total = Help.search_count([])
    check(total > 100, 'entradas de ayuda cargadas', total)
    rmv = Help.search([('key', '=', 'hr.rule.parameter.value')])
    check(rmv and 'UIT' in rmv.name and 'RMV' in rmv.name.upper().replace('REMUNERACIÓN MÍNIMA VITAL', 'RMV'),
          'entrada UIT/RMV apunta al modelo nativo hr.rule.parameter.value', rmv.name)
    for viejo in ('l10n_pe.hr.uit', 'hr.uit', 'l10n_pe.hr.rmv', 'hr.rmv'):
        check(not Help.search_count([('key', '=', viejo)]), 'sin entrada al modelo eliminado %s' % viejo)
    user = usuario_qa('qa20_docs_user', 'Usuario ayuda', ['base.group_user'])
    registro(user, 'usuario interno')
    mapa = Help.with_user(user).get_help_map()   # lo que llama el systray (orm.call get_help_map)
    check('hr.rule.parameter.value' in mapa and mapa['hr.rule.parameter.value']['url'].startswith('http'),
          'get_help_map (systray) responde al usuario interno con la URL de UIT/RMV',
          mapa.get('hr.rule.parameter.value', {}).get('url'))
    check(not Help.with_user(user).has_access('write'), 'el usuario interno no puede editar la ayuda')
    huérfanas = [h.key for h in Help.search([('kind', '=', 'model')])
                 if h.key not in env and h.key in mapa]
    check(not huérfanas, 'ninguna ayuda visible apunta a un modelo inexistente', ', '.join(huérfanas[:15]))
    act = rmv.action_open()
    check(act['type'] == 'ir.actions.act_url' and act['url'].endswith(rmv.path), 'botón «?» abre la guía', act['url'])


def report_base():
    from odoo.tools.image import image_data_uri
    html = str(env['ir.qweb']._render('al_report_base.report_pe_page', {
        'o': company, 'company': company, 'doc_title': qa_name('Documento'), 'doc_number': 'QA-0001',
        'context_timestamp': lambda t: fields.Datetime.context_timestamp(env.user, t),
        'image_data_uri': image_data_uri}))
    check('pe-doc-header' in html and 'class="footer pe-doc-footer"' in html, 'página común: cabecera y pie')
    check(company.vat in html, 'la cabecera muestra el RUC %s' % company.vat)
    check('N.º <span>QA-0001</span>' in html, 'número del documento en la cabecera')
    usan = env['ir.ui.view'].search([('arch_db', 'ilike', 'al_report_base.report_pe_page'), ('type', '=', 'qweb')])
    facturas = usan.filtered(lambda v: 'invoice' in (v.key or '') or 'account' in (v.key or ''))
    aviso('plantillas que usan la cabecera común', '%d: %s' % (len(usan), ', '.join(sorted(usan.mapped('key'))[:12])))
    if not facturas:
        omitido('PDF de factura con cabecera común',
                'ninguna plantilla de factura (account.move) usa al_report_base.report_pe_page; '
                'se validó el PDF de la boleta (al_hr_pe_reports) que sí la usa, ver caso de reportes')


def module_info():
    M = env['ir.module.module']
    mod = M.search([('name', '=', 'al_base_module_info')])
    igual(mod.al_ficha_url, '/al_base_module_info/static/description/index.html', 'ficha del propio módulo')
    check(not M.search([('name', '=', 'base')]).al_ficha_url, '«base» no se considera ficha')
    propios = M.search([('name', '=like', 'al_hr_pe%'), ('state', '=', 'installed')])
    con = propios.filtered('al_ficha_url')
    aviso('fichas de los módulos de planillas', '%d/%d con ficha: %s' % (len(con), len(propios), ', '.join(con.mapped('name'))))
    arch = M.get_view(env.ref('base.module_view_kanban').id, 'kanban')['arch']
    check('al_ficha_url' in arch, 'el kanban de Aplicaciones ofrece el enlace a la ficha')


def mcp():
    from odoo.addons.al_mcp_server.services import tool_executor
    user = usuario_qa('qa20_mcp_lector', 'Lector MCP', ['base.group_user'])
    registro(user, 'usuario interno sin permisos de nómina')
    uenv = env(user=user.id, context=dict(env.context, mcp_scope='read'))
    res = tool_executor.execute_tool(uenv, 'odoo_search_read', {
        'model': 'res.partner', 'domain': [['name', '=like', 'QA20%']], 'fields': ['name'], 'limit': 5})
    check(res.get('records') is not None, 'odoo_search_read de contactos funciona con alcance read',
          'total %s' % res.get('total'))
    try:
        tool_executor.execute_tool(uenv, 'odoo_search_read', {'model': 'hr.payslip', 'fields': ['name'], 'limit': 1})
        check(False, 'leer boletas sin grupo de nómina debe fallar')
    except (AccessError, PermissionError) as exc:
        check(True, 'boletas (hr.payslip) denegadas al usuario sin permisos', type(exc).__name__)
    try:
        tool_executor.execute_tool(uenv, 'odoo_create', {'model': 'res.partner', 'values': {'name': qa_name('MCP')}})
        check(False, 'crear con alcance read debe fallar')
    except (AccessError, PermissionError) as exc:
        check(True, 'odoo_create rechazado con alcance read', '%s: %s' % (type(exc).__name__, exc))
    cnt = tool_executor.execute_tool(uenv, 'odoo_count', {'model': 'res.partner', 'domain': [['name', '=like', 'QA20%']]})
    check(isinstance(cnt, dict), 'odoo_count responde', cnt)


def licencia():
    from odoo.tools.misc import file_path
    mod = env['ir.module.module'].search([('name', '=', 'ol_licencia_perpetua')])
    igual(mod.state, 'installed', 'ol_licencia_perpetua instalado')
    import ast
    manifest = ast.literal_eval(open(file_path('ol_licencia_perpetua/__manifest__.py')).read())
    rep = [a for a in manifest['assets']['web.assets_backend'] if isinstance(a, tuple) and a[0] == 'replace']
    for _op, target, source in rep:
        try:
            file_path(target)
            file_path(source)
            check(True, 'el reemplazo de asset apunta a archivos existentes', '%s ← %s' % (target, source))
        except FileNotFoundError as exc:
            check(False, 'archivo del reemplazo de asset no existe', str(exc))


def preview():
    mod = env['ir.module.module'].search([('name', '=', 'prt_report_attachment_preview')])
    igual(mod.state, 'installed', 'prt_report_attachment_preview instalado')
    from odoo.addons.prt_report_attachment_preview.controllers.report import CxReportController
    from odoo.addons.web.controllers.report import ReportController
    check(issubclass(CxReportController, ReportController), 'el controlador de vista previa extiende /report')
    ctx = CxReportController()._get_extra_context_for_single_record(
        "'Boleta de pago - %s - %s' % (object.employee_id.name, object.date_to)")
    check('object' in ctx, 'evalúa el nombre del PDF con «object» (abre en pestaña con su nombre)', ctx)


def cola():
    partner = env['res.partner'].create({'name': qa_name('Contacto queue_job')})
    registro(partner, 'contacto escrito por un job')
    job = partner.with_delay(description=qa_name('Job de prueba'), priority=5).write({'comment': 'escrito por queue_job'})
    rec = env['queue.job'].search([('uuid', '=', job.uuid)])
    check(rec and rec.state == 'pending', 'job encolado en queue.job (pending)', rec.state)
    check(not partner.comment, 'antes de ejecutarse el contacto no cambia')
    job.perform()
    job.set_done(result='ok QA20')
    job.store()
    rec.invalidate_recordset()
    partner.invalidate_recordset()
    igual(rec.state, 'done', 'job ejecutado con perform → done')
    igual(partner.comment and 'queue_job' in str(partner.comment), True, 'el método diferido se ejecutó (write)')
    registro(rec, 'queue.job')


def ciudad():
    mod = env['ir.module.module'].search([('name', '=', 'al_l10n_pe_city')])
    igual(mod.state, 'uninstalled', 'al_l10n_pe_city no instalado')
    omitido('al_l10n_pe_city', "manifest 'installable': False en Odoo 20 (depende de base_address_extended); no se valida")


caso('al_docs_help', 'Ayuda contextual: entradas, UIT/RMV y respuesta al systray',
     'Un usuario interno pulsa el «?» de la barra superior en la pantalla de UIT/RMV.')
correr(docs_help)
caso('al_report_base', 'Página común de documentos (cabecera y pie)',
     'Se imprime un documento con el formato común de OLANO IT.')
correr(report_base)
caso('al_base_module_info', 'Ficha de módulo en Aplicaciones',
     'El usuario abre Aplicaciones y pulsa «Ver la ficha completa del módulo».')
correr(module_info)
caso('al_mcp_server', 'Herramienta MCP de lectura con usuario no administrador',
     'Un asistente IA conectado con el token de un usuario interno consulta contactos y boletas.')
correr(mcp)
caso('ol_licencia_perpetua', 'Constatar instalación y reemplazo del aviso de suscripción',
     'Se verifica que el módulo está instalado y su asset reemplaza al de web_enterprise.')
correr(licencia)
caso('prt_report_attachment_preview', 'Vista previa de PDF en el navegador',
     'El usuario imprime un PDF y se abre en una pestaña en lugar de descargarse.')
correr(preview)
caso('queue_job', 'Encolar un job con with_delay y ejecutarlo',
     'Un proceso encola una escritura diferida y el job runner la ejecuta.')
correr(cola)
caso('al_l10n_pe_city', 'Módulo no instalable en Odoo 20', 'Solo constatar.')
correr(ciudad)
