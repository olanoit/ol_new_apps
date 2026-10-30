# -*- coding: utf-8 -*-
"""Fase inventario_farmacia: OL-INVENTORY, OL-HEALTH y OL-INDUSTRIES.

Se antepone _qa.py (helpers) y se cierra con _fin.py (JSON + commit).
Todo lo creado lleva el prefijo QA20; los almacenes propios (Q20TA, Q20TB,
Q20FA) aíslan las existencias de las de las otras fases.
"""
import io
import os
import re
from datetime import date, datetime, timedelta
from unittest.mock import patch

from dateutil.relativedelta import relativedelta
from lxml import etree

from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.tools import BinaryBytes

NS = {'cbc': 'urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2',
      'cac': 'urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2'}
# Sin el asistente de confirmación por SMS (stock_sms) al validar entregas.
env = env(context=dict(env.context, skip_sms=True))
MYAPPS = '/home/och/odoo/ce20/myapps/ol_new_apps'
TODAY = fields.Date.context_today(env.user)          # 2026-10-30
MONTH_START = TODAY.replace(day=1)
MONTH_END = MONTH_START + relativedelta(months=1, days=-1)
S = {}   # registros compartidos entre casos




def raises(exc, fn, *args, **kwargs):
    """True si fn lanza exc (dentro de un savepoint, para no ensuciar)."""
    try:
        with env.cr.savepoint():
            fn(*args, **kwargs)
    except exc as e:  # noqa: BLE001
        return str(e) or True
    return False


def xml_text(tree, path):
    nodes = tree.xpath(path, namespaces=NS)
    return nodes[0].text if nodes else None


def get_or_create(model, domain, vals):
    rec = env[model].search(domain, limit=1)
    return rec or env[model].create(vals)


# ======================================================================
#  Preparación común
# ======================================================================
caso('preparacion', 'Datos maestros QA20 de la fase', 'Crear almacenes, categorías y contactos de prueba.')


def _prep():
    Partner = env['res.partner']
    district = env['l10n_pe.res.city.district'].search([], limit=1)
    S['district'] = district
    S['addr_a'] = Partner.create({
        'name': qa_name('Almacén Central Inv'), 'parent_id': company.partner_id.id, 'type': 'other',
        'street': 'Av. Industrial 100 QA20', 'l10n_pe_district': district.id, 'l10n_pe_annex_code': '0091'})
    S['addr_b'] = Partner.create({
        'name': qa_name('Almacén Obra Inv'), 'parent_id': company.partner_id.id, 'type': 'other',
        'street': 'Jr. Las Obras 200 QA20', 'l10n_pe_district': district.id, 'l10n_pe_annex_code': '0092'})
    S['addr_f'] = Partner.create({
        'name': qa_name('Botica Central'), 'parent_id': company.partner_id.id, 'type': 'other',
        'street': 'Av. Salud 300 QA20', 'l10n_pe_district': district.id, 'l10n_pe_annex_code': '0093'})
    Wh = env['stock.warehouse']
    S['wh_a'] = get_or_create('stock.warehouse', [('code', '=', 'Q20TA'), ('company_id', '=', company.id)], {
        'name': qa_name('Inv Central'), 'code': 'Q20TA', 'partner_id': S['addr_a'].id, 'company_id': company.id})
    S['wh_b'] = get_or_create('stock.warehouse', [('code', '=', 'Q20TB'), ('company_id', '=', company.id)], {
        'name': qa_name('Inv Obra'), 'code': 'Q20TB', 'partner_id': S['addr_b'].id, 'company_id': company.id})
    S['wh_f'] = get_or_create('stock.warehouse', [('code', '=', 'Q20FA'), ('company_id', '=', company.id)], {
        'name': qa_name('Botica'), 'code': 'Q20FA', 'partner_id': S['addr_f'].id, 'company_id': company.id})
    for wh in (S['wh_a'], S['wh_b'], S['wh_f']):
        registro(wh, 'almacén propio de la fase')
    S['categ_avco'] = env['product.category'].create({
        'name': qa_name('AVCO tiempo real'), 'property_cost_method': 'average', 'property_valuation': 'real_time'})
    S['categ_avco_per'] = env['product.category'].create({
        'name': qa_name('AVCO kardex'), 'property_cost_method': 'average'})
    S['categ_farma'] = env['product.category'].create({'name': qa_name('Farmacia')})
    S['supplier'] = registro(Partner.create({
        'name': qa_name('Droguería Andina S.A.C.'), 'is_company': True, 'vat': '20100070970',
        'country_id': env.ref('base.pe').id, 'street': 'Av. Proveedores 50',
        'l10n_pe_district': district.id}), 'proveedor')
    S['customer'] = registro(Partner.create({
        'name': qa_name('Clínica San Juan S.A.C.'), 'is_company': True, 'vat': '20131312955',
        'country_id': env.ref('base.pe').id, 'street': 'Calle 1',
        'l10n_pe_district': district.id}), 'cliente')
    S['owner'] = registro(Partner.create({
        'name': qa_name('Alquileres del Sur S.A.C.'), 'is_company': True, 'vat': '20601234565'}), 'propietario')
    S['patient'] = registro(Partner.create({
        'name': qa_name('Mario Cárdenas Ríos'), 'additional_identifiers': {'PE_DNI': '40112233'},
        'street': 'Av. Arequipa 1234', 'phone': '987654321', 'email': 'qa20.paciente@example.com',
        'country_id': env.ref('base.pe').id}), 'paciente (DNI 40112233)')
    S['prescriber'] = registro(Partner.create({
        'name': qa_name('Dra. Elena Vargas'), 'pharmacy_is_prescriber': True,
        'pharmacy_profession': 'physician', 'pharmacy_prescriber_license': '051234'}), 'prescriptor CMP')
    check(all(S[k] for k in ('wh_a', 'wh_b', 'wh_f')), 'almacenes QA20 disponibles',
          ', '.join(S[k].code for k in ('wh_a', 'wh_b', 'wh_f')))


ejecutar(_prep)

# Se crea al principio: pos.config activa productos compartidos y, más tarde,
# choca con las otras fases (SerializationFailure).
caso('preparacion', 'Caja «QA20 Botica» del local Q20FA',
     'Crear el diario de caja, el método Efectivo y el punto de venta del local y abrir la sesión.')


def _prep_pos():
    journal = env['account.journal'].search([('code', '=', 'Q20CJ'), ('company_id', '=', company.id)], limit=1) \
        or env['account.journal'].create({'name': qa_name('Caja Botica'), 'type': 'cash', 'code': 'Q20CJ',
                                          'company_id': company.id})
    cash = env['pos.payment.method'].create({'name': qa_name('Efectivo'), 'type': 'cash', 'journal_id': journal.id,
                                             'company_id': company.id})
    config = env['pos.config'].create({'name': qa_name('Botica'), 'picking_type_id': S['wh_f'].pos_type_id.id,
                                       'payment_method_ids': [Command.set(cash.ids)]})
    config.picking_type_id.use_existing_lots = True
    admin = env.ref('base.user_admin')
    check(admin.has_group('point_of_sale.group_pos_manager'), 'el administrador es responsable del TPV')
    config.with_user(admin).open_ui()
    session = config.current_session_id
    session.with_user(admin).set_opening_control(0, None)
    S['admin'] = admin
    S.update(pos_config=config, cash_pm=cash, session=session)
    registro(config, 'punto de venta')
    registro(session, 'sesión abierta')
    check(session.state == 'opened', 'sesión del TPV abierta', session.state)


ejecutar(_prep_pos)


# ======================================================================
#  OL-INVENTORY
# ======================================================================
def _internal(src, dest, product, qty, owner=None, picking_type=None):
    picking = env['stock.picking'].create({
        'picking_type_id': (picking_type or src.warehouse_id.int_type_id).id,
        'location_id': src.id, 'location_dest_id': dest.id,
        'owner_id': owner.id if owner else False,
        'move_ids': [Command.create({'product_id': product.id, 'product_uom_qty': qty,
                                     'location_id': src.id, 'location_dest_id': dest.id})]})
    picking.action_confirm()
    picking.action_assign()
    picking.move_ids.write({'quantity': qty, 'picked': True})
    picking.button_validate()
    return picking


def _send_guide(picking):
    """Envía la GRE con la firma/SUNAT simuladas (como los tests)."""
    sent = {}

    def sign(record, edi_str):
        sent['xml'] = etree.fromstring(edi_str)
        return {'cdr': b'<cdr/>'}

    picking.l10n_pe_edi_transport_type = '01'
    with patch.object(type(picking), '_l10n_pe_edi_sign', sign), \
            patch.object(type(picking), '_l10n_pe_edi_check_required_data', lambda self: None):
        picking.action_send_delivery_guide()
    return sent.get('xml')


# ---------------------------------------------------------------- al_stock_base
caso('al_stock_base', 'Tipo de operación SUNAT (tabla 12) y página «Logística PE»',
     'Abrir una recepción, una salida y un traslado: la página Logística PE muestra el tipo de '
     'operación de la tabla 12 (02 compra / 01 venta / 21 traslado).')


def _stock_base_t12():
    product = env['product.product'].create({'name': qa_name('Tornillo T12'), 'is_storable': True})
    S['t12_product'] = registro(product)
    wh = S['wh_a']
    supplier_loc = env.ref('stock.stock_location_suppliers')
    customers = env.ref('stock.stock_location_customers')
    receipt = env['stock.picking'].create({
        'picking_type_id': wh.in_type_id.id, 'partner_id': S['supplier'].id,
        'location_id': supplier_loc.id, 'location_dest_id': wh.lot_stock_id.id,
        'move_ids': [Command.create({'product_id': product.id, 'product_uom_qty': 10,
                                     'location_id': supplier_loc.id, 'location_dest_id': wh.lot_stock_id.id})]})
    receipt.action_confirm()
    receipt.move_ids.write({'quantity': 10, 'picked': True})
    receipt.button_validate()
    delivery = env['stock.picking'].create({
        'picking_type_id': wh.out_type_id.id, 'partner_id': S['customer'].id,
        'location_id': wh.lot_stock_id.id, 'location_dest_id': customers.id,
        'move_ids': [Command.create({'product_id': product.id, 'product_uom_qty': 2,
                                     'location_id': wh.lot_stock_id.id, 'location_dest_id': customers.id})]})
    internal = _internal(wh.lot_stock_id, S['wh_b'].lot_stock_id, product, 3)
    registro(receipt | delivery | internal)
    igual(receipt.state, 'done', 'recepción validada')
    igual(receipt.l10n_pe_operation_type, '2', 'recepción: tabla 12 = 02 (compra)')
    igual(delivery.l10n_pe_operation_type, '1', 'salida: tabla 12 = 01 (venta)')
    igual(internal.l10n_pe_operation_type, '21', 'traslado interno: tabla 12 = 21')
    arch = etree.fromstring(env['stock.picking'].get_view(env.ref('stock.view_picking_form').id, 'form')['arch'])
    page = arch.xpath("//page[@name='l10n_pe_stock']")
    check(page and page[0].getparent().tag == 'notebook', 'página «Logística PE» en el formulario')
    check(page and page[0].xpath(".//field[@name='l10n_pe_operation_type']"),
          'el tipo de operación (tabla 12) vive en la página Logística PE')
    visible = [n for n in arch.xpath("//field[@name='l10n_pe_operation_type']")
               if n.get('invisible') not in ('1', 'True')]
    igual(len(visible), 1, 'el campo aparece una sola vez en el formulario')


ejecutar(_stock_base_t12)

caso('al_stock_base', 'Planchas enteras: producto indivisible y redondeo',
     'Crear la unidad «Plancha 2.44×2.15» marcada «Plancha entera» y un tablero con esa unidad: '
     'queda indivisible y una necesidad de 7.3 planchas se entrega como 8.')


def _stock_base_sheet():
    m2 = env.ref('uom.product_uom_square_meter')
    sheet = env['uom.uom'].create({'name': qa_name('Plancha 2.44x2.15'), 'relative_factor': 5.246,
                                   'relative_uom_id': m2.id, 'construction_whole_sheet': True})
    board = env['product.product'].create({'name': qa_name('Tablero melamina 18 mm'), 'is_storable': True,
                                           'uom_id': sheet.id})
    registro(sheet); registro(board)
    check(board.construction_is_indivisible, 'el tablero con unidad «Plancha entera» es indivisible')
    igual(board.construction_sheet_area_m2, 5.246, 'área de la plancha en m² desde el factor', 3)
    tol = company.construction_whole_sheet_tolerance
    igual(board._construction_whole_qty(7.3, company), 8.0 if tol < 0.3 else 7.0,
          'necesidad 7.3 → planchas enteras (tolerancia %s)' % tol)
    igual(board._construction_whole_qty(7.000001, company), 7.0, '7.000001 no sube a 8 (redondeo de la UdM)')
    board.product_tmpl_id.construction_indivisible_mode = 'no'
    check(not board.construction_is_indivisible, 'forzado «No» en el producto anula la regla de la unidad')
    labels = env['stock.picking.type'].with_context(lang='es_PE').get_view(
        env.ref('stock.stock_picking_type_kanban').id, 'kanban')['arch']
    check('Late' not in re.findall(r'>\s*([A-Za-z ]+)\s*<', labels), 'resumen de inventario sin la etiqueta «Late» en inglés')


ejecutar(_stock_base_sheet)

# ------------------------------------------------------ al_l10n_pe_stock_transfer
caso('al_l10n_pe_stock_transfer', 'Traslado interno entre establecimientos con GRE motivo 04',
     'Trasladar 4 andamios de «Inv Central» (anexo 0091) a «Inv Obra» (anexo 0092) y emitir la guía '
     'de remisión: motivo 04, destinatario = remitente.')


def _transfer_internal():
    product = env['product.product'].create({'name': qa_name('Andamio'), 'is_storable': True,
                                             'weight': 10.0, 'categ_id': S['categ_avco_per'].id,
                                             'standard_price': 50.0})
    S['andamio'] = registro(product)
    Quant = env['stock.quant']
    Quant._update_available_quantity(product, S['wh_a'].lot_stock_id, 10.0)
    picking = registro(_internal(S['wh_a'].lot_stock_id, S['wh_b'].lot_stock_id, product, 4.0),
                       'traslado entre establecimientos')
    check(picking.l10n_pe_edi_guide_allowed, 'el traslado admite guía de remisión')
    igual(picking.l10n_pe_edi_origin_address_id, S['addr_a'], 'punto de partida = dirección del almacén origen')
    igual(picking.l10n_pe_edi_destination_address_id, S['addr_b'], 'punto de llegada = dirección del almacén destino')
    tree = _send_guide(picking)
    igual(picking.l10n_pe_edi_reason_for_transfer, '04', 'motivo propuesto 04')
    igual(picking.l10n_pe_edi_status, 'sent', 'GRE marcada como enviada (firma simulada)')
    igual(xml_text(tree, '//cac:Shipment/cbc:HandlingCode'), '04', 'XML: HandlingCode 04')
    igual(xml_text(tree, '//cac:DeliveryCustomerParty//cbc:ID'), company.vat, 'XML: destinatario = RUC de la empresa')
    desp = tree.xpath('//cac:DespatchAddress/cbc:AddressTypeCode', namespaces=NS)[0]
    deliv = tree.xpath('//cac:DeliveryAddress/cbc:AddressTypeCode', namespaces=NS)[0]
    igual((desp.text, desp.get('listID')), ('0091', company.vat), 'XML: partida con anexo 0091 y RUC')
    igual((deliv.text, deliv.get('listID')), ('0092', company.vat), 'XML: llegada con anexo 0092 y RUC')
    omitido('Envío real de la GRE a SUNAT/OSE', 'sin credenciales: firma y envío simulados como en los tests')
    # Mismo establecimiento: sin guía
    shelf = env['stock.location'].create({'name': qa_name('Estante'), 'location_id': S['wh_a'].lot_stock_id.id})
    same = registro(_internal(S['wh_a'].lot_stock_id, shelf, product, 1.0), 'dentro del mismo local')
    check(not same.l10n_pe_edi_guide_allowed, 'traslado dentro del mismo establecimiento no admite guía')
    check(raises(UserError, _send_guide, same), 'enviar la guía de un traslado sin salida del local se bloquea')


ejecutar(_transfer_internal)

caso('al_l10n_pe_stock_transfer', 'Bienes de terceros: guía con propietario y sin valorizar',
     'Recibir 5 andamios en consignación de «Alquileres del Sur» y trasladarlos a la obra: la guía '
     'lleva el propietario en observaciones y los bienes no entran a la valorización ni al kardex.')


def _transfer_third_party():
    product = S.get('andamio') or env['product.product'].create({'name': qa_name('Andamio'), 'is_storable': True, 'weight': 10.0})
    wh = S['wh_a']
    supplier_loc = env.ref('stock.stock_location_suppliers')
    receipt = env['stock.picking'].create({
        'picking_type_id': wh.in_type_id.id, 'partner_id': S['owner'].id, 'owner_id': S['owner'].id,
        'location_id': supplier_loc.id, 'location_dest_id': wh.lot_stock_id.id,
        'move_ids': [Command.create({'product_id': product.id, 'product_uom_qty': 5,
                                     'location_id': supplier_loc.id, 'location_dest_id': wh.lot_stock_id.id})]})
    receipt.action_confirm()
    receipt.move_ids.write({'quantity': 5, 'picked': True})
    receipt.move_line_ids.write({'owner_id': S['owner'].id})
    receipt.button_validate()
    registro(receipt, 'recepción en consignación')
    igual(receipt.state, 'done', 'recepción de terceros validada')
    check(receipt.l10n_pe_has_third_party_goods, 'la recepción se marca «Con bienes de terceros»')
    igual(receipt.move_ids.value, 0.0, 'la recepción de bienes de terceros no se valoriza (value = 0)')
    own_quants = env['stock.quant'].search([('product_id', '=', product.id), ('owner_id', '=', S['owner'].id),
                                            ('location_id', '=', wh.lot_stock_id.id)])
    igual(sum(own_quants.mapped('quantity')), 5.0, 'existencias de terceros separadas por propietario')
    kardex = env['l10n_pe.kardex.line'].search([('move_id', 'in', receipt.move_ids.ids)])
    check(not kardex, 'los bienes de terceros no entran al kardex valorizado')
    picking = registro(_internal(wh.lot_stock_id, S['wh_b'].lot_stock_id, product, 3.0, owner=S['owner']),
                       'traslado de bienes de terceros')
    check(picking.l10n_pe_has_third_party_goods, 'traslado con bienes de terceros')
    igual(picking.l10n_pe_third_party_owner_ids, S['owner'], 'propietario detectado en la transferencia')
    tree = _send_guide(picking)
    check(S['owner'].name in (xml_text(tree, '/*/cbc:Note') or ''), 'XML de la GRE: propietario en las observaciones',
          xml_text(tree, '/*/cbc:Note'))
    html = env['ir.actions.report']._render_qweb_html(
        'al_l10n_pe_delivery_guide_report.action_report_guia_remision', picking.ids)[0].decode()
    check('Propietario de los bienes' in html and S['owner'].name in html, 'PDF de la guía con el propietario')
    own = registro(_internal(wh.lot_stock_id, S['wh_b'].lot_stock_id, product, 1.0, owner=company.partner_id),
                   'propietario = la compañía')
    check(not own.l10n_pe_has_third_party_goods, 'la propia compañía no cuenta como tercero')


ejecutar(_transfer_third_party)

# ------------------------------------------------------------ ol_stock_kardex_pe
caso('ol_stock_kardex_pe', 'Kardex 13.1 valorizado (AVCO) cuadra con el valor de inventario',
     'Comprar 10 u a S/10 y 10 u a S/20, facturar la primera compra (F001), vender 5 y abrir '
     '«Kardex SUNAT» 13.1 de octubre 2026: saldo 15 u × 15 = S/225.')


def _kardex_moves():
    product = env['product.product'].create({
        'name': qa_name('Producto Kardex'), 'default_code': 'QA20-KDX-001', 'is_storable': True,
        'categ_id': S['categ_avco_per'].id, 'l10n_pe_type_of_existence': '1', 'list_price': 100.0,
        'taxes_id': [Command.clear()], 'supplier_taxes_id': [Command.clear()]})
    S['kdx'] = registro(product)
    wh = S['wh_a']

    def buy(qty, price):
        po = env['purchase.order'].create({
            'partner_id': S['supplier'].id, 'picking_type_id': wh.in_type_id.id,
            'order_line': [Command.create({'product_id': product.id, 'product_qty': qty,
                                           'price_unit': price, 'tax_ids': [Command.clear()]})]})
        po.button_confirm()
        po.picking_ids.move_ids.write({'quantity': qty, 'picked': True})
        po.picking_ids.button_validate()
        return po

    po1 = registro(buy(10, 10.0), 'compra 10 @ 10')
    po2 = registro(buy(10, 20.0), 'compra 10 @ 20')
    so = env['sale.order'].create({
        'partner_id': S['customer'].id, 'warehouse_id': wh.id,
        'order_line': [Command.create({'product_id': product.id, 'product_uom_qty': 5,
                                       'price_unit': 100.0, 'tax_ids': [Command.clear()]})]})
    so.action_confirm()
    so.picking_ids.move_ids.write({'quantity': 5, 'picked': True})
    so.picking_ids.button_validate()
    registro(so, 'venta 5')
    S['kdx_so'] = so
    # Factura de proveedor de la primera compra (Tabla 10 del kardex).
    po1.action_create_invoice()
    bill = po1.invoice_ids
    # Diario de compras propio: el compartido choca (SerializationFailure)
    # con las fases que publican facturas a la vez.
    pj = env['account.journal'].search([('code', '=', 'Q20CP'), ('company_id', '=', company.id)], limit=1) \
        or env['account.journal'].create({'name': qa_name('Compras'), 'type': 'purchase', 'code': 'Q20CP',
                                          'company_id': company.id, 'l10n_latam_use_documents': True})
    bill.journal_id = pj
    bill.write({'invoice_date': TODAY,
                'l10n_latam_document_type_id': env.ref('l10n_pe.document_type01').id,
                'l10n_latam_document_number': 'F001-00020123'})
    bill.action_post()
    registro(bill, 'factura de proveedor F001-00020123')
    S['kdx_bill'] = bill
    S['kdx_receipt'] = po1.picking_ids.move_ids
    return product


def _kardex_wizard(**vals):
    return env['l10n_pe.kardex.report.wizard'].create(dict({
        'company_id': company.id, 'date_from': MONTH_START, 'date_to': MONTH_END,
        'report_type': '1301', 'product_ids': [Command.set(S['kdx'].ids)]}, **vals))


def _kardex_blocks(wizard, product):
    return [b for scope in wizard._get_report_data() for b in scope['products'] if b['product'] == product]


def _kardex_131():
    product = _kardex_moves()
    wizard = _kardex_wizard()
    block = _kardex_blocks(wizard, product)[0]
    rows, total = block['lines'], block['total_line']
    igual([r.line_type for r in rows], ['opening', 'move', 'move', 'move'], 'filas: saldo inicial + 3 movimientos')
    opening, in1, in2, out = rows
    igual(opening.operation_type, '16', 'saldo inicial con operación 16')
    igual(in1.operation_type, '02', 'compra: operación 02 (tabla 12)')
    igual((in1.document_type_code, in1.serie, in1.folio), ('01', 'F001', '00020123'),
          'compra: comprobante de la factura (tabla 10)')
    igual(in2.balance_unit_cost, 15.0, 'costo promedio tras la 2.ª compra')
    igual(out.operation_type, '01', 'venta: operación 01')
    igual(out.cost_total_out, 75.0, 'salida 5 × 15 = 75 (AVCO)')
    igual(out.document_type_code, '09', 'venta sin comprobante → guía (09)')
    igual(total.balance_qty, 15.0, 'saldo final 15 u')
    igual(total.balance_value, 225.0, 'saldo final S/ 225')
    p = product.with_context(to_date=fields.Datetime.now())
    igual(total.balance_value, p.total_value, 'saldo del kardex = valor de inventario de Odoo')
    igual(total.balance_qty, p.qty_available, 'saldo del kardex = cantidad a mano')
    rep = registro(env['l10n_pe.kardex.report'].create({
        'company_id': company.id, 'date_from': MONTH_START, 'date_to': MONTH_END,
        'report_type': '1301', 'file_format': 'xlsx', 'product_ids': [Command.set(product.ids)]}),
        'kardex en segundo plano (XLSX)') if 'product_ids' in env['l10n_pe.kardex.report']._fields else None
    from odoo.addons.ol_stock_kardex_pe.reports.kardex_xlsx import build_kardex_xlsx
    content = build_kardex_xlsx(wizard)
    igual(content[:2], b'PK', 'XLSX 13.1 generado')
    html = env['ir.actions.report']._render_qweb_html('ol_stock_kardex_pe.report_kardex', wizard.ids)[0]
    check(b'FORMATO 13.1' in html and b'QA20-KDX-001' in html, 'PDF formato 13.1 con el producto')
    if rep:
        rep._l10n_pe_generate() if hasattr(rep, '_l10n_pe_generate') else None


ejecutar(_kardex_131)

caso('ol_stock_kardex_pe', 'Kardex 12.1 en unidades físicas por almacén',
     'Generar el formato 12.1 agrupado por almacén: solo cantidades, sin costos.')


def _kardex_121():
    product = S['kdx']
    wizard = _kardex_wizard(report_type='1201', group_by_warehouse=True)
    blocks = _kardex_blocks(wizard, product)
    check(blocks, 'hay bloques por almacén')
    igual(sum(b['total_line'].balance_qty for b in blocks), 15.0, 'saldo físico sumado por almacén = 15')
    check(not any(r.cost_total_in or r.cost_total_out for b in blocks for r in b['lines']), '12.1 sin costos')
    lines = env['l10n_pe.kardex.line'].search([('product_id', '=', product.id)], order='date, id')
    igual(len(lines), 3, 'vista SQL del kardex interactivo: 3 movimientos')
    igual(lines[-1:].balance_value, 225.0, 'saldo corrido en el kardex interactivo')


ejecutar(_kardex_121)

caso('ol_stock_kardex_pe', 'TXT del PLE 13.1 con el comprobante del kardex',
     'Generar el libro electrónico 13.1 (PLE) de octubre 2026: la fila de la compra lleva '
     'tipo 01, serie F001 y número 00020123; la unidad es la del producto.')


def _kardex_ple():
    product = S['kdx']
    wizard = env['l10n_pe.stock.ple.wizard'].create({'date_from': MONTH_START, 'date_to': MONTH_END})
    content = wizard._get_ple_report_content('1301')
    rows = [line.split('|') for line in (content or '').split('\n') if line]
    i_date, i_type, i_serie, i_folio = wizard._l10n_pe_kardex_columns()
    receipt = S['kdx_receipt']
    row = next((r for r in rows if r[1] == str(receipt.id).zfill(6)), None)
    check(row, 'fila de la compra en el TXT (CUO = id del movimiento)')
    if row:
        igual((row[i_type], row[i_serie], row[i_folio]), ('01', 'F001', '00020123'), 'TXT: comprobante de la compra')
        igual(row[i_date][3:], TODAY.strftime('%m/%Y'), 'TXT: fecha del periodo en hora de Lima')
        igual(row[i_folio + 3], product.uom_id.l10n_pe_edi_measure_unit_code, 'TXT: unidad del producto (campo 16)')
    mine = [r for r in rows if r[1] in {str(m.id).zfill(6) for m in env['stock.move'].search([('product_id', '=', product.id)])}]
    print('    fila TXT:', '|'.join(row or []))
    igual(len(mine), 3, 'TXT: 3 filas de movimientos del producto QA20')
    check(row and 'QA20KDX001' in row, 'TXT: código del producto saneado (campo 7)')
    igual(wizard._get_stock_valuation(S['categ_avco_per'].id), '1', 'campo 17 (tabla 14): promedio = 1')
    wizard.get_ple_report('1301')
    check(wizard.report_filename and wizard.report_filename.startswith('LE%s' % company.vat),
          'nombre del archivo PLE con el RUC', wizard.report_filename)
    omitido('Presentación del PLE en SUNAT', 'se valida el TXT generado; la carga al PLE es manual/externa')


ejecutar(_kardex_ple)

# ------------------------------------------------------------ al_l10n_pe_mrp_cost
caso('al_l10n_pe_mrp_cost', 'Costo de OF: materiales + MO + CIF con cierre del periodo',
     'Fabricar 2 mesas (4 tableros a S/10, 2 h de corte con un operario a S/10/h), aplicar el costo '
     'hora real del mes (S/15) y repartir una bolsa de CIF de S/300 con el cierre de costos.')


def _mrp_cost():
    acc71 = env['account.account'].search([('company_ids', 'in', company.id), ('code', '=like', '71%')],
                                          order='code', limit=1)
    acc_cif = env['account.account'].search([('company_ids', 'in', company.id), ('code', '=', '6399201')], limit=1) \
        or env['account.account'].create({'name': qa_name('CIF planta'), 'code': '6399201',
                                          'account_type': 'expense', 'company_ids': [Command.set(company.ids)]})
    liability = env['account.account'].search([('company_ids', 'in', company.id), ('account_type', '=', 'liability_current')], limit=1)
    labor_acc = company.l10n_pe_mrp_labor_account_id
    if not labor_acc:
        # Escribir en res.company choca (SerializationFailure) con las otras
        # fases que corren a la vez: se simula el ajuste solo para este caso.
        labor_acc = acc71
        aviso('Contrapartida de MO no configurada en la compañía (Ajustes de Fabricación)',
              'se simula %s durante el cierre, sin escribir en res.company' % acc71.code)
    board = env['product.product'].create({'name': qa_name('Tablero MRP'), 'is_storable': True,
                                           'categ_id': S['categ_avco'].id, 'standard_price': 10.0})
    table = env['product.product'].create({'name': qa_name('Mesa de centro'), 'is_storable': True,
                                           'categ_id': S['categ_avco'].id, 'weight': 12.0})
    wc = env['mrp.workcenter'].create({'name': qa_name('Seccionadora'), 'code': 'Q20S', 'costs_hour': 30.0,
                                       'company_id': company.id})
    bom = env['mrp.bom'].create({
        'product_tmpl_id': table.product_tmpl_id.id, 'product_qty': 1.0,
        'bom_line_ids': [Command.create({'product_id': board.id, 'product_qty': 2.0})],
        'operation_ids': [Command.create({'name': 'Corte', 'workcenter_id': wc.id, 'time_cycle_manual': 60.0})]})
    emp = env['hr.employee'].create({'name': qa_name('Juan Pérez MRP'), 'company_id': company.id, 'hourly_cost': 10.0})
    for r in (board, table, wc, bom, emp):
        registro(r)
    env['stock.quant']._update_available_quantity(board, S['wh_a'].lot_stock_id, 100.0)
    mo = env['mrp.production'].create({'product_id': table.id, 'bom_id': bom.id, 'product_qty': 2.0,
                                       'picking_type_id': S['wh_a'].manu_type_id.id})
    mo.action_confirm()
    wo = mo.workorder_ids
    start = fields.Datetime.now() - timedelta(hours=3)
    env['mrp.workcenter.productivity'].create({
        'workorder_id': wo.id, 'workcenter_id': wc.id, 'employee_id': emp.id,
        'loss_id': env.ref('mrp.block_reason7').id,
        'date_start': start, 'date_end': start + timedelta(hours=2)})
    mo.qty_producing = 2.0
    mo.move_raw_ids.picked = True
    mo.button_mark_done()
    registro(mo, 'orden de fabricación')
    igual(mo.state, 'done', 'OF terminada')
    finished = mo.move_finished_ids.filtered(lambda m: m.product_id == table)
    value_before = finished.value
    # Costo hora real del mes
    Rate = env['l10n_pe.mrp.labor.rate']
    rate = Rate.search([('company_id', '=', company.id), ('date_from', '=', MONTH_START)], limit=1)
    line_vals = {'employee_id': emp.id, 'payroll_cost': 1500.0, 'hours': 100.0}
    if rate:
        aviso('Ya existía un costo hora de octubre', 'se añade la línea del operario QA20 a %s' % rate.display_name)
        rate.line_ids = [Command.create(line_vals)]
        if rate.state != 'applied':
            rate.state = 'applied'
    else:
        rate = Rate.create({'date_from': MONTH_START, 'date_to': MONTH_END, 'line_ids': [Command.create(line_vals)]})
        rate.action_apply()
    registro(rate, 'costo hora real (S/15)')
    igual(rate.state, 'applied', 'costo hora del mes aplicado')
    # Bolsa de CIF con cuenta propia y solo el centro QA20
    pool = registro(env['l10n_pe.mrp.cost.pool'].create({
        'name': qa_name('Gastos de planta'), 'account_prefixes': '6399201', 'split_basis': 'hours',
        'workcenter_ids': [Command.set(wc.ids)], 'counterpart_account_id': acc71.id}), 'bolsa de CIF')
    entry = env['account.move'].create({'move_type': 'entry', 'date': TODAY, 'ref': qa_name('CIF del mes'),
                                        'line_ids': [Command.create({'account_id': acc_cif.id, 'debit': 300.0}),
                                                     Command.create({'account_id': liability.id, 'credit': 300.0})]})
    entry.action_post()
    registro(entry, 'gasto CIF S/300')
    close = env['l10n_pe.mrp.cost.close'].create({'date_from': TODAY, 'date_to': TODAY})
    close.action_compute()
    registro(close, 'cierre de costos del día')
    labor = close.line_ids.filtered(lambda l: l.element == 'labor' and l.production_id == mo)
    overhead = close.line_ids.filtered(lambda l: l.element == 'overhead' and l.production_id == mo)
    igual(sum(labor.mapped('amount')), 10.0, 'ajuste de MO: 2 h × (15 − 10) = 10')
    igual(sum(overhead.mapped('amount')), 300.0, 'CIF repartido a la OF: 300')
    with patch.object(type(company), 'l10n_pe_mrp_labor_account_id', property(lambda self: labor_acc)):
        close.action_validate()
    igual(close.state, 'done', 'cierre validado')
    igual(close.landed_cost_id.state, 'done', 'costo en destino validado')
    finished.invalidate_recordset()
    igual(finished.value, value_before + 310.0, 'valor del producto terminado + 310')
    credit71 = close.landed_cost_id.account_move_id.line_ids.filtered(lambda l: l.account_id == acc71)
    igual(sum(credit71.mapped('credit')), 310.0, 'asiento: abono a la 71 por 310')
    rows = env['l10n_pe.mrp.cost.analysis'].search([('production_id', '=', mo.id)])
    by = {}
    for row in rows:
        by[row.element] = by.get(row.element, 0.0) + row.amount
    igual(by.get('material', 0.0), 40.0, 'análisis: materiales 40')
    igual(by.get('labor', 0.0), 20.0, 'análisis: mano de obra 20')
    igual(by.get('labor_adjustment', 0.0), 10.0, 'análisis: ajuste de MO 10')
    igual(by.get('overhead', 0.0), 300.0, 'análisis: CIF 300')
    check(by.get('machine', 0.0) > 0, 'análisis: costo de máquina', by.get('machine'))


ejecutar(_mrp_cost)

# ------------------------------------------------------------ al_mrp_piece_costing
caso('al_mrp_piece_costing', 'Orden de corte: costo por área, tablero y tapacanto',
     'Fabricar una orden de corte con dos tableros (S/100 y S/200), 3 m de tapacanto y S/10 de costo '
     'extra: cada pieza sale con su costo (lateral 72.37, repisa 35.19, puerta 204.70).')


def _piece_costing():
    categ = env['product.category'].create({'name': qa_name('Piezas de corte'), 'property_cost_method': 'average'})

    def product(name, cost=0.0, **vals):
        return env['product.product'].create(dict({'name': qa_name(name), 'is_storable': True,
                                                   'categ_id': categ.id, 'standard_price': cost}, **vals))
    white = product('Tablero blanco', 100.0, default_code='QA20-BL18')
    oak = product('Tablero roble', 200.0, default_code='QA20-RO18')
    edge = product('Tapacanto blanco', 1.0, default_code='QA20-TBL22', piece_is_edge_banding=True)
    side = product('Lateral', piece_is_cut=True, piece_area_m2=0.5, piece_edge_ml=2.0, piece_material_id=white.id)
    shelf = product('Repisa', piece_is_cut=True, piece_area_m2=0.25, piece_material_id=white.id)
    door = product('Puerta', piece_is_cut=True, piece_area_m2=0.5, piece_edge_ml=1.0, piece_material_id=oak.id)
    main = product('Corte OP 999', piece_area_m2=0.1)
    S['piece'] = {'categ': categ, 'white': white, 'edge': edge}
    stock = S['wh_a'].lot_stock_id
    for comp, qty in ((white, 5), (oak, 5), (edge, 50)):
        env['stock.quant']._update_available_quantity(comp, stock, qty)
    bom = env['mrp.bom'].create({
        'product_tmpl_id': main.product_tmpl_id.id, 'product_qty': 1, 'extra_cost': 10.0,
        'bom_line_ids': [Command.create({'product_id': white.id, 'product_qty': 1}),
                         Command.create({'product_id': oak.id, 'product_qty': 1}),
                         Command.create({'product_id': edge.id, 'product_qty': 3})],
        'byproduct_ids': [Command.create({'product_id': p.id, 'product_qty': 1}) for p in (side, shelf, door)]})
    registro(bom, 'LdM de corte')
    ptype = S['wh_a'].manu_type_id
    mo = env['mrp.production'].create({'product_id': main.id, 'bom_id': bom.id, 'product_qty': 1,
                                       'picking_type_id': ptype.id, 'piece_costing': True})
    mo.action_confirm()
    mo.qty_producing = 1
    mo.move_raw_ids.picked = True
    mo.button_mark_done()
    registro(mo, 'orden de corte con costeo por pieza')
    igual(mo.state, 'done', 'orden de corte terminada')

    def value(p):
        return mo.move_finished_ids.filtered(lambda m: m.product_id == p).value
    igual(value(side), 72.37, 'lateral: 66.67 tablero + 2 tapacanto + 3.70 general')
    igual(value(shelf), 35.19, 'repisa')
    igual(value(door), 204.70, 'puerta (roble)')
    check(abs(value(main) - 0.74) <= 0.011, 'producto principal (área 0.1) ≈ 0.74 (redondeo)', value(main))
    igual(mo.piece_total_cost, 313.0, 'costo total de la orden 313')
    check(abs(mo.piece_balance_diff) <= 1.0, 'cuadre de piezas vs. consumo', mo.piece_balance_diff)
    unbuild = env['mrp.unbuild'].create({'product_id': main.id, 'bom_id': bom.id, 'mo_id': mo.id, 'product_qty': 1})
    check(raises(UserError, unbuild.action_validate), 'no se puede desmontar una orden de corte')


ejecutar(_piece_costing)

caso('al_mrp_piece_costing', 'Importar despiece desde XLSX',
     'Importar el despiece de la obra QA20MC (2 piezas) desde Excel: crea las piezas con área, tablero y tapacanto.')


def _piece_import():
    import openpyxl
    P = S['piece']
    wb = openpyxl.Workbook()
    sh = wb.active
    sh.append(['Obra', 'Mueble terminado', 'Modulo', 'ID pieza', 'Material', 'Largo mm', 'Ancho mm',
               'Area m2', 'Tapacanto codigo', 'Tapacanto ml', 'Cantidad'])
    sh.append(['QA20MC', 'Cocina', 'MA01', '1', 'QA20-BL18', 700, 560, None, 'QA20-TBL22', 1.26, 2])
    sh.append(['QA20MC', 'Cocina', 'MA01', '2', 'QA20-RO18', 700, 300, 0.25, None, 0, 1])
    buf = io.BytesIO()
    wb.save(buf)
    wizard = env['mrp.piece.import.wizard'].create({'file': BinaryBytes(buf.getvalue()), 'file_name': 'despiece.xlsx',
                                                     'categ_id': P['categ'].id})
    wizard.action_import()
    igual(wizard.state, 'done', 'importación terminada')
    check('creadas: 2' in (wizard.summary or ''), 'resumen: 2 piezas creadas', wizard.summary)
    first = env['product.product'].search([('default_code', '=', 'QA20MC-1')])
    registro(first, 'pieza importada')
    check(first.piece_is_cut, 'la pieza queda marcada como pieza de corte')
    igual(first.piece_area_m2, 0.392, 'área calculada 0.7 × 0.56', 3)
    igual(first.piece_material_id, P['white'], 'tablero de la pieza')


ejecutar(_piece_import)


# ======================================================================
#  OL-HEALTH
# ======================================================================
def ingredient(name):
    return env['pharmacy.ingredient']._get_or_create(qa_name(name))


def med(name, ingredients=(), form=None, price=1.0, **vals):
    values = {
        'name': qa_name(name), 'is_storable': True, 'pharmacy_ok': True, 'list_price': price,
        'categ_id': S['categ_farma'].id, 'available_in_pos': True, 'sale_ok': True,
        'taxes_id': [Command.clear()], 'supplier_taxes_id': [Command.clear()],
    }
    if form:
        values['pharmacy_dosage_form_id'] = form.id
    if ingredients:
        values['pharmacy_ingredient_line_ids'] = [Command.create({
            'ingredient_id': ing.id, 'strength': strength, 'strength_unit': 'mg'}) for ing, strength in ingredients]
    values.update(vals)
    return registro(env['product.product'].create(values))


def qa_user(login, groups, name=None):
    user = env['res.users'].search([('login', '=', login)], limit=1)
    if not user:
        user = env['res.users'].create({
            'name': name or qa_name(login), 'login': login, 'company_id': company.id,
            'company_ids': [Command.set(company.ids)],
            'group_ids': [Command.set([env.ref(g).id for g in groups])]})
    return user


TABLET = env.ref('al_pharmacy_base.dosage_form_tablet')
SYRUP = env.ref('al_pharmacy_base.dosage_form_syrup')

caso('preparacion', 'Usuarios y ajustes de farmacia QA20',
     'Crear el técnico y la Q.F. de prueba; categoría Farmacia con salida FEFO.')


def _prep_pharmacy():
    S['tech'] = registro(qa_user('qa20_tecnico_farmacia', ['al_pharmacy_base.group_pharmacy_technician',
                                                           'stock.group_stock_user']), 'técnico de farmacia')
    S['qf'] = registro(qa_user('qa20_qf_farmacia', ['al_pharmacy_base.group_pharmacy_pharmacist',
                                                    'stock.group_stock_user',
                                                    'sales_team.group_sale_salesman_all_leads']), 'Q.F.')
    fefo = env.ref('product_expiry.removal_fefo')
    S['categ_farma'].removal_strategy_id = fefo
    check(fefo, 'estrategia de salida FEFO disponible (product_expiry)')
    S['wh_f'].pharmacy_establishment_code = '0099920'
    S['pharmacist_emp'] = registro(env['hr.employee'].create({
        'name': qa_name('Q.F. Rosa Quispe'), 'pharmacy_role': 'pharmacist', 'pharmacy_license_number': '99345',
        'user_id': S['qf'].id, 'company_id': company.id}), 'Q.F. empleada')


ejecutar(_prep_pharmacy)

# ---------------------------------------------------------------- al_pharmacy_base
caso('al_pharmacy_base', 'Maestro de medicamentos: composición, equivalentes y alternativas',
     'Registrar Apronax (innovador) y tres genéricos de naproxeno 550 mg: el sistema propone primero '
     'los genéricos DCI más baratos y vincula bioequivalentes en ambos sentidos.')


def _pharmacy_master():
    nap = ingredient('Naproxeno sódico')
    para, caf = ingredient('Paracetamol'), ingredient('Cafeína')
    combo = med('Panadol Extra', [(para, 500), (caf, 65)], TABLET, 9.0)
    igual(combo.pharmacy_composition, '%s 500 mg + %s 65 mg' % (para.name, caf.name), 'composición rotulada')
    brand = med('Apronax 550', [(nap, 550)], TABLET, 6.0, pharmacy_product_type='innovator')
    cheap = med('Naproxeno Genfar 550', [(nap, 550)], TABLET, 1.2, pharmacy_product_type='generic_dci')
    dear = med('Naproxeno Portugal 550', [(nap, 550)], TABLET, 2.0, pharmacy_product_type='generic_dci')
    bgen = med('Naprox 550', [(nap, 550)], TABLET, 1.0, pharmacy_product_type='generic_brand')
    other_form = med('Naproxeno suspensión', [(nap, 550)], SYRUP, 8.0, pharmacy_product_type='generic_dci')
    tmpl = brand.product_tmpl_id
    check(cheap.product_tmpl_id in tmpl.pharmacy_equivalent_ids, 'mismo principio, concentración y forma = equivalente')
    check(other_form.product_tmpl_id not in tmpl.pharmacy_equivalent_ids, 'otra forma farmacéutica no es equivalente')
    igual(tmpl._pharmacy_sorted_alternatives().ids,
          (cheap | dear | bgen).product_tmpl_id.ids, 'alternativas: genéricos DCI por precio y luego genérico de marca')
    base = ingredient('Naproxeno base')
    salt = med('Naproxeno 500 base', [(base, 500)], TABLET, 1.8)
    tmpl.pharmacy_bioequivalent_ids = [Command.link(salt.product_tmpl_id.id)]
    igual(salt.product_tmpl_id.pharmacy_bioequivalent_ids, tmpl, 'bioequivalencia recíproca')
    check(salt.product_tmpl_id in tmpl.pharmacy_equivalent_ids, 'el bioequivalente cuenta como equivalente')


ejecutar(_pharmacy_master)

caso('al_pharmacy_base', 'Presentaciones, cuaderno de ocurrencias y turnos del Q.F.',
     'Crear la presentación «Caja x 100» de paracetamol, anotar una ocurrencia en el cuaderno y '
     'registrar un turno de reemplazo del Q.F.')


def _pharmacy_establishment():
    prod = med('Paracetamol 500 caja', price=0.20)
    box = registro(env['pharmacy.presentation'].create({
        'product_tmpl_id': prod.product_tmpl_id.id, 'name': 'Caja', 'factor': 100, 'price': 12.0}))
    igual(box.display_name, 'Caja x 100', 'nombre de la presentación')
    igual(box.uom_id.relative_factor, 100, 'la presentación crea su unidad de 100')
    check(box.uom_id in prod.product_tmpl_id.uom_ids, 'la unidad queda en las unidades del producto')
    igual(box.unit_price, 0.12, 'precio unitario de la caja', 6)
    igual(box._get_quantity_in_presentations(300), 3, '300 unidades = 3 cajas')
    check(raises(ValidationError, env['pharmacy.presentation'].create, {
        'product_tmpl_id': prod.product_tmpl_id.id, 'name': 'Unidad', 'factor': 1, 'price': 0.2}),
        'una presentación de 1 unidad se rechaza')
    entry = registro(env['pharmacy.logbook.entry'].create({
        'entry_type': 'incident', 'description': qa_name('Corte de luz de 20 minutos.'),
        'warehouse_id': S['wh_f'].id}), 'ocurrencia')
    check(entry.name.startswith('OC/'), 'ocurrencia numerada', entry.name)
    check(raises(UserError, entry.write, {'description': 'Otra cosa'}), 'el cuaderno no se edita')
    check(raises(UserError, entry.unlink), 'el cuaderno no se borra')
    now = fields.Datetime.now()
    shift = registro(env['pharmacy.shift'].create({
        'employee_id': S['pharmacist_emp'].id, 'warehouse_id': S['wh_f'].id,
        'date_start': now - timedelta(hours=1), 'date_end': now + timedelta(hours=7), 'role': 'replacement'}))
    log = env['pharmacy.logbook.entry'].search([('entry_type', '=', 'absence'), ('warehouse_id', '=', S['wh_f'].id)],
                                               order='id desc', limit=1)
    check(S['pharmacist_emp'].name in (log.description or ''), 'el turno de reemplazo queda en el cuaderno')
    check(shift.employee_id in env['pharmacy.shift']._get_pharmacists_on_duty(S['wh_f']), 'Q.F. de turno en el local')
    html = env['ir.actions.report']._render_qweb_html('al_pharmacy_base.report_pharmacy_logbook', entry.ids)[0]
    check(b'Corte de luz' in html, 'reporte del cuaderno de ocurrencias')


ejecutar(_pharmacy_establishment)

# ---------------------------------------------------------- al_l10n_pe_pharmacy_edi
caso('al_l10n_pe_pharmacy_edi', 'Presentación en el comprobante electrónico (código SUNAT de la caja)',
     'Facturar 1 «Caja x 100» de amoxicilina: el XML UBL lleva unitCode BX y el precio de la caja.')


def _pharmacy_edi():
    prod = med('Amoxicilina 500 EDI', price=0.15)
    tmpl = prod.product_tmpl_id
    box = registro(env['pharmacy.presentation'].create({'product_tmpl_id': tmpl.id, 'name': 'Caja', 'factor': 100, 'price': 12.0}))
    bottle = env['pharmacy.presentation'].create({'product_tmpl_id': tmpl.id, 'name': 'Frasco', 'factor': 2, 'price': 0.3})
    blister = env['pharmacy.presentation'].create({'product_tmpl_id': tmpl.id, 'name': 'Blíster', 'factor': 10, 'price': 1.5})
    igual(box.uom_id.l10n_pe_edi_measure_unit_code, 'BX', 'Caja → BX (catálogo 03)')
    igual(bottle.uom_id.l10n_pe_edi_measure_unit_code, 'BO', 'Frasco → BO')
    igual(blister.uom_id.l10n_pe_edi_measure_unit_code, 'NIU', 'Blíster → código de la unidad del producto (NIU)')
    tax = env['account.tax'].search([('company_id', '=', company.id), ('type_tax_use', '=', 'sale'),
                                     ('amount', '=', 18), ('amount_type', '=', 'percent'),
                                     ('price_include', '=', False)], limit=1)
    move = env['account.move'].create({
        'move_type': 'out_invoice', 'partner_id': S['customer'].id, 'invoice_date': TODAY,
        'l10n_latam_document_type_id': env.ref('l10n_pe.document_type01').id,
        'invoice_line_ids': [Command.create({'product_id': prod.id, 'product_uom_id': box.uom_id.id,
                                             'price_unit': 12.0, 'quantity': 1, 'tax_ids': [Command.set(tax.ids)]})]})
    move.action_post()
    registro(move, 'factura con presentación caja')
    xml, errors = move._l10n_pe_edi_generate_invoice_bstr()
    check(not errors, 'XML UBL sin errores', errors)
    tree = etree.fromstring(xml if isinstance(xml, bytes) else xml.encode())
    qty = tree.find('.//{*}InvoiceLine/{*}InvoicedQuantity')
    igual(qty.get('unitCode'), 'BX', 'XML: unitCode BX')
    igual(float(qty.text), 1.0, 'XML: cantidad 1 caja')
    igual(float(tree.find('.//{*}InvoiceLine/{*}LineExtensionAmount').text), 12.0, 'XML: valor de la línea 12.00')
    omitido('Envío del CPE a SUNAT/OSE', 'sin credenciales: solo se arma el XML')


ejecutar(_pharmacy_edi)

# ---------------------------------------------------------------- al_pharmacy_stock
def _receipt(product, lot_name, qty=5, days=300, wh=None, partner=None):
    wh = wh or S['wh_f']
    sup = env.ref('stock.stock_location_suppliers')
    picking = env['stock.picking'].create({
        'picking_type_id': wh.in_type_id.id, 'partner_id': (partner or S['supplier']).id,
        'location_id': sup.id, 'location_dest_id': wh.lot_stock_id.id,
        'move_ids': [Command.create({'product_id': product.id, 'product_uom_qty': qty, 'uom_id': product.uom_id.id,
                                     'location_id': sup.id, 'location_dest_id': wh.lot_stock_id.id})]})
    picking.action_confirm()
    if product.tracking == 'lot':
        picking.move_ids.move_line_ids.write({
            'lot_name': lot_name, 'quantity': qty,
            'expiration_date': fields.Datetime.now() + timedelta(days=days)})
    else:
        picking.move_ids.quantity = qty
    picking.move_ids.picked = True
    return picking


def _delivery(product, qty, wh=None, partner=None, origin=False):
    wh = wh or S['wh_f']
    cust = env.ref('stock.stock_location_customers')
    picking = env['stock.picking'].create({
        'picking_type_id': wh.out_type_id.id, 'partner_id': (partner or S['patient']).id, 'origin': origin,
        'location_id': wh.lot_stock_id.id, 'location_dest_id': cust.id,
        'move_ids': [Command.create({'product_id': product.id, 'product_uom_qty': qty, 'uom_id': product.uom_id.id,
                                     'location_id': wh.lot_stock_id.id, 'location_dest_id': cust.id})]})
    picking.action_confirm()
    picking.action_assign()
    return picking


def _lot(product, name, days, qty=10):
    lot = env['stock.lot'].create({'name': name, 'product_id': product.id,
                                   'expiration_date': fields.Datetime.now() + timedelta(days=days)})
    if qty:
        env['stock.quant']._update_available_quantity(product, S['wh_f'].lot_stock_id, qty, lot_id=lot)
    return lot


caso('al_pharmacy_stock', 'Recepción con inspección, cuarentena y salida FEFO',
     'Recibir cetirizina por lote: sin inspección no se valida; con inspección no conforme el lote '
     'queda en cuarentena; al despachar se toma primero el lote que vence antes y nunca el de cuarentena.')


def _pharmacy_stock_reception():
    prod = med('Cetirizina 10 mg', tracking='lot', use_expiration_date=True,
               pharmacy_sanitary_registry='EN-QA2001',
               pharmacy_registry_expiry=TODAY + timedelta(days=700))
    S['cetirizina'] = prod
    pick = _receipt(prod, 'QA20-CTZ01', 10, days=400)
    check(raises(UserError, pick.button_validate), 'sin inspección técnica la recepción no se valida')
    pick.pharmacy_inspection = 'conform'
    pick.button_validate()
    registro(pick, 'recepción conforme')
    igual(pick.state, 'done', 'recepción conforme validada')
    bad = _receipt(prod, 'QA20-CTZ02', 10, days=60)
    bad.write({'pharmacy_inspection': 'nonconform', 'pharmacy_inspection_note': 'Cajas húmedas (QA20).'})
    bad.button_validate()
    registro(bad, 'recepción no conforme')
    lot_bad = bad.move_line_ids.lot_id
    igual(lot_bad.pharmacy_state, 'quarantine', 'recepción no conforme → lote en cuarentena')
    soon = _receipt(prod, 'QA20-CTZ03', 10, days=120)
    soon.pharmacy_inspection = 'conform'
    soon.button_validate()
    registro(soon)
    lot_soon = soon.move_line_ids.lot_id
    out = registro(_delivery(prod, 15), 'despacho FEFO')
    lines = out.move_ids.move_line_ids
    igual(lines.lot_id.mapped('name'), ['QA20-CTZ03', 'QA20-CTZ01'], 'FEFO: primero el lote que vence antes')
    check(lot_bad not in lines.lot_id, 'el lote en cuarentena no se reserva (aunque vence antes)')
    igual(sum(lines.mapped('quantity')), 15.0, 'reserva de 15 en lotes disponibles')
    check(raises(UserError, lot_bad.with_user(S['tech']).action_pharmacy_release), 'el técnico no libera la cuarentena')
    lot_bad.with_user(S['qf']).action_pharmacy_release()
    igual(lot_bad.pharmacy_state, 'available', 'el Q.F. libera el lote')
    prod.pharmacy_registry_expiry = TODAY - timedelta(days=1)
    exp = _receipt(prod, 'QA20-CTZ04', 5)
    exp.pharmacy_inspection = 'conform'
    check(raises(UserError, exp.button_validate), 'registro sanitario vencido: la recepción se bloquea')
    exp.action_cancel()
    prod.pharmacy_registry_expiry = TODAY + timedelta(days=700)
    out.action_cancel()


ejecutar(_pharmacy_stock_reception)

caso('al_pharmacy_stock', 'Alerta sanitaria, baja de lote y canje por vencer',
     'Registrar la alerta DIGEMID de un lote vendido: se bloquea y lista los clientes; dar de baja el '
     'lote bloqueado; cargar al canje con la droguería los lotes por vencer y emitir la nota de crédito.')


def _pharmacy_stock_alert():
    prod = med('Loratadina 10 mg', tracking='lot', use_expiration_date=True, standard_price=0.4,
               seller_ids=[Command.create({'partner_id': S['supplier'].id, 'price': 0.4})])
    lot = _lot(prod, 'QA20-LOR01', 400)
    sale = _delivery(prod, 3)
    sale.move_ids.picked = True
    sale.button_validate()
    registro(sale, 'venta del lote')
    alert = env['pharmacy.alert'].create({
        'name': qa_name('ALERTA 105-2026'),
        'line_ids': [Command.create({'product_tmpl_id': prod.product_tmpl_id.id, 'lot_names': 'QA20-LOR01'})]})
    alert.action_confirm()
    registro(alert, 'alerta sanitaria')
    igual(lot.pharmacy_state, 'blocked', 'el lote alertado queda bloqueado')
    check(S['patient'] in alert.customer_ids, 'la alerta identifica a los clientes que lo compraron')
    check(not _delivery(prod, 1).move_ids.move_line_ids, 'el lote bloqueado no se reserva')
    destruction = env['pharmacy.destruction'].create({
        'warehouse_id': S['wh_f'].id, 'method': 'Operador de residuos sólidos',
        'line_ids': [Command.create({'product_id': prod.id, 'lot_id': lot.id, 'location_id': S['wh_f'].lot_stock_id.id,
                                     'quantity': 7, 'reason': 'damaged'})]})
    destruction.action_validate()
    registro(destruction, 'acta de baja')
    igual(destruction.state, 'done', 'baja validada')
    igual(destruction.scrap_ids.state, 'done', 'desecho de inventario hecho')
    igual(env['stock.quant']._get_available_quantity(prod, S['wh_f'].lot_stock_id, lot_id=lot, strict=True), 0.0,
          'el lote dado de baja queda en 0')
    # Precio por vencer (escala propia de la categoría QA20) y canje
    env['pharmacy.expiry.discount'].create([{'days': 90, 'discount': 20, 'categ_id': S['categ_farma'].id},
                                            {'days': 30, 'discount': 40, 'categ_id': S['categ_farma'].id}])
    near = _lot(prod, 'QA20-LOR45', 45)
    expired = _lot(prod, 'QA20-LORVEN', -5, qty=4)
    _lot(prod, 'QA20-LOR400B', 400)
    igual(near.pharmacy_expiry_discount, 20, 'descuento por vencer a 45 días: 20 %')
    igual(_lot(prod, 'QA20-LOR20', 20, qty=0).pharmacy_expiry_discount, 40, 'a 20 días: 40 %')
    env['stock.lot']._cron_pharmacy_expire_lots()
    igual(expired.pharmacy_state, 'expired', 'el cron marca el lote vencido')
    ret = env['pharmacy.supplier.return'].create({'partner_id': S['supplier'].id, 'warehouse_id': S['wh_f'].id})
    env['pharmacy.supplier.return.load'].create({'return_id': ret.id, 'source': 'seller', 'days': 90}).action_load()
    registro(ret, 'canje con la droguería')
    igual(set(ret.line_ids.lot_id.mapped('name')), {'QA20-LOR45', 'QA20-LORVEN'}, 'canje: lotes por vencer y vencidos')
    igual(ret.amount_total, 14 * 0.4, 'canje valorizado a costo de compra')
    ret.action_send()
    ret.action_accept()
    ret.action_done()
    igual(ret.picking_id.state, 'done', 'devolución a la droguería despachada')
    ret.action_create_refund()
    registro(ret.refund_id, 'nota de crédito del proveedor (borrador)')
    igual((ret.refund_id.move_type, ret.refund_id.state), ('in_refund', 'draft'), 'nota de crédito de proveedor en borrador')
    igual(ret.refund_id.amount_untaxed, 5.6, 'importe de la nota de crédito')


ejecutar(_pharmacy_stock_alert)

# ---------------------------------------------------------- al_pharmacy_prescription
caso('al_pharmacy_prescription', 'Receta: validación, dispensación parcial con genérico y tope',
     'Registrar la receta de Zitromax 500 mg × 6, dispensar 3 del genérico (alternativa) y 3 de la '
     'marca; una 7.ª unidad se rechaza y la dispensación no se puede borrar.')


def _prescription():
    azi = ingredient('Azitromicina')
    brand = med('Zitromax 500', [(azi, 500)], TABLET, 8.0, pharmacy_sale_condition='prescription',
                pharmacy_product_type='innovator')
    generic = med('Azitromicina 500 tableta', [(azi, 500)], TABLET, 1.5, pharmacy_sale_condition='prescription',
                  pharmacy_product_type='generic_dci')
    S['azi'] = (azi, brand, generic)
    Rx = env['pharmacy.prescription']
    empty = Rx.create({'line_ids': [Command.create({'product_id': brand.id, 'quantity': 3})]})
    check(raises(ValidationError, empty.action_validate), 'sin prescriptor ni paciente la receta no se valida')
    empty.unlink() if empty.state == 'draft' and hasattr(empty, 'unlink') else None
    rx = Rx.create({'prescriber_id': S['prescriber'].id, 'patient_id': S['patient'].id,
                    'line_ids': [Command.create({'product_id': brand.id, 'quantity': 6, 'posology': '1 tableta al día'})]})
    rx.action_validate()
    registro(rx, 'receta')
    igual(rx.state, 'valid', 'receta validada')
    first = rx._dispense(generic, 3, source='QA20-TPV/0001')
    check(first.is_alternative, 'el genérico se registra como alternativa')
    igual(rx.state, 'partial', 'dispensación parcial')
    rx._dispense(brand, 3)
    igual(rx.state, 'dispensed', 'receta dispensada por completo')
    check(raises(UserError, rx._dispense, brand, 1), 'no se dispensa más de lo recetado')
    check(raises(UserError, first.unlink), 'la dispensación no se puede borrar')
    html = env['ir.actions.report']._render_qweb_html('al_pharmacy_prescription.report_prescription_book', rx.ids)[0]
    check(b'Libro de recetas' in html, 'libro de recetas imprime')
    S['patient'].pharmacy_allergy_ids = [Command.link(azi.id)]
    rx2 = Rx.create({'prescriber_id': S['prescriber'].id, 'patient_id': S['patient'].id,
                     'line_ids': [Command.create({'product_id': brand.id, 'quantity': 1})]})
    igual(rx2._allergy_warnings(), azi, 'alerta de alergia del paciente al principio activo')
    S['patient'].pharmacy_allergy_ids = [Command.unlink(azi.id)]
    old = Rx.create({'prescriber_id': S['prescriber'].id, 'patient_id': S['patient'].id,
                     'date': TODAY - timedelta(days=40), 'validity_days': 30,
                     'line_ids': [Command.create({'ingredient_id': azi.id, 'quantity': 3})]})
    old.action_validate()
    registro(old, 'receta vencida')
    check(raises(UserError, old._dispense, generic, 1), 'receta vencida no se dispensa')
    Rx._cron_expire_prescriptions()
    igual(old.state, 'expired', 'el cron marca la receta vencida')


ejecutar(_prescription)

# ---------------------------------------------------------- al_l10n_pe_prescriber_cmp
caso('al_l10n_pe_prescriber_cmp', 'Verificación del médico en el CMP',
     'Abrir «Consultar CMP» del médico (portal Conoce a tu médico) y registrar el resultado: solo el Q.F. '
     'puede marcarlo; un médico inhabilitado bloquea la receta.')


def _cmp():
    doc = registro(env['res.partner'].create({
        'name': qa_name('Dr. Jorge Salazar'), 'pharmacy_is_prescriber': True,
        'pharmacy_profession': 'physician', 'pharmacy_prescriber_license': '045678'}))
    action = doc.action_pharmacy_open_cmp()
    check(action['type'] == 'ir.actions.act_url' and 'conoceatumedico.cmp.org.pe' in action['url'],
          'acción abre el portal del CMP', action.get('url'))
    omitido('Consulta real al portal del CMP', 'servicio externo sin API: se verifica la URL generada')
    check(raises(UserError, doc.with_user(S['tech']).action_pharmacy_cmp_enabled), 'el técnico no puede verificar al médico')
    err = raises(AccessError, doc.with_user(S['qf']).action_pharmacy_cmp_enabled)
    check(not err, 'un Q.F. (grupo Químico farmacéutico, sin «Contactos: creación») registra la verificación',
          'AccessError: %s — _pharmacy_set_cmp_status escribe el contacto con los permisos del usuario' % err if err else '')
    qf = S['qf']
    if err:
        qf = qa_user('qa20_qf_contactos', ['al_pharmacy_base.group_pharmacy_pharmacist', 'base.group_partner_manager'])
        registro(qf, 'Q.F. con derecho a editar contactos (para seguir el flujo)')
    doc.with_user(qf).action_pharmacy_cmp_enabled()
    igual(doc.pharmacy_prescriber_status, 'enabled', 'Q.F. registra «Habilitado»')
    igual(doc.pharmacy_prescriber_checked_by_id, qf, 'queda quién verificó')
    doc.with_user(qf).action_pharmacy_cmp_disabled()
    prod = S['azi'][1] if S.get('azi') else med('Losartán 50 CMP')
    rx = env['pharmacy.prescription'].create({'prescriber_id': doc.id, 'patient_id': S['patient'].id,
                                              'line_ids': [Command.create({'product_id': prod.id, 'quantity': 30})]})
    registro(rx, 'receta de médico inhabilitado')
    check(raises(ValidationError, rx.action_validate), 'médico inhabilitado: la receta no se valida')
    nolic = env['res.partner'].create({'name': qa_name('Médico sin CMP'), 'pharmacy_is_prescriber': True})
    check(raises(UserError, nolic.action_pharmacy_open_cmp), 'sin colegiatura no se consulta el CMP')


ejecutar(_cmp)

# ---------------------------------------------------- al_l10n_pe_pharmacy_controlled
def _move_ctrl(product, qty, src, dest, origin=False, partner=False):
    wh = S['wh_f']
    ptype = wh.out_type_id if dest.usage == 'customer' else wh.in_type_id
    p = env['stock.picking'].create({
        'picking_type_id': ptype.id, 'location_id': src.id, 'location_dest_id': dest.id, 'origin': origin,
        'partner_id': partner and partner.id,
        'move_ids': [Command.create({'product_id': product.id, 'product_uom_qty': qty, 'uom_id': product.uom_id.id,
                                     'location_id': src.id, 'location_dest_id': dest.id})]})
    if ptype.code == 'incoming':
        p.pharmacy_inspection = 'conform'
    p.action_confirm()
    p.action_assign()
    p.move_ids.quantity = qty
    p.move_ids.picked = True
    p.with_context(skip_backorder=True).button_validate()
    return registro(p)


caso('al_l10n_pe_pharmacy_controlled', 'Controlados: lista, receta especial y folio de retención',
     'Registrar clonazepam (lista IVA) y morfina (lista IIA): la lista fija la condición y el libro; la '
     'morfina exige receta especial de 3 días y datos del comprador; la receta queda retenida con folio.')


def _controlled_rx():
    clz = med('Clonazepam 2 mg', pharmacy_sanitary_registry='EN-QA2003', pharmacy_control_list='IVA',
              pharmacy_sale_condition='prescription')
    mor = med('Morfina 10 mg/mL', pharmacy_sanitary_registry='EN-QA2004', pharmacy_control_list='IIA',
              pharmacy_sale_condition='controlled')
    S['clz'] = clz
    igual(mor.pharmacy_control_book, 'narcotic', 'lista IIA → libro de estupefacientes')
    igual(clz.pharmacy_control_book, 'psychotropic', 'lista IVA → libro de psicotrópicos')
    check(raises(ValidationError, mor.write, {'pharmacy_sale_condition': 'prescription'}),
          'la lista fija la condición de venta')

    def rx(product, kind='common', **vals):
        r = env['pharmacy.prescription'].create(dict({
            'prescription_type': kind, 'validity_days': 3 if kind == 'special' else 30,
            'prescriber_id': S['prescriber'].id, 'patient_id': S['patient'].id, 'diagnosis': 'Dolor oncológico',
            'line_ids': [Command.create({'product_id': product.id, 'quantity': 10, 'posology': '1 ampolla c/12 h'})]},
            **vals))
        r.action_validate()
        return registro(r)
    check(raises(ValidationError, rx, mor, 'special', validity_days=10), 'receta especial: vigencia máxima 3 días')
    common = rx(mor)
    check(raises(UserError, common._dispense, mor, 2, purchaser_name='X', purchaser_vat='40112233'),
          'lista IIA no se dispensa con receta común')
    special = rx(mor, 'special')
    check(raises(UserError, special._dispense, mor, 2), 'receta especial exige datos del comprador')
    special._dispense(mor, 2, purchaser_name='Ana Cárdenas', purchaser_vat='41223344')
    check(special.retained and (special.retention_folio or '').startswith('F-'), 'receta retenida con folio',
          special.retention_folio)
    igual(special.retention_until, TODAY.replace(year=TODAY.year + 2), 'archivo de la receta retenida por 2 años')


ejecutar(_controlled_rx)

caso('al_l10n_pe_pharmacy_controlled', 'Libro de psicotrópicos y balance trimestral',
     'Comprar 100 clonazepam y 50 metilfenidato, vender 10 y 5 con receta; imprimir el libro del 4.º '
     'trimestre y presentar el balance trimestral justificando una diferencia de stock.')


def _controlled_book():
    clz = S.get('clz') or med('Clonazepam 2 mg B', pharmacy_control_list='IVA', pharmacy_sale_condition='prescription')
    mph = med('Metilfenidato 10 mg', pharmacy_sanitary_registry='EN-QA2005', pharmacy_control_list='IIIA',
              pharmacy_sale_condition='controlled')
    sup, cust, stock = env.ref('stock.stock_location_suppliers'), env.ref('stock.stock_location_customers'), S['wh_f'].lot_stock_id
    _move_ctrl(clz, 100, sup, stock)
    rx = env['pharmacy.prescription'].create({
        'prescriber_id': S['prescriber'].id, 'patient_id': S['patient'].id,
        'line_ids': [Command.create({'product_id': clz.id, 'quantity': 10})]})
    rx.action_validate()
    rx._dispense(clz, 10, source='QA20-TPV/0042')
    registro(rx)
    _move_ctrl(clz, 10, stock, cust, origin='QA20-TPV/0042', partner=S['patient'])
    _move_ctrl(mph, 50, sup, stock)
    _move_ctrl(mph, 5, stock, cust)
    quarter = str((TODAY.month - 1) // 3 + 1)
    report = registro(env['pharmacy.controlled.report'].create({
        'book': 'psychotropic', 'warehouse_id': S['wh_f'].id, 'year': TODAY.year, 'quarter': quarter}))
    igual((report.date_from, report.date_to), (date(2026, 10, 1), date(2026, 12, 31)), 'trimestre 4 de 2026')
    rows = [r for r in report._get_book_lines() if r['product'] == clz]
    igual([(r['qty_in'], r['qty_out'], r['balance']) for r in rows], [(100.0, 0.0, 100.0), (0.0, 10.0, 90.0)],
          'libro: ingreso 100 y salida 10, saldo 90')
    igual(rows[1]['prescription'] if len(rows) > 1 else None, rx.retention_folio, 'la salida lleva el folio de la receta')
    html = env['ir.actions.report']._render_qweb_html('al_l10n_pe_pharmacy_controlled.report_controlled_book', report.ids)[0]
    check(b'Clonazepam' in html, 'libro de control impreso')
    report.action_export_xlsx()
    check(report.file_data, 'libro exportado a XLSX')
    bal = env['pharmacy.controlled.balance'].create({'warehouse_id': S['wh_f'].id, 'book': 'psychotropic',
                                                    'year': TODAY.year, 'quarter': quarter})
    registro(bal, 'balance trimestral')
    igual(bal.deadline, bal.date_to + timedelta(days=15), 'plazo de presentación: 15 días tras el trimestre')
    bal.action_compute()
    line = bal.line_ids.filtered(lambda l: l.product_id == mph)
    igual((line.initial, line.qty_in, line.qty_out, line.final, line.stock_qty), (0.0, 50.0, 5.0, 45.0, 45.0),
          'balance del metilfenidato (IIIA)')
    check(not bal.line_ids.filtered(lambda l: l.product_id == clz), 'la lista IVA no va al balance')
    line.stock_qty = 43.0
    igual(bal.difference_count, 1, 'diferencia con el stock detectada')
    check(raises(UserError, bal.action_submit), 'sin justificar la diferencia no se presenta')
    bal.difference_note = 'Merma autorizada por la DIRIS (acta QA20).'
    bal.action_submit()
    igual(bal.state, 'submitted', 'balance presentado')
    omitido('Presentación del balance ante DIGEMID/DIRIS', 'trámite externo; se valida el acta impresa')


ejecutar(_controlled_book)

# ---------------------------------------------------- al_l10n_pe_pharmacy_digemid
caso('al_l10n_pe_pharmacy_digemid', 'Observatorio de precios: pendientes y archivo de carga',
     'Cambiar el precio de un medicamento con código DIGEMID y exportar el archivo del Observatorio '
     'del local (código 0099920): primera carga completa con precios con IGV.')


def _digemid_obs():
    sita = ingredient('Sitagliptina')
    tax = env['account.tax'].search([('company_id', '=', company.id), ('type_tax_use', '=', 'sale'),
                                     ('amount', '=', 18), ('price_include', '=', False)], limit=1)
    prod = med('Sitagliptina 50 mg', [(sita, 50)], TABLET, 10.0, pharmacy_product_type='generic_dci',
               pharmacy_digemid_code='QA20-DIG-0100', taxes_id=[Command.set(tax.ids)])
    env['stock.quant']._update_available_quantity(prod, S['wh_f'].lot_stock_id, 20)
    Price = env['pharmacy.observatory.price']
    pend = Price.search([('product_tmpl_id', '=', prod.product_tmpl_id.id)])
    igual((pend.state, pend.unit_price), ('pending', 10.0), 'precio pendiente de reportar')
    prod.list_price = 12.0
    pend = Price.search([('product_tmpl_id', '=', prod.product_tmpl_id.id)])
    igual((len(pend), pend.unit_price), (1, 12.0), 'el cambio actualiza el mismo pendiente')
    Export = env['pharmacy.observatory.export']
    wiz = Export.create({'warehouse_id': S['wh_f'].id})
    check(wiz.needs_full, 'primera carga del mes: completa')
    wiz.action_export()
    import openpyxl
    rows = [[c.value for c in r] for r in openpyxl.load_workbook(io.BytesIO(wiz.file_data.content))['Hoja1'].iter_rows()]
    igual(rows[0], ['CodEstab', 'CodProd', 'Precio 1', 'Precio 2'], 'cabecera del formato del Observatorio')
    mine = [r for r in rows[1:] if r[1] == 'QA20-DIG-0100']
    igual(mine, [['0099920', 'QA20-DIG-0100', 14.16, 14.16]], 'fila con precio con IGV (12 × 1.18)')
    igual(pend.state, 'reported', 'el precio pasa a reportado')
    sub = env['pharmacy.observatory.submission'].search([('warehouse_id', '=', S['wh_f'].id)], limit=1)
    registro(sub, 'envío al Observatorio')
    igual((sub.state, sub.mode), ('sent', 'full'), 'envío registrado como carga completa')
    igual(sub.rectification_deadline, company._pharmacy_add_business_days(TODAY, 7), 'plazo de rectificación 7 días hábiles')
    omitido('Carga real en el Observatorio de DIGEMID', 'portal externo sin credenciales; se valida el archivo')
    S['wh_f'].pharmacy_establishment_code = False
    check(raises(UserError, Export.create({'warehouse_id': S['wh_f'].id}).action_export),
          'sin código de establecimiento no se exporta')
    S['wh_f'].pharmacy_establishment_code = '0099920'


ejecutar(_digemid_obs)

caso('al_l10n_pe_pharmacy_digemid', 'Genéricos esenciales y reacción adversa',
     'Calcular el cumplimiento de genéricos del petitorio en el local y registrar una sospecha de '
     'reacción adversa grave (plazo 24 h, queda en el cuaderno).')


def _digemid_generic():
    ing = ingredient('Metformina esencial')
    item = registro(env['pharmacy.essential.medicine'].create({'ingredient_id': ing.id, 'strength': '850 mg',
                                                                'dosage_form_id': TABLET.id}))
    brand = med('Glucophage 850', [(ing, 850)], TABLET, 1.5, pharmacy_product_type='innovator')
    env['stock.quant']._update_available_quantity(brand, S['wh_f'].lot_stock_id, 50)
    gen = med('Metformina 850 tableta', [(ing, 850)], TABLET, 0.3, pharmacy_product_type='generic_dci')
    env['stock.quant']._update_available_quantity(gen, S['wh_f'].lot_stock_id, 100)
    rep = registro(env['pharmacy.generic.compliance'].create({'warehouse_id': S['wh_f'].id}))
    rep.action_compute()
    line = rep.line_ids.filtered(lambda l: l.item_id == item)
    check(line.compliant, 'con stock del genérico DCI el ítem cumple')
    check(gen.product_tmpl_id in line.generic_product_ids, 'el genérico se lista')
    env['stock.quant']._update_available_quantity(gen, S['wh_f'].lot_stock_id, -100)
    rep.action_compute()
    check(not rep.line_ids.filtered(lambda l: l.item_id == item).compliant, 'sin stock del genérico: incumple')
    now = fields.Datetime.now()
    ev = registro(env['pharmacy.adverse.event'].create({
        'patient_initials': 'M.C.R.', 'product_id': gen.id, 'reaction': 'Acidosis láctica',
        'severity': 'serious', 'date_detected': now}))
    igual(ev.deadline, now + timedelta(hours=24), 'grave: notificar en 24 h')
    entry = env['pharmacy.logbook.entry'].search([('res_model', '=', 'pharmacy.adverse.event'), ('res_id', '=', ev.id)])
    igual(entry.entry_type, 'adverse_event', 'la sospecha queda en el cuaderno')
    ev.action_mark_notified()
    igual(ev.state, 'notified', 'marcada como notificada')
    omitido('Notificación a DIGEMID (e-reporting)', 'servicio externo')


ejecutar(_digemid_generic)

# ---------------------------------------------------------------- TPV de botica
import uuid as _uuid



def ui_order(lines, partner=None, payments=None, invoiced=False, **extra):
    """Pedido como lo envía el TPV (create_ui_order_data de los tests).
    lines: [(producto, cantidad, {extra de línea})]."""
    config, session = S['pos_config'], S['session']
    order_lines, total, base = [], 0.0, 0.0
    for product, qty, *rest in lines:
        kw = rest[0] if rest else {}
        price = kw.pop('price_unit', product.lst_price)
        taxes = product.taxes_id.filtered(lambda t: t.company_id == company)
        tv = taxes.compute_all(price, company.currency_id, qty) if taxes else \
            {'total_excluded': price * qty, 'total_included': price * qty}
        order_lines.append(Command.create(dict({
            'price_unit': price, 'product_id': product.id, 'qty': qty,
            'price_subtotal': abs(tv['total_excluded']), 'price_subtotal_incl': abs(tv['total_included']),
            'tax_ids': [Command.set(taxes.ids)]}, **kw)))
        sign = 1 if qty >= 0 else -1
        total += sign * abs(tv['total_included'])
        base += sign * abs(tv['total_excluded'])
    payments = payments or [(S['cash_pm'], total)]
    uid_ = str(_uuid.uuid4())
    data = dict({
        'amount_paid': sum(a for _p, a in payments), 'amount_return': 0, 'amount_tax': total - base,
        'amount_total': total, 'date_order': fields.Datetime.to_string(fields.Datetime.now()),
        'pricelist_id': config.pricelist_id.id, 'name': 'Order %s' % uid_, 'lines': order_lines,
        'partner_id': partner and partner.id, 'session_id': session.id, 'uuid': uid_, 'user_id': S['admin'].id,
        'to_invoice': invoiced, 'fiscal_position_id': False,
        'payment_ids': [Command.create({'amount': a, 'name': fields.Datetime.now(), 'payment_method_id': pm.id})
                        for pm, a in payments]}, **extra)
    env['pos.order'].sync_from_ui([data])
    return registro(env['pos.order'].search([('uuid', '=', uid_)]))


caso('al_pharmacy_pos', 'Dispensación en caja: receta, lote FEFO y rótulo fraccionado',
     'Vender en el TPV 10 cápsulas de amoxicilina (venta con receta) indicando receta 000123, prescriptor '
     'y lote: se crea la receta dispensada; la caja solo ofrece lotes disponibles, el que vence antes primero.')


def _pos_dispense():
    amox = med('Amoxicilina 500 cápsula', price=1.0, tracking='lot', use_expiration_date=True,
               pharmacy_sale_condition='prescription',
               pharmacy_laboratory_id=env['res.partner'].create({'name': qa_name('Laboratorios Andinos')}).id,
               pharmacy_route_id=env.ref('al_pharmacy_base.route_oral').id)
    late = _lot(amox, 'QA20-AMX-LATE', 300, 20)
    soon = _lot(amox, 'QA20-AMX-SOON', 30, 20)
    blocked = _lot(amox, 'QA20-AMX-BLOCK', 10, 5)
    blocked._pharmacy_set_state('blocked', 'Alerta QA20')
    lots = env['pos.order.line'].get_existing_lots(company.id, S['pos_config'].id, amox.id)
    igual([l['name'] for l in lots], ['QA20-AMX-SOON', 'QA20-AMX-LATE'], 'lotes ofrecidos: FEFO y sin el bloqueado')
    env['pharmacy.expiry.discount'].search([('categ_id', '=', S['categ_farma'].id)]) or \
        env['pharmacy.expiry.discount'].create([{'days': 90, 'discount': 20, 'categ_id': S['categ_farma'].id},
                                                {'days': 30, 'discount': 40, 'categ_id': S['categ_farma'].id}])
    info = env['pos.order.line'].get_pharmacy_lot_expiry(S['pos_config'].id, amox.id, ['QA20-AMX-SOON'])
    igual(info['QA20-AMX-SOON']['discount'], 40, 'precio por vencer sugerido en caja (lote a 30 días)')
    order = ui_order([(amox, 10, {'pack_lot_ids': [Command.create({'lot_name': 'QA20-AMX-SOON'})]})],
                     partner=S['patient'], pharmacy_prescription_number='000123',
                     pharmacy_prescriber_name=S['prescriber'].name,
                     pharmacy_prescriber_license=S['prescriber'].pharmacy_prescriber_license)
    check(order.state in ('paid', 'done'), 'venta pagada', order.state)
    rx = order.sudo().pharmacy_prescription_id
    registro(rx, 'receta creada desde el TPV')
    igual(rx.state, 'dispensed', 'receta registrada y dispensada')
    igual(rx.prescriber_id, S['prescriber'], 'prescriptor reconocido por su colegiatura')
    igual(rx.patient_id, S['patient'], 'paciente = cliente de la venta')
    igual(rx.dispensation_ids.lot_id, soon, 'la dispensación guarda el lote')
    igual(rx.dispensation_ids.source, order.name, 'origen de la dispensación = pedido del TPV')
    bad = ui_order([(amox, 2, {'pack_lot_ids': [Command.create({'lot_name': 'QA20-AMX-LATE'})]})],
                   partner=S['patient'], pharmacy_prescription_number='000999')
    entry = env['pharmacy.logbook.entry'].sudo().search([('res_model', '=', 'pos.order'), ('res_id', '=', bad.id)])
    check(not bad.sudo().pharmacy_prescription_id and 'no se pudo registrar' in ' '.join(entry.mapped('description')),
          'receta incompleta: la venta se guarda y el cuaderno lo anota')
    # Venta fraccionada: rótulo con lote y vencimiento
    lora = med('Loratadina 10 fracc', price=0.20, tracking='lot', use_expiration_date=True,
               pharmacy_laboratory_id=amox.pharmacy_laboratory_id.id)
    blister = env['pharmacy.presentation'].create({'product_tmpl_id': lora.product_tmpl_id.id, 'name': 'Blíster',
                                                   'factor': 10, 'price': 2.0})
    _lot(lora, 'QA20-LOR-FR', 400, 30)
    frac = ui_order([(lora, 3, {'pack_lot_ids': [Command.create({'lot_name': 'QA20-LOR-FR'})]})])
    whole = ui_order([(lora, 10, {'pharmacy_presentation_id': blister.id, 'price_unit': 0.2,
                                   'pack_lot_ids': [Command.create({'lot_name': 'QA20-LOR-FR'})]})])
    check(not whole._get_pharmacy_fraction_labels(), 'venta por blíster completo no lleva rótulo')
    labels = frac._get_pharmacy_fraction_labels()
    igual(len(labels), 1, 'venta de unidades sueltas genera un rótulo')
    igual(labels[0]['quantity'], 3, 'rótulo con la cantidad')
    html = env['ir.actions.report']._render_qweb_html('al_pharmacy_pos.report_pos_fraction_label', frac.ids)[0].decode()
    check('QA20-LOR-FR' in html and 'Vence:' in html, 'rótulo impreso con lote y vencimiento')
    omitido('Interfaz del TPV (popup de receta, selección de lote, descuento por vencer)', 'sin navegador: validado en backend')
    S['amox'] = amox


ejecutar(_pos_dispense)

# ---------------------------------------------------------------- al_pharmacy_loyalty
caso('al_pharmacy_loyalty', 'Club de salud: puntos por DNI sin medicamentos con receta',
     'Crear el programa «Club de Salud» (0.25 punto por sol) y vender a la paciente identificada por DNI un '
     'bloqueador (S/8) y amoxicilina con receta (S/40): solo suma 2 puntos; el 10 % de socio no alcanza a la receta.')


def _loyalty():
    otc = med('Bloqueador FPS 50', price=8.0, pharmacy_ok=False, pharmacy_sale_condition='otc')
    rx_prod = med('Amoxicilina 500 club', price=40.0, pharmacy_sale_condition='prescription')
    ctrl = med('Clonazepam 2 club', price=12.0, pharmacy_sale_condition='controlled')
    program = env['loyalty.program'].create({
        'name': qa_name('Club de Salud'), 'program_type': 'loyalty', 'applies_on': 'both', 'trigger': 'auto',
        'rule_ids': [Command.create({'reward_point_mode': 'money', 'reward_point_amount': 0.25})],
        'reward_ids': [Command.create({'reward_type': 'discount', 'discount': 10, 'discount_mode': 'percent',
                                       'discount_applicability': 'order', 'required_points': 1,
                                       'description': 'Descuento de socio'})]})
    registro(program, 'programa de fidelización')
    rule, reward = program.rule_ids, program.reward_ids
    check(program.pharmacy_exclude_prescription, 'exclusión de medicamentos con receta activa por defecto')
    check(otc in rule.valid_product_ids and rx_prod not in rule.valid_product_ids and ctrl not in rule.valid_product_ids,
          'la regla excluye receta y controlados')
    igual(reward.discount_applicability, 'specific', 'descuento «sobre el pedido» pasa a productos específicos')
    igual((otc | rx_prod | ctrl).filtered_domain(reward._get_discount_product_domain()), otc,
          'el descuento solo alcanza al producto de venta libre')
    check(raises(ValidationError, env['loyalty.reward'].create, {
        'program_id': program.id, 'reward_type': 'product', 'reward_product_id': rx_prod.id, 'required_points': 100}),
        'no se puede regalar un medicamento con receta')
    patient = S['patient']
    check(patient.additional_identifiers and patient.additional_identifiers.get('PE_DNI') == '40112233',
          'cliente identificado por DNI', patient.additional_identifiers)
    so = env['sale.order'].create({'partner_id': patient.id, 'order_line': [
        Command.create({'product_id': otc.id, 'product_uom_qty': 1}),
        Command.create({'product_id': rx_prod.id, 'product_uom_qty': 1})]})
    so._update_programs_and_rewards()
    registro(so, 'pedido con puntos')
    pts = so.coupon_point_ids.filtered(lambda p: p.coupon_id.program_id == program)
    igual(sum(pts.mapped('points')), 2.0, 'S/ 8 × 0.25 = 2 puntos; la amoxicilina no suma')
    so.action_confirm()
    card = env['loyalty.card'].search([('program_id', '=', program.id), ('partner_id', '=', patient.id)])
    registro(card, 'tarjeta del socio')
    igual(sum(card.mapped('points')), 2.0, 'puntos acreditados a la tarjeta del DNI 40112233')
    omitido('Canje de puntos en el TPV (tour)', 'interfaz del TPV sin navegador')
    # El programa es «automático» y aplicaría a los pedidos de las otras fases: se archiva.
    program.active = False


ejecutar(_loyalty)

# ---------------------------------------------------------------- al_pharmacy_insurance
caso('al_pharmacy_insurance', 'Convenio con aseguradora: copago en caja y factura a la EPS',
     'Vender paracetamol (10 + IGV) y losartán (20 + IGV) a una afiliada de «Salud Andina EPS»: el convenio '
     'cubre 19.24 y la paciente paga 16.16; al cerrar la caja se liquida y factura lo cubierto a la EPS.')


def _insurance():
    tax = env['account.tax'].create({'name': qa_name('IGV 18 % convenio'), 'amount': 18, 'amount_type': 'percent',
                                     'type_tax_use': 'sale', 'price_include_override': 'tax_excluded',
                                     'company_id': company.id})
    chronic = env['product.category'].create({'name': qa_name('Crónicos'), 'parent_id': S['categ_farma'].id})
    para = med('Paracetamol 500 convenio', price=10.0, pharmacy_ok=False, taxes_id=[Command.set(tax.ids)])
    los = med('Losartán 50 convenio', price=20.0, pharmacy_ok=False, taxes_id=[Command.set(tax.ids)],
              categ_id=chronic.id)
    insurer = registro(env['res.partner'].create({'name': qa_name('Salud Andina EPS S.A.'), 'is_company': True,
                                                  'vat': '20100049008', 'country_id': env.ref('base.pe').id}))
    patient = registro(env['res.partner'].create({'name': qa_name('María Torres Quispe'),
                                                  'additional_identifiers': {'PE_DNI': '44556677'}}))
    sale_journal = env['account.journal'].search([('type', '=', 'sale'), ('company_id', '=', company.id)], limit=1)
    ag = env['pharmacy.insurance.agreement'].create({
        'name': qa_name('Salud Andina · Plan regular'), 'partner_id': insurer.id, 'coverage_percent': 80,
        'copay_amount': 2, 'max_per_sale': 100, 'max_per_month': 30, 'date_start': TODAY - timedelta(days=60),
        'journal_id': sale_journal.id,
        'line_ids': [Command.create({'categ_id': chronic.id, 'coverage_percent': 50})]})
    aff = env['pharmacy.insurance.affiliate'].create({'agreement_id': ag.id, 'partner_id': patient.id,
                                                      'number': 'QA20-SA-001', 'date_start': TODAY - timedelta(days=30)})
    registro(ag, 'convenio')
    registro(aff, 'afiliada')
    ag.action_create_pos_payment_method()
    pm = ag.payment_method_id
    igual(pm.type, 'pay_later', 'método «Convenio» como cuenta de cliente')
    S['pos_config'].payment_method_ids = [Command.link(pm.id)]
    lines = [{'product': p, 'price_unit': pr, 'qty': 1, 'discount': 0.0,
              'taxes': p.taxes_id.filtered(lambda t: t.company_id == company)} for p, pr in ((para, 10.0), (los, 20.0))]
    cov = ag._compute_coverage(patient, lines, TODAY, False)
    check(not cov['error'], 'cobertura sin error', cov['error'])
    igual(cov['amount'], 19.24, 'cubre 80 % de 11.80 + 50 % de 23.60 − copago 2 = 19.24')
    rpc = env['pharmacy.insurance.agreement'].get_pos_coverage(pm.id, S['patient'].id, [
        {'product_id': para.id, 'price_unit': 10.0, 'qty': 1, 'discount': 0, 'tax_ids': tax.ids}], False)
    check('no es afiliado' in (rpc.get('error') or ''), 'en caja: un no afiliado se rechaza', rpc.get('error'))
    check(raises(ValidationError, ui_order, [(para, 1), (los, 1)], partner=patient,
                 payments=[(pm, 25.0), (S['cash_pm'], 10.40)]), 'la caja no acepta cubrir más de lo calculado')
    order = ui_order([(para, 1), (los, 1)], partner=patient, payments=[(pm, 19.24), (S['cash_pm'], 16.16)])
    igual(order.pharmacy_insurance_amount, 19.24, 'pedido con la parte cubierta')
    igual(order.pharmacy_insurance_affiliate_id, aff, 'pedido con la afiliada')
    igual(ag._compute_coverage(patient, lines, TODAY, False)['amount'], 10.76, 'tope mensual: quedan 30 − 19.24')
    S['insurance'] = (ag, order, patient)


ejecutar(_insurance)

# ---------------------------------------------------------------- al_pharmacy_care
caso('al_pharmacy_care', 'Tratamiento crónico: cobertura, recordatorio y WhatsApp',
     'Registrar el tratamiento de enalapril (1 tableta/día) de la paciente; con compras del genérico hace '
     '40 días y de la marca hace 15 (en el TPV) el tratamiento queda cubierto 20 días más y con PDC 67 %.')


def _care():
    ena = ingredient('Enalapril')
    brand = med('Enalabrand 10', [(ena, 10)], TABLET, 1.0, pharmacy_product_type='innovator')
    gen = med('Enalapril 10 tableta', [(ena, 10)], TABLET, 0.3, pharmacy_product_type='generic_dci')
    patient = registro(env['res.partner'].create({
        'name': qa_name('Julia Ramos Quispe'), 'phone': '987 654 321', 'email': 'qa20.julia@example.com',
        'country_id': env.ref('base.pe').id}))
    Treat = env['pharmacy.treatment']
    t = registro(Treat.create({'patient_id': patient.id, 'product_id': brand.id, 'daily_units': 1,
                               'date_start': TODAY - timedelta(days=60)}), 'tratamiento')
    igual(t.refill_state, 'no_purchase', 'sin compras: «sin compra»')

    def pos_purchase(product, qty, days_ago):
        order = env['pos.order'].create({
            'session_id': S['session'].id, 'partner_id': patient.id,
            'lines': [Command.create({'product_id': product.id, 'qty': qty, 'price_unit': 1.0,
                                      'price_subtotal': qty, 'price_subtotal_incl': qty})],
            'amount_tax': 0, 'amount_total': qty, 'amount_paid': qty, 'amount_return': 0,
            'payment_ids': [Command.create({'amount': qty, 'payment_method_id': S['cash_pm'].id})]})
        order.write({'state': 'paid', 'date_order': fields.Datetime.now() - timedelta(days=days_ago)})
        return registro(order)
    pos_purchase(gen, 30, 40)
    pos_purchase(brand, 30, 15)
    t._update_coverage()
    igual(t.last_purchase_date, TODAY - timedelta(days=15), 'última compra hace 15 días')
    igual(t.covered_until, TODAY + timedelta(days=20), 'compras encadenadas: cubierto hasta hoy + 20')
    igual(t.next_refill_date, TODAY + timedelta(days=15), 'próxima reposición (recordatorio 5 días antes)')
    igual(t.refill_state, 'ok', 'estado al día')
    igual(t.pdc, 67.2, 'PDC 41/61 días', 1)
    check(not t.adherent, 'PDC < 80 %: no adherente')
    check(t.whatsapp_url.startswith('https://wa.me/51987654321?text='), 'enlace de WhatsApp con +51', t.whatsapp_url)
    check(raises(ValidationError, Treat.create, {'patient_id': patient.id, 'product_id': brand.id, 'daily_units': 1,
                                                 'date_start': TODAY}), 'un solo tratamiento activo por producto')
    # Tratamiento atrasado → actividad de recordatorio (una sola vez)
    late_p = registro(env['res.partner'].create({'name': qa_name('Pedro Atrasado'), 'phone': '912345678'}))
    late = registro(Treat.create({'patient_id': late_p.id, 'product_id': gen.id, 'daily_units': 1,
                                  'date_start': TODAY - timedelta(days=60)}))
    order = env['pos.order'].create({
        'session_id': S['session'].id, 'partner_id': late_p.id,
        'lines': [Command.create({'product_id': gen.id, 'qty': 30, 'price_unit': 0.3,
                                  'price_subtotal': 9, 'price_subtotal_incl': 9})],
        'amount_tax': 0, 'amount_total': 9, 'amount_paid': 9, 'amount_return': 0,
        'payment_ids': [Command.create({'amount': 9, 'payment_method_id': S['cash_pm'].id})]})
    order.write({'state': 'paid', 'date_order': fields.Datetime.now() - timedelta(days=40)})
    late._update_coverage()
    act_type = env.ref('al_pharmacy_care.mail_activity_type_refill')
    Treat._cron_refill_reminders()
    igual(late.refill_state, 'late', 'tratamiento atrasado')
    igual(late.days_late, 10, '10 días de atraso')
    acts = late.activity_ids.filtered(lambda a: a.activity_type_id == act_type)
    igual(len(acts), 1, 'actividad de recordatorio creada')
    Treat._cron_refill_reminders()
    igual(len(late.activity_ids.filtered(lambda a: a.activity_type_id == act_type)), 1, 'el recordatorio no se repite')
    due = Treat.with_user(S['admin']).get_pos_due_treatments(patient.id, S['pos_config'].id)
    check(isinstance(due, list), 'el TPV consulta los tratamientos por reponer del cliente', due)
    # Desde la receta
    rx = env['pharmacy.prescription'].create({
        'prescriber_id': S['prescriber'].id, 'patient_id': patient.id,
        'line_ids': [Command.create({'product_id': gen.id, 'quantity': 30, 'posology': '1 tableta cada 24 horas'})]})
    rx.action_validate()
    rx.action_create_treatments()
    registro(rx.treatment_ids, 'tratamiento desde receta')
    igual(rx.treatment_ids.posology, '1 tableta cada 24 horas', 'el tratamiento toma la posología de la receta')
    omitido('Envío real del recordatorio por WhatsApp/correo', 'canal externo; se valida enlace y actividad')


ejecutar(_care)

# ---------------------------------------------------------------- al_pharmacy_replenishment
caso('al_pharmacy_replenishment', 'Pedido sugerido por ventas del TPV y faltantes',
     'Vender 30 losartán en caja, anotar faltantes (2 losartán y 1 insulina) y calcular el pedido sugerido '
     'a 15 días: 20 losartán (2 cajas de 10) a la Droguería; crear las órdenes de compra por proveedor.')


def _replenishment():
    units = env.ref('uom.product_uom_unit')
    box10 = env['uom.uom'].create({'name': qa_name('Caja x 10'), 'relative_factor': 10, 'relative_uom_id': units.id})
    vend_b = registro(env['res.partner'].create({'name': qa_name('Distribuidora del Sur'), 'is_company': True}))
    los = med('Losartán 50 reposición', price=0.5, seller_ids=[Command.create({
        'partner_id': S['supplier'].id, 'uom_id': box10.id, 'price': 3.0})])
    ins = med('Insulina NPH reposición', price=30.0, seller_ids=[Command.create({
        'partner_id': S['supplier'].id, 'price': 20.0})])
    env['stock.quant']._update_available_quantity(los, S['wh_f'].lot_stock_id, 35)
    W = env['pharmacy.replenishment.wizard']
    igual(W._suggested_quantity(1.0, 15, 2, 5, 0, 10), 20, 'fórmula: 1/día × 15 + 2 − 5 = 12 → múltiplo de 10 = 20')
    igual(W._suggested_quantity(1.0, 15, 0, 10, 5, 10), 0, 'cubierto por stock y lo que llega: 0')
    ui_order([(los, 30)])
    sh1 = env['pharmacy.shortage'].create({'product_id': los.id, 'quantity': 2, 'warehouse_id': S['wh_f'].id})
    sh2 = env['pharmacy.shortage'].create({'product_id': ins.id, 'quantity': 1, 'warehouse_id': S['wh_f'].id})
    free = env['pharmacy.shortage'].create({'product_id': False, 'quantity': 3, 'warehouse_id': S['wh_f'].id,
                                            'description': qa_name('Colirio lubricante 15 ml')})
    registro(sh1 | sh2 | free, 'faltantes')
    wiz = W.create({'warehouse_id': S['wh_f'].id, 'period_days': 30, 'coverage_days': 15})
    wiz.action_compute()
    line = wiz.line_ids.filtered(lambda l: l.product_id == los)
    igual((line.sold_qty, line.avg_daily, line.shortage_qty, line.stock_qty, line.multiple, line.suggested_qty),
          (30.0, 1.0, 2.0, 5.0, 10.0, 20.0), 'losartán: vendido 30, faltante 2, stock 5 → 20')
    igual(line.vendor_id, S['supplier'], 'proveedor del losartán')
    li = wiz.line_ids.filtered(lambda l: l.product_id == ins)
    igual(li.suggested_qty, 1.0, 'insulina sin ventas: pide el faltante')
    no_vendor = wiz.line_ids.filtered(lambda l: not l.vendor_id)
    if no_vendor:
        print('    sin proveedor (se quitan del sugerido):', no_vendor.product_id.mapped('name'))
        no_vendor.unlink()
    wiz.action_create_purchase()
    orders = env['purchase.order'].search([('origin', '=', 'Sugerido %s' % S['wh_f'].name)])
    registro(orders, 'orden de compra sugerida')
    po_line = orders.order_line.filtered(lambda l: l.product_id == los)
    igual((po_line.uom_id, po_line.product_qty), (box10, 2.0), 'OC en la unidad de compra: 2 cajas de 10')
    igual(set(orders.mapped('state')), {'draft'}, 'OC en borrador')
    igual(sh1.state, 'ordered', 'faltante del losartán pasa a pedido')
    igual(free.state, 'pending', 'faltante sin producto sigue pendiente')


ejecutar(_replenishment)

# ---------------------------------------------------------------- cierre de la caja y liquidación
caso('al_pharmacy_insurance', 'Cierre de caja y liquidación a la aseguradora',
     'Cerrar la caja de la botica (el «Convenio» no queda como cuenta por cobrar de la caja) y crear la '
     'liquidación del mes: factura a la EPS por 19.24 con el DNI y número de afiliado.')


def _close_and_liquidate():
    session = S['session']
    cash_total = sum(session.order_ids.payment_ids.filtered(lambda p: p.payment_method_id == S['cash_pm']).mapped('amount'))
    bank = company.partner_id.bank_ids.filtered(lambda b: not b.allow_out_payment)
    if bank:
        aviso('Cuenta bancaria no confiable en la compañía', '%s (de otra fase) impide a un usuario real publicar el '
              'cierre de caja y las facturas; se cierra como superusuario, que la omite' % ', '.join(bank.mapped('display_name')))
    session.close_session_from_ui({S['cash_pm'].id: cash_total})
    igual(session.state, 'closed', 'sesión de caja cerrada')
    if not S.get('insurance'):
        omitido('Liquidación', 'el caso del convenio no llegó a crear la venta')
        return
    ag, order, patient = S['insurance']
    pm = ag.payment_method_id
    lines = session.sale_move_ids.line_ids
    check(not lines.filtered(lambda l: pm.name in (l.name or '')), 'el cierre no registra cuenta por cobrar del convenio')
    liq = env['pharmacy.insurance.liquidation'].create({'agreement_id': ag.id, 'date_from': MONTH_START,
                                                        'date_to': TODAY})
    liq.action_load_orders()
    igual(liq.order_ids, order, 'la liquidación carga el pedido cubierto')
    liq.action_create_invoice()
    inv = liq.invoice_id
    if inv.partner_bank_id:
        aviso('Cuenta bancaria de la compañía no confiable en la factura',
              '%s (creada por otra fase) bloquea la publicación; se quita de la factura' % inv.partner_bank_id.display_name)
        inv.partner_bank_id = False
    inv.action_post()
    registro(liq, 'liquidación')
    registro(inv, 'factura a la aseguradora')
    igual(inv.move_type, 'out_invoice', 'factura de cliente')
    igual(inv.partner_id, ag.partner_id, 'facturada a la EPS')
    igual(inv.amount_total, 19.24, 'importe = lo cubierto')
    check('DNI 44556677' in inv.invoice_line_ids[0].name and 'QA20-SA-001' in inv.invoice_line_ids[0].name,
          'la línea identifica DNI y número de afiliado', inv.invoice_line_ids[0].name)
    income = lines.filtered(lambda l: l.tax_ids and l.tax_ids[0].amount == 18)
    taxl = lines.filtered(lambda l: l.tax_line_id and l.tax_line_id.amount == 18)
    patient_total = -sum(income.mapped('balance')) - sum(taxl.mapped('balance'))
    check(patient_total >= 16.16 - 0.01, 'el cierre contabiliza la parte del paciente (y otras ventas de la caja)',
          patient_total)
    second = env['pharmacy.insurance.liquidation'].create({'agreement_id': ag.id, 'date_from': MONTH_START,
                                                           'date_to': TODAY})
    second.action_load_orders()
    check(not second.order_ids, 'un pedido se liquida una sola vez')
    second.unlink()
    html = env['ir.actions.report']._render_qweb_html(
        'al_pharmacy_insurance.report_pharmacy_insurance_liquidation', liq.ids)[0]
    check(b'QA20-SA-001' in html, 'reporte de la liquidación')
    omitido('Envío del CPE de la factura a la EPS a SUNAT', 'sin credenciales')


ejecutar(_close_and_liquidate)

# ---------------------------------------------------------------- al_pharmacy_website
caso('al_pharmacy_website', 'Tienda web: genéricos primero, receta adjunta y aprobación del Q.F.',
     'En la tienda, la ficha de Amoxival muestra primero los genéricos; un controlado no entra al carrito; '
     'el pedido con receta queda esperando la imagen y la aprobación del Q.F.; la insulina sale con cadena de frío.')


def _website():
    import base64
    png = base64.b64decode(b'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGNgYGD4DwABBAEAwS2OUAAAAABJRU5ErkJggg==')
    website = company.website_id or env['website'].search([], limit=1)
    amx = ingredient('Amoxicilina web')
    common = {'is_published': True, 'pharmacy_sale_condition': 'prescription'}
    inn = med('Amoxival 500', [(amx, 500)], env.ref('al_pharmacy_base.dosage_form_capsule'), 3.0,
              pharmacy_product_type='innovator', **common)
    gen = med('Amoxicilina 500 web', [(amx, 500)], env.ref('al_pharmacy_base.dosage_form_capsule'), 0.8,
              pharmacy_product_type='generic_dci', **common)
    bgen = med('Amoxil genérico 500', [(amx, 500)], env.ref('al_pharmacy_base.dosage_form_capsule'), 0.5,
               pharmacy_product_type='generic_brand', **common)
    ctrl = med('Clonazepam web', price=1.0, is_published=True, pharmacy_sale_condition='controlled')
    ins = med('Insulina NPH web', price=30.0, is_published=True, pharmacy_sale_condition='otc', pharmacy_cold_chain=True)
    igual(inn.product_tmpl_id._pharmacy_website_alternatives().ids, (gen | bgen).product_tmpl_id.ids,
          'alternativas en la web: genérico DCI y luego de marca')
    cust = registro(env['res.partner'].create({'name': qa_name('Rosa Quispe web'), 'email': 'qa20.rosa@example.com',
                                               'street': 'Av. Arequipa 123', 'city': 'Lima',
                                               'country_id': env.ref('base.pe').id}))

    def web_order(product):
        return registro(env['sale.order'].create({'partner_id': cust.id, 'website_id': website.id,
                                                  'order_line': [Command.create({'product_id': product.id,
                                                                                 'product_uom_qty': 1})]}))
    o_ins = web_order(ins)
    qty, warning = o_ins._verify_updated_quantity(env['sale.order.line'], ctrl.id, 1, ctrl.uom_id.id)
    check(qty == 0 and 'controlado' in (warning or ''), 'un controlado no se agrega al carrito', warning)
    order = web_order(gen)
    check(order.pharmacy_prescription_required, 'el pedido web exige receta')
    igual(order.pharmacy_review_state, 'to_upload', 'estado: por subir receta')
    check(not order._is_cart_ready_for_payment(), 'sin receta no se paga')
    order._pharmacy_save_prescription(png, 'receta.png')
    igual(order.pharmacy_review_state, 'to_review', 'receta subida: por revisar')
    check(raises(UserError, order.action_confirm), 'no se confirma sin la aprobación del Q.F.')
    salesman = qa_user('qa20_vendedor_web', ['base.group_user', 'sales_team.group_sale_salesman'])
    check(raises(UserError, order.with_user(salesman).action_pharmacy_approve_prescription),
          'un vendedor no aprueba recetas')
    order.prepayment_percent = 0.0
    order.with_user(S['qf']).action_pharmacy_approve_prescription()
    igual(order.pharmacy_review_state, 'approved', 'receta aprobada')
    igual(order.pharmacy_reviewer_id, S['qf'], 'queda el Q.F. que revisó')
    igual(order.state, 'sale', 'pedido confirmado tras la aprobación')
    cold = registro(env['sale.order'].create({'partner_id': cust.id, 'warehouse_id': S['wh_f'].id,
                                              'order_line': [Command.create({'product_id': ins.id})]}))
    check(cold.pharmacy_cold_chain, 'pedido con cadena de frío')
    cold.action_confirm()
    pick = cold.picking_ids
    check(pick.pharmacy_cold_chain, 'el albarán hereda la cadena de frío')
    html = env['ir.actions.report']._render_qweb_html('stock.action_report_delivery', pick.ids)[0].decode()
    check('CADENA DE FRÍO' in html, 'el albarán impreso avisa «CADENA DE FRÍO»')
    omitido('Páginas de la tienda y subida de receta por el portal (HTTP)', 'sin navegador; controladores probados en los tests HttpCase')


ejecutar(_website)

# ======================================================================
#  OL-INDUSTRIES
# ======================================================================
INDUSTRY_DIR = os.path.join(MYAPPS, 'OL-INDUSTRIES')
for _mod in sorted(m for m in os.listdir(INDUSTRY_DIR) if m.startswith('al_industry_')):
    caso(_mod, 'Paquete de instalación: módulo y datos declarados presentes',
         'Abrir Aplicaciones y comprobar que el paquete está instalado con sus datos y menús.')

    def _industry(mod=_mod):
        rec = env['ir.module.module'].search([('name', '=', mod)])
        igual(rec.state, 'installed', 'módulo instalado')
        import ast
        manifest = ast.literal_eval(open(os.path.join(INDUSTRY_DIR, mod, '__manifest__.py')).read())
        deps = env['ir.module.module'].search([('name', 'in', manifest.get('depends', []))])
        missing = sorted(set(manifest.get('depends', [])) - set(deps.filtered(lambda m: m.state == 'installed').mapped('name')))
        check(not missing, 'dependencias instaladas (%d)' % len(manifest.get('depends', [])), missing)
        xmlids, menus = [], []
        for fname in manifest.get('data', []):
            path = os.path.join(INDUSTRY_DIR, mod, fname)
            if not fname.endswith('.xml'):
                continue
            tree = etree.parse(path)
            for node in tree.iter('record', 'menuitem', 'template', 'function'):
                if node.get('id'):
                    xid = node.get('id') if '.' in node.get('id') else '%s.%s' % (mod, node.get('id'))
                    xmlids.append(xid)
                    if node.tag == 'menuitem' or node.get('model') == 'ir.ui.menu':
                        menus.append(xid)
        absent = [x for x in xmlids if not env.ref(x, raise_if_not_found=False)]
        check(not absent, 'registros de datos declarados presentes (%d)' % len(xmlids), absent[:10])
        if menus:
            check(all(env.ref(m, raise_if_not_found=False) for m in menus), 'menús declarados presentes (%d)' % len(menus))
        if mod == 'al_industry_base':
            check('al.industry' in env or any(m.startswith('al.industry') for m in env.registry.models),
                  'modelo del asistente de industrias registrado',
                  [m for m in env.registry.models if 'industry' in m])
        omitido('Instalación sobre base vacía y datos demo', 'paquete para bases nuevas: no se reinstala aquí')

    ejecutar(_industry)
