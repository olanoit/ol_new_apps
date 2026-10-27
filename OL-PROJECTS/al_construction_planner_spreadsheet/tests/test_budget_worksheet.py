# -*- coding: utf-8 -*-
"""Hoja de trabajo de presupuesto y plan (requerimiento CREVAL v1.0, 6.7):
plantilla en Presupuestos › cotización, validación, envío al planificador con
el asistente del libro (6.6), costos vigentes y versiones."""
import base64
import json
from io import BytesIO

from odoo.tests import tagged

from odoo.addons.al_construction_planner_suite.tests.test_budget_book_creval import (
    TestBudgetBook,
)


def _sheets_from_book(content):
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(content), data_only=True)
    return {sheet.title: [['' if value is None else value for value in row]
                          for row in sheet.iter_rows(values_only=True)]
            for sheet in book.worksheets}


@tagged('post_install', '-at_install')
class TestBudgetWorksheet(TestBudgetBook):

    def _document(self):
        self.order.action_construction_open_spreadsheet()
        return self.env['documents.document'].search([
            ('construction_sale_order_id', '=', self.order.id),
            ('construction_source_id', '=', False)])

    def test_template_and_folder(self):
        document = self._document()
        self.assertEqual(len(document), 1)
        self.assertEqual(document.construction_sheet_kind, 'budget')
        self.assertEqual(document.folder_id.name, self.order.name)
        self.assertEqual(document.folder_id.folder_id.name, 'Presupuestos')
        data = json.loads(document.spreadsheet_data)
        names = [sheet['name'] for sheet in data['sheets']]
        self.assertEqual(names, ['INSTRUCCIONES', 'OBRA', 'INSUMOS', 'APU', 'APU_DETALLE',
                                 'PRESUPUESTO', 'PLAN', 'LISTAS', 'REFERENCIAS', 'ERRORES'])
        budget = next(sheet for sheet in data['sheets'] if sheet['name'] == 'PRESUPUESTO')
        self.assertEqual(budget['cells']['A1'], 'Ítem')
        # CR-18, c: columnas A a W del libro 6.6; M es el costo unitario.
        self.assertEqual(budget['cells']['F1'], 'Costo unitario sin APU')
        self.assertEqual(budget['cells']['W1'], 'Reserva por holguras')
        self.assertIn('VLOOKUP', budget['cells']['M2'])
        # CR-18, a: ninguna celda se guarda como número.
        for sheet in data['sheets']:
            for ref, content in sheet.get('cells', {}).items():
                self.assertIsInstance(content, str, '%s!%s' % (sheet['name'], ref))
        self.assertEqual(document._get_spreadsheet_metadata()['construction_sheet_kind'],
                         'budget')
        # La obra de la cotización abre la misma hoja.
        self.project.construction_sale_order_id = self.order
        self.project.action_construction_open_spreadsheet()
        self.assertEqual(self.env['documents.document'].search_count([
            ('construction_project_id', '=', self.project.id),
            ('construction_source_id', '=', False)]), 1)

    def test_send_validate_costs_and_versions(self):
        document = self._document()
        sheets = _sheets_from_book(base64.b64decode(self._book()))
        result = document.action_construction_validate(sheets)
        self.assertEqual(result['errors'], [])
        bad = dict(sheets, APU_DETALLE=sheets['APU_DETALLE'] + [['BK-R01', 'NO-EXISTE', '', '',
                                                                 1]])
        errors = document.action_construction_validate(bad)['errors']
        self.assertTrue(any('NO-EXISTE' in error[3] for error in errors))
        action = document.action_construction_send_to_planner(sheets, 'plan')
        wizard = self.env[action['res_model']].browse(action['res_id'])
        self.assertEqual(wizard.state, 'preview')
        self.assertEqual(wizard.spreadsheet_id, document)
        wizard.action_import()
        self.assertEqual(wizard.state, 'done')
        self.assertTrue(self.env['construction.resource.plan'].search(
            [('project_id', '=', self.project.id)]).line_ids)
        # Cada envío guarda una copia numerada.
        copies = self.env['documents.document'].search([('construction_source_id', '=',
                                                         document.id)])
        self.assertEqual(len(copies), 1)
        self.assertIn('v1', copies.name)
        self.assertEqual(document.construction_version, 2)
        document.action_construction_new_version()
        self.assertEqual(document.construction_version, 3)
        costs = document.construction_current_costs(['BK-LAD', 'X'])
        self.assertAlmostEqual(costs['BK-LAD'][0], 0.9661)
        rows = document.construction_odoo_rows()
        self.assertEqual(len(rows['PRESUPUESTO']), 3)
        self.assertEqual(rows['PRESUPUESTO'][0][5], 'Costo unitario sin APU')
        self.assertEqual(rows['PRESUPUESTO'][1][4], 'BK-R01')
        self.assertTrue(all(isinstance(value, str) for row in rows['PRESUPUESTO']
                            for value in row))

    def test_round_trip_keeps_a_foreign_budget(self):
        """CR-18, criterio 3: una cotización que vino de otro sistema (precio
        por partida, sin capas ni recetas) va a la hoja y vuelve igual, con su
        pie (gastos generales, utilidad y descuento)."""
        order = self.env['sale.order'].create({'partner_id': self.customer.id})
        product = self.env['construction.budget.import.wizard'].new(
            {'sale_order_id': order.id})._partida_product()
        Line = self.env['sale.order.line']
        Line.create({'order_id': order.id, 'display_type': 'line_section',
                     'name': 'Estructuras'})
        for item, qty, price in (('02.01', 158.9, 37.15), ('02.02', 12.5, 1210.4),
                                 ('02.03', 3, 845.333)):
            Line.create({'order_id': order.id, 'product_id': product.id,
                         'name': '%s Partida %s (test)' % (item, item),
                         'construction_budget_item': item, 'product_uom_qty': qty,
                         'price_unit': price})
        partidas = order.amount_untaxed
        for name, amount in (('Gastos generales 8 %', partidas * 0.08),
                             ('Utilidad 7 %', partidas * 0.07),
                             ('Descuento 24.9631 %', -partidas * 1.15 * 0.24963145)):
            Line.create({'order_id': order.id, 'product_id': product.id, 'name': name,
                         'product_uom_qty': 1, 'price_unit': order.currency_id.round(amount)})
        total = order.amount_untaxed
        order.action_construction_open_spreadsheet()
        document = self.env['documents.document'].search([
            ('construction_sale_order_id', '=', order.id), ('construction_source_id', '=', False)])
        rows = document.construction_odoo_rows()
        obra = dict(rows['OBRA'])
        self.assertAlmostEqual(float(obra['Gastos generales %']), 8.0, places=3)
        self.assertAlmostEqual(float(obra['Descuento %']), 24.963145, places=3)
        # Valores de la hoja: los datos tal cual y las fórmulas vacías (el
        # servidor las vuelve a calcular).
        budget = [[value if not value.startswith('=') else '' for value in row]
                  for row in rows['PRESUPUESTO']]
        sheets = {'OBRA': rows['OBRA'], 'INSUMOS': [['Código', 'Descripción', 'Tipo de insumo',
                                                     'Unidad']],
                  'APU': [['Código', 'Descripción', 'Unidad']],
                  'APU_DETALLE': [['Código del APU', 'Código del insumo']],
                  'PRESUPUESTO': budget}
        action = document.action_construction_send_to_planner(sheets, 'budget')
        wizard = self.env[action['res_model']].browse(action['res_id'])
        self.assertTrue(wizard.can_send, wizard.preview_html)
        wizard.action_import()
        self.assertAlmostEqual(order.amount_untaxed, total, delta=0.011)
        self.assertEqual(len(order.order_line.filtered(
            lambda l: not l.display_type and l.construction_line_role == 'partida')), 3)



# Las pruebas heredadas ya corren en su clase.
for _name in dir(TestBudgetBook):
    if _name.startswith('test_') and callable(getattr(TestBudgetBook, _name)) \
            and _name not in vars(TestBudgetWorksheet):
        setattr(TestBudgetWorksheet, _name, None)
