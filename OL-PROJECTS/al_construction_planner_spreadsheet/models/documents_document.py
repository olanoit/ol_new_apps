# -*- coding: utf-8 -*-
"""Hoja de trabajo de presupuesto y plan en Hojas de cálculo (requerimiento
CREVAL v1.0, 6.7).

La hoja tiene las mismas hojas y columnas del libro de carga (6.6): se valida,
se envía al planificador (destino Presupuesto o Plan) por el mismo asistente
del libro, trae el costo vigente de los insumos y los datos de Odoo, y guarda
una copia numerada en cada versión y en cada envío. Las obras de mobiliario
llenan su propio maestro con el puente al_construction_planner_spreadsheet_furniture."""
import base64
import json
from io import BytesIO

from markupsafe import Markup

from odoo import fields, models
from odoo.exceptions import UserError

SHEET_KINDS = [
    ('budget', 'Presupuesto y plan'),
    ('furniture', 'Maestro de mobiliario'),
]
# Pestañas de la hoja que van al libro de carga.
BOOK_SHEETS = ('OBRA', 'INSUMOS', 'APU', 'APU_DETALLE', 'PRESUPUESTO', 'PLAN', 'LISTAS')


class DocumentsDocument(models.Model):
    _inherit = 'documents.document'

    construction_project_id = fields.Many2one(
        'project.project', string='Obra (hoja de trabajo)', index='btree_not_null',
        ondelete='set null', readonly=True, copy=False,
        help='Obra cuya hoja de trabajo es esta.')
    construction_sale_order_id = fields.Many2one(
        'sale.order', string='Cotización (hoja de trabajo)', index='btree_not_null',
        ondelete='set null', readonly=True, copy=False)
    construction_sheet_kind = fields.Selection(
        SHEET_KINDS, string='Hoja de trabajo', readonly=True, copy=False)
    construction_version = fields.Integer(string='Versión de la hoja', default=1, copy=False,
                                          readonly=True)
    construction_source_id = fields.Many2one(
        'documents.document', string='Hoja de trabajo vigente', readonly=True, copy=False,
        ondelete='set null', help='En una copia numerada: la hoja de la que se copió.')

    def _get_spreadsheet_metadata(self, access_token=None):
        data = super()._get_spreadsheet_metadata(access_token)
        # sudo: solo se lee el vínculo de la hoja con su obra o cotización
        # para mostrar el menú de la vertical.
        document_sudo = self.sudo()
        data['construction_project_id'] = document_sudo.construction_project_id.id or False
        data['construction_sheet_kind'] = document_sudo.construction_sheet_kind or (
            'furniture' if document_sudo.construction_project_id else False)
        return data

    # ------------------------------------------------------------------
    # Enviar al planificador
    # ------------------------------------------------------------------
    def _construction_order(self):
        self.ensure_one()
        order = self.construction_sale_order_id or \
            self.construction_project_id.construction_sale_order_id
        if not order:
            raise UserError(self.env._(
                'La hoja no tiene cotización: ábrala desde la cotización (botón «Hoja de '
                'trabajo») o indique la orden de venta del contrato de la obra.'))
        return order

    def _construction_book(self, sheets):
        """Libro de carga (6.6) con las pestañas de la hoja."""
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, rows in (sheets or {}).items():
            if name not in BOOK_SHEETS:
                continue
            worksheet = workbook.create_sheet(name)
            for row in rows or []:
                worksheet.append([None if value in ('', False) else value for value in row])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _construction_book_wizard(self, sheets, target='budget'):
        self.ensure_one()
        order = self._construction_order()
        return self.env['construction.budget.book.wizard'].create({
            'sale_order_id': order.id,
            'project_id': (self.construction_project_id or order.project_id).id,
            'target': target,
            'file': base64.b64encode(self._construction_book(sheets)),
            'filename': '%s.xlsx' % self.name,
            'spreadsheet_id': self.id,
        })

    def action_construction_send_to_planner(self, sheets, target='budget'):
        """Lleva la hoja al planificador: el maestro de mobiliario a su
        importador; la de presupuesto y plan al asistente del libro (6.6),
        con su revisión antes de escribir.

        :param sheets: ``{pestaña: [[valor, …], …]}`` con los valores
            calculados, fila por fila (la primera, los encabezados).
        :param target: ``budget`` (presupuesto) o ``plan`` (presupuesto y plan).
        """
        self.ensure_one()
        if not self.construction_project_id and not self.construction_sale_order_id:
            raise UserError(self.env._('Esta hoja no es una hoja de trabajo del planificador.'))
        if self.construction_sheet_kind != 'budget':
            return self._construction_send_furniture(sheets)
        wizard = self._construction_book_wizard(sheets, target)
        wizard.action_preview()
        return {
            'type': 'ir.actions.act_window', 'name': self.env._('Enviar al planificador'),
            'res_model': wizard._name, 'res_id': wizard.id, 'view_mode': 'form',
            'views': [(False, 'form')], 'target': 'new',
        }

    def _construction_send_furniture(self, sheets):
        """Maestro de mobiliario: lo implementa su puente."""
        raise UserError(self.env._(
            'Esta hoja es el maestro de mobiliario: instale el planificador de mobiliario.'))

    def action_construction_validate(self, sheets):
        """«Validar plantilla»: los errores de la hoja como filas de la
        pestaña ERRORES (hoja, fila, columna, mensaje) y los avisos."""
        self.ensure_one()
        wizard = self._construction_book_wizard(sheets)
        _book, errors, warnings = wizard._parse()
        wizard.unlink()
        return {
            'errors': [[sheet, row or '', column or '', message]
                       for sheet, row, column, message in errors],
            'warnings': warnings,
        }

    def construction_current_costs(self, codes):
        """«Actualizar insumos»: {código: (costo vigente, unidad)} del maestro
        de productos para los códigos de la hoja INSUMOS."""
        self.ensure_one()
        company = (self.construction_project_id.company_id or
                   self.construction_sale_order_id.company_id or self.env.company)
        products = self.env['product.product'].with_company(company).search([
            ('default_code', 'in', [str(code) for code in codes or [] if code]),
            ('company_id', 'in', (company.id, False))])
        return {product.default_code: [product.standard_price, product.uom_id.name]
                for product in products}

    def construction_odoo_rows(self):
        """«Traer de Odoo»: las filas de las pestañas con lo que ya tienen la
        cotización y las recetas de la obra (sin encabezados)."""
        self.ensure_one()
        order = self.construction_sale_order_id or \
            self.construction_project_id.construction_sale_order_id
        project = self.construction_project_id or order.project_id
        return self.env['project.project']._construction_budget_rows(order, project)

    def action_construction_new_version(self):
        """«Nueva versión»: una copia numerada de la hoja en la misma carpeta,
        con el enlace en el historial de la cotización o del plan."""
        self.ensure_one()
        return self._construction_save_version(self.env._('Nueva versión de la hoja'))

    def _construction_save_version(self, reason):
        self.ensure_one()
        version = self.construction_version
        copy = self.copy({
            'name': '%s · v%s' % (self.name, version),
            'folder_id': self.folder_id.id,
            'construction_source_id': self.id,
            'construction_version': version,
        })
        self.construction_version = version + 1
        body = Markup('%s: <a href="#" data-oe-model="documents.document" data-oe-id="%s">%s</a>'
                      ) % (reason, copy.id, copy.name)
        order = self.construction_sale_order_id or \
            self.construction_project_id.construction_sale_order_id
        if order:
            order.message_post(body=body)
        plan = self.construction_project_id._construction_current_plan() \
            if self.construction_project_id else False
        if plan:
            plan.message_post(body=body)
        return {'type': 'ir.actions.client', 'tag': 'display_notification', 'params': {
            'type': 'success', 'message': self.env._('Versión %s guardada: %s', version,
                                                     copy.name)}}


class ConstructionBudgetBookWizard(models.TransientModel):
    _inherit = 'construction.budget.book.wizard'

    spreadsheet_id = fields.Many2one('documents.document', string='Hoja de trabajo',
                                     readonly=True)

    def action_import(self):
        res = super().action_import()
        # Cada envío deja una copia numerada de la hoja (6.7).
        if self.spreadsheet_id and self.state == 'done':
            self.spreadsheet_id._construction_save_version(self.env._(
                'Hoja enviada al planificador (%s)', dict(self._fields['target'].selection)[
                    self.target]))
        return res


def spreadsheet_json(data):
    return base64.b64encode(json.dumps(data).encode())
