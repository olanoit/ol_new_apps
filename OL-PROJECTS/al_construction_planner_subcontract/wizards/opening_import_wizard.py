# -*- coding: utf-8 -*-
"""Carga inicial de órdenes de contrata abiertas (requerimiento CREVAL v1.0,
3.14). Para empezar con obras en curso, una plantilla Excel crea las órdenes
con sus condiciones y una valorización de apertura sin efecto contable: guarda
por línea la cantidad acumulada valorizada y, en la orden, lo valorizado,
facturado y pagado, el adelanto otorgado y amortizado y el fondo retenido con
sus vencimientos. Las valorizaciones siguientes parten de ella."""
import io
from collections import defaultdict
from datetime import date, datetime, timedelta

from odoo import Command, fields, models
from odoo.exceptions import UserError

from odoo.addons.al_construction_planner.models.common import norm

SHEETS = {
    'ORDENES': [
        ('ref', 'Referencia', True), ('partner', 'Contratista', True), ('project', 'Obra', True),
        ('condition', 'Condiciones', False), ('specialty', 'Especialidad', False),
        ('advance_pct', 'Adelanto (%)', False), ('guarantee_pct', 'Fondo de garantía (%)', False),
        ('valued', 'Valorizado anterior', False), ('invoiced', 'Facturado anterior', False),
        ('paid', 'Pagado anterior', False), ('advance_granted', 'Adelanto otorgado', False),
        ('advance_amortized', 'Adelanto amortizado', False),
        ('count', 'Nº de valorizados', False), ('cutoff', 'Fecha de corte', True),
    ],
    'LINEAS': [
        ('ref', 'Referencia', True), ('task', 'Partida', True), ('activity', 'Actividad', True),
        ('qty', 'Cantidad contratada', True), ('price', 'Precio', True),
        ('valued_qty', 'Cantidad valorizada', False),
    ],
    'FONDOS': [
        ('ref', 'Referencia', True), ('amount', 'Monto', True), ('date', 'Vencimiento', True),
    ],
}


class ConstructionContractOpeningWizard(models.TransientModel):
    _name = 'construction.contract.opening.wizard'
    _description = 'Carga inicial de órdenes de contrata'

    file = fields.Binary(string='Libro Excel', required=True)
    filename = fields.Char(string='Nombre del archivo')
    log = fields.Text(string='Resultado', readonly=True)

    # ------------------------------------------------------------------
    # Lectura
    # ------------------------------------------------------------------
    def _read(self):
        try:
            import openpyxl
        except ImportError as error:
            raise UserError(self.env._('Falta la librería openpyxl en el servidor.')) from error
        book = openpyxl.load_workbook(io.BytesIO(self.file.content), data_only=True)
        sheets = {norm(name): book[name] for name in book.sheetnames}
        data = {}
        for sheet, columns in SHEETS.items():
            ws = sheets.get(norm(sheet))
            if ws is None:
                if sheet == 'FONDOS':
                    data[sheet] = []
                    continue
                raise UserError(self.env._('Falta la hoja %s en el libro.', sheet))
            rows = list(ws.iter_rows(values_only=True))
            header = [norm(cell) for cell in (rows[0] if rows else [])]
            index = {}
            for key, label, required in columns:
                if norm(label) in header:
                    index[key] = header.index(norm(label))
                elif required:
                    raise UserError(self.env._(
                        'La hoja %(sheet)s no tiene la columna «%(column)s».', sheet=sheet,
                        column=label))
            data[sheet] = [
                {key: row[i] for key, i in index.items()} for row in rows[1:]
                if row and any(cell not in (None, '') for cell in row)]
        return data

    @staticmethod
    def _date(value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return fields.Date.to_date(str(value).strip()) if value else False

    @staticmethod
    def _float(value):
        if value in (None, ''):
            return 0.0
        if isinstance(value, (int, float)):
            return float(value)
        return float(str(value).replace(',', ''))

    def _find(self, model, text, fields_names, extra=None):
        text = str(text or '').strip()
        if not text:
            return self.env[model]
        Model = self.env[model]
        for name in fields_names:
            if name in Model._fields:
                record = Model.search([(name, '=ilike', text)] + (extra or []), limit=1)
                if record:
                    return record
        return Model

    # ------------------------------------------------------------------
    # Carga
    # ------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        data = self._read()
        lines_by_ref = defaultdict(list)
        for row in data['LINEAS']:
            lines_by_ref[str(row['ref']).strip()].append(row)
        funds_by_ref = defaultdict(list)
        for row in data['FONDOS']:
            funds_by_ref[str(row['ref']).strip()].append(row)
        errors, created = [], self.env['purchase.order']
        for row in data['ORDENES']:
            ref = str(row['ref']).strip()
            try:
                with self.env.cr.savepoint():
                    created |= self._import_order(row, lines_by_ref[ref], funds_by_ref[ref])
            except (UserError, ValueError) as error:
                errors.append('%s: %s' % (ref, error.args[0] if error.args else error))
        log = [self.env._('Órdenes creadas: %s.', ', '.join(created.mapped('name')) or '—')]
        if errors:
            log.append(self.env._('Con errores (no se crearon):'))
            log += ['   · %s' % e for e in errors]
        self.log = '\n'.join(log)
        return {'type': 'ir.actions.act_window', 'res_model': self._name, 'res_id': self.id,
                'view_mode': 'form', 'target': 'new'}

    def _import_order(self, row, line_rows, fund_rows):
        ref = str(row['ref']).strip()
        partner = self._find('res.partner', row['partner'], ['vat', 'name'])
        project = self._find('project.project', row['project'], ['construction_code', 'name'],
                             [('is_construction_site', '=', True)])
        if not partner or not project:
            raise UserError(self.env._('No se encuentra el contratista o la obra.'))
        plan = project._construction_current_plan()
        if not plan or plan.state not in ('approved', 'in_progress'):
            raise UserError(self.env._('La obra %s no tiene un plan aprobado.', project.display_name))
        condition = self._find('construction.contract.condition', row.get('condition'), ['name'],
                               [('order_id', '=', False)]) \
            or project.company_id.construction_default_condition_id
        if not condition:
            raise UserError(self.env._('Indique las condiciones de contrata de la orden.'))
        if not line_rows:
            raise UserError(self.env._('La orden no tiene líneas en la hoja LINEAS.'))
        cutoff = self._date(row['cutoff'])
        Wizard = self.env['construction.plan.contract.wizard']
        order = self.env['purchase.order'].create({
            'partner_id': partner.id, 'company_id': plan.company_id.id,
            'partner_ref': ref, 'construction_is_service_order': True,
            'construction_project_id': project.id, 'construction_plan_id': plan.id,
            'construction_specialty': row.get('specialty') or False,
            'origin': self.env._('Carga inicial %s', ref)})
        overrides = {}
        if row.get('advance_pct') not in (None, ''):
            overrides['advance_pct'] = self._float(row['advance_pct'])
        if row.get('guarantee_pct') not in (None, ''):
            overrides['guarantee_pct'] = self._float(row['guarantee_pct'])
        terms = order._construction_apply_terms(condition, overrides)
        taxes_default = terms.scheme_line_ids.filtered(lambda l: l.kind == 'tax') \
            and plan.company_id.account_purchase_tax_id
        valued = {}
        allocations = []
        for line_row in line_rows:
            task = self._find('project.task', line_row['task'], ['name'],
                              [('project_id', '=', project.id)])
            activity = self._find('construction.labor.activity', line_row['activity'],
                                  ['code', 'name'])
            if not task or not activity:
                raise UserError(self.env._(
                    'No se encuentra la partida «%(task)s» o la actividad «%(activity)s».',
                    task=line_row['task'], activity=line_row['activity']))
            plan_lines = plan.line_ids.filtered(
                lambda l: l.task_id == task and l.activity_id == activity
                and l.resource_type == 'contract')
            if not plan_lines:
                raise UserError(self.env._(
                    'El plan no tiene la contrata %(activity)s en %(task)s.',
                    activity=activity.display_name, task=task.display_name))
            product = Wizard._get_service_product(activity)
            qty = self._float(line_row['qty'])
            taxes = product.supplier_taxes_id.filtered(lambda t: t.company_id == plan.company_id) \
                or taxes_default
            po_line = self.env['purchase.order.line'].create({
                'order_id': order.id, 'product_id': product.id,
                'name': '%s · %s' % (task.display_name, activity.display_name),
                'product_qty': qty, 'uom_id': activity.uom_id.id,
                'price_unit': self._float(line_row['price']),
                'tax_ids': [Command.set(taxes.ids)] if taxes else [Command.clear()],
                'date_planned': fields.Datetime.now(),
                'construction_activity_id': activity.id, 'construction_task_id': task.id,
                'analytic_distribution': plan_lines[:1].analytic_distribution or (
                    {str(project.account_id.id): 100.0} if project.account_id else False),
            })
            remaining = qty
            for plan_line in plan_lines:
                share = min(remaining, max(plan_line.qty_planned - plan_line.qty_requested, 0.0))
                if plan_line == plan_lines[-1]:
                    share = remaining
                if share > 0:
                    allocations.append({'plan_line_id': plan_line.id, 'kind': 'service_order',
                                        'purchase_line_id': po_line.id, 'qty_allocated': share})
                    remaining -= share
            plan_lines.with_context(construction_plan_force=True).write(
                {'partner_id': partner.id})
            valued[po_line] = self._float(line_row.get('valued_qty'))
        self.env['construction.resource.plan.allocation'].create(allocations)
        order.button_confirm()
        order.write({
            'construction_opening_valued': self._float(row.get('valued')),
            'construction_opening_invoiced': self._float(row.get('invoiced')),
            'construction_opening_paid': self._float(row.get('paid')),
            'construction_opening_advance_granted': self._float(row.get('advance_granted')),
            'construction_opening_advance_amortized': self._float(row.get('advance_amortized')),
        })
        for fund_row in fund_rows:
            self.env['construction.guarantee.fund'].create({
                'purchase_order_id': order.id, 'amount': self._float(fund_row['amount']),
                'date_maturity': self._date(fund_row['date']), 'is_opening': True})
        opening = self.env['construction.contract.settlement'].create({
            'purchase_order_id': order.id, 'settlement_type': 'opening',
            'period_start': cutoff - timedelta(days=6), 'date_cutoff': cutoff,
            'number': int(self._float(row.get('count'))),
        })
        opening.sudo().write({'line_ids': [Command.create({
            'activity_id': po_line.construction_activity_id.id, 'purchase_line_id': po_line.id,
            'qty_period': qty, 'price_unit': po_line.price_unit, 'retention_pct': 0.0,
        }) for po_line, qty in valued.items() if qty]})
        for po_line, qty in valued.items():
            if qty:
                po_line.qty_received = qty
        opening._terms_recompute()
        opening.with_context(al_construction_approving=True).write(
            {'state': 'approved', 'date_approved': fields.Datetime.now()})
        order.message_post(body=self.env._(
            'Orden cargada con la valorización de apertura %s (sin efecto contable).',
            opening.display_name))
        return order

    def action_download_template(self):
        """Plantilla vacía con las tres hojas."""
        import openpyxl
        book = openpyxl.Workbook()
        book.remove(book.active)
        for sheet, columns in SHEETS.items():
            ws = book.create_sheet(sheet)
            ws.append([label for _key, label, _required in columns])
        buffer = io.BytesIO()
        book.save(buffer)
        attachment = self.env['ir.attachment'].create({
            'name': 'carga_inicial_ordenes_de_contrata.xlsx',
            'raw': buffer.getvalue(),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {'type': 'ir.actions.act_url', 'target': 'self',
                'url': '/web/content/%s?download=true' % attachment.id}
