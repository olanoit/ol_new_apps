# -*- coding: utf-8 -*-
"""Fase 12: «Importar maestro y ETO» (W-15)."""
import base64
import importlib.util
from collections import defaultdict
from pathlib import Path

from odoo.exceptions import AccessError
from odoo.tests import new_test_user, tagged

from odoo.addons.al_construction_planner.wizards.master_import_layout import build_workbook

from .common import PlannerCommon, load_demo

MODULE_DIR = Path(__file__).parents[1]


def _examples():
    """Constantes del script que genera los libros de ejemplo."""
    path = MODULE_DIR / 'tools' / 'generar_ejemplos_importacion.py'
    spec = importlib.util.spec_from_file_location('generar_ejemplos_importacion', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


EX = _examples()
ML = {'Mueble bajo': 'low', 'Mueble alto': 'high', 'Sin ML': 'none'}


@tagged('post_install', '-at_install')
class TestMasterImport(PlannerCommon):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.manager = new_test_user(
            cls.env, 'plan_admin_import', name='Administrador del planificador (test)',
            groups='al_construction_planner.group_planner_manager')
        cls.site = cls.env['project.project'].create({
            'name': 'Obra de ejemplo (test)', 'is_construction_site': True})

    def _wizard(self, content, user=None, project=None, **vals):
        return self.env['construction.master.import.wizard'].with_user(
            user or self.manager).create(dict({
                'project_id': (project or self.site).id, 'file': base64.b64encode(content),
                'filename': 'maestro.xlsx'}, **vals))

    def _example(self, name):
        return (MODULE_DIR / 'static' / 'examples' / name).read_bytes()

    def _import(self, content, **kwargs):
        wizard = self._wizard(content, **kwargs)
        data = wizard._analyze()
        self.assertFalse(data['errors'])
        wizard.action_import()
        return data

    # ------------------------------------------------------------------
    # Ejemplos de static/examples
    # ------------------------------------------------------------------
    def test_examples_are_up_to_date(self):
        """Los libros publicados salen del script (mismas hojas y filas)."""
        from openpyxl import load_workbook
        from io import BytesIO
        master, eto = EX.build()
        for built, name in ((master, EX.MASTER_FILE), (eto, EX.ETO_FILE)):
            published = load_workbook(BytesIO(self._example(name)), read_only=True)
            generated = load_workbook(BytesIO(built), read_only=True)
            self.assertEqual(published.sheetnames, generated.sheetnames)
            for sheet in generated.sheetnames:
                self.assertEqual(list(published[sheet].values), list(generated[sheet].values),
                                 '%s · %s: regenere con tools/generar_ejemplos_importacion.py'
                                 % (name, sheet))

    def test_example_master_eto_and_plan(self):
        """El ejemplo se importa tal cual, con sus advertencias a propósito, y
        el plan generado cuadra con lo calculado desde el libro."""
        wizard = self._wizard(self._example(EX.MASTER_FILE), create_missing_products=True)
        data = wizard._analyze()
        self.assertFalse(data['errors'])
        expected = {'bom_no_stage', 'bom_duplicate', 'product_no_cost'}
        digits = self.env['decimal.precision'].precision_get('Product Unit')
        if digits < 4:
            expected.add('qty_precision')
        self.assertEqual({k for k, v in data['warnings'].items() if v}, expected)
        self.assertEqual(data['warnings']['bom_no_stage'], ['C01 · EJ-MEL-RH'])
        self.assertEqual(data['warnings']['product_no_cost'], ['EJ-PIN'])
        # 0.5 planchas × 195.00 × 6 cocinas C01.
        self.assertAlmostEqual(data['unstaged_amount'], 585.0, places=2)
        self.assertIn('EJ-MEL-RH', wizard.warning_html)
        self.assertIn('Tipologías', wizard.preview_html)
        wizard.action_import()

        typologies = self.env['construction.typology'].search([('project_id', '=', self.site.id)])
        self.assertEqual(sorted(typologies.mapped('code')), ['B01', 'C01', 'C02', 'CL01'])
        by_code = {t.code: t for t in typologies}
        self.assertEqual(by_code['CL01'].family, 'closet')
        self.assertEqual(sum(typologies.mapped('module_count')), 17)
        self.assertAlmostEqual(by_code['C01'].ml_low, 1.85)
        self.assertAlmostEqual(by_code['C02'].ml_high, 2.1)
        tornillo = by_code['C02'].bom_id.bom_line_ids.filtered(
            lambda l: l.product_id.default_code == 'EJ-TOR-450')
        self.assertEqual(tornillo.product_qty, 100)
        self.assertFalse(by_code['C01'].bom_id.bom_line_ids.filtered(
            lambda l: l.product_id.default_code == 'EJ-MEL-RH').construction_consumption_stage)
        tapacanto = self.env['product.product'].search([('default_code', '=', 'EJ-TAP-CAN')])
        self.assertEqual(tapacanto.uom_id, self.meter)
        self.assertEqual(tapacanto.company_id, self.site.company_id or self.env.company)
        Task = self.env['project.task']
        spaces = Task.search([('project_id', '=', self.site.id),
                              ('construction_level', '=', 'space')])
        self.assertEqual(len(spaces), 36)

        eto = self._wizard(self._example(EX.ETO_FILE))
        data = eto._analyze()
        self.assertEqual({k for k, v in data['warnings'].items() if v}, {'eto_unknown_space'})
        eto.action_import()
        modules = Task.search([('project_id', '=', self.site.id),
                               ('construction_level', '=', 'module')])
        self.assertEqual(len(modules), 26)

        plan = self.env['construction.resource.plan'].create({'project_id': self.site.id})
        self.env['construction.plan.generate.wizard'].create(
            {'plan_id': plan.id}).action_generate()
        modules = Task.search([('project_id', '=', self.site.id),
                               ('construction_level', '=', 'module')])
        self.assertEqual(len(modules), 126)

        # Monto de contratas esperado, desde las constantes del ejemplo.
        rates = {code: price for code, _n, _s, _u, _m, _g, price in EX.CATALOG}
        groups = {code: ML.get(group, 'none') for code, _n, _s, _u, measure, group, _p
                  in EX.CATALOG if measure == 'ML'}
        expected_contract = 0.0
        lines = 0
        for floor, apartment, space, typology in EX.tree():
            eto_widths = dict(EX.ETO.get((floor, apartment, space), []))
            mods = [(code, eto_widths.get(code, width), ML[group], activity)
                    for code, _k, width, group, activity in EX.MODULES[typology]]
            for _code, _w, _g, activity in mods:
                expected_contract += rates[activity]
                lines += 1
            for activity, qty in EX.ACTIVITIES[typology]:
                if activity in groups:
                    for _code, width, group, _a in mods:
                        if group == groups[activity] and width:
                            expected_contract += round(width / 1000.0 * rates[activity], 2)
                            lines += 1
                else:
                    expected_contract += round(qty * rates[activity], 2)
                    lines += 1
        contract = plan.line_ids.filtered(lambda l: l.resource_type == 'contract')
        self.assertEqual(len(contract), lines)
        self.assertAlmostEqual(plan.amount_contract, expected_contract, places=2)
        # El ETO baja la instalación al módulo con su ancho real.
        mb02 = modules.filtered(lambda t: t.construction_module_code == 'MB02'
                                and t.construction_apartment_task_id.name == 'Dpto 101')
        install = contract.filtered(lambda l: l.task_id == mb02 and l.stage == 'installation')
        self.assertEqual(install.qty_planned, 0.75)
        # Materiales: una línea por fila de BOM (las repetidas, sumadas) y ambiente.
        material = plan.line_ids.filtered(lambda l: l.resource_type == 'material')
        per_space = {t: len({p for p, _q, _s in rows}) for t, rows in EX.BOM.items()}
        self.assertEqual(len(material), sum(per_space[t] for *_x, t in EX.tree()))
        self.assertEqual(len(material.filtered(lambda l: not l.stage)), 6)
        self.assertTrue(any('Maestro importado' in (body or '')
                            for body in self.site.message_ids.mapped('body')))

    def test_reimport_is_idempotent(self):
        content = self._example(EX.MASTER_FILE)
        self._import(content, create_missing_products=True)
        self._import(self._example(EX.ETO_FILE))

        def snapshot():
            typologies = self.env['construction.typology'].search(
                [('project_id', '=', self.site.id)])
            return (len(typologies), len(typologies.module_line_ids),
                    len(typologies.activity_line_ids), len(typologies.bom_id.bom_line_ids),
                    self.env['project.task'].search_count([('project_id', '=', self.site.id)]),
                    self.env['product.product'].search_count([('default_code', '=like', 'EJ-%')]),
                    self.env['construction.labor.activity'].search_count(
                        [('code', '=like', 'EJ-%')]))

        before = snapshot()
        data = self._import(content, create_missing_products=True)
        self.assertFalse(data['products_to_create'])
        self._import(self._example(EX.ETO_FILE))
        self.assertEqual(snapshot(), before)

    def test_book_replaces_typology_template(self):
        """Lo que trae el libro de una tipología queda exactamente así."""
        self._import(self._example(EX.MASTER_FILE), create_missing_products=True)
        rows = EX.master_rows()
        rows['MODULOS'] = [r for r in rows['MODULOS'] if r['code'] != 'CAMPANA']
        rows['BOM'] = [r for r in rows['BOM'] if r['product_code'] != 'EJ-SIL']
        self._import(build_workbook(rows), create_missing_products=True)
        c01 = self.env['construction.typology'].search(
            [('project_id', '=', self.site.id), ('code', '=', 'C01')])
        self.assertNotIn('CAMPANA', c01.module_line_ids.mapped('code'))
        self.assertNotIn('EJ-SIL', c01.bom_id.bom_line_ids.product_id.mapped('default_code'))

    # ------------------------------------------------------------------
    # Ida y vuelta con el demo del piso 05
    # ------------------------------------------------------------------
    def test_demo_floor_roundtrip(self):
        """El piso 05 del demo, exportado al formato del maestro e importado
        en otra obra, genera el mismo plan (66 módulos, S/ 2,301.53 de
        contratas y S/ 6,642.91 de material al costo de los productos)."""
        demo = load_demo(self.env, 'IMP DEMO')
        rows = defaultdict(list)
        stage_labels = {'production': 'Producción', 'assembly': 'Armado',
                        'installation': 'Instalación', 'finishing': 'Acabado y entrega'}
        for code, typology in demo['typologies'].items():
            rows['TIPOLOGIAS'].append({'code': code, 'name': typology.name, 'family': 'Cocina',
                                       'ml_low': typology.ml_low, 'ml_high': typology.ml_high})
            for module in typology.module_line_ids:
                rows['MODULOS'].append({
                    'typology': code, 'sequence': module.sequence, 'code': module.code,
                    'type': dict(module._fields['module_type'].selection)[module.module_type],
                    'activity': module.assembly_activity_id.code})
            for line in typology.activity_line_ids:
                rows['ACTIVIDADES'].append({'typology': code, 'activity': line.activity_id.code,
                                            'qty': line.qty})
            for line in typology.bom_id.bom_line_ids:
                rows['BOM'].append({
                    'typology': code, 'product_code': line.product_id.default_code,
                    'qty': line.product_qty,
                    'stage': stage_labels.get(line.construction_consumption_stage, '')})
        for space in self.env['project.task'].search([
                ('project_id', '=', demo['project'].id), ('construction_level', '=', 'space')]):
            rows['ARBOL'].append({'floor': space.construction_floor_task_id.name,
                                  'apartment': space.construction_apartment_task_id.name,
                                  'space': space.name,
                                  'typology': space.construction_typology_id.code})
        data = self._import(build_workbook(rows))
        self.assertFalse(data['warnings'].get('product_not_found'))
        # Las tarifas de la obra (P-14) no son parte del maestro: se copian.
        rates = self.env['construction.labor.rate'].search(
            [('project_id', '=', demo['project'].id)])
        rates.copy({'project_id': self.site.id})
        plan = self.env['construction.resource.plan'].create({'project_id': self.site.id})
        self.env['construction.plan.generate.wizard'].create(
            {'plan_id': plan.id}).action_generate()
        for line in plan.line_ids.filtered(lambda l: l.resource_type == 'material'):
            line.price_unit_planned = line.product_id.standard_price
        self.assertEqual(self.env['project.task'].search_count([
            ('project_id', '=', self.site.id), ('construction_level', '=', 'module')]), 66)
        self.assertAlmostEqual(plan.amount_contract, 2301.53, places=2)
        self.assertAlmostEqual(plan.amount_material, 6642.91, places=2)
        self.assertAlmostEqual(plan.amount_total, 8944.44, places=2)

    # ------------------------------------------------------------------
    # Advertencias, errores, permisos y multicompañía
    # ------------------------------------------------------------------
    def test_warnings_and_errors(self):
        self.p_white.default_code = 'TEST-MDP'
        rows = {
            'TIPOLOGIAS': [{'code': 'T1', 'name': 'Cocina T1'}],
            'MODULOS': [
                {'typology': 'T1', 'code': 'M1', 'type': 'Bajo', 'activity': 'ARM-TOR'},
                {'typology': 'T1', 'code': 'M2', 'type': 'Bajo', 'activity': 'NO-EXISTE'},
                {'typology': 'T1', 'code': 'M3', 'type': 'Raro', 'activity': 'ARM-TOR'},
                {'typology': 'T1', 'code': 'M4', 'type': 'Bajo', 'activity': 'COL-PUE'},
                {'typology': 'T9', 'code': 'M1', 'type': 'Bajo', 'activity': 'ARM-TOR'}],
            'BOM': [{'typology': 'T1', 'product_code': 'TEST-MDP', 'qty': 2,
                     'stage': 'Producción'},
                    {'typology': 'T1', 'product_code': 'NO-EXISTE', 'qty': 1,
                     'stage': 'Armado'}],
            'ARBOL': [{'floor': 'P1', 'apartment': 'D1', 'space': 'Cocina', 'typology': 'T7'}],
            'ETO': [{'floor': 'P9', 'apartment': 'D9', 'space': 'Cocina', 'code': 'M1',
                     'width': 600}],
        }
        data = self._wizard(build_workbook(rows), user=self.planner)._analyze()
        self.assertFalse(data['errors'])
        warnings = data['warnings']
        self.assertEqual(warnings['unknown_activity'], ['MODULOS · NO-EXISTE'])
        self.assertEqual(warnings['activity_not_module'], ['M4 · COL-PUE'])
        self.assertEqual(warnings['unknown_typology'], ['MODULOS · T9'])
        self.assertEqual(warnings['product_not_found'], ['NO-EXISTE'])
        self.assertEqual(warnings['tree_unknown_typology'], ['P1 › D1 › Cocina (T7)'])
        self.assertEqual(warnings['eto_unknown_space'], ['P9 › D9 › Cocina'])
        self.assertEqual(len(warnings['invalid_value']), 1)
        self.assertEqual(set(data['modules']['T1']), {'M1'})

        # Columna obligatoria que falta y archivo que no es un libro: errores.
        from openpyxl import Workbook
        from io import BytesIO
        workbook = Workbook()
        workbook.active.title = 'BOM'
        workbook.active.append(['Tipología', 'Cantidad'])
        output = BytesIO()
        workbook.save(output)
        wizard = self._wizard(output.getvalue())
        self.assertIn('Código del producto', wizard._analyze()['errors'][0])
        self.assertTrue(wizard.error_html)
        self.assertTrue(self._wizard(b'no es un libro')._analyze()['errors'])

    def test_planner_imports_without_catalog(self):
        """El planificador importa tipologías, BOM (con sudo justificado) y
        árbol; la hoja CATALOGO solo la aplica un administrador."""
        self.p_white.default_code = 'TEST-MDP'
        rows = {
            'CATALOGO': [{'code': 'X-NUEVA', 'name': 'Nueva', 'stage': 'Armado', 'uom': 'UND'}],
            'TIPOLOGIAS': [{'code': 'T1', 'name': 'Cocina T1'}],
            'MODULOS': [{'typology': 'T1', 'code': 'MB01', 'type': 'Bajo', 'width': 600,
                         'ml_group': 'Mueble bajo', 'activity': 'ARM-TOR'}],
            'ACTIVIDADES': [{'typology': 'T1', 'activity': 'INS-BAJ', 'qty': 0.6}],
            'BOM': [{'typology': 'T1', 'product_code': 'TEST-MDP', 'qty': 2,
                     'stage': 'Producción'}],
            'ARBOL': [{'floor': 'P1', 'apartment': 'D1', 'space': 'Cocina', 'typology': 'T1'}],
        }
        data = self._import(build_workbook(rows), user=self.planner)
        self.assertEqual(data['warnings']['catalog_no_rights'], ['X-NUEVA'])
        self.assertFalse(self.env['construction.labor.activity'].search(
            [('code', '=', 'X-NUEVA')]))
        typology = self.env['construction.typology'].search(
            [('project_id', '=', self.site.id), ('code', '=', 'T1')])
        self.assertEqual(typology.bom_id.bom_line_ids.product_id, self.p_white)
        self.assertEqual(typology.ml_low, 0.6)
        space = self.env['project.task'].search([
            ('project_id', '=', self.site.id), ('construction_level', '=', 'space')])
        self.assertEqual(space.construction_typology_id, typology)
        self.assertEqual(space.construction_floor_task_id.name, 'P1')

    def test_progress_user_cannot_import(self):
        with self.assertRaises(AccessError):
            self._wizard(self._example(EX.ETO_FILE), user=self.reporter)

    def test_multicompany(self):
        """En una obra de otra compañía no se usan las actividades ni los
        productos de esta, y lo creado queda en la compañía de la obra."""
        other = self.env['res.company'].create({'name': 'Otra compañía (test)'})
        self.manager.write({'company_ids': [(4, other.id)]})
        site = self.env['project.project'].with_company(other).create({
            'name': 'Obra de otra compañía (test)', 'is_construction_site': True,
            'company_id': other.id})
        self.p_white.write({'default_code': 'TEST-MDP', 'company_id': self.env.company.id})
        rows = {
            'TIPOLOGIAS': [{'code': 'T1', 'name': 'Cocina T1'}],
            'MODULOS': [{'typology': 'T1', 'code': 'MB01', 'type': 'Bajo',
                         'activity': 'ARM-TOR'}],
            'BOM': [{'typology': 'T1', 'product_code': 'TEST-MDP', 'product_name': 'Melamina',
                     'qty': 2, 'stage': 'Producción'}],
        }
        wizard = self.env['construction.master.import.wizard'].with_user(self.manager) \
            .with_company(other).create({
                'project_id': site.id, 'file': base64.b64encode(build_workbook(rows)),
                'create_missing_products': True})
        data = wizard._analyze()
        self.assertEqual(data['warnings']['unknown_activity'], ['MODULOS · ARM-TOR'])
        self.assertIn('TEST-MDP', data['products_to_create'])
        wizard.action_import()
        typology = self.env['construction.typology'].search([('project_id', '=', site.id)])
        self.assertEqual(typology.company_id, other)
        product = typology.bom_id.bom_line_ids.product_id
        self.assertEqual(product.company_id, other)
        self.assertNotEqual(product, self.p_white)
        self.assertEqual(typology.bom_id.company_id, other)
