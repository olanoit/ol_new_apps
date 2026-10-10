# -*- coding: utf-8 -*-
"""Formato del libro «Maestro de planificación» que lee «Importar maestro y
ETO» (W-15).

Python puro (sin Odoo): lo usan el asistente, para leer el libro y generar la
plantilla vacía, y ``tools/generar_ejemplos_importacion.py``, para generar los
libros de ejemplo de ``static/examples``. Las columnas se reconocen por su
encabezado (sin distinguir mayúsculas ni tildes), en cualquier orden; todas
las hojas son opcionales.
"""
import unicodedata
from io import BytesIO


def norm(text):
    """Texto comparable: sin tildes, en minúsculas y con un solo espacio."""
    text = unicodedata.normalize('NFKD', str(text or ''))
    text = ''.join(c for c in text if not unicodedata.combining(c))
    return ' '.join(text.lower().split())


# Hoja → (descripción, [(clave, encabezado, obligatoria, ayuda)]).
SHEETS = {
    'CATALOGO': ('Actividades de obra (driver de pago). Solo la aplica un administrador del '
                 'planificador.', [
        ('code', 'Código', True, 'Código único de la actividad.'),
        ('name', 'Nombre', True, 'Nombre de la actividad.'),
        ('stage', 'Etapa', True, 'Producción, Armado, Instalación o Acabado y entrega.'),
        ('uom', 'Unidad', True, 'Unidad del driver: UND, ML o el nombre de la unidad.'),
        ('measure', 'Se mide por', False, 'Módulo, ML o Ambiente (por defecto).'),
        ('ml_group', 'Grupo ML', False, 'Mueble bajo o Mueble alto (actividades por ML).'),
        ('price', 'Tarifa base', False, 'Tarifa si no hay una tarifa de la obra o de la contrata.'),
    ]),
    'TIPOLOGIAS': ('Tipologías de ambiente de la obra.', [
        ('code', 'Código', True, 'Código de la tipología (01, C01, CL01…).'),
        ('name', 'Nombre', True, 'Nombre de la tipología.'),
        ('family', 'Familia', False, 'Cocina, Closet, Baño, Lavandería u Otro.'),
        ('product', 'Producto terminado', False,
         'Mueble terminado que fabrica la OF; si no existe se crea con este nombre.'),
        ('ml_low', 'ML bajo', False, 'Solo si los módulos no traen ancho.'),
        ('ml_high', 'ML alto', False, 'Solo si los módulos no traen ancho.'),
    ]),
    'MODULOS': ('Plantilla de módulos de cada tipología.', [
        ('typology', 'Tipología', True, 'Código de la tipología.'),
        ('sequence', 'Orden', False, 'Orden del módulo en la tipología.'),
        ('code', 'Código del módulo', True, 'MB01, MA01, CAMPANA… (el mismo del ETO).'),
        ('type', 'Tipo', True, 'Bajo, Alto, Cajonera, Campana, Repisero u Otro.'),
        ('width', 'Ancho (mm)', False, 'Con ancho, las actividades por ML bajan al módulo.'),
        ('ml_group', 'Grupo ML', False, 'Mueble bajo, Mueble alto o Sin ML.'),
        ('activity', 'Actividad de armado', True, 'Código de una actividad que se mide por módulo.'),
    ]),
    'ACTIVIDADES': ('Actividades por ambiente de cada tipología.', [
        ('typology', 'Tipología', True, 'Código de la tipología.'),
        ('activity', 'Actividad', True, 'Código de la actividad.'),
        ('qty', 'Cantidad por ambiente', True, 'Unidades de driver de un ambiente.'),
    ]),
    'BOM': ('Lista de materiales de un ambiente de cada tipología.', [
        ('typology', 'Tipología', True, 'Código de la tipología.'),
        ('product_code', 'Código del producto', True, 'Referencia interna del producto.'),
        ('product_name', 'Producto', False, 'Nombre (para crear el producto si no existe).'),
        ('qty', 'Cantidad', True, 'Cantidad por ambiente.'),
        ('uom', 'Unidad', False, 'Unidad de la cantidad; vacía: la del producto.'),
        ('stage', 'Etapa de consumo', False,
         'Producción, Armado, Instalación o Acabado y entrega. Sin etapa no se aprueba el plan.'),
        ('cost', 'Costo unitario', False,
         'Costo del maestro: informativo y costo del producto si se crea.'),
    ]),
    'ARBOL': ('Árbol de la obra: piso › departamento › ambiente.', [
        ('floor', 'Piso', True, 'Nombre del piso.'),
        ('apartment', 'Departamento', True, 'Nombre del departamento.'),
        ('space', 'Ambiente', True, 'Nombre del ambiente.'),
        ('typology', 'Tipología', False, 'Código de la tipología del ambiente.'),
    ]),
    'ETO': ('Módulos reales de cada ambiente (código y ancho), del despiece de ingeniería.', [
        ('floor', 'Piso', True, 'Piso del ambiente.'),
        ('apartment', 'Departamento', True, 'Departamento del ambiente.'),
        ('space', 'Ambiente', True, 'Ambiente.'),
        ('code', 'Código del módulo', True, 'El mismo código de la plantilla de la tipología.'),
        ('type', 'Tipo', False, 'Bajo, Alto, Cajonera, Campana, Repisero u Otro.'),
        ('width', 'Ancho (mm)', True, 'Ancho real del módulo.'),
        ('ml_group', 'Grupo ML', False, 'Mueble bajo, Mueble alto o Sin ML.'),
    ]),
}

# Valores de las listas: clave del campo → textos aceptados (además de la clave).
STAGE_VALUES = {
    'production': ('Producción',),
    'assembly': ('Armado',),
    'installation': ('Instalación',),
    'finishing': ('Acabado y entrega', 'Acabado', 'Entrega'),
}
MODULE_TYPE_VALUES = {
    'low': ('Bajo', 'Mueble bajo'),
    'high': ('Alto', 'Mueble alto'),
    'drawer': ('Cajonera',),
    'hood': ('Campana',),
    'shelf': ('Repisero',),
    'other': ('Otro',),
}
ML_GROUP_VALUES = {
    'low': ('Mueble bajo', 'Bajo'),
    'high': ('Mueble alto', 'Alto'),
    'none': ('Sin ML', 'Ninguno', 'No'),
}
FAMILY_VALUES = {
    'kitchen': ('Cocina',),
    'closet': ('Closet', 'Clóset'),
    'bathroom': ('Baño',),
    'laundry': ('Lavandería',),
    'other': ('Otro',),
}
MEASURE_VALUES = {
    'module': ('Módulo',),
    'ml': ('ML', 'Metro lineal', 'Metros lineales'),
    'space': ('Ambiente',),
}


def selection_value(values, text):
    """Clave de la lista cuyo texto coincide con ``text`` (o None)."""
    key = norm(text)
    if not key:
        return None
    for value, labels in values.items():
        if key == norm(value) or key in (norm(label) for label in labels):
            return value
    return None


def sheet_key(title):
    """Hoja del formato que corresponde a una pestaña del libro (o None)."""
    key = norm(title).upper().replace(' ', '_')
    return key if key in SHEETS else None


README = [
    'Maestro de planificación de la obra',
    '',
    'Libro que lee «Importar maestro y ETO» del planificador de obra. Todas las hojas son '
    'opcionales: se importa lo que traiga el libro y el resto de la obra no cambia.',
    'Las columnas se reconocen por su encabezado (fila 1), en cualquier orden. Las filas vacías '
    'se ignoran.',
    'Volver a importar el libro no duplica nada: las tipologías se reconocen por su código, los '
    'módulos por su código, los pisos, departamentos y ambientes por su nombre bajo su nivel '
    'superior y los productos por su referencia interna. La plantilla de módulos, las '
    'actividades y la BOM de cada tipología del libro quedan exactamente como en el libro.',
    'Antes de importar, el asistente muestra lo que se creará y las advertencias, sin escribir '
    'nada.',
]


def build_workbook(rows_by_sheet=None, readme=None):
    """Libro .xlsx (bytes) con la hoja LEEME y las hojas del formato.

    :param rows_by_sheet: ``{hoja: [dict clave → valor]}``; sin él, la
        plantilla vacía con todas las hojas.
    :param readme: líneas adicionales de la hoja LEEME.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    bold = Font(bold=True)
    header_fill = PatternFill('solid', fgColor='DDE4EE')
    workbook = Workbook()
    info = workbook.active
    info.title = 'LEEME'
    lines = README + [''] + list(readme or [])
    lines += ['', 'Hojas y columnas (* obligatoria):']
    for name, (description, columns) in SHEETS.items():
        lines += ['', '%s: %s' % (name, description)]
        lines += ['  %s%s: %s' % (header, ' *' if required else '', help_text)
                  for _key, header, required, help_text in columns]
    for index, line in enumerate(lines, start=1):
        cell = info.cell(row=index, column=1, value=line)
        if index == 1:
            cell.font = Font(bold=True, size=14)
        elif line.endswith(':') or (':' in line and line.split(':')[0] in SHEETS):
            cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    info.column_dimensions['A'].width = 120

    for name, (_description, columns) in SHEETS.items():
        if rows_by_sheet is not None and name not in rows_by_sheet:
            continue
        sheet = workbook.create_sheet(name)
        for col, (_key, header, _required, help_text) in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = bold
            cell.fill = header_fill
            sheet.column_dimensions[cell.column_letter].width = max(14, len(header) + 4)
        for row_index, row in enumerate((rows_by_sheet or {}).get(name, []), start=2):
            for col, (key, _header, _required, _help) in enumerate(columns, start=1):
                value = row.get(key)
                if value not in (None, ''):
                    sheet.cell(row=row_index, column=col, value=value)
        sheet.freeze_panes = 'A2'
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
