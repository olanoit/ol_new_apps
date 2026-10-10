# -*- coding: utf-8 -*-
"""W-15 «Importar maestro y ETO»: tipologías, módulos, actividades, BOM,
árbol de la obra y módulos del ETO desde un libro Excel.

El análisis (``_analyze``) no escribe nada: alimenta la vista previa y las
advertencias. «Importar» vuelve a analizar y aplica (``_apply``), de forma
idempotente. El formato del libro está en ``master_import_layout.py``.
"""
import base64
import logging
from collections import defaultdict
from io import BytesIO

from markupsafe import Markup, escape

from odoo import api, fields, models
from odoo.exceptions import UserError

from .master_import_layout import (
    FAMILY_VALUES, MEASURE_VALUES, ML_GROUP_VALUES, MODULE_TYPE_VALUES, SHEETS, STAGE_VALUES,
    build_workbook, norm, selection_value, sheet_key,
)

_logger = logging.getLogger(__name__)

EXAMPLES_URL = '/al_construction_planner/static/examples/%s'
# Unidades por su abreviatura habitual en los maestros.
UOM_ALIASES = {
    'und': 'uom.product_uom_unit', 'unidad': 'uom.product_uom_unit',
    'unidades': 'uom.product_uom_unit', 'units': 'uom.product_uom_unit',
    'ml': 'uom.product_uom_meter', 'm': 'uom.product_uom_meter',
    'metro': 'uom.product_uom_meter', 'metros': 'uom.product_uom_meter',
}


def _number(value):
    """Número de una celda (acepta coma decimal); None si está vacía o no
    es número."""
    if value in (None, ''):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip().replace(',', '.'))
    except ValueError:
        return None


def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _decimals(value):
    text = ('%.10f' % value).rstrip('0')
    return len(text.split('.')[1]) if '.' in text else 0


class ConstructionMasterImportWizard(models.TransientModel):
    _name = 'construction.master.import.wizard'
    _description = 'Importar maestro y ETO'
    _check_company_auto = True

    project_id = fields.Many2one(
        'project.project', string='Obra', required=True, check_company=True,
        domain=[('is_construction_site', '=', True)])
    company_id = fields.Many2one(related='project_id.company_id', string='Compañía')
    file = fields.Binary(string='Libro Excel', attachment=False)
    filename = fields.Char(string='Nombre del archivo')
    create_missing_products = fields.Boolean(
        string='Crear los productos que no existen',
        help='Crea los productos de la BOM que no existen con su código, nombre, unidad y '
             'costo del libro. Sin marcar, esas filas se omiten y se avisan.')
    preview_html = fields.Html(string='Se importará', compute='_compute_preview', sanitize=False)
    warning_html = fields.Html(string='Advertencias', compute='_compute_preview', sanitize=False)
    error_html = fields.Html(string='Errores', compute='_compute_preview', sanitize=False)

    # ------------------------------------------------------------------
    # Lectura del libro
    # ------------------------------------------------------------------
    def _read_workbook(self):
        """``(hojas, errores)``: ``{hoja: [dict clave → valor con _row]}``."""
        self.ensure_one()
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(base64.b64decode(self.file)), read_only=True,
                                     data_only=True)
        except Exception as exc:  # noqa: BLE001 - cualquier archivo que no sea un libro
            _logger.info('Libro no válido para W-15: %s', exc)
            return {}, [self.env._('El archivo no es un libro Excel (.xlsx) válido.')]
        sheets, errors = {}, []
        for worksheet in workbook.worksheets:
            name = sheet_key(worksheet.title)
            if not name:
                continue
            columns = SHEETS[name][1]
            by_header = {}
            for key, header, _required, _help in columns:
                by_header[norm(header)] = key
                by_header[norm(key)] = key
            rows = worksheet.iter_rows(values_only=True)
            header_row = next(rows, None) or ()
            positions = {}
            for index, header in enumerate(header_row):
                key = by_header.get(norm(header))
                if key and key not in positions:
                    positions[key] = index
            missing = [header for key, header, required, _help in columns
                       if required and key not in positions]
            if missing:
                errors.append(self.env._(
                    'La hoja %(sheet)s no tiene las columnas obligatorias: %(columns)s.',
                    sheet=name, columns=', '.join(missing)))
                continue
            data = []
            for row_number, values in enumerate(rows, start=2):
                row = {key: values[index] if index < len(values) else None
                       for key, index in positions.items()}
                if all(value in (None, '') for value in row.values()):
                    continue
                row['_row'] = row_number
                data.append(row)
            sheets[name] = data
        workbook.close()
        if not sheets and not errors:
            errors.append(self.env._(
                'El libro no tiene ninguna hoja del maestro (%s).', ', '.join(SHEETS)))
        return sheets, errors

    # ------------------------------------------------------------------
    # Análisis (sin escribir)
    # ------------------------------------------------------------------
    def _company(self):
        return self.project_id.company_id or self.env.company

    def _uom(self, text):
        key = norm(text)
        if not key:
            return self.env['uom.uom']
        if key in UOM_ALIASES:
            return self.env.ref(UOM_ALIASES[key], raise_if_not_found=False) or self.env['uom.uom']
        return self.env['uom.uom'].search([('name', '=ilike', _text(text))], limit=1)

    def _analyze(self):
        """Lo que haría la importación y sus advertencias, sin escribir."""
        self.ensure_one()
        result = {'errors': [], 'warnings': defaultdict(list), 'stats': {},
                  'catalog': [], 'typologies': {}, 'modules': {}, 'activities': {}, 'bom': {},
                  'products_to_create': {}, 'tree': [], 'eto': defaultdict(list)}
        if not self.file or not self.project_id:
            return result
        sheets, errors = self._read_workbook()
        result['errors'] = errors
        if errors and not sheets:
            return result
        warn = result['warnings']
        stats = result['stats']
        project = self.project_id
        company = self._company()
        Activity = self.env['construction.labor.activity'].with_context(active_test=False)
        is_manager = self.env.user.has_group('al_construction_planner.group_planner_manager')

        def invalid(sheet, row, column, value):
            warn['invalid_value'].append(self.env._(
                '%(sheet)s fila %(row)s: %(column)s «%(value)s»',
                sheet=sheet, row=row['_row'], column=column, value=_text(value)))

        # Actividades: las del catálogo de la compañía (o compartidas) y las de
        # la hoja CATALOGO.
        activity_domain = [('company_id', 'in', [company.id, False])]
        existing_activities = {}
        for activity in Activity.search(activity_domain, order='company_id'):
            existing_activities.setdefault(activity.code, activity)
        catalog_codes = set()
        rows = sheets.get('CATALOGO', [])
        new = updated = 0
        if rows and not is_manager:
            warn['catalog_no_rights'] = [_text(r['code']) for r in rows]
        for row in rows if is_manager else []:
            code = _text(row['code'])
            stage = selection_value(STAGE_VALUES, row['stage'])
            uom = self._uom(row['uom'])
            measure = selection_value(MEASURE_VALUES, row.get('measure')) or 'space'
            ml_group = selection_value(ML_GROUP_VALUES, row.get('ml_group')) or 'none'
            if not code or not _text(row['name']):
                invalid('CATALOGO', row, self.env._('Código'), code)
                continue
            if not stage:
                invalid('CATALOGO', row, self.env._('Etapa'), row['stage'])
                continue
            if not uom:
                invalid('CATALOGO', row, self.env._('Unidad'), row['uom'])
                continue
            vals = {
                'code': code, 'name': _text(row['name']), 'stage': stage, 'uom_id': uom.id,
                'module_level': measure == 'module', 'ml_based': measure == 'ml',
                'ml_group': ml_group if measure == 'ml' else 'none',
            }
            price = _number(row.get('price'))
            if price is not None:
                vals['default_price'] = price
            record = existing_activities.get(code)
            result['catalog'].append((vals, record))
            catalog_codes.add(code)
            if record:
                updated += 1
            else:
                new += 1
        stats['CATALOGO'] = (len(rows), new, updated)
        activity_codes = set(existing_activities) | catalog_codes
        module_activity_codes = {code for code, activity in existing_activities.items()
                                 if activity.module_level}
        module_activity_codes |= {vals['code'] for vals, _record in result['catalog']
                                  if vals['module_level']}
        module_activity_codes -= {vals['code'] for vals, _record in result['catalog']
                                  if not vals['module_level']}

        # Tipologías: las de la obra y las de la hoja TIPOLOGIAS.
        existing_typologies = {}
        for typology in self.env['construction.typology'].with_context(active_test=False).search(
                [('project_id', '=', project.id)]):
            existing_typologies.setdefault(typology.code, typology)
        rows = sheets.get('TIPOLOGIAS', [])
        new = updated = 0
        for row in rows:
            code = _text(row['code'])
            family = selection_value(FAMILY_VALUES, row.get('family')) or (
                'kitchen' if not _text(row.get('family')) else None)
            if not code or not _text(row['name']):
                invalid('TIPOLOGIAS', row, self.env._('Código'), code)
                continue
            if not family:
                invalid('TIPOLOGIAS', row, self.env._('Familia'), row.get('family'))
                continue
            record = existing_typologies.get(code)
            result['typologies'][code] = {
                'record': record,
                'vals': {'code': code, 'name': _text(row['name']), 'family': family},
                'product': _text(row.get('product')),
                'ml_low': _number(row.get('ml_low')),
                'ml_high': _number(row.get('ml_high')),
            }
            if record:
                updated += 1
            else:
                new += 1
        stats['TIPOLOGIAS'] = (len(rows), new, updated)
        typology_codes = set(existing_typologies) | set(result['typologies'])

        def typology_ok(sheet, row):
            code = _text(row['typology'])
            if code in typology_codes:
                return code
            warn['unknown_typology'].append('%s · %s' % (sheet, code or '—'))
            return None

        # Plantilla de módulos.
        rows = sheets.get('MODULOS', [])
        count = 0
        for row in rows:
            typology = typology_ok('MODULOS', row)
            if not typology:
                continue
            code = _text(row['code'])
            module_type = selection_value(MODULE_TYPE_VALUES, row['type'])
            activity = _text(row['activity'])
            width = _number(row.get('width'))
            ml_group = selection_value(ML_GROUP_VALUES, row.get('ml_group')) or 'none'
            if not code:
                invalid('MODULOS', row, self.env._('Código del módulo'), code)
                continue
            if not module_type:
                invalid('MODULOS', row, self.env._('Tipo'), row['type'])
                continue
            if activity not in activity_codes:
                warn['unknown_activity'].append('MODULOS · %s' % (activity or '—'))
                continue
            if activity not in module_activity_codes:
                warn['activity_not_module'].append('%s · %s' % (code, activity))
                continue
            if width is not None and width < 0:
                invalid('MODULOS', row, self.env._('Ancho (mm)'), row.get('width'))
                continue
            sequence = _number(row.get('sequence'))
            modules = result['modules'].setdefault(typology, {})
            modules[code] = {
                'sequence': int(sequence) if sequence is not None else (len(modules) + 1) * 10,
                'code': code, 'module_type': module_type, 'width_mm': int(round(width or 0)),
                'ml_group': ml_group, '_activity': activity,
            }
            count += 1
        stats['MODULOS'] = (len(rows), count, 0)

        # Actividades por ambiente.
        rows = sheets.get('ACTIVIDADES', [])
        count = 0
        for row in rows:
            typology = typology_ok('ACTIVIDADES', row)
            if not typology:
                continue
            activity = _text(row['activity'])
            qty = _number(row['qty'])
            if activity not in activity_codes:
                warn['unknown_activity'].append('ACTIVIDADES · %s' % (activity or '—'))
                continue
            if qty is None or qty < 0:
                invalid('ACTIVIDADES', row, self.env._('Cantidad por ambiente'), row['qty'])
                continue
            lines = result['activities'].setdefault(typology, {})
            lines[activity] = lines.get(activity, 0.0) + qty
            count += 1
        stats['ACTIVIDADES'] = (len(rows), count, 0)

        # Árbol: se lee antes de la BOM para saber cuántos ambientes tiene
        # cada tipología (monto de la BOM sin etapa en la obra).
        Task = self.env['project.task']
        spaces_in_tree = set()
        rows = sheets.get('ARBOL', [])
        for row in rows:
            path = (_text(row['floor']), _text(row['apartment']), _text(row['space']))
            if not all(path):
                invalid('ARBOL', row, self.env._('Piso, departamento y ambiente'), ' › '.join(path))
                continue
            typology = _text(row.get('typology'))
            if typology and typology not in typology_codes:
                warn['tree_unknown_typology'].append('%s (%s)' % (' › '.join(path), typology))
                typology = ''
            result['tree'].append(path + (typology,))
            spaces_in_tree.add(path)
        existing_spaces = {}
        for space in Task.search([('project_id', '=', project.id),
                                  ('construction_level', '=', 'space')]):
            path = (space.construction_floor_task_id.name, space.construction_apartment_task_id.name,
                    space.name)
            existing_spaces[path] = space
        tree_new = len(spaces_in_tree - set(existing_spaces))
        stats['ARBOL'] = (len(rows), tree_new, len(spaces_in_tree & set(existing_spaces)))
        space_count = defaultdict(int)
        typology_of_path = {path: space.construction_typology_id.code
                            for path, space in existing_spaces.items()}
        typology_of_path.update({row[:3]: row[3] for row in result['tree']})
        for code in typology_of_path.values():
            if code:
                space_count[code] += 1

        # BOM.
        rows = sheets.get('BOM', [])
        Product = self.env['product.product'].with_context(active_test=False)
        codes = {_text(row['product_code']) for row in rows} - {''}
        products = {}
        for product in Product.search([('default_code', 'in', list(codes)),
                                       ('company_id', 'in', [company.id, False])]):
            products.setdefault(product.default_code, product)
        digits = self.env['decimal.precision'].precision_get('Product Unit')
        seen = defaultdict(int)
        count = 0
        for row in rows:
            typology = typology_ok('BOM', row)
            if not typology:
                continue
            code = _text(row['product_code'])
            qty = _number(row['qty'])
            cost = _number(row.get('cost'))
            stage_text = _text(row.get('stage'))
            stage = selection_value(STAGE_VALUES, stage_text) if stage_text else False
            if not code:
                invalid('BOM', row, self.env._('Código del producto'), code)
                continue
            if qty is None or qty <= 0:
                invalid('BOM', row, self.env._('Cantidad'), row['qty'])
                continue
            if stage is None:
                invalid('BOM', row, self.env._('Etapa de consumo'), stage_text)
                continue
            product = products.get(code)
            uom = self._uom(row.get('uom')) if _text(row.get('uom')) else self.env['uom.uom']
            if _text(row.get('uom')) and not uom:
                invalid('BOM', row, self.env._('Unidad'), row.get('uom'))
                continue
            if not product:
                if not self.create_missing_products:
                    warn['product_not_found'].append(code)
                    continue
                if not _text(row.get('product_name')):
                    warn['product_without_name'].append(code)
                    continue
                result['products_to_create'].setdefault(code, {
                    'default_code': code, 'name': _text(row['product_name']),
                    'uom_id': (uom or self.env.ref('uom.product_uom_unit')).id,
                    'standard_price': cost or 0.0,
                })
                product_uom = uom or self.env.ref('uom.product_uom_unit')
                product_cost = cost or 0.0
            else:
                product_uom = product.uom_id
                product_cost = product.standard_price
                if uom and uom != product.uom_id:
                    # Unidades de distinto tipo no se convierten.
                    if uom._has_common_reference(product.uom_id):
                        product_uom = uom
                    else:
                        warn['uom_mismatch'].append('%s (%s ≠ %s)' % (
                            code, uom.name, product.uom_id.name))
            if not (product_cost or cost):
                warn['product_no_cost'].append(code)
            if _decimals(qty) > digits:
                warn['qty_precision'].append('%s · %s %s' % (typology, code, qty))
            key = (typology, code)
            seen[key] += 1
            lines = result['bom'].setdefault(typology, {})
            if code in lines:
                lines[code]['qty'] += qty
                lines[code]['stage'] = lines[code]['stage'] or stage
            else:
                lines[code] = {'qty': qty, 'stage': stage, 'uom': product_uom,
                               'cost': cost if cost is not None else product_cost}
            count += 1
        for (typology, code), times in seen.items():
            if times > 1:
                warn['bom_duplicate'].append(self.env._(
                    '%(typology)s · %(code)s (%(times)s filas)',
                    typology=typology, code=code, times=times))
        unstaged_amount = 0.0
        for typology, lines in result['bom'].items():
            for code, line in lines.items():
                if not line['stage']:
                    warn['bom_no_stage'].append('%s · %s' % (typology, code))
                    unstaged_amount += line['qty'] * (line['cost'] or 0.0) * space_count[typology]
        result['unstaged_amount'] = unstaged_amount
        stats['BOM'] = (len(rows), count, 0)
        stats['PRODUCTOS'] = (0, len(result['products_to_create']), 0)

        # ETO: módulos reales de cada ambiente (del árbol del libro o de la obra).
        rows = sheets.get('ETO', [])
        count = 0
        known_paths = set(existing_spaces) | spaces_in_tree
        for row in rows:
            path = (_text(row['floor']), _text(row['apartment']), _text(row['space']))
            code = _text(row['code'])
            width = _number(row['width'])
            module_type = selection_value(MODULE_TYPE_VALUES, row.get('type')) \
                if _text(row.get('type')) else False
            ml_group = selection_value(ML_GROUP_VALUES, row.get('ml_group')) \
                if _text(row.get('ml_group')) else False
            if path not in known_paths:
                warn['eto_unknown_space'].append(' › '.join(path))
                continue
            if not code:
                invalid('ETO', row, self.env._('Código del módulo'), code)
                continue
            if width is None or width < 0:
                invalid('ETO', row, self.env._('Ancho (mm)'), row['width'])
                continue
            if module_type is None:
                invalid('ETO', row, self.env._('Tipo'), row.get('type'))
                continue
            if ml_group is None:
                invalid('ETO', row, self.env._('Grupo ML'), row.get('ml_group'))
                continue
            result['eto'][path].append({
                'code': code, 'width_mm': int(round(width)), 'module_type': module_type,
                'ml_group': ml_group})
            count += 1
        stats['ETO'] = (len(rows), count, 0)
        return result

    # ------------------------------------------------------------------
    # Vista previa
    # ------------------------------------------------------------------
    def _warning_messages(self):
        return {
            'bom_no_stage': self.env._(
                'Filas de BOM sin etapa de consumo: se importan, pero impiden aprobar el plan '
                'hasta indicar su etapa'),
            'product_not_found': self.env._(
                'Productos que no existen (filas omitidas; marque «Crear los productos que no '
                'existen» o créelos antes)'),
            'product_without_name': self.env._(
                'Productos que no existen y no tienen nombre en el libro (filas omitidas)'),
            'product_no_cost': self.env._(
                'Productos sin costo (el planificador lo escribe con «Aplicar costo»)'),
            'bom_duplicate': self.env._(
                'Productos repetidos en la BOM de una tipología (se suman sus cantidades)'),
            'qty_precision': self.env._(
                'Cantidades con más decimales que la precisión de la unidad (se redondean al '
                'guardar la BOM)'),
            'uom_mismatch': self.env._(
                'Unidades que no corresponden al producto (se usa la del producto)'),
            'unknown_activity': self.env._(
                'Actividades que no existen en el catálogo (filas omitidas)'),
            'activity_not_module': self.env._(
                'Módulos con una actividad de armado que no se mide por módulo (filas omitidas)'),
            'catalog_no_rights': self.env._(
                'La hoja CATALOGO la aplica un administrador del planificador (se omite)'),
            'unknown_typology': self.env._(
                'Filas de una tipología que no existe en la obra ni en la hoja TIPOLOGIAS '
                '(omitidas)'),
            'tree_unknown_typology': self.env._(
                'Ambientes con una tipología que no existe (se crean sin tipología)'),
            'eto_unknown_space': self.env._(
                'Ambientes del ETO que no existen en la obra ni en el árbol del libro (filas '
                'omitidas)'),
            'invalid_value': self.env._('Valores no reconocidos (filas omitidas)'),
        }

    @api.depends('file', 'project_id', 'create_missing_products')
    def _compute_preview(self):
        labels = {
            'CATALOGO': self.env._('Actividades de obra'),
            'TIPOLOGIAS': self.env._('Tipologías'),
            'MODULOS': self.env._('Módulos de la plantilla'),
            'ACTIVIDADES': self.env._('Actividades por ambiente'),
            'BOM': self.env._('Filas de BOM'),
            'PRODUCTOS': self.env._('Productos por crear'),
            'ARBOL': self.env._('Ambientes del árbol'),
            'ETO': self.env._('Módulos del ETO'),
        }
        for wizard in self:
            data = wizard._analyze()
            wizard.error_html = data['errors'] and Markup('<ul class="mb-0">%s</ul>') % Markup(
                '').join(Markup('<li>%s</li>') % error for error in data['errors'])
            if not data['stats']:
                wizard.preview_html = wizard.warning_html = False
                continue
            rows = Markup('').join(
                Markup('<tr><td>%s</td><td class="text-end">%s</td><td class="text-end">%s</td>'
                       '<td class="text-end">%s</td></tr>') % (
                    labels[key], read or '—', new, updated)
                for key, (read, new, updated) in data['stats'].items()
                if read or new or updated)
            wizard.preview_html = Markup(
                '<table class="table table-sm o_main_table"><thead><tr><th>%s</th>'
                '<th class="text-end">%s</th><th class="text-end">%s</th>'
                '<th class="text-end">%s</th></tr></thead><tbody>%s</tbody></table>') % (
                self.env._('Hoja'), self.env._('Filas'), self.env._('Nuevos o importados'),
                self.env._('Actualizados'), rows)
            wizard.warning_html = wizard._warnings_html(data)

    def _warnings_html(self, data):
        items = []
        currency = self._company().currency_id
        for key, message in self._warning_messages().items():
            values = data['warnings'].get(key)
            if not values:
                continue
            names = list(dict.fromkeys(values))
            text = '%s: %s' % (message, ', '.join(names[:8]) + (' …' if len(names) > 8 else ''))
            if key == 'bom_no_stage' and data.get('unstaged_amount'):
                text += self.env._(' (monto en la obra al costo del libro: %(amount)s)',
                                   amount='{:,.2f}'.format(currency.round(data['unstaged_amount'])))
            items.append(text)
        if not items:
            return False
        return Markup('<ul class="mb-0">%s</ul>') % Markup('').join(
            Markup('<li>%s</li>') % escape(item) for item in items)

    # ------------------------------------------------------------------
    # Importación
    # ------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        if not self.file:
            raise UserError(self.env._('Suba el libro del maestro.'))
        data = self._analyze()
        if data['errors']:
            raise UserError('\n'.join(data['errors']))
        summary = self._apply(data)
        # sudo: el planificador puede no tener escritura en la obra; la
        # importación queda en su historial con él como autor.
        project_sudo = self.project_id.sudo()
        project_sudo.message_post(author_id=self.env.user.partner_id.id, body=self.env._(
            'Maestro importado (%(file)s): %(typologies)s tipologías, %(modules)s módulos de '
            'plantilla, %(spaces)s ambientes nuevos y %(eto)s módulos del ETO.',
            file=self.filename or '—', typologies=summary['typologies'],
            modules=summary['modules'], spaces=summary['spaces'], eto=summary['eto']))
        return self.project_id.action_view_construction_typologies()

    def _apply(self, data):
        project = self.project_id
        company = self._company()
        summary = defaultdict(int)

        # Catálogo de actividades (solo administradores: el análisis ya lo filtró).
        Activity = self.env['construction.labor.activity']
        for vals, record in data['catalog']:
            if record:
                record.write(vals)
            else:
                Activity.create(dict(vals, company_id=company.id))
        activities = {}
        for activity in Activity.with_context(active_test=False).search(
                [('company_id', 'in', [company.id, False])], order='company_id'):
            activities.setdefault(activity.code, activity)

        # Productos: sudo porque el planificador no administra el maestro de
        # productos (como «Crear producto desde el plan», W-11); se crean con
        # la compañía de la obra y quedan registrados en el historial.
        Product_sudo = self.env['product.product'].sudo().with_context(active_test=False)
        products = {}
        if data['products_to_create']:
            created = Product_sudo.create([dict(
                vals, type='consu', is_storable=True, company_id=company.id,
                uom_id=vals['uom_id']) for vals in data['products_to_create'].values()])
            products.update({product.default_code: product for product in created})
        codes = {code for lines in data['bom'].values() for code in lines}
        for product in Product_sudo.search([('default_code', 'in', list(codes)),
                                            ('company_id', 'in', [company.id, False])]):
            products.setdefault(product.default_code, product)

        # Tipologías.
        Typology = self.env['construction.typology'].with_context(active_test=False)
        typologies = {t.code: t for t in Typology.search([('project_id', '=', project.id)])}
        for code, info in data['typologies'].items():
            vals = dict(info['vals'])
            record = typologies.get(code)
            if record:
                record.write(vals)
            else:
                record = Typology.create(dict(vals, project_id=project.id))
                typologies[code] = record
            summary['typologies'] += 1
            info['record'] = record

        # Plantilla de módulos: la del libro reemplaza a la de la tipología.
        Module = self.env['construction.typology.module']
        for code, modules in data['modules'].items():
            typology = typologies[code]
            current = {m.code: m for m in typology.module_line_ids}
            for module_code, vals in modules.items():
                vals = dict(vals, assembly_activity_id=activities[vals.pop('_activity')].id)
                if module_code in current:
                    current.pop(module_code).write(vals)
                else:
                    Module.create(dict(vals, typology_id=typology.id))
                summary['modules'] += 1
            if current:
                Module.browse([m.id for m in current.values()]).unlink()

        # ML escritos a mano (tipologías sin anchos en sus módulos).
        for code, info in data['typologies'].items():
            typology = typologies[code]
            if any(typology.module_line_ids.mapped('width_mm')):
                continue
            vals = {key: info[key] for key in ('ml_low', 'ml_high') if info[key] is not None}
            if vals:
                typology.write(vals)

        # Actividades por ambiente.
        TypologyActivity = self.env['construction.typology.activity']
        for code, lines in data['activities'].items():
            typology = typologies[code]
            current = {line.activity_id.code: line for line in typology.activity_line_ids}
            for sequence, (activity_code, qty) in enumerate(lines.items(), start=1):
                vals = {'qty': qty, 'sequence': sequence}
                if activity_code in current:
                    current.pop(activity_code).write(vals)
                else:
                    TypologyActivity.create(dict(vals, typology_id=typology.id,
                                                 activity_id=activities[activity_code].id))
            if current:
                TypologyActivity.browse([line.id for line in current.values()]).unlink()

        # BOM de la tipología: sudo porque el planificador lee las listas de
        # materiales pero no las administra (grupo de usuario de fabricación);
        # la BOM de la tipología es parte del catálogo de la obra que importa.
        for code, lines in data['bom'].items():
            typology = typologies[code]
            info = data['typologies'].get(code, {})
            self._sync_bom(typology, lines, products, info.get('product'), company)

        # Árbol de la obra.
        Task = self.env['project.task']
        nodes = {}
        for task in Task.search([('project_id', '=', project.id),
                                 ('construction_level', 'in', ('floor', 'apartment', 'space'))]):
            nodes[(task.construction_level, task.parent_id.id, task.name)] = task

        def node(name, level, parent):
            key = (level, parent.id if parent else False, name)
            if key not in nodes:
                nodes[key] = Task.create({
                    'name': name, 'project_id': project.id, 'construction_level': level,
                    'parent_id': parent.id if parent else False, 'user_ids': False})
                if level == 'space':
                    summary['spaces'] += 1
            return nodes[key]

        for floor, apartment, space, typology_code in data['tree']:
            space_task = node(space, 'space', node(apartment, 'apartment',
                                                   node(floor, 'floor', None)))
            typology = typologies.get(typology_code) if typology_code else False
            if typology and space_task.construction_typology_id != typology:
                space_task.construction_typology_id = typology

        # ETO: módulos reales de cada ambiente, por código.
        for (floor, apartment, space), modules in data['eto'].items():
            floor_task = nodes.get(('floor', False, floor))
            apartment_task = floor_task and nodes.get(('apartment', floor_task.id, apartment))
            space_task = apartment_task and nodes.get(('space', apartment_task.id, space))
            if not space_task:
                continue
            current = {m.construction_module_code: m for m in Task.search([
                ('parent_id', '=', space_task.id), ('construction_level', '=', 'module')])}
            templates = {m.code: m for m in space_task.construction_typology_id.module_line_ids}
            for vals in modules:
                template = templates.get(vals['code'])
                task_vals = {
                    'construction_width_mm': vals['width_mm'],
                    'construction_module_type': vals['module_type'] or (
                        template.module_type if template else 'other'),
                    'construction_ml_group': vals['ml_group'] or (
                        template.ml_group if template else 'none'),
                }
                task = current.get(vals['code'])
                if task:
                    task.write(task_vals)
                else:
                    Task.with_context(mail_create_nolog=True, tracking_disable=True).create(dict(
                        task_vals, name=vals['code'], project_id=project.id,
                        parent_id=space_task.id, construction_level='module',
                        construction_module_code=vals['code'],
                        construction_unit_state='planned', user_ids=False))
                summary['eto'] += 1
        return summary

    def _sync_bom(self, typology, lines, products, product_name, company):
        """Deja la BOM de la tipología exactamente con ``lines``."""
        typology_sudo = typology.sudo()
        bom_sudo = typology_sudo.bom_id
        if not bom_sudo:
            template_sudo = typology_sudo.product_tmpl_id
            if not template_sudo:
                name = product_name or '%s · %s' % (typology.name, typology.project_id.name)
                Template_sudo = self.env['product.template'].sudo()
                template_sudo = Template_sudo.search([
                    ('name', '=', name), ('company_id', 'in', [company.id, False])], limit=1) \
                    or Template_sudo.create({'name': name, 'type': 'consu', 'is_storable': True,
                                             'company_id': company.id})
            bom_sudo = self.env['mrp.bom'].sudo().create({
                'product_tmpl_id': template_sudo.id, 'product_qty': 1.0,
                'company_id': company.id})
            typology.write({'product_tmpl_id': template_sudo.id, 'bom_id': bom_sudo.id})
        current = {line.product_id.default_code: line for line in bom_sudo.bom_line_ids}
        commands = []
        for code, line in lines.items():
            product = products.get(code)
            if not product:
                continue
            vals = {'product_qty': line['qty'], 'product_uom_id': line['uom'].id,
                    'construction_consumption_stage': line['stage'] or False}
            if code in current:
                commands.append((1, current.pop(code).id, vals))
            else:
                commands.append((0, 0, dict(vals, product_id=product.id)))
        commands += [(2, line.id) for line in current.values()]
        if commands:
            bom_sudo.write({'bom_line_ids': commands})

    # ------------------------------------------------------------------
    # Descargas
    # ------------------------------------------------------------------
    def action_download_template(self):
        """Plantilla vacía con todas las hojas y la hoja LEEME."""
        attachment = self.env['ir.attachment'].create({
            'name': self.env._('Plantilla del maestro de planificación.xlsx'),
            'datas': base64.b64encode(build_workbook()),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'self',
        }

    def action_download_example(self):
        return {'type': 'ir.actions.act_url', 'target': 'self',
                'url': EXAMPLES_URL % 'maestro_planificacion_ejemplo.xlsx'}

    def action_download_eto_example(self):
        return {'type': 'ir.actions.act_url', 'target': 'self',
                'url': EXAMPLES_URL % 'eto_ejemplo.xlsx'}
