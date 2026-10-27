# -*- coding: utf-8 -*-
"""Recetas y árbol de cualquier rubro (v1.7, M17-01 y M17-04).

«Importar recetas y niveles» lee un libro con dos hojas opcionales:

* RECETAS: una fila por recurso de cada plantilla de ejecución (código de la
  receta, tipo de recurso, código del producto o de la actividad, cantidad
  por unidad, etapa, rendimiento…).
* NIVELES: el árbol de la obra con una columna por nivel de su plantilla
  (Disciplina, Subpartida, Partida…), la receta y el metrado de la tarea.

«Copiar desde Work Items» toma los APU del paquete Construction de Odoo
(modelo ``x_work_item`` con sus líneas): el material y el equipo por su
naturaleza de costo, la mano de obra como personal propio o como contrata.
Ambos son idempotentes: la receta se reconoce por su código y queda como en
el origen; las tareas, por su nombre bajo su nivel superior.
"""
import base64
import logging
from io import BytesIO

from markupsafe import Markup

from odoo import Command, api, fields, models
from odoo.exceptions import UserError

from ..models.common import norm

_logger = logging.getLogger(__name__)

RECIPE_COLUMNS = [
    ('recipe_code', 'Código de la receta', True),
    ('recipe_name', 'Nombre de la receta', False),
    ('recipe_uom', 'Unidad de la receta', False),
    ('type', 'Tipo de recurso', True),
    ('code', 'Código del recurso', True),
    ('name', 'Nombre del recurso', False),
    ('qty', 'Cantidad por unidad', True),
    ('uom', 'Unidad', False),
    ('cost', 'Costo unitario', False),
    ('stage', 'Etapa', False),
    ('productivity', 'Rendimiento', False),
    ('crew_size', 'Personas', False),
    # CREVAL v1.0, 6.1 y 6.5: forma de cálculo, porcentaje y moneda del costo.
    ('calc_method', 'Forma de cálculo', False),
    ('calc_pct', 'Porcentaje', False),
    ('currency', 'Moneda', False),
    # CR-05: el nivel (especialidad o sub-especialidad) del % del costo directo.
    ('calc_task', 'Nivel del porcentaje', False),
]
# Forma de cálculo escrita en el libro → código.
CALC_ALIASES = {
    'cantidad': 'qty', 'cantidad por unidad': 'qty', 'cuadrilla': 'crew',
    'cuadrilla / rendimiento': 'crew', 'cuadrilla ÷ rendimiento': 'crew',
    'rendimiento': 'crew', '%mo': 'pct_labor', '% mo': 'pct_labor',
    '% de la mano de obra': 'pct_labor', 'porcentaje de la mano de obra': 'pct_labor',
    '% costo directo': 'pct_group', '% del costo directo': 'pct_group',
    '% del costo directo de un grupo de partidas': 'pct_group',
}
# Unidad que en el APU marca un porcentaje de la mano de obra.
PCT_LABOR_UNITS = ('%mo', '% mo', '%m.o.', '% m.o.')
LEVEL_EXTRA_COLUMNS = [
    ('recipe', 'Receta'),
    ('qty', 'Cantidad'),
]
# Encabezados de libros anteriores que se siguen aceptando.
HEADER_ALIASES = {'NIVELES': {'Metrado': 'qty'}}
# Naturaleza de costo del Work Item → comportamiento del recurso.
WORK_ITEM_NATURES = {'material': 'material', 'equipment': 'equipment', 'labour': 'labor'}


def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _number(value):
    if value in (None, ''):
        return None
    try:
        return float(str(value).replace(',', '.')) if isinstance(value, str) else float(value)
    except ValueError:
        return None


class ConstructionRecipeImportWizard(models.TransientModel):
    _name = 'construction.recipe.import.wizard'
    _description = 'Importar recetas y niveles'
    _check_company_auto = True

    project_id = fields.Many2one(
        'project.project', string='Obra', check_company=True,
        domain="[('is_construction_site', '=', True)]",
        help='Para la hoja NIVELES y para recetas propias de la obra. Vacía: recetas de la '
             'compañía.')
    company_id = fields.Many2one(
        'res.company', string='Compañía', required=True, default=lambda self: self.env.company)
    source = fields.Selection(
        [('excel', 'Libro Excel'), ('work_item', 'Work Items del paquete Construction')],
        string='Origen', default='excel', required=True)
    file = fields.Binary(string='Libro Excel')
    filename = fields.Char(string='Nombre del archivo')
    labor_behavior = fields.Selection(
        [('labor', 'Personal propio (cuadrilla)'), ('contract', 'Contrata')],
        string='Mano de obra como', default='labor', required=True,
        help='La mano de obra del APU entra como horas de cuadrilla propia o como contrata.')
    # CREVAL v1.0, 6.4: la modalidad reemplaza a «Mano de obra como»; la
    # mano de obra entra como personal propio y la modalidad decide al
    # generar el plan si va a una contrata.
    execution_mode = fields.Selection(
        [('own', 'Personal propio'), ('labor_contract', 'Contrata solo de mano de obra'),
         ('full_contract', 'Contrata a todo costo')],
        string='Modalidad de las recetas', default='own', required=True,
        help='Modalidad por defecto de las recetas importadas; se cambia por partida.')
    contract_activity_id = fields.Many2one(
        'construction.labor.activity', string='Actividad de la contrata',
        help='La que paga a la contrata en las modalidades de contrata.')
    default_stage_id = fields.Many2one(
        'construction.stage', string='Etapa por defecto',
        help='La de los recursos sin etapa en el origen.')
    work_item_scope = fields.Selection(
        [('templates', 'Plantillas de Work Item'), ('project', 'Los de la venta de la obra')],
        string='Work Items', default='templates', required=True)
    work_item_available = fields.Boolean(compute='_compute_work_item_available')
    result_html = fields.Html(string='Resultado', readonly=True, sanitize=False)

    @api.onchange('project_id')
    def _onchange_project_apu_source(self):
        """La fuente del APU de la obra (CREVAL v1.0, 6.2)."""
        if self.project_id.construction_apu_source == 'excel':
            self.source = 'excel'
        elif self.project_id.construction_apu_source == 'work_item' and self.work_item_available:
            self.source = 'work_item'
            self.work_item_scope = 'project'

    def _compute_work_item_available(self):
        for wizard in self:
            wizard.work_item_available = 'x_work_item' in self.env

    # ------------------------------------------------------------------
    # Libro
    # ------------------------------------------------------------------
    def _level_lines(self):
        template = self.project_id.construction_level_template_id
        return template.line_ids.sorted(lambda l: (l.sequence, l.id)) if template else \
            self.env['construction.level.template.line']

    def _read_workbook(self):
        self.ensure_one()
        content = self.with_context(bin_size=False).file
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(base64.b64decode(content)), read_only=True,
                                     data_only=True)
        except Exception as exc:  # noqa: BLE001 - cualquier archivo que no sea un libro
            _logger.info('Libro de recetas no válido: %s', exc)
            raise UserError(self.env._('El archivo no es un libro Excel (.xlsx) válido.')) from exc
        level_columns = [(str(line.id), line.name, False) for line in self._level_lines()]
        layouts = {
            'RECETAS': RECIPE_COLUMNS,
            'NIVELES': level_columns + [(key, header, False) for key, header in
                                        LEVEL_EXTRA_COLUMNS],
        }
        sheets = {}
        for worksheet in workbook.worksheets:
            name = norm(worksheet.title).upper()
            if name not in layouts:
                continue
            by_header = {norm(header): key for key, header, _required in layouts[name]}
            for header, key in HEADER_ALIASES.get(name, {}).items():
                by_header.setdefault(norm(header), key)
            rows = worksheet.iter_rows(values_only=True)
            positions = {}
            for index, header in enumerate(next(rows, None) or ()):
                key = by_header.get(norm(header))
                if key and key not in positions:
                    positions[key] = index
            missing = [header for key, header, required in layouts[name]
                       if required and key not in positions]
            if missing:
                raise UserError(self.env._(
                    'La hoja %(sheet)s no tiene las columnas obligatorias: %(columns)s.',
                    sheet=name, columns=', '.join(missing)))
            data = []
            for row_number, values in enumerate(rows, start=2):
                row = {key: values[index] if index < len(values) else None
                       for key, index in positions.items()}
                if all(value in (None, '') for value in row.values()):
                    continue
                row['_row'] = row_number
                data.append(row)
            sheets[name] = data
        workbook.close()
        if not sheets:
            raise UserError(self.env._('El libro no tiene las hojas RECETAS ni NIVELES.'))
        return sheets

    def _uom(self, text):
        text = _text(text)
        if not text:
            return self.env['uom.uom']
        # Tabla de equivalencias de la empresa (CREVAL v1.0, 6.4).
        alias_uom = self.env['construction.uom.alias']._find(text, self.company_id)
        if alias_uom:
            return alias_uom
        aliases = {'und': 'uom.product_uom_unit', 'm2': 'uom.product_uom_square_meter',
                   'm²': 'uom.product_uom_square_meter', 'ml': 'uom.product_uom_meter',
                   'm': 'uom.product_uom_meter', 'kg': 'uom.product_uom_kgm',
                   'h': 'uom.product_uom_hour', 'hh': 'uom.product_uom_hour'}
        if norm(text) in aliases:
            return self.env.ref(aliases[norm(text)])
        return self.env['uom.uom'].search([('name', '=ilike', text)], limit=1)

    def _resource_type(self, text):
        Type = self.env['construction.resource.type']
        text = _text(text)
        for rtype in Type.search([]):
            if norm(text) in (norm(rtype.code), norm(rtype.name)):
                return rtype
        return Type

    def _stage_code(self, text):
        text = _text(text)
        if not text:
            return self.default_stage_id.code or False
        return self.env['construction.stage']._find_by_text(text)

    # ------------------------------------------------------------------
    # Recetas
    # ------------------------------------------------------------------
    def _activity_for(self, product, uom, stage):
        """Actividad (driver) de la mano de obra del APU: una por producto."""
        Activity = self.env['construction.labor.activity']
        code = product.default_code or 'MO-%s' % product.id
        activity = Activity.search([('code', '=', code),
                                    ('company_id', 'in', (self.company_id.id, False))], limit=1)
        if not activity:
            stage = stage or next(iter(self.env['construction.stage']._field_codes()), False)
            if not stage:
                raise UserError(self.env._(
                    'Elija la etapa por defecto: la mano de obra %s necesita una actividad con '
                    'etapa.', product.display_name))
            activity = Activity.create({
                'code': code, 'name': product.name, 'uom_id': uom.id, 'stage': stage,
                'company_id': self.company_id.id,
                'product_id': product.id if product.type == 'service' else False,
            })
        return activity

    def _resolve_code(self, code):
        """(actividad, producto) del código de un recurso del libro: se busca
        en las actividades (por su código o el de su producto) y en los
        productos (CREVAL v1.0, CR-04)."""
        Activity = self.env['construction.labor.activity']
        if not code:
            return Activity, self.env['product.product']
        company_domain = [('company_id', 'in', (self.company_id.id, False))]
        activity = Activity.search([('code', '=', code)] + company_domain, limit=1) \
            or Activity.search([('product_id.default_code', '=', code)] + company_domain,
                               limit=1)
        product = self.env['product.product'].search([('default_code', '=', code)], limit=1) \
            or activity.product_id
        return activity, product

    def _upsert_recipes(self, recipes, source, source_refs=None):
        """Crea o actualiza las recetas ``{código: {'name', 'uom', 'product',
        'lines': [vals de línea]}}``: las líneas quedan como en el origen."""
        Template = self.env['construction.execution.template']
        created = updated = 0
        result = Template
        scope = self._recipe_scope_values()
        for code, data in recipes.items():
            domain = [('code', '=', code), ('company_id', '=', self.company_id.id)] + [
                (name, '=', value) for name, value in scope.items()]
            template = Template.with_context(active_test=False).search(domain, limit=1)
            vals = {'name': data.get('name') or code, 'source': source,
                    'execution_mode': self.execution_mode,
                    'contract_activity_id': self.contract_activity_id.id,
                    'source_ref': (source_refs or {}).get(code, False),
                    'line_ids': [Command.clear()] + [Command.create(line)
                                                     for line in data['lines']]}
            if data.get('uom'):
                vals['uom_id'] = data['uom'].id
            if data.get('product'):
                vals['product_id'] = data['product'].id
            if template:
                template.write(vals)
                updated += 1
            else:
                template = Template.create(dict(vals, code=code, company_id=self.company_id.id,
                                                **scope))
                created += 1
            result |= template
        return result, created, updated

    def _recipe_scope_values(self):
        """Dónde viven las recetas que se importan: la obra o, sin ella, la
        biblioteca de la compañía. El puente de ventas agrega la cotización
        de origen."""
        return {'project_id': self.project_id.id or False}

    def _resolve_calc_tasks(self, templates, recipes):
        """«Nivel del porcentaje» de las líneas por % del costo directo: el
        nivel de la obra cuyo nombre empieza con ese código (CR-05)."""
        errors = []
        Task = self.env['project.task']
        for template in templates:
            pending = recipes.get(template.code, {}).get('calc_tasks')
            if not pending:
                continue
            lines = template.line_ids.sorted('id')
            for index, (code, row) in pending.items():
                task = Task
                if self.project_id:
                    candidates = Task.search([
                        ('project_id', '=', self.project_id.id),
                        '|', ('name', '=ilike', code), ('name', '=ilike', '%s %%' % code)])
                    task = candidates.sorted(lambda t: len(t.construction_path or ''))[:1]
                if not task:
                    errors.append(self.env._(
                        'RECETAS fila %(row)s: no hay un nivel %(code)s en la obra para el '
                        'porcentaje (elija la obra en el asistente).', row=row, code=code))
                    continue
                lines[index].write({'calc_scope': 'task', 'calc_task_id': task.id})
        return errors

    def _recipes_from_rows(self, rows):
        Product = self.env['product.product']
        recipes, errors = {}, []
        for row in rows:
            code = _text(row.get('recipe_code'))
            recipe = recipes.setdefault(code, {'name': _text(row.get('recipe_name')),
                                               'uom': self._uom(row.get('recipe_uom')),
                                               'lines': [], 'calc_tasks': {}})
            rtype = self._resource_type(row.get('type'))
            qty = _number(row.get('qty'))
            stage = self._stage_code(row.get('stage'))
            if not rtype or qty is None or stage is None:
                errors.append(self.env._('RECETAS fila %(row)s: tipo, cantidad o etapa no válidos.',
                                         row=row['_row']))
                continue
            resource_code = _text(row.get('code'))
            vals = {'resource_type_id': rtype.id, 'resource_type': rtype.behavior, 'qty': qty,
                    'stage': stage, 'price_unit': _number(row.get('cost')) or 0.0,
                    'productivity': _number(row.get('productivity')) or 0.0,
                    'crew_size': _number(row.get('crew_size')) or 1.0,
                    'note': _text(row.get('name')) or False}
            uom_text = norm(_text(row.get('uom')))
            method = CALC_ALIASES.get(norm(_text(row.get('calc_method'))), False)
            if not method and uom_text in PCT_LABOR_UNITS:
                # «%MO»: la cantidad del APU es el porcentaje.
                method = 'pct_labor'
                vals['calc_pct'] = qty
            if method:
                vals['calc_method'] = method
                pct = _number(row.get('calc_pct'))
                if pct is not None:
                    vals['calc_pct'] = pct
            level_code = _text(row.get('calc_task'))
            if method == 'pct_group' and level_code:
                # Se resuelve al final: la hoja NIVELES puede crear el nivel.
                recipe['calc_tasks'][len(recipe['lines'])] = (level_code, row['_row'])
            currency_text = _text(row.get('currency'))
            if currency_text:
                currency = self.env['res.currency'].with_context(active_test=False).search(
                    [('name', '=ilike', currency_text)], limit=1)
                if currency:
                    vals['line_currency_id'] = currency.id
            uom = self._uom(row.get('uom')) if uom_text not in PCT_LABOR_UNITS else \
                self.env['uom.uom']
            if method in ('pct_labor', 'pct_group') and rtype.behavior in ('contract', 'labor'):
                # Las herramientas por porcentaje no se pagan ni se compran.
                vals.update(resource_type='equipment', resource_type_id=self.env[
                    'construction.resource.type']._default_for('equipment').id)
            activity, product = self._resolve_code(resource_code)
            if vals['resource_type'] in ('contract', 'labor'):
                # El código puede ser el de la actividad o el de su producto;
                # la línea lleva los dos (CR-04).
                if not activity and product:
                    activity = self._activity_for(product, uom or product.uom_id, stage)
                if not activity:
                    errors.append(self.env._(
                        'RECETAS fila %(row)s: el código %(code)s no es una actividad ni un '
                        'producto.', row=row['_row'], code=resource_code))
                    continue
                vals.update(activity_id=activity.id,
                            product_id=(activity.product_id or product).id or False,
                            product_uom_id=(uom or activity.uom_id).id)
            elif method in ('pct_labor', 'pct_group'):
                # Herramientas y gastos por porcentaje: sin producto comprable.
                product = product or Product
                vals.update(product_id=product.id or False,
                            product_uom_id=(uom or self.env.ref('uom.product_uom_unit')).id)
            else:
                if not product:
                    errors.append(self.env._(
                        'RECETAS fila %(row)s: el producto %(code)s no existe.',
                        row=row['_row'], code=resource_code))
                    continue
                vals.update(product_id=product.id, product_uom_id=(uom or product.uom_id).id)
            recipe['lines'].append(vals)
        return recipes, errors

    # ------------------------------------------------------------------
    # Niveles
    # ------------------------------------------------------------------
    def _import_levels(self, rows):
        """Crea las tareas que faltan, nivel por nivel, y pone la receta y la
        cantidad presupuestada en la última columna con nombre de cada fila."""
        if not self.project_id.construction_level_template_id:
            raise UserError(self.env._(
                'La obra %s no tiene plantilla de niveles: la hoja NIVELES la necesita.',
                self.project_id.display_name))
        Task = self.env['project.task'].with_context(mail_create_nolog=True,
                                                    tracking_disable=True)
        Template = self.env['construction.execution.template']
        levels = self._level_lines()
        cache = {(task.parent_id.id, task.construction_level_line_id.id, task.name): task
                 for task in Task.search([('project_id', '=', self.project_id.id),
                                          ('construction_level_line_id', '!=', False)])}
        created, errors = 0, []
        for row in rows:
            parent = Task
            for line in levels:
                name = _text(row.get(str(line.id)))
                if not name:
                    break
                key = (parent.id or False, line.id, name)
                task = cache.get(key)
                if not task:
                    task = Task.create({'name': name, 'project_id': self.project_id.id,
                                        'construction_level_line_id': line.id,
                                        'parent_id': parent.id or False, 'user_ids': False})
                    cache[key] = task
                    created += 1
                parent = task
            if not parent:
                continue
            code = _text(row.get('recipe'))
            if code:
                recipe = Template.search([('code', '=', code),
                                          ('company_id', '=', self.company_id.id),
                                          ('project_id', 'in', (self.project_id.id, False))],
                                         order='project_id', limit=1)
                if not recipe:
                    errors.append(self.env._('NIVELES fila %(row)s: la receta %(code)s no existe.',
                                             row=row['_row'], code=code))
                    continue
                parent.construction_execution_template_id = recipe
            qty = _number(row.get('qty'))
            if qty is not None:
                parent.construction_qty = qty
        return created, errors

    # ------------------------------------------------------------------
    # Work Items
    # ------------------------------------------------------------------
    def _read_work_items(self):
        """Los Work Items como datos: ``[{'ref', 'name', 'uom', 'product',
        'lines': [{'product', 'qty', 'cost', 'nature'}]}]``."""
        if 'x_work_item' not in self.env:
            raise UserError(self.env._(
                'Esta base no tiene los Work Items del paquete Construction de Odoo.'))
        WorkItem = self.env['x_work_item']
        if self.work_item_scope == 'project':
            if not self.project_id:
                raise UserError(self.env._('Elija la obra cuya venta trae los Work Items.'))
            domain = [('x_sale_order_line_id.order_id.project_id', '=', self.project_id.id)]
        else:
            domain = [('x_is_template', '=', True)]
        items = []
        for item in WorkItem.search(domain):
            product = item.x_product_id
            items.append({
                'ref': product.default_code or item.x_name, 'name': item.x_name,
                'uom': item.x_unit_custom_id or product.uom_id, 'product': product,
                'source_ref': 'x_work_item,%s' % item.id,
                'lines': [{'product': line.x_product_id, 'qty': line.x_quantity,
                           'cost': line.x_unit_cost, 'nature': line.x_cost_nature,
                           'uom': line.x_product_id.uom_id}
                          for line in item.x_work_item_line_ids],
            })
        return items

    def _recipes_from_work_items(self, items):
        """Recetas desde los Work Items leídos (``_read_work_items``)."""
        Type = self.env['construction.resource.type']
        stage = self.default_stage_id.code or False
        recipes, refs = {}, {}
        for item in items:
            lines = []
            for data in item['lines']:
                behavior = WORK_ITEM_NATURES.get(data.get('nature'), 'material')
                if behavior == 'labor':
                    behavior = self.labor_behavior
                rtype = Type._default_for(behavior)
                vals = {'resource_type_id': rtype.id, 'resource_type': behavior,
                        'qty': data['qty'], 'price_unit': data.get('cost') or 0.0,
                        'stage': stage, 'product_uom_id': data['uom'].id}
                # La forma de cálculo «%MO» del APU se conserva (CREVAL v1.0,
                # 6.2): el porcentaje de la mano de obra de la receta.
                if norm(data['uom'].name) in PCT_LABOR_UNITS:
                    vals.update(calc_method='pct_labor', calc_pct=data['qty'],
                                resource_type='equipment',
                                resource_type_id=Type._default_for('equipment').id)
                    behavior = 'equipment'
                if behavior in ('contract', 'labor'):
                    activity = self._activity_for(data['product'], data['uom'], stage)
                    vals['activity_id'] = activity.id
                else:
                    vals['product_id'] = data['product'].id
                lines.append(vals)
            recipes[item['ref']] = {'name': item['name'], 'uom': item['uom'],
                                    'product': item['product'], 'lines': lines}
            refs[item['ref']] = item.get('source_ref')
        return recipes, refs

    # ------------------------------------------------------------------
    # Acciones
    # ------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        messages, errors = [], []
        if self.source == 'work_item':
            recipes, refs = self._recipes_from_work_items(self._read_work_items())
            templates, created, updated = self._upsert_recipes(recipes, 'work_item', refs)
            # Vínculo con el Work Item para sincronizar en los dos sentidos
            # (v1.8, M18-13).
            for template in templates:
                model, _sep, item_id = (template.source_ref or '').partition(',')
                if model == 'x_work_item' and item_id.isdigit():
                    template._link_work_item(int(item_id))
            messages.append(self.env._('Recetas desde Work Items: %(new)s nuevas y %(upd)s '
                                       'actualizadas.', new=created, upd=updated))
        else:
            if not self.file:
                raise UserError(self.env._('Adjunte el libro Excel.'))
            sheets = self._read_workbook()
            if sheets.get('RECETAS'):
                recipes, recipe_errors = self._recipes_from_rows(sheets['RECETAS'])
                errors += recipe_errors
                refs = dict.fromkeys(recipes, self.filename or False)
                templates, created, updated = self._upsert_recipes(recipes, 'excel', refs)
                messages.append(self.env._('Recetas: %(new)s nuevas y %(upd)s actualizadas.',
                                           new=created, upd=updated))
            if sheets.get('NIVELES'):
                created, level_errors = self._import_levels(sheets['NIVELES'])
                errors += level_errors
                messages.append(self.env._('Niveles: %s tareas nuevas.', created))
            if sheets.get('RECETAS'):
                errors += self._resolve_calc_tasks(templates, recipes)
        body = Markup('').join(Markup('<p>%s</p>') % message for message in messages)
        if errors:
            body += Markup('<ul class="text-warning">%s</ul>') % Markup('').join(
                Markup('<li>%s</li>') % error for error in errors[:50])
        self.result_html = body
        return {
            'type': 'ir.actions.act_window', 'res_model': self._name, 'res_id': self.id,
            'view_mode': 'form', 'target': 'new', 'name': self.env._('Importar recetas y niveles'),
        }

    def action_download_template(self):
        """Libro vacío con las hojas RECETAS y NIVELES (columnas de la plantilla
        de niveles de la obra)."""
        self.ensure_one()
        from openpyxl import Workbook
        from openpyxl.styles import Font

        workbook = Workbook()
        recipes = workbook.active
        recipes.title = 'RECETAS'
        levels = workbook.create_sheet('NIVELES')
        headers = {
            recipes: [header for _key, header, _required in RECIPE_COLUMNS],
            levels: self._level_lines().mapped('name') + [h for _k, h in LEVEL_EXTRA_COLUMNS],
        }
        for sheet, names in headers.items():
            for col, header in enumerate(names, start=1):
                cell = sheet.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                sheet.column_dimensions[cell.column_letter].width = max(14, len(header) + 4)
        output = BytesIO()
        workbook.save(output)
        attachment = self.env['ir.attachment'].create({
            'name': 'recetas_y_niveles.xlsx', 'datas': base64.b64encode(output.getvalue()),
            'res_model': self._name, 'res_id': self.id,
        })
        return {'type': 'ir.actions.act_url', 'url': '/web/content/%s?download=true'
                % attachment.id, 'target': 'self'}
