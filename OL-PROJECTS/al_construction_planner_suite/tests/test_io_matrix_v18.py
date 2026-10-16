# -*- coding: utf-8 -*-
"""Panel de recursos por tipo (v1.8, M18-03) y tabla insumo-producto
(M18-04) sobre el piso 05 del demo."""
import base64
import io

from odoo.tests import tagged

from .test_schedule import TestSchedule


@tagged('post_install', '-at_install')
class TestIoMatrixV18(TestSchedule):

    def test_panel_by_resource_type(self):
        """Al marcar el piso, un bloque por tipo de recurso con su subtotal en
        el orden del catálogo, y los totales coinciden con el plan del piso."""
        summary = self.plan.get_selection_summary([str(self.floor.id)])
        groups = summary['groups']
        lines = self.plan.line_ids.filtered(lambda l: l.floor_task_id == self.floor)
        self.assertAlmostEqual(sum(g['amount'] for g in groups),
                               round(sum(lines.mapped('amount_planned')), 2), places=2)
        self.assertAlmostEqual(summary['total'], round(sum(lines.mapped('amount_planned')), 2))
        types = self.env['construction.resource.type'].browse([g['type_id'] for g in groups])
        self.assertEqual(types.mapped('sequence'), sorted(types.mapped('sequence')))
        material = next(g for g in groups if g['resource_type'] == 'material')
        contract = next(g for g in groups if g['resource_type'] == 'contract')
        for group in (material, contract):
            expected = lines.filtered(lambda l, g=group: l.resource_type_id.id == g['type_id'])
            self.assertAlmostEqual(group['amount'],
                                   round(sum(expected.mapped('amount_planned')), 2), places=2)
            self.assertAlmostEqual(sum(stage['amount'] for stage in group['stages']),
                                   group['amount'], places=2)
        # Con el filtro de etapas, el panel solo suma esas etapas.
        filtered = self.plan.get_selection_summary(
            [str(self.floor.id)], {'stages': ['furniture_installation']})
        for group in filtered['groups']:
            self.assertTrue(all(stage['stage'] in ('installation', 'finishing', 'cleaning')
                                for stage in group['stages']))

    def test_io_matrix_rows_columns_and_totals(self):
        options = self.plan.get_io_matrix_options()
        self.assertIn('space_task_id', [level['field'] for level in options['levels']])
        matrix = self.plan.get_io_matrix([str(self.floor.id)], {}, {'level': 'space_task_id'})
        lines = self.plan.line_ids.filtered(lambda l: l.floor_task_id == self.floor)
        self.assertEqual({row['id'] for row in matrix['rows']}, set(self.spaces.ids))
        self.assertAlmostEqual(matrix['total'], round(sum(lines.mapped('amount_planned')), 2),
                               places=2)
        self.assertAlmostEqual(sum(matrix['row_totals'].values()), matrix['total'], places=2)
        # Materiales de producción bajo «1 Producción › 1.1 Producción de
        # piezas», contratas de instalación bajo «2 Instalación › 2.1».
        by_stage = {}
        for column in matrix['columns']:
            by_stage.setdefault(column['stage'], set()).add(column['type_label'])
        self.assertIn('Material', by_stage['production'])
        self.assertIn('Contrata', by_stage['installation'])
        production = next(c for c in matrix['columns'] if c['stage'] == 'production')
        self.assertEqual(production['parent'], 'furniture_production')
        groups = {group['key']: group for group in matrix['groups']}
        self.assertEqual(groups['parent:furniture_production']['label'], '1 Producción')
        stage_group = groups['stage:production']
        self.assertEqual(stage_group['parent'], 'parent:furniture_production')
        # Las columnas van en el orden del árbol de etapas.
        order = [column['stage'] for column in matrix['columns']]
        self.assertLess(order.index('production'), order.index('installation'))
        # Una fila por departamento: mismas sumas.
        by_unit = self.plan.get_io_matrix([str(self.floor.id)], {},
                                          {'level': 'apartment_task_id'})
        self.assertEqual(len(by_unit['rows']), 8)
        self.assertAlmostEqual(by_unit['total'], matrix['total'], places=2)
        # En cantidad y lo comprometido.
        qty = self.plan.get_io_matrix([str(self.floor.id)], {},
                                      {'level': 'space_task_id', 'measure': 'qty'})
        self.assertAlmostEqual(qty['total'], round(sum(lines.mapped('qty_planned')), 4), places=2)
        committed = self.plan.get_io_matrix([str(self.floor.id)], {},
                                            {'level': 'space_task_id', 'value': 'committed'})
        self.assertAlmostEqual(committed['total'], 0.0)

    def test_io_matrix_excel(self):
        import openpyxl
        action = self.plan.action_export_io_matrix(
            [str(self.floor.id)], {}, {'level': 'space_task_id'})
        attachment = self.env['ir.attachment'].browse(int(action['url'].split('/')[-1].split('?')[0]))
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(attachment.datas)))
        sheet = book.active
        values = [cell.value for cell in sheet[1]]
        self.assertIn('1 Producción', values)
        last = sheet.max_row
        self.assertEqual(sheet.cell(last, 1).value, 'Total')
        matrix = self.plan.get_io_matrix([str(self.floor.id)], {}, {'level': 'space_task_id'})
        self.assertAlmostEqual(sheet.cell(last, sheet.max_column).value, matrix['total'], places=2)


for _name in dir(TestSchedule):
    if _name.startswith('test_') and callable(getattr(TestSchedule, _name)) \
            and _name not in vars(TestIoMatrixV18):
        setattr(TestIoMatrixV18, _name, None)
