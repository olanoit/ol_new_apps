# -*- coding: utf-8 -*-
"""Excel de carga de APU y presupuesto (requerimiento CREVAL v1.0, 6.6 y 6.8)
con datos de demostración: un muro de ladrillo con operario por cuadrilla,
herramientas al 5 % de la mano de obra y ladrillo a 0.9661."""
import base64
from io import BytesIO

from odoo.exceptions import UserError
from odoo.tests import tagged
from odoo.tests.common import TransactionCase
from odoo.tools import BinaryBytes


@tagged('post_install', '-at_install')
class TestBudgetBook(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        env = cls.env
        env.ref('product.decimal_price').digits = 4
        env['tier.definition'].search([('model', '=', 'construction.resource.plan')]).active = False
        cls.stage = env['construction.stage'].create({
            'name': 'Muros (test libro)', 'code': 'book_test_stage', 'sequence': 930})
        template = env['construction.level.template'].create({
            'name': 'Especialidad, sub-especialidad y partida (test libro)', 'line_ids': [
                (0, 0, {'sequence': 1, 'name': 'Especialidad', 'role': 'group'}),
                (0, 0, {'sequence': 2, 'name': 'Sub-especialidad', 'role': 'group'}),
                (0, 0, {'sequence': 3, 'name': 'Partida', 'role': 'space'})]})
        cls.project = env['project.project'].create({
            'name': 'OBR-03 Oficina de demostración (test libro)', 'is_construction_site': True,
            'construction_level_template_id': template.id, 'allow_billable': True,
            'construction_apu_source': 'excel'})
        if not cls.project.account_id:
            cls.project._create_analytic_account()
        cls.customer = env['res.partner'].create({'name': 'Cliente del libro (test)'})
        cls.order = env['sale.order'].create({'partner_id': cls.customer.id,
                                              'project_id': cls.project.id})

    def _book(self, qty_second=50.0):
        from openpyxl import Workbook
        # Los precios en la moneda de la compañía: el test no mide la conversión.
        currency = self.env.company.currency_id.name
        book = Workbook()
        sheet = book.active
        sheet.title = 'OBRA'
        for row in (('Obra', 'OBR-03'), ('Tipo de contrato', 'Suma alzada'),
                    ('Gastos generales %', 8), ('Utilidad %', 7)):
            sheet.append(row)
        sheet = book.create_sheet('INSUMOS')
        sheet.append(['Código', 'Descripción', 'Tipo de insumo', 'Unidad', 'Precio sin IGV',
                      'Moneda'])
        sheet.append(['BK-LAD', 'Ladrillo KK 18 huecos (test)', 'Material', 'und', 0.9661, currency])
        sheet.append(['BK-OPE', 'Operario (test)', 'Mano de obra', 'hh', 23.36, currency])
        sheet.append(['BK-HER', 'Herramientas manuales (test)', 'Herramienta %MO', '%MO', None,
                      currency])
        sheet = book.create_sheet('APU')
        sheet.append(['Código', 'Descripción', 'Unidad', 'Rendimiento por día',
                      'Horas por jornada', 'Modalidad por defecto', 'Etapa'])
        sheet.append(['BK-R01', 'Muro de ladrillo KK (test)', 'm2', 19, 8, 'Personal propio',
                      'book_test_stage'])
        sheet = book.create_sheet('APU_DETALLE')
        sheet.append(['Código del APU', 'Código del insumo', 'Cuadrilla', '% de la mano de obra',
                      'Cantidad por unidad', 'Precio manual'])
        sheet.append(['BK-R01', 'BK-OPE', 1, None, None, None])
        sheet.append(['BK-R01', 'BK-LAD', None, None, 54, None])
        sheet.append(['BK-R01', 'BK-HER', None, 5, None, None])
        sheet = book.create_sheet('PRESUPUESTO')
        sheet.append(['Ítem', 'Especialidad', 'Sub-especialidad', 'Partida', 'Código del APU',
                      'Unidad', 'Modalidad', 'Metrado real', 'Holgura de metrado (%)'])
        sheet.append(['01.01', 'Arquitectura', 'Muros', 'Muro del eje A', 'BK-R01', 'm2', None,
                      100, 10])
        sheet.append(['01.02', 'Arquitectura', 'Muros', 'Muro del eje B', 'BK-R01', 'm2', None,
                      qty_second, None])
        sheet = book.create_sheet('PLAN')
        sheet.append(['Ítem', 'Partida', 'Modalidad final', 'Contratista sugerido', 'Inicio',
                      'Fin'])
        sheet.append(['01.02', 'Muro del eje B', 'Personal propio', None, '2026-11-02',
                      '2026-11-20'])
        sheet = book.create_sheet('LISTAS')
        sheet.append(['Tipo de insumo', None, 'Modalidad', None, 'Escrito como', 'Unidad'])
        sheet.append(['Material', None, 'Personal propio', None, 'hh', 'Hours'])
        sheet.append([None, None, None, None, 'und', 'Units'])
        sheet.append([None, None, None, None, 'm2', 'm²'])
        output = BytesIO()
        book.save(output)
        return BinaryBytes(output.getvalue())

    def _wizard(self, target='budget', **vals):
        return self.env['construction.budget.book.wizard'].create(dict({
            'sale_order_id': self.order.id, 'target': target, 'file': self._book(),
            'filename': 'libro.xlsx'}, **vals))

    def test_budget_and_plan(self):
        wizard = self._wizard('plan')
        wizard.action_preview()
        self.assertIn('3 insumos, 1 APU y 2 partidas', wizard.preview_html)
        wizard.action_import()
        recipe = self.env['construction.execution.template'].search([('code', '=', 'BK-R01')])
        # Mano de obra 1 × 8 ÷ 19 × 23.36 = 9.8358; ladrillo 54 × 0.9661 =
        # 52.1694; herramientas 5 % de la mano de obra = 0.4918.
        unit = 8 / 19 * 23.36 * 1.05 + 54 * 0.9661
        self.assertAlmostEqual(recipe.amount_unit, round(unit, 2))
        lines = self.order.order_line.filtered(
            lambda l: not l.display_type and l.construction_line_role == 'partida')
        self.assertEqual(len(lines), 2)
        first = lines.filtered(lambda l: l.construction_budget_item == '01.01')
        self.assertAlmostEqual(first.construction_qty_real, 100.0)
        self.assertAlmostEqual(first.product_uom_qty, 110.0)
        self.assertAlmostEqual(first.construction_cost_apu, unit, places=4)
        self.assertEqual(self.order.order_line.filtered(
            lambda l: l.display_type == 'line_section').name, 'Arquitectura')
        overhead = self.order.order_line.filtered(lambda l: l.construction_line_role == 'overhead')
        self.assertAlmostEqual(overhead.price_unit, sum(lines.mapped('price_subtotal')) * 0.08,
                               places=2)
        self.assertEqual(self.order.construction_contract_type, 'lump_sum')
        # Plan: niveles con la receta y el metrado real, sin las líneas del pie.
        plan = self.env['construction.resource.plan'].search([('project_id', '=', self.project.id)])
        self.assertEqual(len(plan), 1)
        tasks = plan.line_ids.task_id
        self.assertEqual(len(tasks), 2)
        self.assertEqual(set(tasks.mapped('construction_qty')), {100.0, 50.0})
        second = tasks.filtered(lambda t: t.construction_qty == 50.0)
        self.assertEqual(str(second.date_deadline)[:10], '2026-11-20')
        self.assertEqual(set(plan.line_ids.mapped('resource_type')),
                         {'labor', 'material', 'equipment'})
        # Volver a enviar no duplica.
        count = len(self.order.order_line)
        plan_lines = len(plan.line_ids)
        again = self._wizard('plan')
        again.action_preview()
        again.action_import()
        self.assertEqual(len(self.order.order_line), count)
        self.assertEqual(len(plan.line_ids), plan_lines)
        self.assertEqual(self.env['construction.execution.template'].search_count(
            [('code', '=', 'BK-R01')]), 1)
        self.assertTrue(self.env['construction.uom.alias'].search([('name', '=', 'hh')]))

    def test_errors_block(self):
        wizard = self._wizard()
        content = wizard.file.content
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(content))
        book['APU_DETALLE'].append(['BK-R01', 'NO-EXISTE', None, None, 1, None])
        book['INSUMOS'].append(['BK-LAD', 'Repetido', 'Material', 'und', 1,
                                self.env.company.currency_id.name])
        output = BytesIO()
        book.save(output)
        wizard.file = BinaryBytes(output.getvalue())
        wizard.action_preview()
        self.assertIn('NO-EXISTE no existe en la hoja INSUMOS', wizard.preview_html)
        self.assertIn('BK-LAD está repetido', wizard.preview_html)
        with self.assertRaises(UserError):
            wizard.action_import()

    def _edited_book(self, edit, order=None):
        """El libro de ``_book`` cambiado por ``edit(book)``."""
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(self._book().content))
        edit(book)
        output = BytesIO()
        book.save(output)
        return BinaryBytes(output.getvalue())

    def _send(self, file, target='budget', order=None):
        wizard = self.env['construction.budget.book.wizard'].create({
            'sale_order_id': (order or self.order).id, 'target': target, 'file': file,
            'filename': 'libro.xlsx'})
        wizard.action_preview()
        return wizard

    def test_partida_without_apu_needs_its_cost(self):
        """CR-18, D1: sin APU, el costo lo escribe el presupuestador en
        «Costo unitario sin APU»; sin él, error y no se envía."""
        def no_cost(book):
            book['PRESUPUESTO'].append(['01.03', 'Arquitectura', 'Muros', 'Limpieza final',
                                        None, 'und', None, 1])
        wizard = self._send(self._edited_book(no_cost))
        self.assertIn('Partida 01.03: sin código de APU ni costo unitario', wizard.preview_html)
        self.assertFalse(wizard.can_send)
        with self.assertRaises(UserError):
            wizard.action_import()

        def with_cost(book):
            sheet = book['PRESUPUESTO']
            sheet.cell(row=1, column=10, value='Costo unitario sin APU')
            sheet.append(['01.03', 'Arquitectura', 'Muros', 'Limpieza final', None, 'und', None,
                          1, None, 350.25])
        wizard = self._send(self._edited_book(with_cost))
        self.assertTrue(wizard.can_send)
        wizard.action_import()
        line = self.order.order_line.filtered(lambda l: l.construction_budget_item == '01.03')
        self.assertAlmostEqual(line.price_unit, 350.25)
        self.assertAlmostEqual(line.construction_cost_apu, 350.25)
        self.assertFalse(line.construction_execution_template_id)
        first = self.order.order_line.filtered(lambda l: l.construction_budget_item == '01.01')
        self.assertEqual(first.construction_execution_template_id.code, 'BK-R01')

    def test_odoo_unit_names_are_valid(self):
        """CR-18, b: el nombre de una unidad de Odoo vale sin equivalencia."""
        hour = self.env.ref('uom.product_uom_hour')

        def odoo_units(book):
            sheet = book['INSUMOS']
            sheet.cell(row=3, column=4, value=hour.with_context(lang='en_US').name)
            # Sin la equivalencia «hh» de LISTAS.
            book['LISTAS'].cell(row=2, column=5, value=None)
            book['LISTAS'].cell(row=2, column=6, value=None)
        wizard = self._send(self._edited_book(odoo_units))
        self.assertNotIn('no tiene equivalencia', wizard.preview_html or '')
        book, errors, _warnings = wizard._parse()
        self.assertEqual(errors, [])
        self.assertEqual(book['insumos']['BK-OPE']['uom'], hour)

    def test_zero_direct_cost_blocks_sending(self):
        """CR-18, e: con costo directo 0 no se envía."""
        def zero(book):
            sheet = book['PRESUPUESTO']
            sheet.delete_rows(2, 2)
            sheet.cell(row=1, column=10, value='Costo unitario sin APU')
            sheet.append(['01.09', 'Arquitectura', 'Muros', 'Partida sin costo', None, 'und',
                          None, 1, None, 0])
        wizard = self._send(self._edited_book(zero))
        self.assertIn('El costo directo de la hoja es 0', wizard.preview_html)
        self.assertFalse(wizard.can_send)

    def test_footer_recalculated_on_sending(self):
        """CR-18, D2: gastos generales y utilidad sobre las partidas; el
        descuento sobre partidas más ellos; cada monto redondeado una vez."""
        def footer(book):
            sheet = book['OBRA']
            sheet.append(['Descuento %', 24.963145])
            sheet.cell(row=3, column=2, value=10)
        self._send(self._edited_book(footer)).action_import()
        lines = self.order.order_line.filtered(lambda l: not l.display_type)
        role = {r: lines.filtered(lambda l, r=r: l.construction_line_role == r)
                for r in ('partida', 'overhead', 'profit', 'discount')}
        partidas = sum(role['partida'].mapped('price_subtotal'))
        currency = self.order.currency_id
        self.assertEqual(role['overhead'].price_unit, currency.round(partidas * 0.10))
        self.assertEqual(role['profit'].price_unit, currency.round(partidas * 0.07))
        before = partidas + role['overhead'].price_subtotal + role['profit'].price_subtotal
        self.assertEqual(role['discount'].price_unit, currency.round(-before * 0.24963145))
        self.assertIn('24.9631 %', role['discount'].name)
        self.assertAlmostEqual(self.order.amount_untaxed,
                               before + role['discount'].price_subtotal, places=2)

    def test_preview_lists_each_partida(self):
        """CR-18, e: la vista previa da precio y metrado actuales y nuevos;
        en rojo lo que baja a 0 o cambia más del porcentaje."""
        self._send(self._book()).action_import()

        def cheaper(book):
            book['PRESUPUESTO'].cell(row=3, column=8, value=10)
        wizard = self._send(self._edited_book(cheaper))
        self.assertIn('Muro del eje A', wizard.preview_html)
        self.assertIn('Muro del eje B', wizard.preview_html)
        self.assertIn('1 en rojo', wizard.preview_html)

    def test_confirmed_order_keeps_its_offer(self):
        """CR-18, D3: confirmada, «presupuesto» no se envía; «presupuesto y
        plan» actualiza costos y plan sin tocar precio ni metrado ofertado."""
        self._send(self._book()).action_import()
        self.order.action_confirm()
        lines = self.order.order_line.filtered(lambda l: not l.display_type)
        offer = {line: (line.price_unit, line.product_uom_qty) for line in lines}
        total = self.order.amount_untaxed
        wizard = self._send(self._book())
        self.assertIn('La cotización está confirmada', wizard.preview_html)
        self.assertFalse(wizard.can_send)

        def more(book):
            book['PRESUPUESTO'].cell(row=2, column=8, value=120)
        wizard = self._send(self._edited_book(more), target='plan')
        self.assertIn('Va por cambio de obra', wizard.preview_html)
        wizard.write({'replan_reason': 'Metrado medido (test)'})
        wizard.action_import()
        for line, (price, qty) in offer.items():
            self.assertEqual((line.price_unit, line.product_uom_qty), (price, qty))
        self.assertEqual(self.order.amount_untaxed, total)
        first = lines.filtered(lambda l: l.construction_budget_item == '01.01')
        self.assertEqual(first.construction_qty_real, 120.0)

    def test_quote_without_site_keeps_its_recipes(self):
        """CR-18, D4: sin obra, el envío guarda las recetas en la cotización
        (no en la biblioteca) y la obra las recibe al crear sus niveles."""
        Template = self.env['construction.execution.template']
        library = Template.create({'code': 'BK-R01', 'name': 'Receta de la biblioteca (test)',
                                   'company_id': self.env.company.id})
        order = self.env['sale.order'].create({'partner_id': self.customer.id})
        self._send(self._book(), order=order).action_import()
        own = order.construction_recipe_ids
        self.assertEqual(own.mapped('code'), ['BK-R01'])
        self.assertNotEqual(own, library)
        self.assertEqual(library.name, 'Receta de la biblioteca (test)')
        self.assertFalse(library.line_ids)
        site = self.env['project.project'].create({
            'name': 'OBR-04 Obra desde la cotización (test)', 'is_construction_site': True})
        order.project_id = site
        copy = Template.search([('project_id', '=', site.id), ('code', '=', 'BK-R01')])
        self.assertEqual(len(copy), 1)
        self.assertEqual(len(copy.line_ids), len(own.line_ids))
        partida = order.order_line.filtered(lambda l: l.construction_budget_item == '01.01')
        self.assertEqual(partida.construction_execution_template_id, copy)
        # Una sola vez: volver a asignar la obra no duplica.
        order.project_id = site
        self.assertEqual(Template.search_count(
            [('project_id', '=', site.id), ('code', '=', 'BK-R01')]), 1)

    def test_template(self):
        from openpyxl import load_workbook
        content = self.env['construction.budget.book.wizard']._template_workbook()
        book = load_workbook(BytesIO(content))
        self.assertEqual(book.sheetnames, ['INSTRUCCIONES', 'OBRA', 'INSUMOS', 'APU',
                                           'APU_DETALLE', 'PRESUPUESTO', 'LISTAS'])
