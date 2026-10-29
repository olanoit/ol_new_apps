# -*- coding: utf-8 -*-
from io import BytesIO

from odoo import models
from odoo.exceptions import UserError
from odoo.tools import BinaryBytes

from odoo.addons.al_construction_planner_furniture.wizards.master_import_layout import sheet_key


class DocumentsDocument(models.Model):
    _inherit = 'documents.document'

    def _construction_send_furniture(self, sheets):
        """Lleva las pestañas del maestro al asistente «Importar maestro y
        ETO» de la obra de la hoja (v1.5, M15-15)."""
        self.ensure_one()
        project = self.construction_project_id
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, rows in (sheets or {}).items():
            if not sheet_key(name):
                continue
            worksheet = workbook.create_sheet(name[:31])
            for row in rows or []:
                worksheet.append([None if value in ('', False) else value for value in row])
        if not workbook.worksheets:
            raise UserError(self.env._('La hoja no tiene ninguna pestaña del maestro.'))
        output = BytesIO()
        workbook.save(output)
        wizard = self.env['construction.master.import.wizard'].create({
            'project_id': project.id,
            'file': BinaryBytes(output.getvalue()),
            'filename': '%s.xlsx' % self.name,
            # Los productos nuevos de la hoja (PRODUCTOS NUEVOS) se crean.
            'create_missing_products': True,
        })
        return {
            'type': 'ir.actions.act_window',
            'name': self.env._('Importar maestro y ETO'),
            'res_model': wizard._name,
            'res_id': wizard.id,
            'view_mode': 'form',
            'views': [(False, 'form')],
            'target': 'new',
        }
