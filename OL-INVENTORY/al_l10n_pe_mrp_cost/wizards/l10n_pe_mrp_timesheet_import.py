# -*- coding: utf-8 -*-
import io
from datetime import date, datetime, time, timedelta

import pytz

from odoo import _, fields, models
from odoo.exceptions import UserError

PE_TZ = pytz.timezone('America/Lima')
#: Columnas del Excel (la primera fila son los títulos).
COLUMNS = ('orden', 'operacion', 'empleado', 'fecha', 'inicio', 'horas')


class L10nPeMrpTimesheetImport(models.TransientModel):
    """Importa el tareo de planta como registros de tiempo de las órdenes.

    Columnas: Orden | Operación | Empleado | Fecha | Inicio | Horas.

    * Orden: referencia de la orden de fabricación (WH/MO/00012).
    * Operación: nombre de la operación de la orden o de su centro de
      producción; vacía si la orden tiene una sola.
    * Empleado: DNI (número de identificación), código de barras o nombre.
    * Inicio: hora (opcional, 08:00 si falta). Horas: decimal (1,5 = 1 h 30).

    Todo o nada: si una fila tiene un error no se importa ninguna.
    """
    _name = 'l10n_pe.mrp.timesheet.import'
    _description = 'Importar tareo de fabricación'

    file = fields.Binary(string='Archivo Excel', required=True)
    filename = fields.Char()
    company_id = fields.Many2one('res.company', required=True, default=lambda self: self.env.company)

    def action_import(self):
        self.ensure_one()
        rows = self._l10n_pe_read_rows()
        errors, vals_list = [], []
        productive = self.env['mrp.workcenter.productivity.loss'].search(
            [('loss_type', '=', 'productive')], limit=1)
        for number, row in rows:
            try:
                vals_list.append(self._l10n_pe_row_vals(row, productive))
            except UserError as error:
                errors.append(_('Fila %(row)s: %(error)s', row=number, error=error.args[0]))
        if errors:
            raise UserError('\n'.join(errors))
        if not vals_list:
            raise UserError(_('El archivo no tiene filas de tareo.'))
        self.env['mrp.workcenter.productivity'].create(vals_list)
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'type': 'success',
                'message': _('%s registros de tiempo importados.', len(vals_list)),
                'next': {'type': 'ir.actions.act_window_close'},
            },
        }

    def _l10n_pe_read_rows(self):
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(io.BytesIO(self.file.content), read_only=True, data_only=True)
        except Exception as error:
            raise UserError(_('No se pudo leer el Excel: %s', error)) from error
        sheet = workbook.worksheets[0]
        rows = []
        for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value not in (None, '') for value in values):
                continue
            rows.append((number, dict(zip(COLUMNS, list(values) + [None] * len(COLUMNS)))))
        return rows

    def _l10n_pe_row_vals(self, row, productive):
        production = self.env['mrp.production'].search([
            ('name', '=', str(row['orden'] or '').strip()), ('company_id', '=', self.company_id.id)], limit=1)
        if not production:
            raise UserError(_('no existe la orden «%s».', row['orden']))
        workorder = self._l10n_pe_workorder(production, str(row['operacion'] or '').strip())
        employee = self._l10n_pe_employee(str(row['empleado'] or '').strip())
        day = row['fecha']
        if isinstance(day, datetime):
            day = day.date()
        if not isinstance(day, date):
            try:
                day = datetime.strptime(str(day).strip(), '%d/%m/%Y').date()
            except ValueError:
                raise UserError(_('fecha «%s» no válida (dd/mm/aaaa).', row['fecha'])) from None
        start = row['inicio'] or time(8, 0)
        if isinstance(start, datetime):
            start = start.time()
        if not isinstance(start, time):
            try:
                start = datetime.strptime(str(start).strip(), '%H:%M').time()
            except ValueError:
                raise UserError(_('hora de inicio «%s» no válida (hh:mm).', row['inicio'])) from None
        try:
            hours = float(str(row['horas']).replace(',', '.'))
        except (TypeError, ValueError):
            raise UserError(_('horas «%s» no válidas.', row['horas'])) from None
        if hours <= 0 or hours > 24:
            raise UserError(_('las horas deben estar entre 0 y 24.'))
        date_start = PE_TZ.localize(datetime.combine(day, start)).astimezone(pytz.utc).replace(tzinfo=None)
        return {
            'workorder_id': workorder.id,
            'workcenter_id': workorder.workcenter_id.id,
            'employee_id': employee.id,
            'company_id': production.company_id.id,
            'loss_id': productive.id,
            'date_start': date_start,
            'date_end': date_start + timedelta(hours=hours),
            'description': _('Tareo importado'),
        }

    def _l10n_pe_workorder(self, production, operation):
        workorders = production.workorder_ids
        if operation:
            workorders = workorders.filtered(
                lambda w: operation.lower() in (w.name.lower(), w.workcenter_id.name.lower(),
                                                (w.workcenter_id.code or '').lower()))
        if len(workorders) != 1:
            raise UserError(_('la orden %(mo)s no tiene una sola operación «%(op)s».',
                              mo=production.name, op=operation or _('(vacía)')))
        return workorders

    def _l10n_pe_employee(self, key):
        # sudo: el DNI y el código de barras son campos privados del empleado
        # (grupo de Empleados) y quien importa el tareo es de Fabricación;
        # solo se usan para encontrar al empleado y no se devuelven.
        employees_sudo = self.env['hr.employee'].sudo()
        company = [('company_id', 'in', (False, self.company_id.id))]
        employee_sudo = employees_sudo.search(
            company + ['|', ('identification_id', '=', key), ('barcode', '=', key)], limit=2)
        if not employee_sudo:
            employee_sudo = employees_sudo.search(company + [('name', '=ilike', key)], limit=2)
        if len(employee_sudo) != 1:
            raise UserError(_('empleado «%s» no encontrado o repetido.', key))
        return self.env['hr.employee'].browse(employee_sudo.id)
